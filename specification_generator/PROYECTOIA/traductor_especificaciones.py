"""Excel editable, instantánea válida y consultas de traducción en memoria."""
from __future__ import annotations

import hashlib
import io
import json
import os
import re
import tempfile
import threading
import unicodedata
from datetime import datetime, timedelta, timezone
from functools import lru_cache
from pathlib import Path
from rutas_listas import carpeta_listas
from rutas_proyecto import resolver_ruta

BASE = carpeta_listas()
ARCHIVO_EXCEL = "TRADUCCIONES_ESPECIFICACIONES.xlsx"
ARCHIVO_ESTADO = ".estado_traductor.json"
ARCHIVO_HISTORIAL = "CAMBIOS_TRADUCTOR.txt"


def clave(texto):
    return " ".join(unicodedata.normalize("NFC", texto).upper().split())


class DiccionarioTraductor:
    def __init__(self, registros):
        mapas = {"es": {}, "en": {}}
        for fila in registros:
            if not fila["activo"]:
                continue
            mapas["es"][clave(fila["es"])] = fila["en"]
            # Las variantes comparten inglés; la primera define la inversa.
            mapas["en"].setdefault(clave(fila["en"]), fila["es"])
        self.mapas = mapas
        self.patrones = {}
        for idioma, mapa in mapas.items():
            palabras = sorted(mapa, key=lambda p: (-len(p), p))
            alternativas = [r"\s+".join(re.escape(parte) for parte in p.split()) for p in palabras]
            self.patrones[idioma] = re.compile(r"(?<!\w)(?:" + "|".join(alternativas) + r")(?!\w)", re.I) if palabras else None
        self.traducir = lru_cache(maxsize=4096)(self._traducir)

    def _traducir(self, texto, origen, destino):
        if origen not in self.mapas or destino not in self.mapas:
            raise ValueError("El idioma debe ser es o en.")
        if origen == destino or not texto.strip():
            return texto, True
        patron = self.patrones[origen]
        # Una sola pasada sobre el original: ninguna traducción se retraduce.
        salida = patron.sub(lambda m: self.mapas[origen][clave(m.group())], texto) if patron else texto
        resto = patron.sub("", texto) if patron else texto
        resto = re.sub(r"\b(?:ASTM\s+[A-Z]\s*\d+(?:-\d+)?|INS-[A-Z0-9-]+|N/A|NA)\b", "", resto, flags=re.I)
        seguro = re.fullmatch(r"[\d\s.,+−\-–—/<>=≤≥±%():;°]*", resto) is not None
        return salida, seguro


def validar_registros(registros):
    if not isinstance(registros, list):
        raise ValueError("El diccionario debe contener una lista de equivalencias.")
    encontrados = {}
    for numero, fila in enumerate(registros, 2):
        if not isinstance(fila, dict) or not all(isinstance(fila.get(campo), str) and fila[campo].strip() for campo in ("es", "en")) or not isinstance(fila.get("activo"), bool):
            raise ValueError(f"Fila {numero}: completa ESPAÑOL, INGLÉS y ACTIVO.")
        fuente = clave(fila["es"])
        if fuente in encontrados:
            raise ValueError(f"Fila {numero}: ESPAÑOL duplicado ({fila['es']}); primera aparición en fila {encontrados[fuente]}.")
        encontrados[fuente] = numero
    return registros


def leer_excel(contenido):
    from openpyxl import load_workbook
    libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=False)
    try:
        hoja = libro["Traducciones"]
        filas = hoja.iter_rows(values_only=True, max_col=4)
        cabecera = next(filas, ())
        if tuple(clave(str(c or "")) for c in cabecera) != ("ESPAÑOL", "INGLÉS", "ACTIVO", "OBSERVACIONES"):
            raise ValueError("La hoja Traducciones debe tener ESPAÑOL, INGLÉS, ACTIVO y OBSERVACIONES.")
        registros = []
        primeras_filas = {}
        for numero, valores in enumerate(filas, 2):
            if all(v is None or str(v).strip() == "" for v in valores):
                continue
            es, en, activo, observaciones = valores
            if not isinstance(es, str) or not isinstance(en, str) or es.startswith("=") or en.startswith("="):
                raise ValueError(f"Fila {numero}: las traducciones deben ser texto, sin fórmulas.")
            if not es.strip() or not en.strip():
                raise ValueError(f"Fila {numero}: completa ESPAÑOL e INGLÉS.")
            fuente = clave(es)
            if fuente in primeras_filas:
                raise ValueError(f"Fila {numero}: ESPAÑOL duplicado ({es}); primera aparición en fila {primeras_filas[fuente]}.")
            primeras_filas[fuente] = numero
            activado = clave(str(activo or ""))
            if activado not in {"SI", "SÍ", "NO", "TRUE", "FALSE", "1", "0"}:
                raise ValueError(f"Fila {numero}: ACTIVO debe ser SI o NO.")
            registros.append({"es": es.strip(), "en": en.strip(), "activo": activado in {"SI", "SÍ", "TRUE", "1"}, "observaciones": str(observaciones or "").strip()})
        return validar_registros(registros)
    except KeyError as exc:
        raise ValueError("Falta la hoja Traducciones en el Excel.") from exc
    finally:
        libro.close()


def diferencias(anteriores, actuales):
    antes = {clave(r["es"]): r for r in anteriores}
    despues = {clave(r["es"]): r for r in actuales}
    cambios = []
    for k in despues.keys() - antes.keys():
        cambios.append(f"AGREGADA: {despues[k]['es']} → {despues[k]['en']}")
    for k in antes.keys() - despues.keys():
        cambios.append(f"ELIMINADA: {antes[k]['es']} → {antes[k]['en']}")
    for k in antes.keys() & despues.keys():
        if antes[k] != despues[k]:
            cambios.append(f"MODIFICADA: {antes[k]} → {despues[k]}")
    if anteriores != actuales and not cambios:
        cambios.append("ORDEN MODIFICADO: actualizado el orden de preferencia para las traducciones inversas.")
    return sorted(cambios)


class ServicioTraductor:
    def __init__(self, carpeta=BASE):
        self.carpeta = resolver_ruta(carpeta)
        self.lock = threading.RLock()
        self.estado = {}
        try:
            estado = json.loads((self.carpeta / ARCHIVO_ESTADO).read_text(encoding="utf-8"))
            if isinstance(estado, dict) and estado.get("version") == 1:
                validar_registros(estado["registros"])
                self.estado = estado
        except (OSError, ValueError, KeyError, TypeError):
            pass
        self.diccionario = DiccionarioTraductor(self.estado.get("registros", []))

    def configurar_carpeta(self, carpeta):
        carpeta = resolver_ruta(carpeta)
        with self.lock:
            if carpeta != self.carpeta.resolve():
                nuevo = ServicioTraductor(carpeta)
                self.carpeta = carpeta
                self.estado = nuevo.estado
                if nuevo.estado:
                    self.diccionario = nuevo.diccionario

    def actualizar(self):
        """Ejecutar en un hilo de trabajo; activa el nuevo diccionario al terminar."""
        with self.lock:
            contenido = (self.carpeta / ARCHIVO_EXCEL).read_bytes()
            huella = hashlib.sha256(contenido).hexdigest()
            if huella == self.estado.get("sha256_excel"):
                return {"cambios": 0, "equivalencias": len(self.diccionario.mapas["es"])}
            registros = leer_excel(contenido)
            nuevo = DiccionarioTraductor(registros)
            cambios = diferencias(self.estado.get("registros", []), registros)
            fecha = datetime.now(timezone(timedelta(hours=-6))).isoformat(timespec="seconds")
            estado = {"version": 1, "sha256_excel": huella, "ultima_comprobacion": fecha,
                      "ultima_modificacion": fecha if cambios else self.estado.get("ultima_modificacion", fecha),
                      "registros": registros, "ultimos_cambios": cambios if cambios else self.estado.get("ultimos_cambios", [])}
            archivos = {self.carpeta / ARCHIVO_ESTADO: json.dumps(estado, ensure_ascii=False, indent=2) + "\n"}
            if cambios:
                historial = self.carpeta / ARCHIVO_HISTORIAL
                previo = historial.read_text(encoding="utf-8") if historial.exists() else ""
                archivos[historial] = previo + f"\n[{fecha}] {len(cambios)} cambios\n" + "\n".join(cambios) + "\n"
            self._guardar(archivos)
            self.estado = estado
            self.diccionario = nuevo
            return {"cambios": len(cambios), "equivalencias": len(nuevo.mapas["es"])}

    def guardar_registros(self, registros, huella_inicial):
        """Guarda la edición del panel Métodos y actualiza el estado e historial."""
        from openpyxl import load_workbook
        validar_registros(registros)
        with self.lock:
            ruta = self.carpeta / ARCHIVO_EXCEL
            anterior = ruta.read_bytes()
            if hashlib.sha256(anterior).hexdigest() != huella_inicial:
                raise ValueError("El Excel cambió desde que abriste el editor. Cierra y vuelve a abrir la lista para incorporar esos cambios.")
            libro = load_workbook(io.BytesIO(anterior))
            try:
                hoja = libro["Traducciones"]
                if hoja.max_row > 1:
                    hoja.delete_rows(2, hoja.max_row - 1)
                for fila in registros:
                    hoja.append([fila["es"], fila["en"], "SI" if fila["activo"] else "NO", fila.get("observaciones", "")])
                hoja.auto_filter.ref = hoja.dimensions
                salida = io.BytesIO()
                libro.save(salida)
            finally:
                libro.close()
            # Valida el Excel producido antes de reemplazar el archivo editable.
            leer_excel(salida.getvalue())
            self._guardar({ruta: salida.getvalue()})
            try:
                resultado = self.actualizar()
                return {**resultado, "sha256_excel": self.estado["sha256_excel"]}
            except Exception:
                self._guardar({ruta: anterior})
                raise

    @staticmethod
    def _guardar(archivos):
        temporales, respaldos, escritos = {}, {}, []
        try:
            for ruta, texto in archivos.items():
                respaldos[ruta] = ruta.read_bytes() if ruta.exists() else None
                with tempfile.NamedTemporaryFile(mode="wb", dir=ruta.parent, delete=False, prefix=".traductor_", suffix=".tmp") as archivo:
                    temporales[ruta] = Path(archivo.name)
                    archivo.write(texto.encode("utf-8") if isinstance(texto, str) else texto)
            for ruta, temporal in temporales.items():
                os.replace(temporal, ruta)
                escritos.append(ruta)
        except Exception:
            for ruta in reversed(escritos):
                if respaldos[ruta] is None:
                    ruta.unlink(missing_ok=True)
                else:
                    ruta.write_bytes(respaldos[ruta])
            raise
        finally:
            for temporal in temporales.values():
                temporal.unlink(missing_ok=True)


SERVICIO = ServicioTraductor(carpeta_listas())


def traducir(texto, origen, destino):
    return SERVICIO.diccionario.traducir(texto, origen, destino)


def actualizar_traductor():
    SERVICIO.configurar_carpeta(carpeta_listas())
    return SERVICIO.actualizar()
