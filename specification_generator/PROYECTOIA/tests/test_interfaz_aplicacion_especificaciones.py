import unittest
import json
import queue
import sqlite3
import tempfile
from datetime import date
from pathlib import Path
from unittest.mock import Mock, patch

from docx import Document

from generador_especs_hii import (
    _escribir_notas,
    _llenar_pruebas,
    _separar_notas_de_pruebas,
)
from interfaz_aplicacion_especificaciones import (
    CAMPO_PERIODICIDAD,
    GeneradorEspecificacion,
    LEYENDA_PERIODICIDAD,
    NOTAS_PLANTILLA,
    VALORES_ELABORO_APROBO,
    _normalizar_periodicidad,
    _nombres_maestros_para_metodo,
    _opciones_maestro,
    _part_id_canonico,
    _part_id_existe,
    _part_id_por_nombre,
    _prueba_cercana,
    _sugerencias_nombre_producto,
    _sugerencias_part_id,
    _traducir_nombre_prueba,
    _unidad_sugerida,
    _valores_control_predeterminados,
    _valor_reanalisis_editor,
    _valor_reanalisis_json,
    _valores_prueba_predeterminados,
    _SalidaCarga,
    _procesar_carga_especificaciones,
    _datos_generador_desde_json,
    _integrar_edicion_json,
    _actualizar_comentarios_celda,
    _agregar_comentario_a_celda,
    _guardar_celda_con_comentario,
    crear_documento_alineado_manual,
    crear_datos_especificacion_nueva,
    crear_json_temporal_comparacion,
    _leer_filas_maestro_editable,
    _guardar_cambios_maestro,
    _consultar_tabla_historico,
    _exportar_tabla_historico_excel,
    alinear_filas_generador,
)


class RevisionAutomaticaTests(unittest.TestCase):
    def editor(self):
        editor = GeneradorEspecificacion.__new__(GeneradorEspecificacion)
        editor.datos_originales = {"control_documental": {"fecha": "24-02-12"}}
        editor._originales_idioma = {}
        editor._contenido_inicial_revision = ("original",)
        editor._contenido_para_revision = Mock(return_value=("editado",))
        editor._revision_inicial = "01"
        editor._revision_aplicada = False
        editor._cambio_alineacion_revision = False
        editor._resultados_revision = queue.Queue()
        editor._consultas_revision = {"P1"}
        editor._revisiones_historial = {"P1": "08"}
        editor.variables_control = {
            campo: Mock(get=Mock(return_value=valor))
            for campo, valor in (("codigo_unico_producto", "P1"), ("revision", "01"), ("fecha", "24-02-12"))}
        return editor

    def test_modificar_actualiza_revision_y_fecha_del_sistema_una_sola_vez(self):
        editor = self.editor()
        with patch("interfaz_aplicacion_especificaciones.date") as reloj:
            reloj.today.return_value = date(2026, 10, 14)
            self.assertTrue(editor._actualizar_revision_automatica())
            reloj.today.return_value = date(2026, 10, 15)
            self.assertTrue(editor._actualizar_revision_automatica())
        editor.variables_control["revision"].set.assert_called_once_with("08")
        editor.variables_control["fecha"].set.assert_called_once_with("14/10/26")
        self.assertEqual(editor.datos_originales["control_documental"]["fecha"], "24-02-12")

    def test_sin_cambios_o_especificacion_nueva_conserva_fecha_y_revision(self):
        for nueva in (False, True):
            with self.subTest(nueva=nueva):
                editor = self.editor()
                if nueva:
                    editor.datos_originales = None
                    editor._contenido_para_revision.return_value = editor._contenido_inicial_revision
                else:
                    editor._contenido_para_revision.return_value = editor._contenido_inicial_revision
                self.assertTrue(editor._actualizar_revision_automatica())
                editor.variables_control["revision"].set.assert_not_called()
                editor.variables_control["fecha"].set.assert_not_called()

    def test_fecha_cambia_cuando_llega_el_resultado_del_historial(self):
        editor = self.editor()
        editor._revisiones_historial.clear()
        self.assertFalse(editor._actualizar_revision_automatica())
        editor.variables_control["fecha"].set.assert_not_called()
        editor._resultados_revision.put(("P1", "08"))
        with patch("interfaz_aplicacion_especificaciones.date") as reloj:
            reloj.today.return_value = date(2026, 10, 14)
            self.assertTrue(editor._actualizar_revision_automatica())
        editor.variables_control["fecha"].set.assert_called_once_with("14/10/26")


class DatosNuevaEspecificacionTests(unittest.TestCase):
    def test_reanalisis_guarda_meses_y_reabre_sin_duplicar_unidad(self):
        for entrada in ("06", "06 meses", " 06 MESES ", "06 months", " 06 MONTHS "):
            with self.subTest(entrada=entrada):
                self.assertEqual(_valor_reanalisis_json(entrada), "06 meses")
                self.assertEqual(_valor_reanalisis_json(entrada, "en"), "06 months")
                self.assertEqual(_valor_reanalisis_editor(entrada), "06")
                contenido = _datos_generador_desde_json({"FECHA DE REANALISIS": entrada})
                self.assertEqual(contenido["fecha_reanalisis"], "06")
        self.assertEqual(_valor_reanalisis_json(""), "")

    def setUp(self):
        self.maestro = [
            {
                "METODO": "ASTM D445 / ASTM D7042",
                "NOMBRE IQM": "Viscosidad cinemática",
                "NOMBRE IQM INGLES": "Kinematic viscosity",
                "UNIDADES": "mm²/s",
                "UNITS": "mm2/s",
            },
            {
                "METODO": "ASTM D92",
                "NOMBRE IQM": "Punto de inflamación",
                "NOMBRE IQM INGLES": "Flash point",
                "UNIDADES": "°C",
                "UNITS": "°F",
            },
            {
                "METODO": "ASTM D92",
                "NOMBRE IQM": "Prueba de temperatura",
                "NOMBRE IQM INGLES": "Temperature test",
                "UNIDADES": "°C",
                "UNITS": "°F",
            },
        ]

    def test_sugiere_nombres_unidades_y_metodos_de_la_lista_segun_idioma(self):
        self.assertEqual(
            _opciones_maestro(self.maestro, "prueba", "es", "visco"),
            ("Viscosidad cinemática",),
        )
        self.assertEqual(
            _opciones_maestro(self.maestro, "prueba", "en", "kinematic"),
            ("Kinematic viscosity",),
        )
        self.assertEqual(
            _opciones_maestro(self.maestro, "unidades", "es", "°"),
            ("°C",),
        )
        self.assertEqual(
            _opciones_maestro(self.maestro, "unidades", "en", "°"),
            ("°F",),
        )
        self.assertEqual(
            _opciones_maestro(self.maestro, "metodo", "es", "ASTM D445"),
            ("ASTM D445 / ASTM D7042",),
        )

    def test_sugiere_el_metodo_maestro_en_ingles_para_pruebas_en_ingles(self):
        maestro_bilingue = [
            {
                "METODO": "MÉTODO INTERNO",
                "MÉTODO INGLÉS": "INTERNAL METHOD",
                "NOMBRE IQM": "Método interno",
                "NOMBRE IQM INGLES": "Internal test",
                "UNIDADES": "mg/kg",
                "UNITS": "mg/kg",
            }
        ]

        self.assertEqual(
            _opciones_maestro(maestro_bilingue, "metodo", "en", "internal"),
            ("INTERNAL METHOD",),
        )
        self.assertEqual(
            _opciones_maestro(maestro_bilingue, "metodo", "es", "interno"),
            ("MÉTODO INTERNO",),
        )
        self.assertEqual(
            _valores_prueba_predeterminados(
                maestro_bilingue, "Internal test", "en"
            ),
            {"metodo": "INTERNAL METHOD", "unidades": "mg/kg"},
        )

    def test_astm_sugiere_el_nombre_maestro_en_el_idioma_elegido(self):
        self.assertEqual(
            _nombres_maestros_para_metodo(
                self.maestro, "ASTM D445", "es"
            ),
            ("Viscosidad cinemática",),
        )
        self.assertEqual(
            _nombres_maestros_para_metodo(
                self.maestro, "ASTM D445", "en"
            ),
            ("Kinematic viscosity",),
        )

    def test_no_ofrece_nombres_fuera_de_las_coincidencias_de_metodo(self):
        self.assertEqual(
            _opciones_maestro(
                self.maestro,
                "prueba",
                "en",
                "point",
                _nombres_maestros_para_metodo(
                    self.maestro, "ASTM D445", "en"
                ),
            ),
            (),
        )

    def test_metodo_astm_conjugado_requiere_coincidencia_completa(self):
        self.assertEqual(
            _opciones_maestro(
                self.maestro,
                "metodo",
                "es",
                "ASTM D445 / ASTM D7042",
            ),
            ("ASTM D445 / ASTM D7042",),
        )
        self.assertEqual(
            _nombres_maestros_para_metodo(
                self.maestro, "ASTM D445 / ASTM D7042", "es"
            ),
            ("Viscosidad cinemática",),
        )

    def test_sugerencias_grises_y_unidad_celsius(self):
        self.assertEqual(
            _prueba_cercana(self.maestro, "en", "flash poin"),
            "Flash point",
        )
        self.assertEqual(
            _unidad_sugerida(
                self.maestro,
                "Temperature test",
                "ASTM D92",
                "en",
            ),
            "°C",
        )

    def test_prueba_exacta_propone_metodo_y_unidad_por_defecto(self):
        self.assertEqual(
            _valores_prueba_predeterminados(
                self.maestro, "Viscosidad cinemática", "es"
            ),
            {"metodo": "ASTM D445 / ASTM D7042", "unidades": "mm²/s"},
        )
        self.assertEqual(
            _valores_prueba_predeterminados(
                self.maestro, "Punto de inflamación", "es"
            ),
            {"metodo": "ASTM D92", "unidades": "°C"},
        )
        self.assertEqual(
            _valores_prueba_predeterminados(
                self.maestro, "Kinematic viscosity", "en"
            ),
            {"metodo": "ASTM D445 / ASTM D7042", "unidades": "mm2/s"},
        )
        self.assertEqual(
            _valores_prueba_predeterminados(
                self.maestro, "Flash point", "en"
            ),
            {"metodo": "ASTM D92", "unidades": "°C"},
        )
        self.assertEqual(
            _valores_prueba_predeterminados(
                self.maestro, "Nombre libre", "es"
            ),
            {},
        )

    def test_carga_ejecuta_extraccion_antes_de_alineacion(self):
        orden = []
        with (
            patch(
                "gestor_especificaciones.procesar",
                side_effect=lambda base: orden.append(
                    ("extraccion", base)
                ) or {"nuevos": 1},
            ),
            patch(
                "alinear_especificaciones_json.ejecutar",
                side_effect=lambda base: orden.append(
                    ("alineacion", base)
                ) or {"archivos": 1},
            ),
        ):
            resultado = _procesar_carga_especificaciones(Path("."))

        self.assertEqual(
            [registro[0] for registro in orden],
            ["extraccion", "alineacion"],
        )
        self.assertEqual(resultado["extraccion"], {"nuevos": 1})
        self.assertEqual(resultado["alineacion"], {"archivos": 1})

    def test_procesamiento_envia_a_ver_mas_mensajes_del_proceso(self):
        mensajes = []
        with (
            patch(
                "gestor_especificaciones.procesar",
                side_effect=lambda _base: print("PROCESANDO: archivo.docx")
                or {"nuevos": 1},
            ),
            patch(
                "alinear_especificaciones_json.ejecutar",
                side_effect=lambda _base: print("ALINEADO: archivo.json")
                or {"archivos": 1},
            ),
        ):
            resultado = _procesar_carga_especificaciones(
                Path("."),
                al_mensaje=mensajes.append,
            )

        self.assertEqual(resultado["extraccion"], {"nuevos": 1})
        self.assertEqual(resultado["alineacion"], {"archivos": 1})
        self.assertIn("Iniciando extracción de especificaciones.", mensajes)
        self.assertIn("PROCESANDO: archivo.docx", mensajes)
        self.assertIn("Iniciando alineación de especificaciones.", mensajes)
        self.assertIn("ALINEADO: archivo.json", mensajes)

    def test_salida_carga_emite_lineas_completas_y_residuo(self):
        mensajes = []
        salida = _SalidaCarga(mensajes.append, None)

        self.assertEqual(
            salida.write("línea uno\nlínea"),
            len("línea uno\nlínea"),
        )
        salida.write(" dos")
        salida.flush()

        self.assertEqual(mensajes, ["línea uno", "línea dos"])

    def test_falla_de_extraccion_no_omite_el_intento_de_alinear(self):
        with (
            patch(
                "gestor_especificaciones.procesar",
                side_effect=RuntimeError("falló extracción"),
            ),
            patch(
                "alinear_especificaciones_json.ejecutar",
                return_value={"archivos": 0},
            ) as alinear,
        ):
            resultado = _procesar_carga_especificaciones(Path("."))

        self.assertEqual(resultado["error_extraccion"], "falló extracción")
        alinear.assert_called_once_with(Path("."))

    def test_prepara_json_alineado_para_editar_con_valores_de_vista(self):
        origen = {
            "producto": {"nombre": "Aceite viejo", "nombre_fuente": "Aceite fuente"},
            "documento": {"idioma": "en", "fecha_documento_normalizada": "2026-10-01"},
            "control_documental": {"revision": "02", "codigo_unico_producto": "P-1"},
            "FECHA DE REANALISIS": "12/10/26",
            "pruebas": [
                {
                    "prueba": "Nombre sin limpiar",
                    "unidades": "unidad fuente",
                    "metodo": "método fuente",
                    CAMPO_PERIODICIDAD: "QIP",
                }
            ],
            "alineacion_metodos_lab": {
                "detalle": [
                    {
                        "valores_alineados": {
                            "prueba": "Nombre alineado",
                            "unidades": "mm²/s",
                            "metodo": "ASTM D445",
                            "especificacion": "10 MIN",
                            "valor_tipico": "12",
                        }
                    }
                ]
            },
            "notas": ["Nota existente"],
            "Desarrollador": "Equipo de Desarrollo",
        }

        cargados = _datos_generador_desde_json(origen)

        self.assertEqual(cargados["nombre_producto"], "Aceite viejo")
        self.assertEqual(cargados["idioma"], "en")
        self.assertEqual(cargados["control"]["revision"], "02")
        self.assertEqual(
            cargados["pruebas"],
            [
                {
                    "prueba": "Nombre alineado",
                    "unidades": "mm²/s",
                    "metodo": "ASTM D445",
                    "especificacion": "10 MIN",
                    "valor_tipico": "12",
                    "periodicidad": "QIP",
                }
            ],
        )
        self.assertEqual(cargados["notas"], ["Nota existente"])
        self.assertEqual(cargados["desarrollador"], "Equipo de Desarrollo")

    def test_integrar_edicion_preserva_metadatos_y_fecha_reanalisis(self):
        origen = {
            "producto": {"nombre": "Anterior", "nombre_fuente": "Original"},
            "documento": {
                "idioma": "es",
                "fecha_documento_normalizada": "2026-10-01",
            },
            "FECHA DE REANALISIS": "12/10/26",
            "campo_futuro": {"valor": 1},
            "Desarrollador": "Anterior",
        }
        actualizado = {
            "esquema": "especificacion_producto_es_v3",
            "producto": {"nombre": "Nuevo"},
            "documento": {
                "idioma": "en",
                "fecha_documento_normalizada": "",
            },
            "control_documental": {"revision": "03"},
            "pruebas": [],
            "notas": [],
            "alineacion_metodos_lab": {"detalle": []},
            "Desarrollador": "Nuevo",
        }

        resultado = _integrar_edicion_json(origen, actualizado)

        self.assertEqual(resultado["producto"]["nombre"], "Nuevo")
        self.assertEqual(resultado["producto"]["nombre_fuente"], "Original")
        self.assertEqual(resultado["documento"]["idioma"], "en")
        self.assertEqual(
            resultado["documento"]["fecha_documento_normalizada"],
            "2026-10-01",
        )
        self.assertEqual(resultado["FECHA DE REANALISIS"], "12/10/26")
        self.assertEqual(resultado["campo_futuro"], {"valor": 1})
        self.assertEqual(resultado["Desarrollador"], "Nuevo")

    def test_cliente_ausente_recibe_guion_y_su_borrado_se_conserva(self):
        for control in ({}, {"nombre_cliente": " ", "revision_cliente": None}):
            with self.subTest(control=control):
                original = {"control_documental": control, "FECHA DE REANALISIS": "12"}
                contenido = _datos_generador_desde_json(original)
                self.assertEqual(contenido["control"]["nombre_cliente"], "-")
                self.assertEqual(contenido["control"]["revision_cliente"], "-")
                self.assertEqual(contenido["fecha_reanalisis"], "12")
                contenido["control"].update(nombre_cliente="", revision_cliente="")
                actualizado = _integrar_edicion_json(original, {"control_documental": contenido["control"], "FECHA DE REANALISIS": "06"})
                self.assertEqual(actualizado["control_documental"]["nombre_cliente"], "")
                self.assertEqual(actualizado["control_documental"]["revision_cliente"], "")
                self.assertEqual(actualizado["FECHA DE REANALISIS"], "06")

    def test_traduce_nombres_de_prueba_al_cambiar_idioma(self):
        self.assertEqual(
            _traducir_nombre_prueba(
                self.maestro,
                "Punto de inflamación",
                "es",
                "en",
                "ASTM D92",
            ),
            "Flash point",
        )
        self.assertEqual(
            _traducir_nombre_prueba(
                self.maestro,
                "Kinematic viscosity",
                "en",
                "es",
                "ASTM D445 / ASTM D7042",
            ),
            "Viscosidad cinemática",
        )

    def test_campos_control_tienen_predeterminados_del_dia_y_guiones(self):
        valores = _valores_control_predeterminados(date(2026, 10, 12))
        self.assertEqual(valores["fecha"], "12/10/26")
        self.assertEqual(valores["folio_vap"], "-")
        self.assertEqual(valores["nombre_cliente"], "-")
        self.assertEqual(valores["revision_cliente"], "-")

    def test_periodicidad_requiere_de_uno_a_tres_simbolos_validos(self):
        self.assertEqual(_normalizar_periodicidad("◊, ▶ I", 1), "◊▶I")
        self.assertEqual(_normalizar_periodicidad("QIP", 1), "QIP")
        self.assertEqual(_normalizar_periodicidad("", 1), "▶")
        self.assertEqual(_normalizar_periodicidad("  ", 1), "▶")
        for valor in ("◊ ▶ C I", "◊ ▶ C I P", "X", "◊ ◊"):
            with self.subTest(valor=valor), self.assertRaises(ValueError):
                _normalizar_periodicidad(valor, 1)
        self.assertEqual(len(LEYENDA_PERIODICIDAD["es"]), 8)
        self.assertEqual(len(LEYENDA_PERIODICIDAD["en"]), 8)

    def test_sugerencias_part_id_limitadas_a_cinco_y_validacion_exacta(self):
        registros = [
            {"part_id": f"AB-{i}", "descripcion": f"Aceite {i}"}
            for i in range(8)
        ]
        self.assertEqual(len(_sugerencias_part_id(registros, "AB")), 5)
        self.assertEqual(_sugerencias_part_id(registros, "Aceite 7"), ("AB-7",))
        self.assertTrue(_part_id_existe("ab-2", registros))
        self.assertEqual(_part_id_canonico("ab-2", registros), "AB-2")
        self.assertFalse(_part_id_existe("aceite 2", registros))

    def test_sugerencias_y_resolucion_del_catalogo_producto_part_id(self):
        registros = [
            {"part_id": "P-100", "descripcion": "Aceite Hidráulico"},
            {"part_id": "P-200", "descripcion": "Aceite Hidráulico"},
            {"part_id": "P-300", "descripcion": "Grasa EP"},
        ]
        self.assertEqual(
            _sugerencias_nombre_producto(registros, "aceite"),
            ("Aceite Hidráulico",),
        )
        self.assertEqual(_part_id_por_nombre("Grasa EP", registros), "P-300")
        self.assertEqual(_part_id_por_nombre("Aceite Hidráulico", registros), "")
        self.assertEqual(_part_id_por_nombre("Producto libre", registros), "")

    def test_crea_json_espanol_con_detalle_y_campos_futuros_vacios(self):
        datos = crear_datos_especificacion_nueva(
            "Aceite de prueba",
            "es",
            {"revision": "01", "codigo_unico_producto": "P-100"},
            [
                {
                    "prueba": "Viscosidad",
                    "unidades": "mm²/s",
                    "metodo": "ASTM D445 / ASTM D7042",
                    "especificacion": "28,8 a 35,2",
                    "valor_tipico": "30,23",
                    "periodicidad": "◊, ▶",
                }
            ],
            ["Nota de prueba"],
            "Investigación y desarrollo",
        )

        self.assertEqual(datos["documento"]["idioma"], "es")
        self.assertEqual(datos["producto"]["nombre"], "Aceite de prueba")
        self.assertEqual(datos["FECHA DE REANALISIS"], "")
        self.assertEqual(datos["Desarrollador"], "Investigación y desarrollo")
        self.assertEqual(datos["pruebas"][0]["PERIODIDICAD DE ANALISIS"], "◊▶")
        self.assertEqual(
            datos["alineacion_metodos_lab"]["detalle"][0]["valores_alineados"][
                "metodo"
            ],
            "ASTM D445 / ASTM D7042",
        )
        self.assertEqual(datos["notas"], ["Nota de prueba"])
        self.assertEqual(
            datos["control_documental"]["elaboro_y_aprobo"],
            VALORES_ELABORO_APROBO["es"],
        )

    def test_crea_json_ingles_con_estructura_de_fila_compartida(self):
        datos = crear_datos_especificacion_nueva(
            "Test oil",
            "en",
            {"revision": "02", "codigo_unico_producto": "P-200"},
            [{"prueba": "Kinematic viscosity", "periodicidad": "P"}],
            [],
        )

        self.assertEqual(datos["documento"]["idioma"], "en")
        self.assertEqual(
            datos["notas"],
            [NOTAS_PLANTILLA["en"]],
        )
        self.assertEqual(
            datos["control_documental"]["elaboro_y_aprobo"],
            VALORES_ELABORO_APROBO["en"],
        )
        self.assertEqual(datos["pruebas"][0]["prueba"], "Kinematic viscosity")
        self.assertEqual(
            datos["alineacion_metodos_lab"]["detalle"][0]["valores_alineados"][
                "prueba"
            ],
            "Kinematic viscosity",
        )

    def test_crea_json_espanol_con_notas_de_plantilla_por_defecto(self):
        datos = crear_datos_especificacion_nueva(
            "Producto",
            "es",
            {},
            [{"prueba": "Color", "periodicidad": "I"}],
            [],
        )
        self.assertEqual(datos["notas"], [NOTAS_PLANTILLA["es"]])

    def test_el_generador_omite_filas_de_encabezado_en_espanol_e_ingles(self):
        filas = [
            {
                "prueba": "PRUEBA",
                "metodo": "MÉTODO",
                "especificacion": "ESPECIFICACIÓN",
            },
            {
                "prueba": "TEST",
                "metodo": "METHOD",
                "especificacion": "SPECIFICATION",
            },
            {
                "prueba": "PRUEBA MÉTODO",
                "metodo": "",
                "especificacion": "ESPECIFICACIÓN",
            },
            {
                "prueba": "Viscosidad",
                "metodo": "ASTM D445",
                "especificacion": "MÍNIMO",
            },
        ]

        validas, notas = _separar_notas_de_pruebas(filas)

        self.assertEqual([fila["prueba"] for fila in validas], ["Viscosidad"])
        self.assertEqual(notas, [])

    def test_exige_nombre_de_producto_y_al_menos_una_prueba_nombrada(self):
        with self.assertRaisesRegex(ValueError, "nombre del producto"):
            crear_datos_especificacion_nueva("", "es", {}, [{"prueba": "Color"}], [])
        with self.assertRaisesRegex(ValueError, "al menos una prueba"):
            crear_datos_especificacion_nueva("Producto", "es", {}, [], [])
        with self.assertRaisesRegex(ValueError, "nombre de la prueba"):
            crear_datos_especificacion_nueva(
                "Producto", "es", {}, [{"prueba": "", "unidades": "kg"}], []
            )
        resultado = crear_datos_especificacion_nueva(
            "Producto",
            "es",
            {},
            [
                {"prueba": "", "unidades": "", "metodo": ""},
                {"prueba": "Color", "periodicidad": "C"},
            ],
            [],
        )
        self.assertEqual(len(resultado["pruebas"]), 1)

    def test_alinear_generador_reusa_reglas_y_conserva_superindices(self):
        maestro = [
            {
                "METODO": "ASTM D445",
                "NOMBRE IQM": "Viscosidad cinemática",
                "NOMBRE IQM INGLES": "Kinematic viscosity",
                "UNIDADES": "cSt, mm²/s",
                "UNITS": "mm²/s, cSt",
            }
        ]
        pruebas = [
            {
                "prueba": "Viscosidad cinemática ¹",
                "unidades": "mm²/s",
                "metodo": "ASTM D445 ²",
                "especificacion": "30 max ³",
                "valor_tipico": "typical",
                "periodicidad": "▶",
            }
        ]

        alineadas, conteo = alinear_filas_generador(pruebas, maestro, "es")

        self.assertEqual(conteo, {"alineados": 1, "revision": 0, "sin_coincidencia": 0})
        self.assertEqual(alineadas[0]["prueba"], "Viscosidad cinemática ¹")
        self.assertEqual(alineadas[0]["metodo"], "ASTM D445 ²")
        self.assertEqual(alineadas[0]["unidades"], "mm²/s")
        self.assertEqual(alineadas[0]["especificacion"], "30 MÁXIMO ³")
        self.assertEqual(alineadas[0]["valor_tipico"], "TYPICAL")
        self.assertEqual(alineadas[0]["periodicidad"], "▶")
        self.assertEqual(pruebas[0]["especificacion"], "30 max ³")

    def test_alinear_generador_no_cuenta_filas_vacias_con_periodicidad(self):
        fila_vacia = {
            "prueba": "",
            "unidades": "",
            "metodo": "",
            "especificacion": "",
            "valor_tipico": "",
            "periodicidad": "▶",
        }

        filas, conteo = alinear_filas_generador(
            [fila_vacia],
            [{"METODO": "ASTM D445", "NOMBRE IQM": "Color"}],
            "es",
        )

        self.assertEqual(filas, [fila_vacia])
        self.assertEqual(conteo, {"alineados": 0, "revision": 0, "sin_coincidencia": 0})


class DocumentoAlineadoManualTests(unittest.TestCase):
    def test_especificacion_comparte_numeracion_y_conserva_comentarios(self):
        pruebas = [{"prueba": "Viscosidad", "metodo": "ASTM D445", "especificacion": "30 max"}]
        notas = _agregar_comentario_a_celda(pruebas, "Nota base", 0, "especificacion", "30 max")
        notas = notas.replace("¹ ", "¹ Nota de especificación")
        notas = _actualizar_comentarios_celda(pruebas, notas, "metodo", 0)
        notas = notas.replace("¹ ", "¹ Nota de método")
        notas = _actualizar_comentarios_celda(pruebas, notas, "prueba", 0)
        self.assertEqual(pruebas[0], {
            "prueba": "Viscosidad ¹", "metodo": "ASTM D445 ²", "especificacion": "30 max ³",
        })
        self.assertEqual(notas, "Nota base\n¹ \n² Nota de método\n³ Nota de especificación")
        notas_editadas = _guardar_celda_con_comentario(pruebas, notas, 0, "especificacion", "40 max ³")
        self.assertEqual(notas_editadas, notas)
        self.assertEqual(pruebas[0]["especificacion"], "40 max ³")
        notas = _guardar_celda_con_comentario(pruebas, notas, 0, "metodo", "ASTM D445")
        self.assertEqual(pruebas[0]["especificacion"], "40 max ²")
        self.assertEqual(notas, "Nota base\n¹ \n² Nota de especificación")

    def test_no_permite_comentarios_en_especificacion_vacia(self):
        with self.assertRaisesRegex(ValueError, "Agrega primero una especificación"):
            _actualizar_comentarios_celda([{"especificacion": ""}], "", "especificacion", 0)

    def test_comentarios_metodo_se_numeran_segun_orden_de_filas(self):
        pruebas = [
            {"metodo": "ASTM D445"},
            {"metodo": "ASTM D7042"},
        ]

        notas = _actualizar_comentarios_celda(
            pruebas,
            "Nota existente",
            "metodo",
            1,
        )
        notas = notas.replace("¹ ", "¹ Comentario segunda fila")
        notas = _actualizar_comentarios_celda(pruebas, notas, "metodo", 0)

        self.assertEqual(pruebas[0]["metodo"], "ASTM D445 ¹")
        self.assertEqual(pruebas[1]["metodo"], "ASTM D7042 ²")
        self.assertEqual(
            notas,
            "Nota existente\n¹ \n² Comentario segunda fila",
        )

    def test_comentarios_de_prueba_y_metodo_se_numeran_en_orden_de_fila(self):
        pruebas = [
            {"prueba": "Viscosidad", "metodo": "ASTM D445"},
        ]

        notas = _actualizar_comentarios_celda(pruebas, "", "metodo", 0)
        notas = _actualizar_comentarios_celda(
            pruebas,
            notas.replace("¹ ", "¹ Nota del método"),
            "prueba",
            0,
        )

        self.assertEqual(pruebas[0]["prueba"], "Viscosidad ¹")
        self.assertEqual(pruebas[0]["metodo"], "ASTM D445 ²")
        self.assertEqual(notas, "¹ \n² Nota del método")

    def test_agregar_comentario_a_prueba_no_reemplaza_su_valor_con_el_metodo(self):
        pruebas = [
            {"prueba": "ASTM D2265", "metodo": "ASTM D2265"},
        ]

        notas = _agregar_comentario_a_celda(
            pruebas,
            "",
            0,
            "prueba",
            "ASTM D2265",
        )

        self.assertEqual(pruebas[0]["prueba"], "ASTM D2265 ¹")
        self.assertEqual(pruebas[0]["metodo"], "ASTM D2265")
        self.assertEqual(notas, "¹ ")

    def test_borrar_superindice_quita_comentario_y_renumera_los_siguientes(self):
        pruebas = [
            {"prueba": "Prueba 1 ¹", "metodo": ""},
            {"prueba": "Prueba 2 ²", "metodo": ""},
        ]
        notas = "¹ Comentario que se elimina\n² Comentario que permanece"

        notas = _guardar_celda_con_comentario(
            pruebas,
            notas,
            0,
            "prueba",
            "Prueba 1",
        )

        self.assertEqual(pruebas[0]["prueba"], "Prueba 1")
        self.assertEqual(pruebas[1]["prueba"], "Prueba 2 ¹")
        self.assertEqual(notas, "¹ Comentario que permanece")

    def test_borrar_metodo_quita_su_comentario_y_renumera_los_restantes(self):
        pruebas = [
            {"prueba": "Prueba 1", "metodo": "Método 1"},
            {"prueba": "Prueba 2", "metodo": "Método 2"},
        ]
        notas = _actualizar_comentarios_celda(pruebas, "", "metodo", 0)
        notas = _actualizar_comentarios_celda(
            pruebas,
            notas.replace("¹ ", "¹ Comentario a quitar"),
            "metodo",
            1,
        )
        notas = notas.replace("² ", "² Comentario que queda")

        pruebas[0]["metodo"] = ""
        notas = _actualizar_comentarios_celda(
            pruebas,
            notas,
            remover_numero="1",
        )

        self.assertEqual(pruebas[0]["metodo"], "")
        self.assertEqual(pruebas[1]["metodo"], "Método 2 ¹")
        self.assertEqual(notas, "¹ Comentario que queda")

    def test_borrar_fila_quita_todos_sus_comentarios_en_una_sola_renumeracion(self):
        pruebas = [
            {"prueba": "Prueba 1", "metodo": "Método 1"},
            {"prueba": "Prueba 2", "metodo": "Método 2"},
        ]
        notas = _actualizar_comentarios_celda(pruebas, "", "prueba", 0)
        notas = _actualizar_comentarios_celda(pruebas, notas, "metodo", 0)
        notas = _actualizar_comentarios_celda(pruebas, notas, "metodo", 1)
        notas = notas.replace("¹ ", "¹ Comentario prueba 1")
        notas = notas.replace("² ", "² Comentario método 1")
        notas = notas.replace("³ ", "³ Comentario método 2")

        marcadores_eliminados = {"1", "2"}
        del pruebas[0]
        notas = _actualizar_comentarios_celda(
            pruebas,
            notas,
            remover_numero=marcadores_eliminados,
        )

        self.assertEqual(pruebas[0]["metodo"], "Método 2 ¹")
        self.assertEqual(notas, "¹ Comentario método 2")

    def test_comentario_metodo_se_escribe_en_nueva_linea_de_notas_word(self):
        documento = Document()
        tabla = documento.add_table(rows=3, cols=2)
        tabla.rows[2].cells[0].text = "Notas: Nota de plantilla"

        _escribir_notas(tabla, ["¹ Comentario del método", "Nota adicional"])

        self.assertEqual(
            tabla.rows[2].cells[0].text,
            "Notas: Nota de plantilla\n¹ Comentario del método\nNota adicional",
        )

    def test_periodicidad_json_no_se_escribe_en_la_tabla_word(self):
        documento = Document()
        tabla = documento.add_table(rows=3, cols=5)
        tabla.rows[2].cells[0].text = "Notas:"
        _llenar_pruebas(
            tabla,
            ["prueba", "unidades", "metodo", "especificacion", "valor_tipico"],
            [
                {
                    "prueba": "Viscosidad",
                    "unidades": "mm2/s",
                    "metodo": "ASTM D445 / ASTM D7042",
                    "especificacion": "28.8-35.2",
                    "valor_tipico": "30.23",
                    CAMPO_PERIODICIDAD: "◊▶",
                }
            ],
        )

        valores_word = [celda.text for celda in tabla.rows[1].cells]
        self.assertEqual(valores_word[0], "Viscosidad")
        self.assertNotIn("◊", " ".join(valores_word))

    def test_crea_detalle_con_los_valores_editados_en_el_json(self):
        datos = {
            "producto": {"nombre": "Producto de prueba"},
            "pruebas": [
                {
                    "prueba": "COLOR NUEVO",
                    "unidades": "N/A",
                    "metodo": "USP",
                    "especificacion": "AMBAR",
                    "valor_tipico": "",
                }
            ],
        }

        resultado = crear_documento_alineado_manual(datos)

        self.assertEqual(datos.get("alineacion_metodos_lab"), None)
        self.assertEqual(
            resultado["alineacion_metodos_lab"]["detalle"][0]["valores_alineados"],
            {
                "prueba": "COLOR NUEVO",
                "unidades": "N/A",
                "metodo": "USP",
                "especificacion": "AMBAR",
                "valor_tipico": "",
            },
        )
        self.assertEqual(
            resultado["alineacion_metodos_lab"]["detalle"][0]["estado"],
            "EDICION_MANUAL",
        )

    def test_rechaza_lista_pruebas_con_elementos_invalidos(self):
        with self.assertRaisesRegex(ValueError, "'pruebas' debe ser una lista de objetos"):
            crear_documento_alineado_manual(
                {"producto": {"nombre": "Producto"}, "pruebas": ["no es una fila"]}
            )

    def test_conserva_valores_alineados_existentes_distintos_de_pruebas_raiz(self):
        datos = {
            "producto": {"nombre": "Producto de prueba"},
            "pruebas": [{"prueba": "Nombre extraído", "metodo": "VISUAL"}],
            "alineacion_metodos_lab": {
                "detalle": [
                    {
                        "valores_alineados": {
                            "prueba": "Nombre maestro",
                            "unidades": "N/A",
                            "metodo": "USP",
                            "especificacion": "AMBAR",
                            "valor_tipico": "",
                        }
                    }
                ]
            },
        }

        resultado = crear_documento_alineado_manual(datos)

        self.assertEqual(
            resultado["alineacion_metodos_lab"]["detalle"][0]["valores_alineados"]["prueba"],
            "Nombre maestro",
        )

    def test_no_fabrica_valores_alineados_desde_pruebas_raiz_si_hay_auditoria_antigua(self):
        with self.assertRaisesRegex(ValueError, "no tiene valores_alineados completos"):
            crear_documento_alineado_manual(
                {
                    "producto": {"nombre": "Producto"},
                    "pruebas": [{"prueba": "Nombre extraído"}],
                    "alineacion_metodos_lab": {"detalle": [{"nombre_original": "Nombre extraído"}]},
                }
            )

    def test_rechaza_campos_con_objetos_no_serializables_en_la_tabla(self):
        with self.assertRaisesRegex(ValueError, "pruebas\\[1\\]\\.metodo"):
            crear_documento_alineado_manual(
                {
                    "producto": {"nombre": "Producto"},
                    "pruebas": [{"prueba": "Color", "metodo": {"valor": "USP"}}],
                }
            )


class ComparacionTemporalWordTests(unittest.TestCase):
    def test_extrae_word_a_json_temporal_sin_escribir_registro(self):
        with tempfile.TemporaryDirectory() as temporal:
            base = Path(temporal)
            origen = base / "comparacion.docx"
            documento = Document()
            documento.sections[0].header.paragraphs[0].text = "Hydraulic Oil"
            tabla = documento.add_table(rows=2, cols=5)
            for celda, encabezado in zip(
                tabla.rows[0].cells,
                ("TEST", "UNITS", "METHOD", "SPECIFICATION", "TYPICAL VALUE"),
            ):
                celda.text = encabezado
            for celda, valor in zip(
                tabla.rows[1].cells,
                ("Kinematic viscosity", "mm2/s", "ASTM D445", "10 MIN", "12.5"),
            ):
                celda.text = valor
            documento.save(origen)

            directorio_temporal = base / "temporal"
            directorio_temporal.mkdir()
            salida = crear_json_temporal_comparacion(origen, directorio_temporal)
            datos = json.loads(salida.read_text(encoding="utf-8"))

            self.assertEqual(datos["producto"]["nombre"], "Hydraulic Oil")
            self.assertEqual(datos["documento"]["idioma"], "en")
            self.assertEqual(datos["pruebas"][0]["prueba"], "Kinematic viscosity")
            self.assertEqual(datos["pruebas"][0]["metodo"], "ASTM D445")
            self.assertFalse((base / "especificaciones.db").exists())
            self.assertTrue(origen.is_file())


class ListaMaestraEditableTests(unittest.TestCase):
    def test_lee_y_guarda_todos_los_campos_sin_quitar_hojas(self):
        from openpyxl import Workbook, load_workbook

        with tempfile.TemporaryDirectory() as temporal:
            ruta = Path(temporal) / "maestro.xlsx"
            libro = Workbook()
            hoja = libro.active
            hoja.title = "METODOS"
            hoja.append(
                [
                    "METODO",
                    "MÉTODO INGLÉS",
                    "NOMBRE IQM",
                    "NOMBRE IQM INGLES",
                    "UNIDADES",
                    "UNITS",
                    "ACREDITACION NOMBRE",
                ]
            )
            hoja.append(
                [
                    "ASTM D445",
                    "ASTM D445",
                    "Viscosidad",
                    "Viscosity",
                    "mm²/s",
                    "mm2/s",
                    "ACREDITADO",
                ]
            )
            libro.create_sheet("RESUMEN").append(["Conservado"])
            libro.save(ruta)
            libro.close()

            filas = _leer_filas_maestro_editable(ruta)
            self.assertEqual(filas[0][0], 2)
            self.assertEqual(
                filas[0][1],
                (
                    "ASTM D445",
                    "ASTM D445",
                    "Viscosidad",
                    "Viscosity",
                    "mm²/s",
                    "mm2/s",
                    "ACREDITADO",
                ),
            )

            valores_actualizados = (
                "ASTM D445 / ASTM D7042",
                "ASTM D445 / ASTM D7042",
                "Viscosidad cinemática",
                "Kinematic viscosity",
                "cSt",
                "cSt",
                "NO ACREDITADO",
            )
            _guardar_cambios_maestro(ruta, {2: valores_actualizados})

            filas_actualizadas = _leer_filas_maestro_editable(ruta)
            self.assertEqual(filas_actualizadas, [(2, valores_actualizados)])
            libro_guardado = load_workbook(ruta, read_only=True)
            self.assertIn("RESUMEN", libro_guardado.sheetnames)
            self.assertEqual(
                libro_guardado["RESUMEN"]["A1"].value,
                "Conservado",
            )
            libro_guardado.close()


class VisorHistoricoTests(unittest.TestCase):
    def test_consulta_tabla_pagina_y_busqueda_sin_modificar_base(self):
        with tempfile.TemporaryDirectory() as temporal:
            ruta = Path(temporal) / "especificaciones.db"
            conexion = sqlite3.connect(ruta)
            conexion.execute(
                "CREATE TABLE versiones (id INTEGER PRIMARY KEY, producto TEXT)"
            )
            conexion.executemany(
                "INSERT INTO versiones (producto) VALUES (?)",
                [
                    ("Aceite hidráulico",),
                    ("Grasa industrial",),
                    ("Aceite motor",),
                    ("Aceite =2+2",),
                ],
            )
            conexion.executemany(
                "INSERT INTO versiones (producto) VALUES (?)",
                [("Aceite de prueba",) for _ in range(125)],
            )
            conexion.commit()
            conexion.close()

            columnas, filas, total = _consultar_tabla_historico(
                ruta,
                "versiones",
                limite=2,
                desplazamiento=1,
            )
            self.assertEqual(columnas, ("id", "producto"))
            self.assertEqual(filas, [("2", "Grasa industrial"), ("3", "Aceite motor")])
            self.assertEqual(total, 129)

            columnas, filas, total = _consultar_tabla_historico(
                ruta,
                "versiones",
                busqueda="ACEITE",
            )
            self.assertEqual(len(filas), 100)
            self.assertEqual(total, 128)

            from openpyxl import load_workbook

            destino = Path(temporal) / "exportado.xlsx"
            cantidad = _exportar_tabla_historico_excel(
                ruta,
                "versiones",
                destino,
                busqueda="ACEITE",
            )
            self.assertEqual(cantidad, 128)
            libro = load_workbook(destino, data_only=False, read_only=False)
            hoja = libro["versiones"]
            self.assertEqual(hoja.max_row, 129)
            self.assertEqual(hoja["B4"].value, "Aceite =2+2")
            self.assertEqual(hoja["B4"].data_type, "s")
            self.assertEqual(hoja.freeze_panes, "A2")
            self.assertTrue(hoja.auto_filter.ref)
            libro.close()

            with self.assertRaises(ValueError):
                _consultar_tabla_historico(
                    ruta,
                    "versiones; DROP TABLE versiones",
                )
            conexion = sqlite3.connect(ruta)
            self.assertEqual(
                conexion.execute("SELECT COUNT(*) FROM versiones").fetchone()[0],
                129,
            )
            conexion.close()


if __name__ == "__main__":
    unittest.main()
