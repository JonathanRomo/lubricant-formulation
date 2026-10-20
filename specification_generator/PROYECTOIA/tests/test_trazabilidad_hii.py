import io
import json
import os
import sqlite3
import tempfile
import unittest
from datetime import date, datetime
from dataclasses import replace
from contextlib import closing
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch
from uuid import UUID
from zipfile import ZipFile

from lxml import etree as ET
from openpyxl import Workbook, load_workbook
from openpyxl.packaging.custom import StringProperty

import reconstruir_base_datos as recuperacion
import generador_hii as trazabilidad
from alinear_especificaciones_json import procesar_json
from generador_hii import GeneradorHII, construir_excel, huella_datos, preparar_datos, seleccionar_plantilla
from interfaz_seleccion_especificaciones import cargar_catalogo
from rutas_proyecto import BASE


MAESTRO = {"METODO": "ASTM D445", "NOMBRE IQM": "Viscosidad cinemática",
           "NOMBRE IQM INGLES": "Kinematic viscosity", "UNIDADES": "cSt", "UNITS": "mm2/s",
           "ACREDITACION NOMBRE": "ACREDITADO", "FILA_MAESTRO": "2"}


def documento():
    return {"producto": {"nombre": "Producto"}, "documento": {"idioma": "es"},
            "control_documental": {"codigo_unico_producto": "20100001", "revision": "01", "nombre_cliente": "Cliente"},
            "notas": ["Nota"], "pruebas": [{"numero": 1, "prueba": "Viscosidad cinemática", "metodo": "ASTM D445",
                                           "unidades": "cSt", "especificacion": "10", "periodicidad": "A"}]}


class TrazabilidadTests(unittest.TestCase):
    def setUp(self):
        temporal = tempfile.TemporaryDirectory()
        self.addCleanup(temporal.cleanup)
        self.base = Path(temporal.name).resolve()
        self.origen = self.base / "ESPECIFICACIONES_JSON" / "producto.json"
        self.especificacion = self.base / "ESPECIFICACIONES_JSON_ALINEADAS" / "producto.json"
        self.destino = self.base / "HII" / "producto.xlsm"
        self.origen.parent.mkdir()
        self.destino.parent.mkdir()
        self.origen.write_text(json.dumps(documento()), encoding="utf-8")
        procesar_json(self.origen, self.especificacion, [MAESTRO])
        libro = Workbook()
        libro.active["A1"] = "Prueba"
        libro.custom_doc_props.append(StringProperty(name="PropiedadExistente", value="Conservar"))
        salida = io.BytesIO()
        libro.save(salida)
        libro.close()
        self.contenido = salida.getvalue()

    def leer(self):
        return json.loads(self.especificacion.read_text(encoding="utf-8"))

    def test_preparacion_de_seleccion_lee_json_una_sola_vez(self):
        leer = Path.read_text
        with patch.object(Path, "read_text", autospec=True, side_effect=leer) as lectura:
            datos, vista = trazabilidad.preparar_seleccion(self.especificacion)
        self.assertEqual(sum(c.args[0] == self.especificacion for c in lectura.call_args_list), 1)
        self.assertEqual(datos.nombre, vista.nombre_producto)
        self.assertEqual(huella_datos(datos), huella_datos(preparar_datos(self.especificacion)))

    def test_cache_plantilla_se_reutiliza_sin_contaminar_salidas_y_se_invalida(self):
        plantilla = self.base / "modelo.xlsx"
        plantilla.write_bytes(self.contenido)
        trazabilidad._cargar_plantilla.cache_clear()
        self.addCleanup(trazabilidad._cargar_plantilla.cache_clear)
        datos = replace(preparar_datos(self.especificacion), reanalisis="12 meses")
        with patch("generador_hii.load_workbook", wraps=load_workbook) as cargar:
            primera = construir_excel(plantilla, datos)
            segunda = construir_excel(plantilla, replace(datos, reanalisis="06 meses"))
            self.assertEqual(construir_excel(plantilla, datos), primera)
            self.assertEqual(cargar.call_count, 1)
            libro = load_workbook(plantilla)
            libro.active["A1"] = "Plantilla modificada"
            libro.save(plantilla)
            libro.close()
            nueva = construir_excel(plantilla, datos)
            self.assertEqual(cargar.call_count, 2)
        for contenido, esperado in ((primera, "12 meses"), (segunda, "06 meses")):
            libro = load_workbook(io.BytesIO(contenido))
            self.assertEqual(libro.active["D33"].value, esperado)
            libro.close()
        libro = load_workbook(io.BytesIO(nueva))
        self.assertEqual(libro.active["A1"].value, "Plantilla modificada")
        libro.close()

    def test_plano_excel_no_crea_widgets_y_conserva_celdas_del_archivo(self):
        datos = preparar_datos(self.especificacion)
        with patch("generador_hii.tk.PhotoImage", side_effect=AssertionError("Tk no permitido")), patch(
            "generador_hii.tk.Canvas", side_effect=AssertionError("Tk no permitido")
        ):
            region, ordenes = trazabilidad.preparar_dibujo_excel(self.contenido, datos)
        self.assertGreater(region[2], 0)
        self.assertGreater(region[3], 0)
        self.assertTrue(any(tipo == "text" and opciones["text"] == "Prueba" for tipo, _, opciones in ordenes))

    def guardar(self):
        huella = huella_datos(preparar_datos(self.especificacion))
        return trazabilidad.guardar_con_historial(self.destino, self.especificacion, self.contenido, huella)

    def test_cinco_campos_id_compartido_y_guardados_independientes(self):
        self.assertEqual(self.leer()["hii_generadas"], [])
        antes = preparar_datos(self.especificacion)
        primero = self.guardar()
        self.assertEqual(set(primero), {"id", "archivo", "fecha_generacion", "version_generador", "hash_especificacion"})
        UUID(primero["id"])
        self.assertIsNotNone(datetime.fromisoformat(primero["fecha_generacion"]).tzinfo)
        self.assertEqual((self.especificacion.parent / primero["archivo"]).resolve(), self.destino)
        self.assertEqual(primero["hash_especificacion"], huella_datos(antes))
        libro = load_workbook(self.destino)
        self.assertEqual(libro.custom_doc_props["HII_ID"].value, primero["id"])
        libro.close()
        segundo = self.guardar()
        self.assertNotEqual(primero["id"], segundo["id"])
        self.assertEqual(self.leer()["hii_generadas"], [primero, segundo])
        self.assertEqual(huella_datos(antes), huella_datos(preparar_datos(self.especificacion, date(2020, 1, 1))))

    def test_hash_cambia_con_valores_efectivos_y_no_con_auditoria(self):
        antes = huella_datos(preparar_datos(self.especificacion))
        datos = self.leer()
        datos["alineacion_metodos_lab"]["fecha"] = "otra fecha"
        datos["alineacion_metodos_lab"]["detalle"][0]["valores_alineados"]["valor_tipico"] = "99"
        self.especificacion.write_text(json.dumps(datos), encoding="utf-8")
        self.assertEqual(antes, huella_datos(preparar_datos(self.especificacion)))
        datos["alineacion_metodos_lab"]["detalle"][0]["valores_alineados"]["especificacion"] = "20"
        self.especificacion.write_text(json.dumps(datos), encoding="utf-8")
        self.assertNotEqual(antes, huella_datos(preparar_datos(self.especificacion)))

    def test_realinear_desde_original_y_en_sitio_conserva_historial(self):
        registro = self.guardar()
        procesar_json(self.origen, self.especificacion, [MAESTRO])
        self.assertEqual(self.leer()["hii_generadas"], [registro])
        procesar_json(self.especificacion, self.especificacion, [MAESTRO])
        self.assertEqual(self.leer()["hii_generadas"], [registro])

    def test_traslado_reajusta_rutas_y_fusiona_historial_reciente(self):
        registro = self.guardar()
        datos_cargados = self.leer()
        otro = self.guardar()
        trazabilidad.conservar_historial(datos_cargados, self.especificacion, self.especificacion)
        self.assertEqual(datos_cargados["hii_generadas"], [registro, otro])
        copia = self.base / "subcarpeta" / "copia.json"
        trazabilidad.conservar_historial(datos_cargados, self.especificacion, copia)
        for fila in datos_cargados["hii_generadas"]:
            self.assertEqual((copia.parent / fila["archivo"]).resolve(), self.destino)

    def test_preserva_partes_y_propiedades_al_reemplazar_id(self):
        salida = io.BytesIO()
        with ZipFile(io.BytesIO(self.contenido)) as original, ZipFile(salida, "w") as modificado:
            for info in original.infolist():
                modificado.writestr(info, original.read(info))
            modificado.writestr("xl/vbaProject.bin", b"macros originales")
            modificado.writestr("xl/media/imagen.png", b"imagen original")
        una = trazabilidad.incrustar_identificador(salida.getvalue(), "primero")
        dos = trazabilidad.incrustar_identificador(una, "segundo")
        with ZipFile(io.BytesIO(salida.getvalue())) as original, ZipFile(io.BytesIO(dos)) as final:
            for nombre in original.namelist():
                if nombre not in {"_rels/.rels", "[Content_Types].xml", "docProps/custom.xml"}:
                    self.assertEqual(original.read(nombre), final.read(nombre))
            propiedades = ET.fromstring(final.read("docProps/custom.xml"))
            self.assertEqual({p.get("name"): p[0].text for p in propiedades},
                             {"HII_ID": "segundo", "PropiedadExistente": "Conservar"})

    def test_error_excel_no_registra_y_error_json_informa_guardado_parcial(self):
        antes = self.especificacion.read_bytes()
        with patch.object(trazabilidad.os, "replace", side_effect=PermissionError("Excel abierto")):
            with self.assertRaisesRegex(OSError, "Excel abierto"):
                self.guardar()
        self.assertFalse(self.destino.exists())
        self.assertEqual(self.especificacion.read_bytes(), antes)
        reemplazar = os.replace

        def fallar_json(origen, destino):
            if destino == self.especificacion:
                raise PermissionError("JSON bloqueado")
            reemplazar(origen, destino)

        with patch.object(trazabilidad.os, "replace", side_effect=fallar_json):
            with self.assertRaisesRegex(OSError, "El Excel se guardó.*HII_ID="):
                self.guardar()
        self.assertTrue(self.destino.is_file())
        self.assertEqual(self.especificacion.read_bytes(), antes)
        self.assertFalse(list(self.base.rglob("*.tmp")))

    def test_historial_invalido_no_se_sobrescribe_al_guardar_o_alinear(self):
        datos = self.leer()
        datos["hii_generadas"] = "historial dañado"
        self.especificacion.write_text(json.dumps(datos), encoding="utf-8")
        antes = self.especificacion.read_bytes()
        with self.assertRaises(ValueError):
            self.guardar()
        with self.assertRaises(ValueError):
            procesar_json(self.origen, self.especificacion, [MAESTRO])
        self.assertEqual(self.especificacion.read_bytes(), antes)
        self.assertFalse(self.destino.exists())

    def test_reconstruye_sqlite_sin_modificar_json_incluso_si_falta_excel(self):
        self.guardar()
        ultimo = self.guardar()
        datos = self.leer()
        datos["hii_generadas"][0]["fecha_generacion"] = "2026-10-13T10:00:00-06:00"
        ultimo["fecha_generacion"] = "2026-10-14T10:00:00-06:00"
        datos["hii_generadas"][1] = ultimo
        datos["hii_generadas"].reverse()  # La última entrada no necesariamente es la más reciente.
        datos["hii_generadas"].append({"archivo": "incompleto.xlsm"})
        self.especificacion.write_text(json.dumps(datos), encoding="utf-8")
        antes = self.especificacion.read_bytes()
        # Simular un traslado completo de la carpeta portable y un Excel ausente.
        nueva = self.base / "trasladado"
        nueva.mkdir()
        self.especificacion.parent.rename(nueva / self.especificacion.parent.name)
        self.destino.parent.rename(nueva / "HII")
        self.origen.parent.rename(nueva / self.origen.parent.name)
        (nueva / "HII" / self.destino.name).unlink()
        resultado = recuperacion.reconstruir(nueva)
        self.assertEqual(resultado["hii_generadas"], 1)
        self.assertTrue(any("no localizada" in a for a in resultado["advertencias"]))
        self.assertTrue(any("incompleto" in a for a in resultado["advertencias"]))
        with closing(sqlite3.connect(nueva / "especificaciones.db")) as con:
            self.assertEqual(con.execute("SELECT name FROM sqlite_master WHERE name LIKE 'hii_%'").fetchall(), [])
            self.assertEqual(con.execute("PRAGMA integrity_check").fetchone()[0], "ok")
        self.assertEqual((nueva / "ESPECIFICACIONES_JSON_ALINEADAS/producto.json").read_bytes(), antes)

    def test_sqlite_danado_no_impide_guardar_excel_y_json(self):
        generador = SimpleNamespace(registrar_hii=Mock(side_effect=sqlite3.DatabaseError("base dañada")))
        datos = preparar_datos(self.especificacion)
        aviso = GeneradorHII.guardar_archivo(generador, self.destino, self.especificacion, datos, self.contenido)
        self.assertIsNone(aviso)
        generador.registrar_hii.assert_not_called()
        self.assertTrue(self.destino.exists())
        self.assertEqual(len(self.leer()["hii_generadas"]), 1)

    def test_botones_individual_y_lote_registran_sus_archivos(self):
        generador = GeneradorHII.__new__(GeneradorHII)
        generador._cerrada = False
        generador.datos = preparar_datos(self.especificacion)
        generador.excel = self.contenido
        generador.actual = self.especificacion
        generador.formatos = self.base / "formatos"
        generador.ventana = Mock()
        generador.estado = Mock()
        generador.registrar_hii = Mock()
        generador.resolver_idioma = Mock(return_value=self.especificacion)
        generador.catalogo = SimpleNamespace(especificaciones=[SimpleNamespace(ruta=self.especificacion)])
        generador.marcadas = {self.especificacion}
        generador.ambos_idiomas = Mock(get=Mock(return_value=False))
        with patch("generador_hii.filedialog") as dialogos, patch("generador_hii.messagebox") as mensajes, \
             patch("generador_hii.seleccionar_plantilla", return_value=self.base / "formato.xlsm"), patch("generador_hii.construir_excel", return_value=self.contenido):
            dialogos.asksaveasfilename.return_value = str(self.destino)
            dialogos.askdirectory.return_value = str(self.destino.parent)
            generador.guardar()
            generador.guardar_seleccionadas()
            mensajes.showerror.assert_not_called()
            mensajes.showwarning.assert_not_called()
        historial = self.leer()["hii_generadas"]
        self.assertEqual(len(historial), 2)
        generador.registrar_hii.assert_not_called()
        for registro in historial:
            archivo = self.especificacion.parent / registro["archivo"]
            libro = load_workbook(archivo)
            self.assertEqual(libro.custom_doc_props["HII_ID"].value, registro["id"])
            libro.close()

    def test_rechaza_vista_previa_desactualizada_antes_de_escribir(self):
        generador = SimpleNamespace(registrar_hii=Mock())
        datos = preparar_datos(self.especificacion)
        actual = self.leer()
        actual["producto"]["nombre"] = "Producto corregido"
        self.especificacion.write_text(json.dumps(actual), encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "cambió"):
            GeneradorHII.guardar_archivo(generador, self.destino, self.especificacion, datos, self.contenido)
        generador.registrar_hii.assert_not_called()
        self.assertFalse(self.destino.exists())

    def generador_bilingue(self):
        ingles = self.especificacion.with_name("producto_ENG.json")
        datos = self.leer()
        datos["documento"]["idioma"] = "en"
        ingles.write_text(json.dumps(datos), encoding="utf-8")
        formatos = self.base / "formatos"
        formatos.mkdir()
        (formatos / "91400005 vigente.xlsx").write_bytes(self.contenido)
        libro = Workbook()
        libro.template = True
        libro.save(formatos / "91400019 vigente.xltx")
        libro.close()
        generador = GeneradorHII.__new__(GeneradorHII)
        generador._cerrada = False
        generador.datos = preparar_datos(self.especificacion)
        generador.excel = self.contenido
        generador.actual = generador.origen_actual = self.especificacion
        generador.idioma_objetivo = None
        generador.base_datos = self.base / "sin_usar.db"
        generador.formatos = formatos
        generador.estado, generador.ventana = Mock(), Mock()
        generador.registrar_hii = Mock()
        generador.ambos_idiomas = Mock(get=Mock(return_value=True))
        generador.catalogo = cargar_catalogo(self.especificacion.parent)
        generador.marcadas = {self.especificacion, ingles}
        return generador, ingles

    def test_ambos_individual_y_lote_respetan_extensiones_sin_duplicar(self):
        generador, ingles = self.generador_bilingue()
        with patch("generador_hii.filedialog") as dialogos, patch("generador_hii.messagebox") as mensajes, \
             patch("generador_hii.conectar"), patch("generador_hii.vincular"), \
             patch("generador_hii.construir_excel", side_effect=lambda plantilla, datos: plantilla.read_bytes()):
            dialogos.askdirectory.return_value = str(self.destino.parent)
            mensajes.askyesno.return_value = True
            generador.guardar()
            generador.guardar_seleccionadas()
            mensajes.showerror.assert_not_called()
        generador.registrar_hii.assert_not_called()
        for ruta, extension in ((self.especificacion, ".xlsx"), (ingles, ".xltx")):
            historial = json.loads(ruta.read_text(encoding="utf-8"))["hii_generadas"]
            self.assertEqual(len(historial), 2)
            excel = (ruta.parent / historial[-1]["archivo"]).resolve()
            self.assertEqual(excel.suffix, extension)
            libro = load_workbook(excel)
            self.assertEqual(libro.custom_doc_props["HII_ID"].value, historial[-1]["id"])
            libro.close()

    def test_cancelar_reemplazo_de_ambos_no_modifica_archivos_ni_historial(self):
        generador, ingles = self.generador_bilingue()
        destino = self.destino.parent / "producto HII ESPANOL.xlsx"
        destino.write_bytes(b"Excel anterior")
        antes = {p: p.read_bytes() for p in (self.especificacion, ingles, destino)}
        with patch("generador_hii.filedialog") as dialogos, patch("generador_hii.messagebox") as mensajes, \
             patch("generador_hii.conectar"), patch("generador_hii.vincular"):
            dialogos.askdirectory.return_value = str(self.destino.parent)
            mensajes.askyesno.return_value = False
            generador.guardar()
            generador.registrar_hii.assert_not_called()
        self.assertEqual({p: p.read_bytes() for p in antes}, antes)

    def test_seleccion_de_plantilla_reconoce_formatos_y_rechaza_ambiguedad(self):
        formatos = self.base / "formatos"
        formatos.mkdir()
        plantilla = formatos / "91400005 vigente.XLSX"
        plantilla.write_bytes(self.contenido)
        self.assertEqual(seleccionar_plantilla(formatos, "ESPANOL"), plantilla)
        (formatos / "91400005 copia.xltx").write_bytes(self.contenido)
        with self.assertRaises(ValueError):
            seleccionar_plantilla(formatos, "ESPANOL")

    def test_construye_hii_xlsx_y_xltx_con_datos_y_propiedad_de_trazabilidad(self):
        try:
            original = seleccionar_plantilla(BASE / "FORMATOS", "ESPANOL")
        except ValueError:
            self.skipTest("Plantilla local no disponible")
        for extension, es_plantilla in ((".xlsx", False), (".xltx", True)):
            with self.subTest(extension=extension):
                plantilla = self.base / ("91400005 formato" + extension)
                libro = load_workbook(original)
                libro.template = es_plantilla
                libro.save(plantilla)
                libro.close()
                datos = preparar_datos(self.especificacion)
                contenido = construir_excel(plantilla, datos)
                salida = trazabilidad.incrustar_identificador(contenido, "formato-" + extension)
                final = load_workbook(io.BytesIO(salida))
                self.assertEqual(final.template, es_plantilla)
                self.assertEqual(final.worksheets[0]["F15"].value, datos.pruebas[0]["especificacion"])
                self.assertEqual(final.custom_doc_props["HII_ID"].value, "formato-" + extension)
                final.close()

    def test_plantillas_reales_conservan_macros_e_imagenes(self):
        for idioma in ("ESPANOL", "INGLES"):
            with self.subTest(idioma=idioma):
                try:
                    plantilla = seleccionar_plantilla(BASE / "FORMATOS", idioma)
                except ValueError:
                    self.skipTest("Plantillas locales no disponibles")
                datos = preparar_datos(self.especificacion)
                datos.idioma = idioma
                contenido = construir_excel(plantilla, datos)
                salida = trazabilidad.incrustar_identificador(contenido, "prueba-plantilla")
                with ZipFile(plantilla) as original, ZipFile(io.BytesIO(salida)) as final:
                    self.assertIsNone(final.testzip())
                    conservadas = [n for n in original.namelist() if n.startswith("xl/media/") or n.endswith("vbaProject.bin")]
                    self.assertTrue(conservadas)
                    for nombre in conservadas:
                        self.assertEqual(original.read(nombre), final.read(nombre))
                libro = load_workbook(io.BytesIO(salida))
                self.assertEqual(libro.custom_doc_props["HII_ID"].value, "prueba-plantilla")
                libro.close()


class CompatibilidadPlantillasHII(unittest.TestCase):
    @unittest.skipUnless(os.name == "nt", "Las métricas reales usan las fuentes de Windows")
    def test_altura_comentarios_mide_anchos_de_letras_y_saltos(self):
        medir = trazabilidad._altura_texto_comentarios
        estrecho = medir("i" * 100, 250, 8, "Arial", True)
        ancho = medir("W" * 100, 250, 8, "Arial", True)
        self.assertGreater(ancho, estrecho)
        self.assertGreater(medir("Primera\nSegunda\nTercera", 250, 8, "Arial"),
                           medir("Primera Segunda Tercera", 250, 8, "Arial"))
        self.assertLessEqual(medir("Nota breve", 250, 8, "Arial"), 18)

    def test_notas_crecen_reducen_fuente_y_conservan_bordes_en_plantillas_reales(self):
        for idioma in ("ESPANOL", "INGLES"):
            plantilla = seleccionar_plantilla(BASE / "FORMATOS", idioma)
            original = load_workbook(plantilla)
            hoja = original.active
            fuente_original = hoja["D21"].font.sz
            comun = next(linea.strip() for linea in hoja["D21"].value.splitlines()
                         if trazabilidad.es_nota_formato(linea))
            altura_original = sum(hoja.row_dimensions[r].height or hoja.sheet_format.defaultRowHeight
                                  for r in range(21, 29))
            alturas_originales = {r: hoja.row_dimensions[r].height for r in range(20, 36)}
            original.close()
            for cantidad in (0, 3, 4, 12):
                with self.subTest(idioma=idioma, cantidad=cantidad):
                    notas = [(f"{i + 1}. " + "Observaciones del usuario con varias palabras. " * 3).strip()
                             for i in range(cantidad)]
                    # La nota común puede llegar antes y repetida desde la especificación.
                    comentarios = "\n".join([comun, *notas, comun])
                    pruebas = [{"numero": str(i + 1), "prueba": "Prueba"} for i in range(6)]
                    datos = trazabilidad.DatosHII("Producto", "123", idioma, pruebas, comentarios, "12 meses", "19/10/26")
                    contenido = construir_excel(plantilla, datos)
                    final = load_workbook(io.BytesIO(contenido))
                    s = final.active
                    celda = s["D24"]  # Tres filas insertadas por las seis pruebas.
                    esperado = "\n\n".join(["\n".join(notas), comun]) if notas else comun
                    self.assertEqual(celda.value, esperado)
                    self.assertEqual(celda.font.sz, fuente_original - (1 if cantidad > 3 else 0))
                    self.assertTrue(celda.alignment.wrap_text)
                    altura = sum(s.row_dimensions[r].height or s.sheet_format.defaultRowHeight
                                 for r in range(24, 32))
                    self.assertGreaterEqual(altura, altura_original)
                    if cantidad == 12:
                        self.assertGreater(altura, altura_original)
                        self.assertGreater(s.row_dimensions[24].height, alturas_originales[21])
                    for r, altura_previa in alturas_originales.items():
                        if r != 21:
                            self.assertEqual(s.row_dimensions[r + 3].height, altura_previa,
                                             f"Se alteró la altura original de la fila {r}")
                    for r in range(15, 21):
                        self.assertEqual(s[f"F{r}"].border.right.style, "thin")
                    self.assertEqual(s["D34"].value, "12 meses")
                    self.assertIn("D24:F31", str(s.merged_cells))
                    final.close()
                    trazabilidad.preparar_dibujo_excel(contenido, datos)
            with self.assertRaisesRegex(ValueError, "fila libre"):
                construir_excel(plantilla, replace(datos, comentarios="Nota extensa " * 2000))
            # Cada exportación parte de la fuente original, sin reducciones acumuladas.
            final_original = load_workbook(plantilla)
            self.assertEqual(final_original.active["D21"].font.sz, fuente_original)
            final_original.close()

    def test_expansion_respeta_comentarios_y_reanalisis_en_ambos_dis_enos(self):
        for capacidad in (3, 6):
            for idioma in ("ESPANOL", "INGLES"):
                with self.subTest(capacidad=capacidad, idioma=idioma), tempfile.TemporaryDirectory() as carpeta:
                    plantilla = Path(carpeta) / "modelo.xlsx"
                    libro = Workbook()
                    hoja = libro.active
                    for i in range(capacidad):
                        hoja.cell(15 + i, 1, i + 1)
                    comentarios = 20 if capacidad == 3 else 23
                    contenido = comentarios + 1
                    reanalisis = 30 if capacidad == 3 else 32
                    hoja.cell(comentarios, 4, "COMMENTS:" if idioma == "INGLES" else "COMENTARIOS:")
                    hoja.merge_cells(start_row=contenido, end_row=reanalisis - 2, start_column=4, end_column=6)
                    hoja.cell(contenido, 4, "Nota original")
                    hoja.cell(reanalisis, 4, "REANALYSIS DATE")
                    hoja.cell(reanalisis + 1, 4, "X months")
                    libro.save(plantilla)
                    libro.close()
                    for cantidad in (1, capacidad, 12):
                        pruebas = [{"numero": str(i + 1), "prueba": f"Prueba {i + 1}"} for i in range(cantidad)]
                        datos = trazabilidad.DatosHII("Producto", "123", idioma, pruebas, "Nota nueva", "12 meses", "19/10/26")
                        salida = construir_excel(plantilla, datos)
                        final = load_workbook(io.BytesIO(salida))
                        extra = max(0, cantidad - capacidad)
                        self.assertEqual(final.active.cell(contenido + extra, 4).value, "Nota nueva")
                        self.assertEqual(final.active.cell(reanalisis + 1 + extra, 4).value, "12 meses")
                        for i in range(cantidad):
                            self.assertEqual(final.active.cell(15 + i, 2).value, f"Prueba {i + 1}")
                        final.close()


if __name__ == "__main__":
    unittest.main()
