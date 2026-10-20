import json
import sqlite3
import tempfile
import unittest
from contextlib import closing
from pathlib import Path
from unittest.mock import patch

from docx import Document
from openpyxl import Workbook

import gestor_especificaciones as gestor


class ActualizacionExtraccionTests(unittest.TestCase):
    def setUp(self):
        temporal = tempfile.TemporaryDirectory()
        self.addCleanup(temporal.cleanup)
        self.base = Path(temporal.name)
        self.fuente = self.base / gestor.ENTRADA / "producto.docx"
        self.fuente.parent.mkdir()
        libro = Workbook()
        libro.active.append(["ID", "Description"])
        libro.active.append(["20100001", "PRODUCTO"])
        libro.save(self.base / gestor.CATALOGO)
        libro.close()
        self.documento()

    def documento(self, valor="50 MÁXIMO"):
        doc = Document()
        doc.add_paragraph("PRODUCTO")
        tabla = doc.add_table(rows=2, cols=4)
        for c, t in zip(tabla.rows[0].cells, ("PRUEBA", "UNIDADES", "MÉTODO", "ESPECIFICACIÓN")):
            c.text = t
        for c, t in zip(tabla.rows[1].cells, ("VISCOSIDAD", "cSt", "D445", valor)):
            c.text = t
        control = doc.add_table(rows=3, cols=2)
        for fila, (k, v) in zip(control.rows, (("Fecha", "01/01/2020"), ("Revisión", "00"), ("Nombre Cliente", "CLIENTE"))):
            fila.cells[0].text, fila.cells[1].text = k, v
        doc.save(self.fuente)

    def ejecutar(self, lector="anterior"):
        with patch.object(gestor, "VERSION_LECTOR", lector), patch("builtins.print"):
            return gestor.procesar(self.base)

    def json(self):
        return next((self.base / gestor.SALIDA).glob("*.JSON"))

    def leer(self):
        return json.loads(self.json().read_text(encoding="utf-8"))

    def editar(self, datos):
        self.json().write_text(json.dumps(datos, ensure_ascii=False), encoding="utf-8")

    def versiones(self):
        with closing(sqlite3.connect(self.base / gestor.DB)) as con:
            return con.execute("SELECT id,revision_json,archivo_json,version_lector FROM versiones ORDER BY id").fetchall()

    def test_lector_nuevo_actualiza_sin_crear_revision_y_luego_reutiliza_cache(self):
        original = gestor.extraer

        def antiguo(*args, **kwargs):
            datos = original(*args, **kwargs)
            datos["pruebas"][0]["prueba"] = "LECTURA ANTIGUA"
            return datos

        with patch.object(gestor, "extraer", side_effect=antiguo):
            self.ejecutar()
        ruta = self.json()
        previo = self.leer()
        version = self.versiones()[0]
        huella = gestor.sha256(self.fuente)
        resultado = self.ejecutar("nuevo")
        actual = self.leer()
        self.assertEqual(resultado["reextraidos"], 1)
        self.assertEqual((resultado["nuevos"], resultado["actualizados"]), (0, 0))
        self.assertEqual(actual["pruebas"][0]["prueba"], "VISCOSIDAD")
        self.assertEqual(actual["fecha_creacion_json"], previo["fecha_creacion_json"])
        self.assertEqual(actual["revision_json"], previo["revision_json"])
        self.assertEqual(self.json(), ruta)
        self.assertEqual(len(self.versiones()), 1)
        self.assertEqual(self.versiones()[0][:3], version[:3])
        self.assertEqual(gestor.sha256(self.fuente), huella)
        with patch.object(gestor, "extraer", side_effect=AssertionError("No releer")):
            self.assertEqual(self.ejecutar("nuevo")["ignorados"], 1)

    def test_conserva_ediciones_manuales_y_campos_ajenos_a_extraccion(self):
        self.ejecutar()
        datos = self.leer()
        datos["producto"]["nombre"] = "NOMBRE MANUAL"
        datos["control_documental"]["nombre_cliente"] = "CLIENTE MANUAL"
        datos["pruebas"][0]["especificacion"] = "99 MANUAL"
        datos["notas"] = ["Nota manual"]
        datos["mi_campo"] = {"conservar": True}
        self.editar(datos)
        resultado = self.ejecutar("nuevo")
        self.assertEqual(resultado["ediciones_conservadas"], 1)
        for campo in ("producto", "control_documental", "pruebas", "notas", "mi_campo"):
            self.assertEqual(self.leer()[campo], datos[campo])
        self.ejecutar("nuevo")
        self.assertEqual(self.leer()["producto"]["nombre"], "NOMBRE MANUAL")
        self.assertEqual(len(self.versiones()), 1)

    def test_respeta_una_eliminacion_manual(self):
        self.ejecutar()
        datos = self.leer()
        del datos["control_documental"]["nombre_cliente"]
        self.editar(datos)
        self.ejecutar("nuevo")
        self.assertNotIn("nombre_cliente", self.leer()["control_documental"])

    def test_json_heredado_sin_base_rellena_vacios_y_conserva_valores_existentes(self):
        self.ejecutar()
        datos = self.leer()
        datos["notas"] = ["Nota heredada manual"]
        datos["pruebas"][0]["prueba"] = ""
        self.editar(datos)
        with closing(sqlite3.connect(self.base / gestor.DB)) as con:
            con.execute("UPDATE versiones SET version_lector=NULL,base_extraccion_json=NULL")
            con.commit()
        self.ejecutar("nuevo")
        self.assertEqual(self.leer()["pruebas"][0]["prueba"], "VISCOSIDAD")
        self.assertEqual(self.leer()["notas"], ["Nota heredada manual"])
        self.assertEqual(len(self.versiones()), 1)

    def test_repara_json_ausente_sin_crear_version(self):
        self.ejecutar()
        previo = self.leer()
        ruta = self.json()
        ruta.unlink()
        resultado = self.ejecutar()
        self.assertEqual(resultado["reextraidos"], 1)
        self.assertEqual(len(self.versiones()), 1)
        self.assertEqual(self.leer()["fecha_creacion_json"], previo["fecha_creacion_json"])

    def test_relectura_fallida_conserva_word_json_y_version_anterior(self):
        self.ejecutar()
        json_antes = self.json().read_bytes()
        fuente_antes = self.fuente.read_bytes()
        with patch.object(gestor, "extraer", side_effect=ValueError("Estructura ambigua")):
            resultado = self.ejecutar("nuevo")
        self.assertEqual(resultado["pendientes_actualizacion"], 1)
        self.assertEqual(resultado["movidos_revisar"], 0)
        self.assertEqual(self.json().read_bytes(), json_antes)
        self.assertEqual(self.fuente.read_bytes(), fuente_antes)
        self.assertEqual(self.versiones()[0][3], "anterior")

    def test_cambio_de_catalogo_no_inventa_revision(self):
        self.ejecutar()
        libro = Workbook()
        libro.active.append(["ID", "Description"])
        libro.active.append(["20100001", "PRODUCTO"])
        libro.active.append(["20100002", "OTRO PRODUCTO"])
        libro.save(self.base / gestor.CATALOGO)
        libro.close()
        self.assertEqual(self.ejecutar()["reextraidos"], 1)
        self.assertEqual(len(self.versiones()), 1)

    def test_cambio_real_del_word_si_crea_version_documental(self):
        self.ejecutar()
        self.documento("60 MÁXIMO")
        self.assertEqual(self.ejecutar()["actualizados"], 1)
        self.assertEqual(len(self.versiones()), 2)

    def test_cambio_de_version_de_aplicacion_no_invalida_el_lector(self):
        self.ejecutar()
        with patch.object(gestor, "APP_VERSION", "999"), patch.object(gestor, "extraer", side_effect=AssertionError("No releer")):
            resultado = self.ejecutar()
        self.assertEqual(resultado["ignorados"], 1)

    def test_fallo_al_publicar_actualizacion_no_dana_extraccion_existente(self):
        import cache_extraccion

        self.ejecutar()
        datos = self.leer()
        datos["producto"]["nombre"] = "NOMBRE MANUAL"
        self.editar(datos)
        original = self.json().read_bytes()
        escribir = cache_extraccion.escribir_atomico
        llamadas = 0

        def fallar_una_vez(*args):
            nonlocal llamadas
            llamadas += 1
            if llamadas == 1:
                raise OSError("No se pudo publicar")
            return escribir(*args)

        with patch.object(cache_extraccion, "escribir_atomico", side_effect=fallar_una_vez):
            resultado = self.ejecutar("nuevo")
        self.assertEqual(resultado["pendientes_actualizacion"], 1)
        self.assertEqual(self.json().read_bytes(), original)
        self.assertEqual(self.versiones()[0][3], "anterior")
        self.assertTrue(self.fuente.exists())


if __name__ == "__main__":
    unittest.main()
