from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document

import gestor_especificaciones as gestor


class IdiomaContenidoWordTests(unittest.TestCase):
    def test_lectura_xml_no_confunde_etiquetas_de_tabla_con_texto(self):
        with tempfile.TemporaryDirectory() as base:
            ruta = Path(base) / "texto.docx"
            with ZipFile(ruta, "w") as archivo:
                archivo.writestr("word/document.xml", '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                                 '<w:tbl>IGNORAR<w:tr><w:tc><w:p><w:r><w:t>CLEAR</w:t></w:r></w:p>'
                                 '</w:tc></w:tr></w:tbl></w:document>')
            self.assertEqual(gestor.texto_ooxml(ruta), "CLEAR")

    def documento(self, base, valores, ingles=False, otro_texto="", tipico="", prueba="Viscosidad"):
        ruta = Path(base) / "original.docx"
        doc = Document()
        doc.add_paragraph(otro_texto)
        tabla = doc.add_table(rows=1, cols=4)
        encabezados = ("TEST", "METHOD", "SPECIFICATION", "TYPICAL VALUES") if ingles else (
            "PRUEBA", "METODO", "ESPECIFICACION", "VALOR TIPICO")
        for celda, encabezado in zip(tabla.rows[0].cells, encabezados):
            celda.text = encabezado
        for valor in valores:
            fila = tabla.add_row()
            for celda, texto in zip(fila.cells, (prueba, "ASTM D445", valor, tipico)):
                celda.text = texto
        doc.save(ruta)
        return ruta

    def test_dos_palabras_inglesas_prevalecen_sobre_formato_espanol(self):
        for valores in (("CLEAR AND BRIGHT",), ("PASS", "REPORT"), ("10 MINIMUM", "20 MAXIMUM"),
                        ("light amber", "typical"), ("10 mínimum", "20 máximum")):
            with self.subTest(valores=valores), tempfile.TemporaryDirectory() as base:
                ruta = self.documento(base, valores)
                original = ruta.read_bytes()
                self.assertTrue(gestor.contiene_marcadores_ingles(ruta))
                self.assertGreaterEqual(len(gestor.marcadores_ingles_especificacion(ruta)), 2)
                self.assertEqual(ruta.read_bytes(), original)

    def test_una_palabra_repetida_no_basta_ni_coinciden_fragmentos(self):
        for valores in (("PASS", "PASS", "PASS"), ("REPORT",), ("CLEARANCE", "BRIGHTNESS"),
                        ("CLARO Y BRILLANTE",), ("10 MIN", "20 MAX")):
            with self.subTest(valores=valores), tempfile.TemporaryDirectory() as base:
                ruta = self.documento(base, valores)
                self.assertFalse(gestor.contiene_marcadores_ingles(ruta))

    def test_no_cuenta_palabras_fuera_de_la_columna_especificacion(self):
        with tempfile.TemporaryDirectory() as base:
            ruta = self.documento(base, ("CUMPLE",), otro_texto="CLEAR BRIGHT AMBER",
                                   tipico="MINIMUM MAXIMUM", prueba="PASS REPORT")
            self.assertEqual(gestor.marcadores_ingles_especificacion(ruta), ())
            self.assertFalse(gestor.contiene_marcadores_ingles(ruta))

    def test_conserva_deteccion_anterior_por_etiquetas_inglesas(self):
        with tempfile.TemporaryDirectory() as base:
            ruta = self.documento(base, ("10",), ingles=True)
            self.assertTrue(gestor.contiene_marcadores_ingles(ruta))

    def test_columnas_con_celdas_fusionadas_y_notas_no_cuentan_como_especificacion(self):
        with tempfile.TemporaryDirectory() as base:
            ruta = Path(base) / "fusionada.docx"
            doc = Document()
            tabla = doc.add_table(rows=3, cols=5)
            tabla.cell(0, 0).merge(tabla.cell(0, 1)).text = "PRUEBA"
            for columna, valor in ((2, "METODO"), (3, "ESPECIFICACION"), (4, "VALOR TIPICO")):
                tabla.cell(0, columna).text = valor
            tabla.cell(1, 3).text = "CLEAR BRIGHT"
            tabla.cell(2, 0).merge(tabla.cell(2, 4)).text = "Notes: PASS REPORT"
            doc.save(ruta)
            self.assertEqual(gestor.marcadores_ingles_especificacion(ruta), ("CLEAR", "BRIGHT"))
            self.assertTrue(gestor.contiene_marcadores_ingles(ruta))

    def test_cache_evita_releer_xml_en_la_misma_carga(self):
        with tempfile.TemporaryDirectory() as base:
            ruta = self.documento(base, ("CLEAR BRIGHT",))
            with gestor.lectura_word_compartida(), patch.object(
                gestor.ElementTree, "fromstring", wraps=gestor.ElementTree.fromstring
            ) as parsear:
                self.assertTrue(gestor.contiene_marcadores_ingles(ruta))
                self.assertTrue(gestor.contiene_marcadores_ingles(ruta))
                parsear.assert_called_once()


if __name__ == "__main__":
    unittest.main()
