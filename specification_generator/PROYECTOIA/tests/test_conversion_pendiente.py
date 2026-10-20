"""Regresiones de conversión con datos sintéticos y directorios temporales."""
import sqlite3
import subprocess
import tempfile
import unittest
from contextlib import closing
from pathlib import Path
from unittest.mock import Mock, patch

from docx import Document
from openpyxl import Workbook

import gestor_especificaciones as gestor


class ConversionPendienteTests(unittest.TestCase):
    def setUp(self):
        temporal = tempfile.TemporaryDirectory()
        self.addCleanup(temporal.cleanup)
        self.base = Path(temporal.name)
        self.entrada = self.base / gestor.ENTRADA
        self.entrada.mkdir()
        libro = Workbook()
        libro.active.append(["ID", "Description"])
        libro.active.append(["20100001", "PRODUCTO"])
        libro.save(self.base / gestor.CATALOGO)
        libro.close()

    def documento(self, nombre="producto.docx", encabezado="PRUEBA"):
        ruta = self.entrada / nombre
        doc = Document()
        doc.sections[0].header.paragraphs[0].text = "PRODUCTO"
        tabla = doc.add_table(rows=2, cols=4)
        for celda, texto in zip(tabla.rows[0].cells, (encabezado, "MÉTODO", "UNIDADES", "ESPECIFICACIÓN")):
            celda.text = texto
        for celda, texto in zip(tabla.rows[1].cells, ("VISCOSIDAD", "D445", "cSt", "50 MÁXIMO")):
            celda.text = texto
        doc.save(ruta)
        return ruta

    def registros(self, tabla):
        with closing(sqlite3.connect(self.base / gestor.DB)) as con:
            return con.execute(f"SELECT * FROM {tabla}").fetchall()

    def test_fallo_conversion_no_aparta_ni_crea_historial_y_continua_con_docx(self):
        originales = []
        for nombre in ("a.doc", "b.doc", "c.docm"):
            ruta = self.entrada / nombre
            ruta.write_bytes(nombre.encode())
            originales.append((ruta, ruta.read_bytes()))
        self.documento("z.docx")
        preparar = gestor.preparar

        def convertir(ruta, temporal):
            if ruta.suffix != ".docx":
                raise gestor.ErrorConversionWord("Word no está disponible")
            return preparar(ruta, temporal)

        progreso = Mock()
        with patch.object(gestor, "preparar", side_effect=convertir), patch("builtins.print") as mensajes:
            resultado = gestor.procesar(self.base, al_progreso=progreso)
        self.assertEqual(resultado["pendientes_conversion"], 3)
        self.assertEqual((resultado["errores"], resultado["movidos_revisar"], resultado["nuevos"]), (0, 0, 1))
        self.assertEqual(len(self.registros("archivos")), 1)
        self.assertEqual(len(self.registros("versiones")), 1)
        self.assertFalse((self.base / gestor.REVISAR).exists())
        for ruta, contenido in originales:
            self.assertEqual(ruta.read_bytes(), contenido)
        self.assertEqual(progreso.call_count, 4)
        avisos = [c for c in mensajes.call_args_list if str(c.args[0]).startswith("CONVERSIÓN PENDIENTE")]
        self.assertEqual(len(avisos), 1)

    def test_siguiente_carga_reintenta_y_crea_una_sola_version(self):
        fuente = self.entrada / "pendiente.doc"
        fuente.write_bytes(b"documento antiguo")
        modelo = self.documento()
        modelo = modelo.rename(self.base / "modelo.docx")
        with patch.object(gestor, "preparar", side_effect=gestor.ErrorConversionWord("no disponible")), patch("builtins.print"):
            primero = gestor.procesar(self.base)
        self.assertEqual(primero["pendientes_conversion"], 1)
        self.assertEqual(self.registros("archivos"), [])
        self.assertEqual(self.registros("versiones"), [])
        with patch.object(gestor, "preparar", return_value=(modelo, True)), patch("builtins.print"):
            segundo = gestor.procesar(self.base)
            tercero = gestor.procesar(self.base)
        self.assertEqual((segundo["nuevos"], tercero["ignorados"]), (1, 1))
        self.assertEqual(len(self.registros("versiones")), 1)
        self.assertTrue(fuente.exists())

    def test_estructura_invalida_si_se_aparta_para_revision(self):
        ruta = self.entrada / "invalido.docx"
        doc = Document()
        doc.add_paragraph("PRODUCTO")
        doc.save(ruta)
        with patch("builtins.print"):
            resultado = gestor.procesar(self.base)
        self.assertEqual((resultado["errores"], resultado["movidos_revisar"]), (1, 1))
        self.assertTrue((self.base / gestor.REVISAR / ruta.name).exists())
        self.assertEqual(self.registros("versiones"), [])

    def test_alias_nombre_prueba_conserva_los_nombres(self):
        ruta = self.documento(encabezado="NOMBRE DE LA PRUEBA")
        datos = gestor.extraer(ruta)
        self.assertEqual(datos["pruebas"][0]["prueba"], "VISCOSIDAD")

    def test_reintenta_inicio_solo_una_vez_y_siguientes_no_reinician_word(self):
        def fallar(*args):
            if gestor._ERROR_INICIO_WORD is not None:
                raise RuntimeError("inicio ya fallido")
            gestor._ERROR_INICIO_WORD = "inicio bloqueado"
            raise TimeoutError("inicio bloqueado")

        with (
            patch.object(gestor.sys, "platform", "win32"),
            patch.object(gestor, "_ERROR_INICIO_WORD", None),
            patch.object(gestor, "_PROCESOS_WORD_PREEXISTENTES", None),
            patch.object(gestor, "_convertir_word_com", side_effect=fallar) as convertir,
            patch.object(gestor.shutil, "which", return_value=None),
        ):
            with self.assertRaises(gestor.ErrorConversionWord):
                gestor.convertir_word_legacy(self.entrada / "a.doc", self.base)
            self.assertEqual(convertir.call_count, 2)
            with self.assertRaises(gestor.ErrorConversionWord):
                gestor.convertir_word_legacy(self.entrada / "b.doc", self.base)
            self.assertEqual(convertir.call_count, 3)

    def test_inicio_transitorio_se_recupera_en_el_reintento(self):
        def convertir(origen, destino):
            if mock.call_count == 1:
                gestor._ERROR_INICIO_WORD = "inicio bloqueado"
                raise TimeoutError("inicio bloqueado")
            destino.write_bytes(b"convertido")

        with (
            patch.object(gestor.sys, "platform", "win32"),
            patch.object(gestor, "_ERROR_INICIO_WORD", None),
            patch.object(gestor, "_PROCESOS_WORD_PREEXISTENTES", None),
            patch.object(gestor, "_convertir_word_com", side_effect=convertir) as mock,
        ):
            resultado = gestor.convertir_word_legacy(self.entrada / "a.doc", self.base)
        self.assertEqual(resultado.read_bytes(), b"convertido")
        self.assertEqual(mock.call_count, 2)

    def test_fallo_libreoffice_tambien_deja_conversion_pendiente(self):
        for error in (OSError("no disponible"), subprocess.TimeoutExpired("soffice", 30)):
            with (
                self.subTest(error=type(error).__name__),
                patch.object(gestor.sys, "platform", "linux"),
                patch.object(gestor.shutil, "which", return_value="soffice"),
                patch.object(gestor.subprocess, "run", side_effect=error),
                self.assertRaises(gestor.ErrorConversionWord),
            ):
                gestor.convertir_word_legacy(self.entrada / "a.doc", self.base)

    def test_inspeccion_reconoce_embedding_y_restablece_referencia_entre_cargas(self):
        respuesta = Mock(returncode=0, stdout="[100,321]", stderr="")
        with patch.object(gestor.subprocess, "run", return_value=respuesta) as consultar:
            self.assertEqual(gestor._procesos_word_automatizacion(), {100, 321})
        self.assertIn("-embedding", consultar.call_args.args[0][-1])
        with patch.object(gestor, "_PROCESOS_WORD_PREEXISTENTES", {100}):
            gestor._restablecer_error_inicio_word()
            self.assertIsNone(gestor._PROCESOS_WORD_PREEXISTENTES)

    def test_interfaz_informa_pendientes_sin_presentarlos_como_errores(self):
        from interfaz_aplicacion_especificaciones import AplicacionEspecificaciones

        app = AplicacionEspecificaciones.__new__(AplicacionEspecificaciones)
        app.raiz = Mock()
        resultado = {"extraccion": {"pendientes_conversion": 3, "errores": 0},
                     "alineacion": {"archivos": 0, "errores": 0}}
        with patch("interfaz_aplicacion_especificaciones.messagebox.showwarning") as mostrar:
            app._mostrar_resultado_carga(resultado, 3, 3, self.entrada)
        titulo, mensaje = mostrar.call_args.args
        self.assertEqual(titulo, "Carga terminada con pendientes")
        self.assertIn("3 documento(s) pendientes de conversión", mensaje)
        self.assertIn("Se conservaron en su carpeta", mensaje)


if __name__ == "__main__":
    unittest.main()
