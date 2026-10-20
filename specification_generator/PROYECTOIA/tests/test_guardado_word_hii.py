import json
from pathlib import Path
import queue
import tempfile
from threading import get_ident
import unittest
from unittest.mock import Mock, patch

from docx import Document

import generador_hii as modulo


def crear_formato(ruta, ingles=False):
    documento = Document()
    documento.sections[0].header.paragraphs[0].text = "ESPECIFICACIÓN NOMBRE"
    documento.sections[0].footer.paragraphs[0].text = "Dirección"
    documento.add_paragraph("FORMATO INGLÉS" if ingles else "FORMATO ESPAÑOL")
    control = documento.add_table(rows=4, cols=2)
    for fila, etiqueta in zip(control.rows, ("PART ID", "REVISION", "DATE", "CUSTOMER NAME")):
        fila.cells[0].text = etiqueta
    pruebas = documento.add_table(rows=3, cols=5)
    encabezados = ("TEST", "UNITS", "METHOD", "SPECIFICATION", "TYPICAL VALUE") if ingles else (
        "PRUEBA", "UNIDADES", "MÉTODO", "ESPECIFICACIÓN", "VALOR TÍPICO")
    for celda, etiqueta in zip(pruebas.rows[0].cells, encabezados):
        celda.text = etiqueta
    pruebas.rows[2].cells[0].text = "Notes:" if ingles else "Notas:"
    documento.save(ruta)


class GuardadoWordHIITests(unittest.TestCase):
    def setUp(self):
        temporal = tempfile.TemporaryDirectory()
        self.addCleanup(temporal.cleanup)
        self.base = Path(temporal.name)
        self.formatos = self.base / "formatos"
        self.formatos.mkdir()
        self.rutas = {}
        self.plantillas = {}
        for idioma, ingles in (("ESPANOL", False), ("INGLES", True)):
            plantilla = self.formatos / ("formato (ENG).docx" if ingles else "formato.docx")
            crear_formato(plantilla, ingles)
            self.plantillas[idioma] = plantilla
            valores = {"prueba": "VISCOSITY" if ingles else "VISCOSIDAD", "metodo": "ASTM D445",
                       "unidades": "cSt", "especificacion": "20 MAX", "valor_tipico": "15"}
            datos = {
                "producto": {"nombre": "PRODUCT EN" if ingles else "PRODUCTO ES"},
                "documento": {"idioma": "en" if ingles else "es"},
                "control_documental": {"codigo_unico_producto": "20100001", "revision": "07",
                                       "fecha": "13/10/26", "nombre_cliente": "CLIENTE"},
                "pruebas": [{**valores, "especificacion": "VALOR ORIGINAL"}],
                "notas": ["NOTE EN" if ingles else "NOTA ES"],
                "alineacion_metodos_lab": {"detalle": [{"valores_alineados": valores}]},
            }
            ruta = self.base / f"producto_{idioma}.json"
            ruta.write_text(json.dumps(datos), encoding="utf-8")
            self.rutas[idioma] = ruta

    def generador(self, ambos=False, idioma="ESPANOL"):
        g = modulo.GeneradorHII.__new__(modulo.GeneradorHII)
        g.formatos = self.formatos
        g.datos = modulo.preparar_datos(self.rutas[idioma])
        g.actual = self.rutas[idioma]
        g.origen_actual = self.rutas["ESPANOL"]
        g.idioma_objetivo = idioma
        g.ambos_idiomas = Mock(get=Mock(return_value=ambos))
        g.resolver_idioma = Mock(side_effect=lambda origen: self.rutas[g.idioma_objetivo])
        g._disponibles = Mock(return_value={"ESPANOL", "INGLES"})
        g._guardando_word = False
        g._cola_word = queue.Queue()
        g.ventana = Mock()
        g.guardar_word_boton = Mock()
        g._programar = Mock()
        g.estado = Mock()
        hilo_principal = get_ident()
        g.estado.configure.side_effect = lambda **_: self.assertEqual(get_ident(), hilo_principal)
        g.guardar_word_boton.configure.side_effect = lambda **_: self.assertEqual(get_ident(), hilo_principal)
        return g

    def terminar(self, g):
        self.assertTrue(g._guardando_word)
        resultado = g._cola_word.get(timeout=10)
        g._cola_word.put(resultado)
        g._recibir_word()
        self.assertFalse(g._guardando_word)
        return resultado

    def verificar_word(self, ruta, ingles):
        documento = Document(ruta)
        self.assertIn("FORMATO INGLÉS" if ingles else "FORMATO ESPAÑOL", documento.paragraphs[0].text)
        self.assertIn("PRODUCT EN" if ingles else "PRODUCTO ES", documento.sections[0].header.paragraphs[0].text)
        control, pruebas = documento.tables
        self.assertEqual([fila.cells[1].text for fila in control.rows], ["20100001", "07", "13/10/26", "CLIENTE"])
        self.assertEqual(pruebas.rows[1].cells[0].text, "VISCOSITY" if ingles else "VISCOSIDAD")
        self.assertEqual(pruebas.rows[1].cells[3].text, "20 MAX")
        self.assertEqual(pruebas.rows[1].cells[4].text, "15")
        self.assertIn("NOTE EN" if ingles else "NOTA ES", pruebas.rows[-1].cells[0].text)

    def test_guarda_idioma_elegido_con_formato_word_y_control_sin_incrementar(self):
        g = self.generador(idioma="INGLES")
        destino = self.base / "salida.docx"
        originales = {ruta: ruta.read_bytes() for ruta in (*self.rutas.values(), *self.plantillas.values())}
        with patch.object(modulo, "filedialog") as dialogos, patch.object(modulo, "messagebox") as mensajes:
            dialogos.asksaveasfilename.return_value = str(destino)
            g.guardar_word()
            g.guardar_word_boton.configure.assert_called_with(state="disabled")
            guardadas, errores = self.terminar(g)
            self.assertEqual(guardadas, [destino])
            self.assertEqual(errores, [])
            mensajes.showerror.assert_not_called()
            dialogos.askdirectory.assert_not_called()
        self.verificar_word(destino, ingles=True)
        self.assertEqual(originales, {ruta: ruta.read_bytes() for ruta in originales})
        g.guardar_word_boton.configure.assert_called_with(state="normal")

    def test_ambos_exporta_version_y_formato_de_cada_idioma(self):
        g = self.generador(ambos=True, idioma="INGLES")
        carpeta = self.base / "salidas"
        carpeta.mkdir()
        with patch.object(modulo, "filedialog") as dialogos, patch.object(modulo, "messagebox") as mensajes:
            dialogos.askdirectory.return_value = str(carpeta)
            g.guardar_word()
            guardadas, errores = self.terminar(g)
            self.assertEqual(len(guardadas), 2)
            self.assertEqual(errores, [])
            mensajes.showerror.assert_not_called()
            dialogos.asksaveasfilename.assert_not_called()
        self.assertEqual(g.idioma_objetivo, "INGLES")
        for ruta in guardadas:
            self.verificar_word(ruta, ingles="INGLES" in ruta.name)

    def test_cancelar_dialogos_no_inicia_exportacion(self):
        for ambos in (True, False):
            with self.subTest(ambos=ambos):
                g = self.generador(ambos=ambos)
                with patch.object(modulo, "filedialog") as dialogos, patch.object(modulo.threading, "Thread") as hilo:
                    dialogos.askdirectory.return_value = ""
                    dialogos.asksaveasfilename.return_value = ""
                    g.guardar_word()
                    hilo.assert_not_called()
                self.assertFalse(g._guardando_word)

    def test_no_hay_formato_del_idioma_informa_antes_de_pedir_destino(self):
        self.plantillas["INGLES"].unlink()
        g = self.generador(idioma="INGLES")
        with patch.object(modulo, "filedialog") as dialogos, patch.object(modulo, "messagebox") as mensajes:
            g.guardar_word()
            mensajes.showerror.assert_called_once()
            self.assertIn("INGLES", mensajes.showerror.call_args.args[1])
            dialogos.asksaveasfilename.assert_not_called()
        self.assertFalse(g._guardando_word)

    def test_no_permite_reemplazar_plantilla_ni_extension_incorrecta(self):
        for destino in (self.plantillas["ESPANOL"], self.base / "salida.xlsm"):
            with self.subTest(destino=destino):
                g = self.generador()
                with patch.object(modulo, "filedialog") as dialogos, patch.object(modulo, "messagebox") as mensajes:
                    dialogos.asksaveasfilename.return_value = str(destino)
                    g.guardar_word()
                    mensajes.showerror.assert_called_once()
                self.assertFalse(g._guardando_word)

    def test_rechazar_reemplazo_de_ambos_conserva_documento(self):
        g = self.generador(ambos=True)
        destino = self.base / f"{self.rutas['ESPANOL'].stem} ESPANOL.docx"
        destino.write_bytes(b"Anterior")
        with patch.object(modulo, "filedialog") as dialogos, patch.object(modulo, "messagebox") as mensajes:
            dialogos.askdirectory.return_value = str(self.base)
            mensajes.askyesno.return_value = False
            g.guardar_word()
            mensajes.askyesno.assert_called_once()
        self.assertFalse(g._guardando_word)
        self.assertEqual(destino.read_bytes(), b"Anterior")

    def test_error_de_generacion_conserva_salida_anterior_y_reactiva_boton(self):
        g = self.generador()
        destino = self.base / "salida.docx"
        destino.write_bytes(b"Anterior")
        with patch.object(modulo, "filedialog") as dialogos, patch.object(modulo, "messagebox") as mensajes, patch.object(
            modulo, "generar_documento", side_effect=ValueError("No se pudo completar")
        ):
            dialogos.asksaveasfilename.return_value = str(destino)
            g.guardar_word()
            guardadas, errores = self.terminar(g)
            self.assertEqual(guardadas, [])
            self.assertIn("No se pudo completar", errores[0])
            mensajes.showerror.assert_called_once()
        self.assertEqual(destino.read_bytes(), b"Anterior")
        self.assertEqual(list(self.base.glob("word_*")), [])
        g.guardar_word_boton.configure.assert_called_with(state="normal")

    def test_preparar_word_no_muta_original_ni_usa_valores_sin_alinear(self):
        ruta = self.rutas["ESPANOL"]
        original = ruta.read_bytes()
        datos = modulo.preparar_datos_word(ruta)
        self.assertEqual(datos["pruebas"][0]["especificacion"], "20 MAX")
        self.assertEqual(ruta.read_bytes(), original)

    def test_plantillas_reales_de_ambos_idiomas(self):
        try:
            formatos = modulo.detectar_formatos_word(modulo.BASE / "FORMATOS")
            plantillas = {idioma: modulo.seleccionar_plantilla_word(formatos, idioma)
                          for idioma in ("ESPANOL", "INGLES")}
        except (OSError, ValueError):
            self.skipTest("Plantillas Word locales no disponibles")
        for idioma in ("ESPANOL", "INGLES"):
            with self.subTest(idioma=idioma):
                plantilla = plantillas[idioma]
                datos = modulo.preparar_datos_word(self.rutas[idioma])
                destino = self.base / f"real_{idioma}.docx"
                guardadas, errores = modulo.generar_archivos_word([(self.rutas[idioma], destino, datos, plantilla)])
                self.assertEqual(errores, [])
                self.assertEqual(guardadas, [destino])
                documento = Document(destino)
                texto = "\n".join(celda.text for tabla in documento.tables for fila in tabla.rows for celda in fila.cells)
                self.assertIn("20 MAX", texto)
                self.assertNotIn("VALOR ORIGINAL", texto)


if __name__ == "__main__":
    unittest.main()
