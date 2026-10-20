import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml

import gestor_especificaciones as gestor


class EstructurasWordTests(unittest.TestCase):
    def setUp(self):
        temporal = tempfile.TemporaryDirectory()
        self.addCleanup(temporal.cleanup)
        self.ruta = Path(temporal.name) / "especificacion.docx"

    def pruebas(self, contenedor, ingles=False, titulo=False, cantidad=1):
        tabla = contenedor.add_table(rows=1, cols=4)
        if titulo:
            tabla.rows[0].cells[0].merge(tabla.rows[0].cells[3]).text = "PRUEBAS DEL PRODUCTO"
            encabezado = tabla.add_row()
        else:
            encabezado = tabla.rows[0]
        etiquetas = ("TEST", "UNITS", "TEST METHOD", "SPECIFICATION") if ingles else (
            "PRUEBA", "UNIDADES", "MÉTODO", "ESPECIFICACIÓN")
        for c, t in zip(encabezado.cells, etiquetas):
            c.text = t
        for i in range(cantidad):
            fila = tabla.add_row()
            for c, t in zip(fila.cells, (f"VISCOSIDAD {i}", "mm²/s", "ASTM D445 / ASTM D7042", "15.86 TYPICAL")):
                c.text = t
        return tabla

    def control(self, doc, pares):
        tabla = doc.add_table(rows=0, cols=2)
        for etiqueta, valor in pares:
            fila = tabla.add_row()
            fila.cells[0].text = etiqueta
            fila.cells[1].text = valor
        return tabla

    def leer(self, doc, **kwargs):
        doc.save(self.ruta)
        return gestor.extraer(self.ruta, **kwargs)

    def test_tabla_anidada_y_encabezado_en_segunda_fila(self):
        doc = Document()
        doc.add_paragraph("ESPECIFICACIÓN")
        doc.add_paragraph("PRODUCTO ANTIGUO")
        contenedor = doc.add_table(rows=1, cols=1)
        self.pruebas(contenedor.cell(0, 0), titulo=True)
        datos = self.leer(doc)
        self.assertEqual(datos["producto"]["nombre"], "PRODUCTO ANTIGUO")
        self.assertEqual(len(datos["pruebas"]), 1)
        self.assertEqual(datos["pruebas"][0]["fila_origen"], 3)
        self.assertEqual(datos["pruebas"][0]["metodo"], "ASTM D445 / ASTM D7042")

    def test_nombre_en_cuadro_de_texto(self):
        doc = Document()
        p = doc.add_paragraph()
        p._p.append(parse_xml('''<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
          <w:pict><w:txbxContent><w:p><w:r><w:t>PRODUCTO EN CUADRO</w:t></w:r></w:p></w:txbxContent></w:pict>
          </w:r>'''))
        self.pruebas(doc)
        self.assertEqual(self.leer(doc)["producto"]["nombre"], "PRODUCTO EN CUADRO")

    def test_r01_conserva_nota_y_no_confunde_fecha_del_formato(self):
        doc = Document()
        doc.add_paragraph("PRODUCTO ANTIGUO")
        self.pruebas(doc, cantidad=5)
        doc.add_paragraph("◊Característica crítica de control para estadístico.")
        self.control(doc, (("Elaboró:", "Validación de Producto"), ("Folio APQP", "VAP12K331"),
                           ("Fecha", "27/11/12"), ("Revisión", "00"),
                           ("Nombre Cliente", "CLIENTE")))
        tabla = doc.sections[0].footer.add_table(rows=2, cols=5, width=6000000)
        for c, t in zip(tabla.rows[0].cells, ("CÓDIGO", "REVISIÓN", "FECHA DE REV.", "RETENER", "RESPONSABLE")):
            c.text = t
        for c, t in zip(tabla.rows[1].cells, ("R01-PRO-VAP-01", "00", "07/04/2010", "1 AÑO", "VALIDACIÓN")):
            c.text = t
        datos = self.leer(doc)
        self.assertEqual(len(datos["pruebas"]), 5)
        self.assertEqual(datos["control_documental"]["elaboro_y_aprobo"], "Validación de Producto")
        self.assertEqual(datos["control_documental"]["fecha"], "27/11/12")
        self.assertEqual(datos["documento"]["codigo_formato"], "R01-PRO-VAP-01")
        self.assertEqual(datos["documento"]["revision_formato"], "00")
        self.assertEqual(datos["notas"], ["◊Característica crítica de control para estadístico."])

    def test_ingles_apostrofes_y_formulario_de_aceptacion(self):
        doc = Document()
        doc.add_paragraph("SPECIFICATION")
        doc.add_paragraph("PRODUCTO EN INGLÉS")
        self.pruebas(doc, ingles=True, cantidad=12)
        self.control(doc, (("COMPANY:", ""), ("NAME:", ""), ("DATE:", "")))
        self.control(doc, (("ELABORATED AND\nAPPROVED BY:", "RESEARCH AND DEVELOPMENT"),
                           ("PART ID:", "20300017"), ("DATE:", "08/06/2021"),
                           ("REVISION:", "2021"), ("CUSTOMER’S NAME:", "CLIENTE INGLÉS"),
                           ("CUSTOMER’S REVISION NUMBER:", "----")))
        datos = self.leer(doc, idioma="en")
        self.assertEqual(datos["control_documental"]["fecha"], "08/06/2021")
        self.assertEqual(datos["control_documental"]["nombre_cliente"], "CLIENTE INGLÉS")
        self.assertEqual(datos["control_documental"]["revision_cliente"], "----")
        self.assertEqual(len(datos["pruebas"]), 12)
        self.assertEqual(datos["pruebas"][0]["especificacion"], "15.86 TYPICAL")
        self.assertEqual(datos["pruebas"][0]["valor_tipico"], "")
        self.assertTrue(gestor.contiene_marcadores_ingles(self.ruta))

    def test_codigo_de_formato_numerico_es_distinto_del_producto(self):
        doc = Document()
        doc.add_paragraph("PRODUCTO")
        self.pruebas(doc, cantidad=13)
        self.control(doc, (("Código Único de\nProducto", "20500194"), ("Fecha", "08/01/2024"),
                           ("Revisión", "2023")))
        doc.sections[0].footer.paragraphs[0].text = "CÓDIGO:91600060 REV. 00"
        datos = self.leer(doc)
        self.assertEqual(datos["control_documental"]["codigo_unico_producto"], "20500194")
        self.assertEqual(datos["control_documental"]["revision"], "2023")
        self.assertEqual(datos["documento"]["codigo_formato"], "91600060")
        self.assertEqual(datos["documento"]["revision_formato"], "00")

    def test_no_acepta_una_tabla_sin_columna_de_nombres(self):
        doc = Document()
        doc.add_paragraph("PRODUCTO")
        tabla = self.pruebas(doc)
        tabla.cell(0, 0).text = "DESCONOCIDA"
        with self.assertRaisesRegex(ValueError, "nombre y especificacion"):
            self.leer(doc)

    def test_no_acepta_filas_sin_nombre_de_prueba(self):
        doc = Document()
        doc.add_paragraph("PRODUCTO")
        tabla = self.pruebas(doc)
        tabla.cell(1, 0).text = ""
        with self.assertRaisesRegex(ValueError, "sin nombre"):
            self.leer(doc)

    def test_catalogo_rechaza_dos_nombres_igualmente_validos(self):
        doc = Document()
        doc.add_paragraph("PRODUCTO A")
        doc.add_paragraph("PRODUCTO B")
        self.pruebas(doc)
        catalogo = [{"descripcion_clave": "PRODUCTO A"}, {"descripcion_clave": "PRODUCTO B"}]
        with self.assertRaisesRegex(ValueError, "ambiguo"):
            self.leer(doc, catalogo=catalogo)

    def test_catalogo_distingue_titulo_de_nombre_institucional(self):
        doc = Document()
        doc.sections[0].header.paragraphs[0].text = "Chemical & Schutz High Performance Lubricants"
        doc.add_paragraph("ESPECIFICACIÓN")
        doc.add_paragraph("PRODUCTO A")
        self.pruebas(doc)
        datos = self.leer(doc, catalogo=[{"descripcion_clave": "PRODUCTO A"}])
        self.assertEqual(datos["producto"]["nombre"], "PRODUCTO A")

    def test_aliases_ingleses_con_resultados_tipicos_y_tipo_de_prueba(self):
        doc = Document()
        doc.add_paragraph("PRODUCTO")
        tabla = doc.add_table(rows=2, cols=4)
        for c, t in zip(tabla.rows[0].cells, ("TEST DESCRIPTION", "TEST TYPE", "ASTM METHOD", "TYPICAL RESULTS")):
            c.text = t
        for c, t in zip(tabla.rows[1].cells, ("Viscosity", "B", "D445", "403.1")):
            c.text = t
        prueba = self.leer(doc)["pruebas"][0]
        self.assertEqual(prueba["prueba"], "Viscosity")
        self.assertEqual(prueba["tipo_prueba"], "B")
        self.assertEqual(prueba["valor_tipico"], "403.1")
        self.assertEqual(prueba["especificacion"], "")

    def test_columnas_prueba_duplicadas_no_sobrescriben_los_nombres(self):
        doc = Document()
        doc.add_paragraph("PRODUCTO")
        tabla = doc.add_table(rows=2, cols=5)
        for c, t in zip(tabla.rows[0].cells, ("PRUEBA", "UNIDADES", "METODO", "ESPECIFICACION", "PRUEBA")):
            c.text = t
        for c, t in zip(tabla.rows[1].cells, ("VISCOSIDAD", "cSt", "D445", "10 MINIMO", "15")):
            c.text = t
        with self.assertRaisesRegex(ValueError, "duplicadas"):
            self.leer(doc)

    def test_notas_plural_y_encabezado_repetido_no_son_pruebas(self):
        doc = Document()
        doc.add_paragraph("PRODUCTO")
        tabla = self.pruebas(doc)
        for c, t in zip(tabla.add_row().cells, ("PRUEBA", "UNIDADES", "METODO", "ESPECIFICACION")):
            c.text = t
        fila = tabla.add_row()
        fila.cells[0].merge(fila.cells[3]).text = "NOTAS: Característica controlada."
        datos = self.leer(doc)
        self.assertEqual(len(datos["pruebas"]), 1)
        self.assertEqual(datos["notas"], ["Característica controlada."])

    def test_codigo_de_producto_incrustado_en_tabla_no_es_una_prueba(self):
        doc = Document()
        doc.add_paragraph("PRODUCTO")
        tabla = self.pruebas(doc)
        fila = tabla.add_row()
        fila.cells[0].text = "PRODUCT CODE"
        fila.cells[1].text = "21800030"
        datos = self.leer(doc)
        self.assertEqual(len(datos["pruebas"]), 1)
        self.assertEqual(datos["control_documental"]["codigo_unico_producto"], "21800030")

    def test_control_de_cambios_usa_texto_vigente(self):
        doc = Document()
        doc.add_paragraph("PRODUCTO")
        tabla = self.pruebas(doc)
        celda = tabla.cell(1, 3)
        celda.text = ""
        celda.paragraphs[0]._p.append(parse_xml('''<w:del xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
          <w:r><w:delText>288 - 352</w:delText></w:r></w:del>'''))
        celda.paragraphs[0]._p.append(parse_xml('''<w:ins xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
          <w:r><w:t>28,8 – 35,2</w:t></w:r></w:ins>'''))
        datos = self.leer(doc)
        self.assertEqual(datos["pruebas"][0]["especificacion"], "28,8 – 35,2")

    def test_codigo_vacio_incrustado_no_reemplaza_part_id(self):
        doc = Document()
        doc.add_paragraph("PRODUCTO")
        tabla = self.pruebas(doc)
        fila = tabla.add_row()
        fila.cells[0].text = "CÓDIGO DE PRODUCTO"
        fila.cells[1].text = "----"
        self.control(doc, (("Part ID", "21800030"), ("Fecha", "08/06/2021")))
        datos = self.leer(doc)
        self.assertEqual(datos["control_documental"]["codigo_unico_producto"], "21800030")
        self.assertEqual(len(datos["pruebas"]), 1)

    def test_subtitulo_de_grupo_se_conserva_como_nota(self):
        doc = Document()
        doc.add_paragraph("PRODUCTO")
        tabla = self.pruebas(doc)
        fila = tabla.add_row()
        fila.cells[0].text = "Carbon Type Analysis, %"
        fila = tabla.add_row()
        for c, t in zip(fila.cells, ("Aromatic (Ca)", "%", "D2140", "4 MAXIMUM")):
            c.text = t
        datos = self.leer(doc)
        self.assertEqual(len(datos["pruebas"]), 2)
        self.assertIn("Carbon Type Analysis, %", datos["notas"])


if __name__ == "__main__":
    unittest.main()
