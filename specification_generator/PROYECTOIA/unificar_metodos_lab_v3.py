#!/usr/bin/env python3
from __future__ import annotations
import argparse, hashlib, json, re, shutil, sys, unicodedata
from collections import Counter
from datetime import datetime
from pathlib import Path

VERSION="3.1.0"; CARPETA="METODOS_LAB"; SALIDA="METODOS_LAB_UNIFICADO.xlsx"
BITACORA="CAMBIOS_METODOS_LAB.txt"; ESTADO=".estado_metodos_lab.json"
PYTHON_REQUERIDO=(3,11)
FUENTES={"METODOS ASTM FEUM USP NO ACREDITADOS REV2":"NORMAL","METODOS 3RIA":"NORMAL","LISTAS DE METODOS ACREDITADOS":"ACREDITADOS"}
COLUMNAS=["METODO","MÉTODO INGLÉS","NOMBRE IQM","NOMBRE IQM INGLES","UNIDADES","UNITS","ACREDITACION NOMBRE"]
CLAVE_FILA=("METODO","NOMBRE IQM","NOMBRE IQM INGLES","UNIDADES","UNITS","ACREDITACION NOMBRE")

def cargar_openpyxl():
    try: import openpyxl
    except ModuleNotFoundError as exc:
        if exc.name!="openpyxl":raise
        raise RuntimeError("Falta openpyxl. Ejecuta .venv\\Scripts\\python.exe -m pip install -r requirements_metodos_lab.txt") from exc
    global Workbook,load_workbook,Alignment,Font,PatternFill,Table,TableStyleInfo,get_column_letter
    from openpyxl import Workbook,load_workbook
    from openpyxl.styles import Alignment,Font,PatternFill
    from openpyxl.worksheet.table import Table,TableStyleInfo
    from openpyxl.utils import get_column_letter

def txt(v):
    if v is None:return ""
    if isinstance(v,float) and v.is_integer():v=int(v)
    s=str(v).replace("\xa0"," ").replace("\r"," ").replace("\n"," ").replace("\t"," ")
    return re.sub(r"\s+"," ","".join(c for c in s if c.isprintable())).strip()
def key(v):
    s="".join(c for c in unicodedata.normalize("NFD",txt(v).upper()) if unicodedata.category(c)!="Mn")
    return re.sub(r"\s+"," ",re.sub(r"[^A-Z0-9]+"," ",s)).strip()
def clean(v,asterisco=False):
    s=txt(v)
    if asterisco:s=re.sub(r"^\s*\*+\s*","",s)
    return "NA" if key(s) in {"NA","N A","NO APLICA","NO APLICABLE"} else s.strip(" ;")
def digest(p):
    h=hashlib.sha256()
    with p.open("rb") as f:
        for b in iter(lambda:f.read(1048576),b""):h.update(b)
    return h.hexdigest()
def archivos(carpeta):
    m={k:[] for k in FUENTES}
    for p in carpeta.iterdir():
        if p.is_file() and not p.name.startswith("~$") and p.suffix.lower() in {".xlsx",".xlsm"} and key(p.stem) in m:m[key(p.stem)].append(p)
    errores=[]; out=[]
    for n,t in FUENTES.items():
        if len(m[n])==0:errores.append(f"Falta: {n}.xlsx")
        elif len(m[n])>1:errores.append(f"Duplicado: {n}")
        else:out.append((m[n][0],t))
    if errores:raise FileNotFoundError("\n".join(errores))
    return out
def header_row(ws):
    best=(0,1)
    for r in range(1,min(ws.max_row,30)+1):
        ks={key(ws.cell(r,c).value) for c in range(1,ws.max_column+1)}
        score=5*bool({"METODO","METHOD"}&ks)+sum(x in ks for x in ("NOMBRE","NAME","NOMBRE IQM","NOMBRE IQM INGLES","UNIDADES","UNITS"))
        if score>best[0]:best=(score,r)
    if best[0]<6:raise ValueError("encabezado no reconocido")
    return best[1]
def leer(p,tipo):
    wb=load_workbook(p,read_only=True,data_only=True,keep_links=False); out=[]
    for ws in wb.worksheets:
        try:r0=header_row(ws)
        except ValueError:continue
        hs=[key(ws.cell(r0,c).value) for c in range(1,ws.max_column+1)]
        for r in range(r0+1,ws.max_row+1):
            d={h:txt(ws.cell(r,c+1).value) for c,h in enumerate(hs) if h}
            metodo=clean(d.get("METODO") or d.get("METHOD"))
            metodo_ingles=clean(d.get("METODO INGLES") or d.get("ENGLISH METHOD") or d.get("METHOD ENGLISH") or d.get("METODO EN INGLES"))
            if tipo=="ACREDITADOS":nom,eng=d.get("NOMBRE",""),d.get("NAME","")
            else:nom,eng=d.get("NOMBRE IQM",""),d.get("NOMBRE IQM INGLES","")
            if not any((metodo,txt(nom),txt(eng))):continue
            out.append({"METODO":metodo,"MÉTODO INGLÉS":metodo_ingles,"_METODO_INGLES_FUENTE":bool(metodo_ingles),"NOMBRE IQM":clean(nom,True),"NOMBRE IQM INGLES":clean(eng,True),"UNIDADES":clean(d.get("UNIDADES")),"UNITS":clean(d.get("UNITS")),"ACREDITACION NOMBRE":"ACREDITADO" if "*" in txt(nom) else "NO ACREDITADO","_ORIGEN":p.name})
    wb.close();return out

def claves_fila(fila):
    return tuple(key(fila.get(columna)) for columna in CLAVE_FILA)

def metodos_ingles_existentes(path):
    if not path.is_file():return {},{}
    wb=load_workbook(path,read_only=True,data_only=True,keep_links=False)
    try:
        if "METODOS" not in wb.sheetnames:return {},{}
        ws=wb["METODOS"]
        encabezados={key(ws.cell(1,c).value):c for c in range(1,ws.max_column+1)}
        if not all(key(columna) in encabezados for columna in CLAVE_FILA):return {},{}
        col_ingles=encabezados.get(key("MÉTODO INGLÉS"))
        if col_ingles is None:return {},{}
        resultado={};por_nombre={}
        for r in range(2,ws.max_row+1):
            fila={columna:txt(ws.cell(r,encabezados[key(columna)]).value) for columna in CLAVE_FILA}
            metodo_ingles=clean(ws.cell(r,col_ingles).value)
            if metodo_ingles:
                resultado.setdefault(claves_fila(fila),[]).append(metodo_ingles)
                nombre_ingles=key(fila["NOMBRE IQM INGLES"])
                if nombre_ingles:por_nombre.setdefault(nombre_ingles,[]).append(metodo_ingles)
        return resultado,por_nombre
    finally:wb.close()

def conservar_metodos_ingles(filas,existentes,por_nombre=None):
    for fila in filas:
        if not fila.pop("_METODO_INGLES_FUENTE",False):
            traducciones=existentes.get(claves_fila(fila),[])
            if traducciones:
                fila["MÉTODO INGLÉS"]=traducciones.pop(0)
            else:
                nombre_ingles=key(fila.get("NOMBRE IQM INGLES"))
                alternativas=[]
                if nombre_ingles and por_nombre:
                    alternativas=[
                        metodo
                        for nombre,metodos in por_nombre.items()
                        if nombre in nombre_ingles or nombre_ingles in nombre
                        for metodo in metodos
                    ]
                traducciones_unicas=tuple(dict.fromkeys(alternativas))
                fila["MÉTODO INGLÉS"]=(
                    traducciones_unicas[0]
                    if len(traducciones_unicas)==1
                    else fila["METODO"]
                )
def crear(path,filas):
    wb=Workbook();ws=wb.active;ws.title="METODOS";ws.append(COLUMNAS)
    for r in filas:ws.append([r[c] for c in COLUMNAS])
    blue=PatternFill("solid",fgColor="1F4E78");white=Font(color="FFFFFF",bold=True)
    for c in ws[1]:c.fill=blue;c.font=white;c.alignment=Alignment(horizontal="center",wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for c in row:c.alignment=Alignment(vertical="top",wrap_text=True)
        row[6].fill=PatternFill("solid",fgColor="E2F0D9" if row[6].value=="ACREDITADO" else "FCE4D6")
    for i,w in enumerate((25,25,52,52,18,18,23),1):ws.column_dimensions[get_column_letter(i)].width=w
    ws.freeze_panes="A2";ws.sheet_view.showGridLines=False
    if ws.max_row>1:
        tab=Table(displayName="TablaMetodos",ref=f"A1:G{ws.max_row}");tab.tableStyleInfo=TableStyleInfo(name="TableStyleMedium2",showRowStripes=True);ws.add_table(tab)
    rs=wb.create_sheet("RESUMEN");rs.append(["CONCEPTO","VALOR"]);rs.append(["Version",VERSION]);rs.append(["Fecha",datetime.now().isoformat(timespec="seconds")]);rs.append(["Total filas",len(filas)])
    for n,c in Counter(r["_ORIGEN"] for r in filas).items():rs.append([n,c])
    temp=path.with_suffix(".tmp.xlsx");wb.save(temp);shutil.move(temp,path)
def ejecutar(base):
    carpeta=base/CARPETA
    if not carpeta.exists():raise FileNotFoundError(f"No existe la carpeta {carpeta}")
    ep=carpeta/ESTADO;ant=json.loads(ep.read_text(encoding="utf-8")) if ep.exists() else {"archivos":{}}
    filas=[];actual={"version":VERSION,"archivos":{}};mensajes=[]
    metodos_ingles,metodos_por_nombre=metodos_ingles_existentes(carpeta/SALIDA)
    for p,t in archivos(carpeta):
        datos=leer(p,t);filas+=datos;firma={"hash":digest(p),"filas":len(datos)};actual["archivos"][p.name]=firma;viejo=ant.get("archivos",{}).get(p.name)
        if viejo is None:mensajes.append(f"{p.name}: registrado por primera vez con {len(datos)} filas utiles.")
        elif viejo.get("hash")!=firma["hash"]:mensajes.append(f"{p.name}: modificado. Filas utiles {viejo.get('filas',0)} -> {len(datos)}.")
    conservar_metodos_ingles(filas,metodos_ingles,metodos_por_nombre)
    crear(carpeta/SALIDA,filas)
    if mensajes:
        with (carpeta/BITACORA).open("a",encoding="utf-8") as f:f.write(f"\n[{datetime.now().isoformat(timespec='seconds')}]\n"+"\n".join(mensajes)+"\n")
    ep.write_text(json.dumps(actual,ensure_ascii=False,indent=2),encoding="utf-8")
    print(f"GENERADO: {carpeta/SALIDA}\nFILAS EN METODOS: {len(filas)}\nCAMBIOS REGISTRADOS: {len(mensajes)}")
def main():
    if sys.version_info[:2]!=PYTHON_REQUERIDO:
        raise SystemExit(f"Se requiere Python 3.11. Detectado {sys.version_info.major}.{sys.version_info.minor}.")
    p=argparse.ArgumentParser();p.add_argument("--base",type=Path,default=Path(__file__).resolve().parent);a=p.parse_args()
    try:cargar_openpyxl();ejecutar(a.base.resolve());return 0
    except Exception as e:print(f"ERROR: {e}",file=sys.stderr);return 1
if __name__=="__main__":raise SystemExit(main())
