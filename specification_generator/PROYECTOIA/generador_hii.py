#!/usr/bin/env python3
"""Genera HII bilingües preservando macros, objetos y formato del XLSM original."""
from __future__ import annotations

import argparse
import copy
import hashlib
import io
import json
import math
import os
import queue
import re
import sqlite3
import sys
import tempfile
import textwrap
import threading
import tkinter as tk
import tkinter.font as tkfont
from concurrent.futures import CancelledError
from contextlib import contextmanager
from dataclasses import asdict, dataclass, replace
from datetime import date, datetime
from fractions import Fraction
from functools import lru_cache
from pathlib import Path
from threading import Event
from tkinter import filedialog, messagebox, ttk as ttk_nativo
from uuid import uuid4
from zipfile import ZipFile, ZIP_DEFLATED

from lxml import etree as ET
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

import estilo_interfaz as ui
from estilo_interfaz import ttk
from generador_especs_hii import clave_texto, leer_idioma_datos, detectar_json, _separar_notas_de_pruebas
from interfaz_seleccion_especificaciones import (
    preparar_vista_previa_alineada, cargar_catalogo, buscar_especificaciones,
    encabezados_tabla, crear_divisor_redimensionable, Catalogo, Especificacion,
)

from rutas_proyecto import BASE, resolver_ruta, rutas_configuradas, ruta_portatil, reubicar_ruta_heredada
from especificaciones_wip import es_wip

if __name__ == "__main__":
    # Las llamadas de recuperación reutilizan este módulo al ejecutar el archivo directamente.
    sys.modules.setdefault("generador_hii", sys.modules[__name__])

NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"}
FECHA = re.compile(r"DD\s*/\s*MM\s*/\s*YY", re.I)
REFERENCIA = re.compile(r"(?<![\w])([$]?[A-Z]{1,3}[$]?)(\d+)(?!\w)")
COLUMNAS = ("numero", "prueba", "metodo", "periodicidad", "unidades", "especificacion")
FILAS_POR_BLOQUE = 100
EXTENSIONES_HII = {".xlsm", ".xlsx", ".xltm", ".xltx"}


# Historial y trazabilidad de los archivos HII

VERSION_GENERADOR = "1.1.0"
NS_CUSTOM = "http://schemas.openxmlformats.org/officeDocument/2006/custom-properties"
NS_VT = "http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"
NS_RELS = "http://schemas.openxmlformats.org/package/2006/relationships"
NS_TYPES = "http://schemas.openxmlformats.org/package/2006/content-types"
TIPO_CUSTOM = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties"


def hash_especificacion(contenido: dict) -> str:
    """Recibe los valores efectivos de la HII, sin fecha ni historial."""
    canonico = json.dumps(contenido, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(canonico.encode("utf-8")).hexdigest()


def ruta_relativa(archivo: Path, carpeta: Path) -> str:
    try:
        return Path(os.path.relpath(archivo.resolve(), carpeta.resolve())).as_posix()
    except ValueError:
        # Windows no admite rutas relativas entre unidades diferentes.
        return archivo.resolve().as_posix()


def leer_historial(datos: dict) -> list[dict]:
    if not isinstance(datos, dict):
        raise ValueError("El JSON debe contener un objeto; se conserva el archivo anterior.")
    historial = datos.get("hii_generadas", [])
    if not isinstance(historial, list) or any(not isinstance(r, dict) for r in historial):
        raise ValueError("'hii_generadas' debe ser una lista de registros; se conserva el archivo anterior.")
    if any("id" in r and not isinstance(r["id"], str) for r in historial):
        raise ValueError("El ID de una HII debe ser texto; se conserva el archivo anterior.")
    return copy.deepcopy(historial)


def conservar_historial(datos: dict, origen: Path, destino: Path) -> None:
    """Fusiona el historial actual del destino y reajusta referencias al cambiar de carpeta."""
    historial = leer_historial(datos)
    for registro in historial:
        if isinstance(registro.get("archivo"), str) and registro["archivo"]:
            registro["archivo"] = ruta_relativa(origen.parent / registro["archivo"], destino.parent)
    if destino.is_file():
        actuales = leer_historial(json.loads(destino.read_text(encoding="utf-8-sig")))
        por_id = {r.get("id"): i for i, r in enumerate(historial) if r.get("id")}
        for registro in actuales:
            indice = por_id.get(registro.get("id"))
            if indice is not None:
                historial[indice] = registro
            elif registro not in historial:
                historial.append(registro)
                if registro.get("id"):
                    por_id[registro["id"]] = len(historial) - 1
    datos["hii_generadas"] = historial


def crear_registro(destino: Path, especificacion: Path, huella: str) -> dict:
    return {
        "id": str(uuid4()),
        "archivo": ruta_relativa(destino, especificacion.parent),
        "fecha_generacion": datetime.now().astimezone().isoformat(timespec="seconds"),
        "version_generador": VERSION_GENERADOR,
        "hash_especificacion": huella,
    }


def escribir_atomico(ruta: Path, contenido: bytes) -> None:
    """Reemplaza solo al terminar de escribir; no trunca un archivo existente si falla."""
    temporal = None
    try:
        with tempfile.NamedTemporaryFile(dir=ruta.parent, prefix=f".{ruta.name}.", suffix=".tmp", delete=False) as archivo:
            temporal = Path(archivo.name)
            archivo.write(contenido)
            archivo.flush()
            os.fsync(archivo.fileno())
        os.replace(temporal, ruta)
    finally:
        if temporal is not None:
            temporal.unlink(missing_ok=True)


def incrustar_identificador(contenido: bytes, identificador: str) -> bytes:
    """Añade HII_ID como propiedad de Excel conservando macros y demás partes."""
    with ZipFile(io.BytesIO(contenido)) as origen:
        partes = {i.filename: origen.read(i) for i in origen.infolist()}
        propiedades = (ET.fromstring(partes["docProps/custom.xml"]) if "docProps/custom.xml" in partes
                       else ET.Element(f"{{{NS_CUSTOM}}}Properties", nsmap={None: NS_CUSTOM, "vt": NS_VT}))
        for propiedad in list(propiedades):
            if propiedad.get("name") == "HII_ID":
                propiedades.remove(propiedad)
        pid = max([1, *(int(p.get("pid", "1")) for p in propiedades)]) + 1
        propiedad = ET.SubElement(propiedades, f"{{{NS_CUSTOM}}}property", name="HII_ID", pid=str(pid),
                                 fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}")
        ET.SubElement(propiedad, f"{{{NS_VT}}}lpwstr").text = identificador
        relaciones = ET.fromstring(partes["_rels/.rels"])
        relacion = next((r for r in relaciones if r.get("Type") == TIPO_CUSTOM), None)
        if relacion is None:
            usados = {r.get("Id") for r in relaciones}
            indice = 1
            while f"rId{indice}" in usados:
                indice += 1
            relacion = ET.SubElement(relaciones, f"{{{NS_RELS}}}Relationship", Id=f"rId{indice}", Type=TIPO_CUSTOM)
        relacion.set("Target", "docProps/custom.xml")
        tipos = ET.fromstring(partes["[Content_Types].xml"])
        tipo = next((t for t in tipos if t.get("PartName") == "/docProps/custom.xml"), None)
        if tipo is None:
            tipo = ET.SubElement(tipos, f"{{{NS_TYPES}}}Override", PartName="/docProps/custom.xml")
        tipo.set("ContentType", "application/vnd.openxmlformats-officedocument.custom-properties+xml")
        for nombre, raiz in (("docProps/custom.xml", propiedades), ("_rels/.rels", relaciones), ("[Content_Types].xml", tipos)):
            partes[nombre] = ET.tostring(raiz, encoding="UTF-8", xml_declaration=True)
        salida = io.BytesIO()
        with ZipFile(salida, "w", compression=ZIP_DEFLATED) as archivo:
            archivo.comment = origen.comment
            for info in origen.infolist():
                archivo.writestr(info, partes.pop(info.filename))
            for nombre, datos in partes.items():
                archivo.writestr(nombre, datos)
    return salida.getvalue()


def guardar_con_historial(destino: Path, especificacion: Path, contenido: bytes, huella: str) -> dict:
    """Publica el Excel primero y registra después. Si falla el JSON, el Excel conserva su ID."""
    datos = json.loads(especificacion.read_text(encoding="utf-8-sig"))
    historial = leer_historial(datos)  # Validar antes de reemplazar el Excel.
    registro = crear_registro(destino, especificacion, huella)
    excel = incrustar_identificador(contenido, registro["id"])
    datos["hii_generadas"] = [*historial, registro]
    json_actualizado = (json.dumps(datos, ensure_ascii=False, indent=2) + "\n").encode("utf-8")
    escribir_atomico(destino, excel)
    try:
        escribir_atomico(especificacion, json_actualizado)
    except OSError as exc:
        raise OSError(f"El Excel se guardó en {destino}, con HII_ID={registro['id']}, "
                      f"pero no se pudo actualizar el historial del JSON: {exc}") from exc
    return registro

# Registro SQLite y relaciones entre versiones

def _base_registro(con):
    archivo = con.execute("PRAGMA database_list").fetchone()[2]
    return Path(archivo).parent if archivo else BASE


def _clave_ruta(con, ruta):
    return ruta_portatil(ruta, _base_registro(con))


def _ruta_registro(con, ruta):
    return resolver_ruta(ruta, _base_registro(con))


def _migrar_rutas(con):
    """Convierte registros anteriores una sola vez, conservando sus enlaces y sus HII."""
    if con.execute("SELECT 1 FROM hii_estado WHERE clave='rutas_portatiles'").fetchone():
        return
    base = _base_registro(con)
    cambios = {}
    for fila in con.execute("SELECT * FROM hii_especificaciones").fetchall():
        anterior = fila["ruta"]
        nueva = ruta_portatil(reubicar_ruta_heredada(anterior, base), base)
        if nueva != anterior:
            cambios[anterior] = nueva
            con.execute("INSERT OR IGNORE INTO hii_especificaciones VALUES (?, ?, ?, ?, ?)",
                        (nueva, fila["codigo"], fila["revision"], fila["cliente"], fila["idioma"]))
    # Se crean primero los destinos de las claves foráneas; los cambios son atómicos.
    for tabla in ("hii_traducciones", "hii_enlaces_automaticos"):
        for origen, destino in con.execute(f"SELECT origen, destino FROM {tabla}").fetchall():
            pareja = (cambios.get(origen, origen), cambios.get(destino, destino))
            if pareja != (origen, destino):
                con.execute(f"INSERT OR IGNORE INTO {tabla} VALUES (?, ?)", pareja)
                con.execute(f"DELETE FROM {tabla} WHERE origen=? AND destino=?", (origen, destino))
    for anterior, nueva in cambios.items():
        con.execute("UPDATE hii_generadas SET especificacion=? WHERE especificacion=?", (nueva, anterior))
        con.execute("DELETE FROM hii_especificaciones WHERE ruta=?", (anterior,))
    for fila in con.execute("SELECT * FROM hii_generadas").fetchall():
        anterior = fila["ruta"]
        nueva = ruta_portatil(reubicar_ruta_heredada(anterior, base), base)
        if nueva != anterior:
            con.execute("INSERT OR IGNORE INTO hii_generadas VALUES (?, ?, ?, ?)",
                        (nueva, fila["especificacion"], fila["idioma"], fila["fecha"]))
            con.execute("DELETE FROM hii_generadas WHERE ruta=?", (anterior,))
    con.execute("INSERT OR IGNORE INTO hii_estado VALUES ('rutas_portatiles', '1')")


@contextmanager
def conectar(ruta, *, recuperar=True):
    ruta = resolver_ruta(ruta)
    if recuperar:
        from reconstruir_base_datos import asegurar_base
        asegurar_base(ruta.parent, destino=ruta)
    con = sqlite3.connect(ruta, timeout=30)
    try:
        con.row_factory = sqlite3.Row
        con.execute("PRAGMA busy_timeout=30000")
        con.execute("PRAGMA foreign_keys=ON")
        modo_journal = con.execute("PRAGMA journal_mode").fetchone()[0]
        if str(modo_journal).lower() != "wal":
            con.execute("PRAGMA journal_mode=WAL")
        con.executescript("""
            CREATE TABLE IF NOT EXISTS hii_especificaciones (
                ruta TEXT PRIMARY KEY, codigo TEXT NOT NULL, revision TEXT NOT NULL,
                cliente TEXT NOT NULL, idioma TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS hii_traducciones (
                origen TEXT NOT NULL REFERENCES hii_especificaciones(ruta),
                destino TEXT NOT NULL REFERENCES hii_especificaciones(ruta),
                PRIMARY KEY(origen, destino));
            CREATE TABLE IF NOT EXISTS hii_enlaces_automaticos (
                origen TEXT NOT NULL, destino TEXT NOT NULL,
                PRIMARY KEY(origen, destino));
            CREATE TABLE IF NOT EXISTS hii_generadas (
                ruta TEXT PRIMARY KEY, especificacion TEXT NOT NULL REFERENCES hii_especificaciones(ruta),
                idioma TEXT NOT NULL, fecha TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS hii_estado (
                clave TEXT PRIMARY KEY, valor TEXT NOT NULL);
            CREATE VIEW IF NOT EXISTS hii_versiones_idiomas AS
                SELECT a.ruta AS hii_origen, a.idioma AS idioma_origen,
                       b.ruta AS hii_traduccion, b.idioma AS idioma_traduccion,
                       a.especificacion AS especificacion_origen,
                       b.especificacion AS especificacion_traduccion
                FROM hii_generadas a JOIN hii_generadas b
                JOIN hii_traducciones t ON
                    (a.especificacion=t.origen AND b.especificacion=t.destino) OR
                    (a.especificacion=t.destino AND b.especificacion=t.origen);
        """)
        with con:
            _migrar_rutas(con)
            yield con
    finally:
        con.close()


def indexar(con, ruta):
    ruta = _ruta_registro(con, ruta)
    datos = json.loads(ruta.read_text(encoding="utf-8-sig"))
    control = datos.get("control_documental", {})
    fila = (_clave_ruta(con, ruta), clave_texto(str(control.get("codigo_unico_producto") or "")),
            clave_texto(str(control.get("revision") or "")),
            clave_texto(str(control.get("nombre_cliente") or "")), leer_idioma_datos(datos))
    con.execute("""INSERT INTO hii_especificaciones VALUES (?, ?, ?, ?, ?)
        ON CONFLICT(ruta) DO UPDATE SET codigo=excluded.codigo,
        revision=excluded.revision, cliente=excluded.cliente, idioma=excluded.idioma""", fila)
    return fila


def candidatas(con, origen, idioma, rutas_disponibles):
    fila = indexar(con, origen)
    disponibles = {_clave_ruta(con, p) for p in rutas_disponibles}
    explicitas = con.execute("""SELECT DISTINCT e.* FROM hii_especificaciones e
        JOIN hii_traducciones t ON (e.ruta=t.destino AND t.origen=?)
          OR (e.ruta=t.origen AND t.destino=?) WHERE e.idioma=?""",
        (fila[0], fila[0], idioma)).fetchall()
    validas = [_ruta_registro(con, e["ruta"]) for e in explicitas if e["ruta"] in disponibles
               and (e["codigo"], e["cliente"]) == (fila[1], fila[3])]
    if validas:
        return validas
    if not fila[1] or not fila[2]:
        return []
    return [_ruta_registro(con, e["ruta"]) for e in con.execute("""SELECT ruta FROM hii_especificaciones
        WHERE codigo=? AND revision=? AND cliente=? AND idioma=? ORDER BY ruta""",
        (*fila[1:4], idioma)) if e["ruta"] in disponibles]


def vincular(con, origen, destino, permitir_revision_distinta=False):
    a, b = indexar(con, origen), indexar(con, destino)
    compatible = (a[1], a[3]) == (b[1], b[3]) and (permitir_revision_distinta or a[2] == b[2])
    if not compatible or a[4] == b[4]:
        raise ValueError("Las traducciones deben tener el mismo código, revisión y cliente, y distinto idioma.")
    con.execute("""INSERT OR IGNORE INTO hii_traducciones SELECT ?, ?
        WHERE NOT EXISTS (SELECT 1 FROM hii_traducciones WHERE origen=? AND destino=?)""",
        (a[0], b[0], b[0], a[0]))
    con.execute("DELETE FROM hii_enlaces_automaticos WHERE (origen=? AND destino=?) OR (origen=? AND destino=?)",
                (a[0], b[0], b[0], a[0]))
    exportar_relaciones(con, [origen, destino])


def exportar_relaciones(con, rutas):
    """Referencias relativas al JSON actual; admite más de una traducción confirmada."""
    modificadas = set()
    for ruta in rutas:
        ruta = _ruta_registro(con, ruta)
        clave = _clave_ruta(con, ruta)
        datos = json.loads(ruta.read_text(encoding="utf-8-sig"))
        relaciones = {}
        filas = con.execute("""SELECT DISTINCT e.* FROM hii_especificaciones e
            JOIN hii_traducciones t ON (t.origen=? AND t.destino=e.ruta)
                OR (t.destino=? AND t.origen=e.ruta) ORDER BY e.ruta""", (clave, clave))
        actual = con.execute("SELECT * FROM hii_especificaciones WHERE ruta=?", (clave,)).fetchone()
        for fila in filas:
            destino = _ruta_registro(con, fila["ruta"])
            if not destino.is_file() or not actual or fila["idioma"] == actual["idioma"]:
                continue
            if (fila["codigo"], fila["cliente"]) != (actual["codigo"], actual["cliente"]):
                continue
            relaciones.setdefault(fila["idioma"], []).append({
                "archivo_json": Path(os.path.relpath(destino, ruta.parent)).as_posix(),
                "codigo": fila["codigo"], "revision": fila["revision"], "cliente": fila["cliente"],
            })
        if not relaciones and "versiones_idioma" not in datos:
            continue
        if datos.get("versiones_idioma") == relaciones:
            continue
        datos["versiones_idioma"] = relaciones
        temporal = None
        try:
            with tempfile.NamedTemporaryFile(mode="w", encoding="utf-8", dir=ruta.parent,
                prefix=".relaciones_", suffix=".tmp", delete=False) as archivo:
                temporal = Path(archivo.name)
                json.dump(datos, archivo, ensure_ascii=False, indent=2)
                archivo.write("\n")
            os.replace(temporal, ruta)
        finally:
            if temporal is not None:
                temporal.unlink(missing_ok=True)
        modificadas.add(ruta)
    return modificadas


def sincronizar_idiomas(con, carpeta):
    """Autoenlaza solo grupos con una única versión por idioma y revisión."""
    from collections import defaultdict
    grupos = defaultdict(list)
    rutas = []
    for ruta in sorted(_ruta_registro(con, carpeta).rglob("*")):
        if not ruta.is_file() or ruta.suffix.lower() != ".json" or ruta.name.startswith("."):
            continue
        try:
            fila = indexar(con, ruta)
        except (OSError, ValueError, AttributeError):
            continue
        rutas.append(ruta.resolve())
        if fila[1] and fila[2] and fila[4] in {"ESPANOL", "INGLES"}:
            grupos[fila[1:4]].append(fila)
    disponibles = {_clave_ruta(con, r) for r in rutas}
    for a, b in con.execute("SELECT origen, destino FROM hii_enlaces_automaticos").fetchall():
        if a in disponibles or b in disponibles:
            con.execute("DELETE FROM hii_traducciones WHERE origen=? AND destino=?", (a, b))
            con.execute("DELETE FROM hii_enlaces_automaticos WHERE origen=? AND destino=?", (a, b))
    for grupo in grupos.values():
        es = [r for r in grupo if r[4] == "ESPANOL"]
        en = [r for r in grupo if r[4] == "INGLES"]
        if len(es) == len(en) == 1:
            pareja = (es[0][0], en[0][0])
            existe = con.execute("SELECT 1 FROM hii_traducciones WHERE (origen=? AND destino=?) OR (origen=? AND destino=?)",
                                 (*pareja, pareja[1], pareja[0])).fetchone()
            if not existe:
                con.execute("INSERT INTO hii_traducciones VALUES (?, ?)", pareja)
                con.execute("INSERT INTO hii_enlaces_automaticos VALUES (?, ?)", pareja)
    return exportar_relaciones(con, rutas)


def registrar(con, excel, especificacion, idioma, fecha):
    indexar(con, especificacion)
    con.execute("""INSERT INTO hii_generadas VALUES (?, ?, ?, ?)
        ON CONFLICT(ruta) DO UPDATE SET especificacion=excluded.especificacion,
        idioma=excluded.idioma, fecha=excluded.fecha""",
        (_clave_ruta(con, excel), _clave_ruta(con, especificacion), idioma, fecha))

# Catalogo en cache y disponibilidad de idiomas

VERSION_CACHE = 1
NOMBRE_CACHE = ".cache_catalogo_hii"


def _serializar(especificacion: Especificacion, carpeta: Path) -> dict:
    datos = asdict(especificacion)
    datos.pop("ruta")
    datos["versiones_idioma"] = [ruta_relativa(ruta, carpeta) for ruta in especificacion.versiones_idioma]
    return datos


def _restaurar(datos: dict, ruta: Path, carpeta: Path) -> Especificacion:
    if not isinstance(datos, dict):
        raise ValueError("Entrada de caché inválida")
    datos = dict(datos)
    for campo in ("nombre_producto", "codigo", "cliente", "revision", "busqueda", "idioma", "version_alineador"):
        if not isinstance(datos.get(campo), str):
            raise ValueError("Metadatos de caché inválidos")
    for campo in ("edicion_bilingue", "en_progreso"):
        if not isinstance(datos.get(campo), bool):
            raise ValueError("Estado de caché inválido")
    referencias = datos.get("versiones_idioma")
    if not isinstance(referencias, list) or any(not isinstance(r, str) for r in referencias):
        raise ValueError("Referencias de caché inválidas")
    datos["versiones_idioma"] = tuple((carpeta / r).resolve() for r in referencias)
    return Especificacion(ruta=ruta, **datos)


def cargar_catalogo_hii(carpeta: Path, *, cache: Path | None = None, cancelar: Event | None = None) -> Catalogo:
    """Relee solo JSON nuevos/modificados; una caché ilegible o no escribible no impide abrir."""
    if cancelar is not None and cancelar.is_set():
        raise CancelledError()
    carpeta = carpeta.resolve()
    cache = cache if cache is not None else carpeta / NOMBRE_CACHE
    anteriores = {}
    cache_valida = False
    try:
        guardados = json.loads(cache.read_text(encoding="utf-8"))
        if guardados.get("version") == VERSION_CACHE and isinstance(guardados.get("archivos"), dict):
            anteriores = guardados["archivos"]
            cache_valida = True
    except (OSError, UnicodeError, ValueError, AttributeError):
        pass
    especificaciones, errores, actuales = [], [], {}
    for ruta in detectar_json(carpeta):
        if cancelar is not None and cancelar.is_set():
            raise CancelledError()
        relativo = ruta.relative_to(carpeta).as_posix()
        try:
            info = ruta.stat()
            firma = [info.st_mtime_ns, info.st_size, es_wip(ruta)]
        except OSError as exc:
            errores.append(f"{ruta.name}: {exc}")
            continue
        anterior = anteriores.get(relativo)
        if isinstance(anterior, dict) and anterior.get("firma") == firma:
            try:
                especificaciones.append(_restaurar(anterior.get("datos"), ruta, carpeta))
                actuales[relativo] = anterior
                continue
            except (ValueError, TypeError, OSError):
                pass
        catalogo = cargar_catalogo(carpeta, archivos=[ruta])
        errores.extend(catalogo.errores)
        if catalogo.especificaciones:
            especificacion = catalogo.especificaciones[0]
            especificaciones.append(especificacion)
            try:
                info = ruta.stat()
                if firma == [info.st_mtime_ns, info.st_size, es_wip(ruta)]:
                    actuales[relativo] = {"firma": firma, "datos": _serializar(especificacion, carpeta)}
            except OSError:
                pass
    if cancelar is not None and cancelar.is_set():
        raise CancelledError()
    if not cache_valida or actuales != anteriores:
        try:
            contenido = json.dumps({"version": VERSION_CACHE, "archivos": actuales}, ensure_ascii=False).encode("utf-8")
            escribir_atomico(cache, contenido)
        except OSError:
            pass  # La caché es opcional, incluso en carpetas de solo lectura.
    return Catalogo(tuple(especificaciones), tuple(errores))


def candidatas_catalogo(especificaciones, origen: Especificacion, idioma: str) -> list[Path]:
    """Respeta enlaces explícitos (también de distinta revisión); conserva las ambigüedades."""
    codigo, revision, cliente = map(clave_texto, (origen.codigo, origen.revision, origen.cliente))
    compatibles = [e for e in especificaciones if e.idioma == idioma and
                   tuple(map(clave_texto, (e.codigo, e.cliente))) == (codigo, cliente)]
    explicitas = [e.ruta for e in compatibles if e.ruta in origen.versiones_idioma or origen.ruta in e.versiones_idioma]
    if explicitas:
        return sorted(set(explicitas))
    if not codigo or not revision:
        return []
    return sorted({e.ruta for e in compatibles if clave_texto(e.revision) == revision})


def disponibilidad_catalogo(especificaciones) -> dict[Path, set[str]]:
    """Disponibilidad de todas las versiones en una pasada, sin leer archivos ni SQLite."""
    por_ruta = {e.ruta: e for e in especificaciones}
    claves = {e.ruta: tuple(map(clave_texto, (e.codigo, e.revision, e.cliente))) for e in especificaciones}
    grupos = {(claves[e.ruta], e.idioma) for e in especificaciones if all(claves[e.ruta][:2])}
    disponibles = {e.ruta: {e.idioma} | {i for i in ("ESPANOL", "INGLES") if
                    all(claves[e.ruta][:2]) and (claves[e.ruta], i) in grupos} for e in especificaciones}
    for origen in especificaciones:
        for ruta in origen.versiones_idioma:
            destino = por_ruta.get(ruta)
            if destino and (claves[origen.ruta][0], claves[origen.ruta][2]) == (claves[ruta][0], claves[ruta][2]):
                disponibles[origen.ruta].add(destino.idioma)
                disponibles[ruta].add(origen.idioma)
    return disponibles

# Zoom local de la seccion HII

def fuente_escalada(base, factor):
    familia, tamano, *estilo = base
    magnitud = max(1, round(abs(tamano) * factor))
    return (familia, -magnitud if tamano < 0 else magnitud, *estilo)


def _fuente(widget, valor):
    fuente = tkfont.Font(root=widget, font=valor or "TkDefaultFont")
    return tuple(fuente.actual(c) for c in ("family", "size", "weight", "slant"))


def aplicar_zoom_canvas(canvas, factor):
    anterior = getattr(canvas, "_hii_factor", 1.0)
    if factor == anterior:
        return
    if not hasattr(canvas, "_hii_textos"):
        canvas._hii_textos = {
            item: (_fuente(canvas, canvas.itemcget(item, "font")), float(canvas.itemcget(item, "width") or 0))
            for item in canvas.find_all() if canvas.type(item) == "text"
        }
    vista_x, vista_y = canvas.xview()[0], canvas.yview()[0]
    canvas.scale("all", 0, 0, factor / anterior, factor / anterior)
    for item, (fuente, ancho) in canvas._hii_textos.items():
        canvas.itemconfigure(item, font=fuente_escalada(fuente, factor), width=ancho * factor)
    proporcion = Fraction(str(factor)).limit_denominator(10)
    canvas._hii_imagenes_zoom = []
    for item, original in getattr(canvas, "_hii_imagenes_base", []):
        imagen = (original if factor == 1 else
                  original.zoom(proporcion.numerator).subsample(proporcion.denominator))
        canvas._hii_imagenes_zoom.append(imagen)
        canvas.itemconfigure(item, image=imagen)
    canvas.configure(scrollregion=tuple(valor * factor for valor in canvas._hii_region))
    canvas.xview_moveto(vista_x)
    canvas.yview_moveto(vista_y)
    canvas._hii_factor = factor


class ZoomHII:
    def __init__(self, contenido):
        self.contenido = contenido

    def aplicar(self, factor):
        self._aplicar(self.contenido, factor)

    def _aplicar(self, widget, factor):
        if isinstance(widget, ui.Boton):
            if not hasattr(widget, "_hii_fuente"):
                valor = widget.cget("font")
                widget._hii_fuente = (tuple(valor) if isinstance(valor, (tuple, list)) else
                                      (valor.cget("family"), valor.cget("size"), valor.cget("weight")))
                widget._hii_alto = widget.cget("height")
                widget._hii_ancho = widget.cget("width")
            ancho = (max(110, len(str(widget.cget("text"))) * 7 + 24)
                     if widget._ancho_automatico else widget._hii_ancho)
            widget.configure(font=fuente_escalada(widget._hii_fuente, factor),
                             width=round(ancho * factor), height=round(widget._hii_alto * factor))
            return  # CTk contiene widgets internos; el botón se escala como una unidad.
        if isinstance(widget, ttk_nativo.Treeview):
            estilo = ttk_nativo.Style(widget)
            if not hasattr(widget, "_hii_tabla"):
                original = widget.cget("style") or "Treeview"
                widget._hii_tabla = (
                    _fuente(widget, estilo.lookup(original, "font")),
                    _fuente(widget, estilo.lookup(original + ".Heading", "font")),
                    int(estilo.lookup(original, "rowheight") or 22),
                    {c: (widget.column(c, "width"), widget.column(c, "minwidth")) for c in widget["columns"]},
                )
            fuente, encabezado, alto, columnas = widget._hii_tabla
            nombre = f"HII{id(widget)}.Treeview"
            estilo.configure(nombre, font=fuente_escalada(fuente, factor), rowheight=round(alto * factor))
            estilo.configure(nombre + ".Heading", font=fuente_escalada(encabezado, factor))
            widget.configure(style=nombre)
            for columna, (ancho, minimo) in columnas.items():
                widget.column(columna, width=round(ancho * factor), minwidth=round(minimo * factor))
        elif isinstance(widget, (ttk_nativo.Label, ttk_nativo.Entry, tk.Text)):
            if not hasattr(widget, "_hii_fuente"):
                widget._hii_fuente = _fuente(widget, widget.cget("font"))
            widget.configure(font=fuente_escalada(widget._hii_fuente, factor))
        elif isinstance(widget, ttk_nativo.Checkbutton):
            estilo = ttk_nativo.Style(widget)
            if not hasattr(widget, "_hii_fuente"):
                widget._hii_fuente = _fuente(widget, estilo.lookup(widget.cget("style") or "TCheckbutton", "font"))
            nombre = f"HII{id(widget)}.TCheckbutton"
            estilo.configure(nombre, font=fuente_escalada(widget._hii_fuente, factor),
                             indicatorsize=max(8, round(14 * factor)))
            widget.configure(style=nombre)
        elif isinstance(widget, ttk_nativo.Scrollbar):
            estilo = ttk_nativo.Style(widget)
            nombre = f"HII{id(widget)}.{str(widget.cget('orient')).capitalize()}.TScrollbar"
            estilo.configure(nombre, width=max(8, round(14 * factor)))
            widget.configure(style=nombre)
        elif isinstance(widget, tk.Scrollbar):
            if not hasattr(widget, "_hii_ancho"):
                widget._hii_ancho = widget.winfo_pixels(widget.cget("width"))
            widget.configure(width=max(8, round(widget._hii_ancho * factor)))
        elif isinstance(widget, tk.Canvas) and hasattr(widget, "_hii_region"):
            aplicar_zoom_canvas(widget, factor)
        for hijo in widget.winfo_children():
            self._aplicar(hijo, factor)


def texto(valor) -> str:
    return "" if valor is None else str(valor).strip()


@dataclass
class DatosHII:
    nombre: str
    codigo: str
    idioma: str
    pruebas: list[dict]
    comentarios: str
    reanalisis: str
    fecha: str


def huella_datos(datos: DatosHII) -> str:
    contenido = asdict(datos)
    contenido.pop("fecha")
    contenido["pruebas"] = [{campo: texto(prueba.get(campo)) for campo in COLUMNAS} for prueba in datos.pruebas]
    return hash_especificacion(contenido)


def comentarios_simples(notas: list[str]) -> str:
    # TODO: asignar aquí comportamientos especiales a TODO lo relacionado con
    # comentarios (referencias, superíndices, reglas por prueba y por idioma).
    # Por ahora únicamente conserva el texto, sin interpretar ni traducirlo.
    return "\n".join(dict.fromkeys(texto(nota) for nota in notas if texto(nota)))


def leer_reanalisis(datos: dict, notas: list[str]) -> str:
    for objeto in (datos, datos.get("control_documental", {}), datos.get("producto", {})):
        if isinstance(objeto, dict):
            for campo, valor in objeto.items():
                if clave_texto(campo) in {"FECHA DE REANALISIS", "FECHA REANALISIS",
                                         "REANALYSIS DATE", "REANALISIS", "REANALYSIS"}:
                    if texto(valor):
                        return texto(valor)
    for nota in notas:
        coincidencia = re.search(
            r"(?:fecha\s+(?:de\s+)?rean[aá]lisis|rean[aá]lisis|reanalysis\s+date)\s*[:：]\s*([^\n]+)",
            texto(nota), re.I)
        if coincidencia:
            return coincidencia.group(1).strip()
    return ""


def preparar_datos(ruta: Path, fecha: date | None = None, *, datos=None, vista=None) -> DatosHII:
    if datos is None:
        datos = json.loads(ruta.read_text(encoding="utf-8-sig"))
    if vista is None:
        vista = preparar_vista_previa_alineada(ruta, datos=datos)
    idioma = leer_idioma_datos(datos)
    if idioma not in {"ESPANOL", "INGLES"}:
        raise ValueError("La especificación debe indicar idioma español o inglés.")
    codigo = texto(datos.get("control_documental", {}).get("codigo_unico_producto"))
    if not codigo:
        raise ValueError("Falta control_documental.codigo_unico_producto.")
    pruebas = []
    originales = {texto(p.get("numero", i)): p for i, p in enumerate(datos.get("pruebas", []), 1)}
    for i, detalle in enumerate(datos["alineacion_metodos_lab"]["detalle"], 1):
        valores = dict(detalle["valores_alineados"])
        original = originales.get(texto(detalle.get("numero_prueba", i)), {})
        periodicidad = next((texto(objeto.get(campo))
                            for objeto in (valores, original)
                            for campo in ("periodicidad", "PERIODIDICAD DE ANALISIS", "periodicidad_analisis")
                            if texto(objeto.get(campo))), "▶")
        valores["periodicidad"] = periodicidad
        pruebas.append(valores)
    pruebas, _ = _separar_notas_de_pruebas(pruebas)
    for i, prueba in enumerate(pruebas, 1):
        prueba["numero"] = str(i)
    if not pruebas:
        raise ValueError("La especificación alineada no contiene pruebas.")
    return DatosHII(vista.nombre_producto, codigo, idioma, pruebas,
                    comentarios_simples(list(vista.notas)),
                    leer_reanalisis(datos, list(vista.notas)),
                    (fecha or date.today()).strftime("%d/%m/%y"))


def seleccionar_plantilla(carpeta: Path, idioma: str) -> Path:
    codigo = "91400019" if idioma == "INGLES" else "91400005"
    candidatas = sorted(p for p in carpeta.iterdir() if p.is_file() and p.name.startswith(codigo)
                        and p.suffix.lower() in EXTENSIONES_HII)
    if len(candidatas) != 1:
        raise ValueError(f"Se requiere una sola plantilla HII {codigo} compatible (.xlsm, .xlsx, .xltm o .xltx) en {carpeta}.")
    return candidatas[0]


def desplazar_referencia(valor: str, extra: int) -> str:
    return REFERENCIA.sub(lambda m: m[1] + str(int(m[2]) + (extra if int(m[2]) >= 21 else 0)), valor)


def escribir_celda(hoja, referencia: str, valor: str) -> None:
    fila_numero = int(re.search(r"\d+", referencia)[0])
    filas = hoja.find("s:sheetData", NS)
    fila = next((f for f in filas if int(f.get("r")) == fila_numero), None)
    if fila is None:
        fila = ET.SubElement(filas, f"{{{NS['s']}}}row", r=str(fila_numero))
        filas[:] = sorted(filas, key=lambda f: int(f.get("r")))
    celda = next((c for c in fila if c.get("r") == referencia), None)
    if celda is None:
        celda = ET.SubElement(fila, f"{{{NS['s']}}}c", r=referencia)
        fila[:] = sorted(fila, key=lambda c: c.get("r").rstrip("0123456789"))
    for hijo in list(celda):
        celda.remove(hijo)
    celda.set("t", "inlineStr")
    cadena = ET.SubElement(celda, f"{{{NS['s']}}}is")
    t = ET.SubElement(cadena, f"{{{NS['s']}}}t")
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = valor


def preparar_seleccion(ruta):
    contenido = json.loads(ruta.read_text(encoding="utf-8-sig"))
    vista = preparar_vista_previa_alineada(ruta, datos=contenido)
    return preparar_datos(ruta, datos=contenido, vista=vista), vista


def _firma_plantilla(ruta):
    estado = ruta.stat()
    return estado.st_size, estado.st_mtime_ns, estado.st_ctime_ns


@lru_cache(maxsize=4)
def _cargar_plantilla(ruta, firma):
    contenido = ruta.read_bytes()
    if _firma_plantilla(ruta) != firma:
        raise ValueError("La plantilla cambió mientras se leía. Vuelve a intentarlo.")
    modelo_libro = load_workbook(io.BytesIO(contenido), read_only=False)
    try:
        modelo_hoja = modelo_libro.worksheets[0]
        anchos = tuple(modelo_hoja.column_dimensions[get_column_letter(c)].width for c in range(1, 7))
        nombre_hoja = modelo_hoja.title
    finally:
        modelo_libro.close()
    with ZipFile(io.BytesIO(contenido)) as archivo:
        partes = {info.filename: archivo.read(info.filename) for info in archivo.infolist()}
        infos = tuple(archivo.infolist())
    return partes, infos, anchos, nombre_hoja


_CERROJO_PLANTILLA = threading.Lock()
_TURNO_EXCEL = threading.Semaphore(1)


def construir_excel(plantilla: Path, datos: DatosHII) -> bytes:
    """Edita copias de las partes necesarias; la plantilla preparada permanece intacta."""
    extra = max(0, len(datos.pruebas) - 6)
    plantilla = Path(plantilla).resolve()
    with _CERROJO_PLANTILLA:
        originales, infos, anchos, nombre_hoja = _cargar_plantilla(plantilla, _firma_plantilla(plantilla))
    partes = dict(originales)
    hoja_nombre = "xl/worksheets/sheet1.xml"
    hoja = ET.fromstring(partes[hoja_nombre])
    estilos = ET.fromstring(partes["xl/styles.xml"])
    formatos_celda = estilos.find("s:cellXfs", NS)
    estilos_envueltos = {}
    if extra:
        filas = hoja.find("s:sheetData", NS)
        modelo = copy.deepcopy(next(f for f in filas if f.get("r") == "20"))
        for elemento in hoja.iter():
            for atributo in ("r", "ref", "sqref"):
                if elemento.get(atributo):
                    valor = elemento.get(atributo)
                    elemento.set(atributo, str(int(valor) + extra) if atributo == "r" and valor.isdigit() and int(valor) >= 21
                                 else desplazar_referencia(valor, extra))
            if ET.QName(elemento).localname in {"f", "formula1", "formula2"} and elemento.text:
                elemento.text = desplazar_referencia(elemento.text, extra)
        for numero in range(21, 21 + extra):
            fila = copy.deepcopy(modelo)
            fila.set("r", str(numero))
            for celda in fila:
                celda.set("r", re.sub(r"\d+", str(numero), celda.get("r")))
            filas.append(fila)
        filas[:] = sorted(filas, key=lambda f: int(f.get("r")))
    for indice in range(max(6, len(datos.pruebas))):
        prueba = datos.pruebas[indice] if indice < len(datos.pruebas) else {}
        for col, campo in enumerate(COLUMNAS, 1):
            escribir_celda(hoja, f"{get_column_letter(col)}{15 + indice}", texto(prueba.get(campo)))
        fila = next(f for f in hoja.find("s:sheetData", NS) if f.get("r") == str(15 + indice))
        for celda in fila:
            if re.match(r"^[A-F]\d+$", celda.get("r", "")):
                estilo = int(celda.get("s", "0"))
                if estilo not in estilos_envueltos:
                    nuevo = copy.deepcopy(formatos_celda[estilo])
                    alineacion = nuevo.find("s:alignment", NS)
                    if alineacion is None:
                        alineacion = ET.SubElement(nuevo, f"{{{NS['s']}}}alignment")
                    alineacion.set("wrapText", "1")
                    nuevo.set("applyAlignment", "1")
                    estilos_envueltos[estilo] = len(formatos_celda)
                    formatos_celda.append(nuevo)
                celda.set("s", str(estilos_envueltos[estilo]))
        lineas = max(max(1, sum(max(1, math.ceil(len(linea) / max(4, ancho - 2)))
                                   for linea in texto(prueba.get(campo)).split("\n")))
                     for ancho, campo in zip(anchos, COLUMNAS))
        fila.set("ht", str(max(float(fila.get("ht", "15")), lineas * 12 + 6)))
        fila.set("customHeight", "1")
    comentario = "D24" if datos.idioma == "INGLES" else "D25"
    # Comentarios: mantener toda lógica especial centralizada en comentarios_simples.
    if datos.comentarios:
        escribir_celda(hoja, desplazar_referencia(comentario, extra), datos.comentarios)
    escribir_celda(hoja, f"D{33 + extra}", datos.reanalisis)
    partes[hoja_nombre] = ET.tostring(hoja, xml_declaration=True, encoding="UTF-8")
    formatos_celda.set("count", str(len(formatos_celda)))
    partes["xl/styles.xml"] = ET.tostring(estilos, xml_declaration=True, encoding="UTF-8")
    for nombre, contenido in list(partes.items()):
        if nombre.startswith("xl/drawings/") and nombre.endswith(".xml") and "/_rels/" not in nombre:
            raiz = ET.fromstring(contenido)
            for ancla in raiz:
                nodos = ancla.findall(".//a:t", NS)
                for nodo in nodos:
                    nodo.text = FECHA.sub(datos.fecha, nodo.text or "")
                origen = ancla.find("xdr:from", NS)
                fila = int(origen.find("xdr:row", NS).text) if origen is not None else -1
                columna = int(origen.find("xdr:col", NS).text) if origen is not None else -1
                reemplazo = datos.nombre if (fila == 6 and columna == 1) else datos.codigo if (fila == 9 and columna == 1) else None
                if reemplazo is not None and nodos:
                    nodos[0].text = reemplazo
                    for nodo in nodos[1:]:
                        nodo.text = ""
                for nodo in ancla.findall(".//xdr:row", NS):
                    if int(nodo.text) >= 20:
                        nodo.text = str(int(nodo.text) + extra)
            partes[nombre] = ET.tostring(raiz, xml_declaration=True, encoding="UTF-8")
        elif extra and (nombre.startswith("xl/tables/") or nombre == "xl/workbook.xml") and nombre.endswith(".xml"):
            raiz = ET.fromstring(contenido)
            for elemento in raiz.iter():
                if elemento.get("ref"):
                    elemento.set("ref", desplazar_referencia(elemento.get("ref"), extra))
                if ET.QName(elemento).localname == "definedName" and elemento.text:
                    elemento.text = desplazar_referencia(elemento.text, extra)
            partes[nombre] = ET.tostring(raiz, xml_declaration=True, encoding="UTF-8")
    # Limita la impresión a la HII, sin incluir el catálogo auxiliar de la plantilla inglesa.
    libro = ET.fromstring(partes["xl/workbook.xml"])
    definidos = libro.find("s:definedNames", NS)
    if definidos is None:
        definidos = ET.Element(f"{{{NS['s']}}}definedNames")
        posicion = list(libro).index(libro.find("s:sheets", NS)) + 1
        libro.insert(posicion, definidos)
    area = next((n for n in definidos if n.get("name") == "_xlnm.Print_Area" and n.get("localSheetId") == "0"), None)
    if area is None:
        area = ET.SubElement(definidos, f"{{{NS['s']}}}definedName", name="_xlnm.Print_Area", localSheetId="0")
    area.text = "'" + nombre_hoja.replace("'", "''") + f"'!$A$1:$F${43 + extra}"
    partes["xl/workbook.xml"] = ET.tostring(libro, xml_declaration=True, encoding="UTF-8")
    # Reemplazar fechas también en cadenas compartidas, encabezados y otros objetos XML.
    for nombre, contenido in list(partes.items()):
        if nombre.endswith((".xml", ".vml")) and b"DD/MM/YY" in contenido:
            partes[nombre] = contenido.replace(b"DD/MM/YY", datos.fecha.encode())
    salida = io.BytesIO()
    with ZipFile(salida, "w") as archivo:
        for info in infos:
            archivo.writestr(copy.copy(info), partes[info.filename])
    return salida.getvalue()


class _PlanoExcel:
    def __init__(self):
        self.ordenes = []

    def create_rectangle(self, *coordenadas, **opciones):
        self.ordenes.append(("rectangle", coordenadas, opciones))

    def create_text(self, *coordenadas, **opciones):
        self.ordenes.append(("text", coordenadas, opciones))


def preparar_dibujo_excel(contenido, datos):
    """Calcula el plano de la hoja sin crear widgets ni imágenes Tk."""
    canvas = _PlanoExcel()
    libro = load_workbook(io.BytesIO(contenido), data_only=True)
    hoja = libro.worksheets[0]
    filas = 43 + max(0, len(datos.pruebas) - 6)
    xs, ys = [20], [20]
    for col in range(1, 7):
        xs.append(xs[-1] + max(45, hoja.column_dimensions[get_column_letter(col)].width * 7))
    for fila in range(1, filas + 1):
        alto = (hoja.row_dimensions[fila].height or 15) * 1.33
        if 15 <= fila < 15 + len(datos.pruebas):
            alto = max(alto, max(len(textwrap.wrap(texto(hoja.cell(fila, col).value),
                                   max(5, int((xs[col] - xs[col-1]) / 6)))) for col in range(1, 7)) * 14 + 6)
        ys.append(ys[-1] + alto)
    canvas.create_rectangle(20, 20, xs[-1], ys[-1], fill="white", outline="white")
    cubiertas = set()
    for combinado in hoja.merged_cells.ranges:
        for fila in range(combinado.min_row, combinado.max_row + 1):
            for col in range(combinado.min_col, combinado.max_col + 1):
                if (fila, col) != (combinado.min_row, combinado.min_col):
                    cubiertas.add((fila, col))
    for fila in range(1, filas + 1):
        for col in range(1, 7):
            if (fila, col) in cubiertas:
                continue
            celda = hoja.cell(fila, col)
            combinado = next((m for m in hoja.merged_cells.ranges if celda.coordinate in m), None)
            ultima_col = min(6, combinado.max_col) if combinado else col
            ultima_fila = min(filas, combinado.max_row) if combinado else fila
            x, y, x2, y2 = xs[col-1], ys[fila-1], xs[ultima_col], ys[ultima_fila]
            relleno = "white"
            if celda.fill.fgColor.type == "rgb" and celda.fill.patternType == "solid":
                relleno = "#" + celda.fill.fgColor.rgb[-6:]
            borde = any(getattr(celda.border, lado).style for lado in ("left", "right", "top", "bottom"))
            canvas.create_rectangle(x, y, x2, y2, fill=relleno, outline="#666" if borde else relleno)
            if celda.value is not None:
                canvas.create_text(x + 4, y + 3, text=texto(celda.value), anchor="nw", width=max(10, x2-x-8),
                                   font=("Arial", 9, "bold" if celda.font.bold else "normal"))
    with ZipFile(io.BytesIO(contenido)) as z:
        for nombre in z.namelist():
            if not (nombre.startswith("xl/drawings/") and nombre.endswith(".xml") and "/_rels/" not in nombre):
                continue
            for ancla in ET.fromstring(z.read(nombre)):
                origen = ancla.find("xdr:from", NS)
                if origen is None:
                    continue
                col, fila = int(origen.find("xdr:col", NS).text), int(origen.find("xdr:row", NS).text)
                textos = ancla.findall(".//a:t", NS)
                if col >= 6 or fila >= filas:
                    continue
                xoff = int(origen.find("xdr:colOff", NS).text) / 9525
                yoff = int(origen.find("xdr:rowOff", NS).text) / 9525
                destino = ancla.find("xdr:to", NS)
                ancho = max(60, xs[-1]-xs[col]-xoff)
                if destino is not None:
                    col_fin = int(destino.find("xdr:col", NS).text)
                    if col_fin <= 6:
                        ancho = max(30, xs[col_fin] + int(destino.find("xdr:colOff", NS).text) / 9525 - xs[col] - xoff)
                if textos:
                    valor = "\n".join(t.text for t in textos if t.text)
                    canvas.create_text(xs[col] + xoff, ys[fila] + yoff, text=valor, anchor="nw",
                                       width=ancho, font=("Arial", 10, "bold"))
                imagen = ancla.find(".//a:blip", NS)
                if imagen is not None:
                    import posixpath
                    relacion = imagen.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
                    ruta_rels = posixpath.join(posixpath.dirname(nombre), "_rels", posixpath.basename(nombre) + ".rels")
                    rels = ET.fromstring(z.read(ruta_rels))
                    destino_imagen = next(r.get("Target") for r in rels if r.get("Id") == relacion)
                    ruta_imagen = (destino_imagen.lstrip("/") if destino_imagen.startswith("/") else
                                   posixpath.normpath(posixpath.join(posixpath.dirname(nombre), destino_imagen)))
                    canvas.ordenes.append(("image", (xs[col] + xoff, ys[fila] + yoff),
                                           {"contenido": z.read(ruta_imagen), "ancho": ancho}))
    canvas._hii_region = (0, 0, xs[-1] + 20, ys[-1] + 20)
    libro.close()
    return canvas._hii_region, tuple(canvas.ordenes)


def _trabajar_excel(trabajo, salida):
    """Solo datos y archivos: nunca llama a Tk ni ejecuta callbacks de la interfaz."""
    try:
        with _TURNO_EXCEL:
            if trabajo["cancelar"].is_set():
                raise CancelledError()
            contenido = trabajo["contenido"]
            if contenido is None:
                contenido = construir_excel(trabajo["plantilla"], trabajo["datos"])
            if trabajo["cancelar"].is_set():
                raise CancelledError()
            dibujo = preparar_dibujo_excel(contenido, trabajo["datos"]) if trabajo["dibujar"] else None
            if _firma_plantilla(trabajo["plantilla"]) != trabajo["firma"]:
                raise ValueError("La plantilla cambió durante la preparación. Vuelve a intentarlo.")
            resultado = contenido, dibujo
        salida.put((trabajo["revision"], True, resultado))
    except Exception as exc:
        salida.put((trabajo["revision"], False, exc))


def _procesar_trabajos_excel(entrada, salida):
    while True:
        trabajo = entrada.get()
        if trabajo is None:
            return
        _trabajar_excel(trabajo, salida)


class GeneradorHII:
    def __init__(self, padre, carpeta_formatos=None, carpeta_json=None, app=None, base_datos=BASE / "especificaciones.db"):
        self.app = app
        rutas = rutas_configuradas(BASE)
        self.base_datos = resolver_ruta(base_datos, BASE)
        self.idioma_objetivo = None
        self.origen_actual = None
        self.editores_traduccion = {}
        self.formatos = resolver_ruta(carpeta_formatos if carpeta_formatos is not None else rutas["formatos"], BASE)
        self.carpeta_json = resolver_ruta(carpeta_json if carpeta_json is not None else rutas["json_alineadas"], BASE)
        self.excel = None
        self.imagenes = []
        self._vista_actual = None
        self._revision_excel = 0
        self._revision_dibujo = 0
        self._cola_excel = queue.Queue()
        self._pendientes_excel = queue.Queue()
        self._hilo_excel = None
        self._trabajos_excel = {}
        self._dibujo_excel = None
        self._firma_excel = None
        self.datos = None
        self.marcadas = set()
        self.compacta = False
        self.actual = None
        self.vistas_visibles = {"tabla": True, "excel": False}
        self._vista_ampliada = None
        self._vistas_anteriores = None
        self._vistas_renderizadas = set()
        self.zoom_factor = 1.0
        self.catalogo = Catalogo((), ())
        self.visibles = []
        self._cerrada = False
        self._tareas = {}
        self._revision_lista = 0
        self._cargando_catalogo = False
        self._recarga_pendiente = False
        self._acciones_tras_catalogo = []
        self._cola_catalogo = queue.Queue()
        self._cancelar_catalogo = threading.Event()
        self.ventana = ui.Ventana(padre)
        self.ventana.bind("<Destroy>", self._al_destruir, add="+")
        self.ventana.title("Generar HII — especificación alineada y Excel")
        self.ventana.geometry("1400x850")
        self.ventana.minsize(1000, 650)
        self.ventana.configure(background="#edf2f7")
        cabecera = tk.Frame(self.ventana, background="#0752b9", padx=18, pady=8)
        cabecera.pack(fill="x")
        tk.Label(cabecera, text="GENERAR HII", background="#0752b9", foreground="white",
                 font=("Segoe UI", 17, "bold")).pack(side="left")
        self.zoom_mas = ttk.Button(cabecera, text="+", width=3, fondo="#0752b9", command=lambda: self.cambiar_zoom(.1))
        self.zoom_mas.pack(side="right", padx=(3, 0))
        self.zoom_porcentaje = tk.Label(cabecera, text="100%", background="#0752b9", foreground="white", font=("Segoe UI", 10))
        self.zoom_porcentaje.pack(side="right", padx=6)
        self.zoom_menos = ttk.Button(cabecera, text="−", width=3, fondo="#0752b9", command=lambda: self.cambiar_zoom(-.1))
        self.zoom_menos.pack(side="right", padx=(3, 0))
        self.cuerpo = ttk.Frame(self.ventana)
        self.cuerpo.pack(fill="both", expand=True)
        self.zoom_contenido = ZoomHII(self.cuerpo)
        barra = ttk.Frame(self.cuerpo, padding=8)
        barra.pack(fill="x")
        self.editar_seleccion = ttk.Button(barra, text="Ampliar", command=self.alternar_lista)
        self.editar_seleccion.pack(side="left")
        self.lote_boton = ttk.Button(barra, text="Guardar seleccionadas", command=self.guardar_seleccionadas, state="disabled")
        self.lote_boton.pack(side="left", padx=8)
        self.guardar_boton = ttk.Button(barra, text="Guardar HII", command=self.guardar, state="disabled")
        self.guardar_boton.pack(side="left", padx=8)
        self.estado = ttk.Label(barra, text="Marca una especificación de la lista.")
        self.estado.pack(side="left")
        idiomas = ttk.Frame(self.cuerpo, padding=(8, 0, 8, 8))
        idiomas.pack(fill="x")
        ttk.Label(idiomas, text="Generar versión:").pack(side="left", padx=(0, 8))
        self.botones_idioma = {}
        for idioma, etiqueta in (("ESPANOL", "Español"), ("INGLES", "Inglés")):
            boton = ttk.Button(idiomas, text=etiqueta, command=lambda i=idioma: self.generar_idioma(i), state="disabled")
            boton.pack(side="left", padx=3)
            self.botones_idioma[idioma] = boton
        self.ambos_idiomas = tk.BooleanVar(master=self.ventana, value=False)
        self.ambos_check = ttk.Checkbutton(idiomas, text="Ambos", variable=self.ambos_idiomas,
                                           command=self._cambiar_ambos, state="disabled")
        self.ambos_check.pack(side="left", padx=(10, 0))
        self.etiqueta_idioma = ttk.Label(idiomas, text="Selecciona una especificación.")
        self.etiqueta_idioma.pack(side="left", padx=12)
        self.contenido = ttk.Panedwindow(self.cuerpo, orient="horizontal")
        self.contenido.pack(fill="both", expand=True, padx=18, pady=(0, 16))
        self.busqueda_panel = ttk.Frame(self.contenido, padding=12)
        self.previa_panel = ttk.Frame(self.contenido)
        self.contenido.add(self.busqueda_panel, weight=2)
        self.contenido.add(self.previa_panel, weight=3)
        ttk.Label(self.busqueda_panel, text="Buscar especificación", foreground="#163b70",
                  font=("Segoe UI", 12, "bold")).pack(anchor="w", pady=(0, 8))
        self.consulta = tk.StringVar()
        ttk.Entry(self.busqueda_panel, textvariable=self.consulta).pack(fill="x", pady=(0, 8))
        self.consulta.trace_add("write", lambda *_: self.programar_busqueda())
        marco_lista = ttk.Frame(self.busqueda_panel)
        marco_lista.pack(fill="both", expand=True)
        self.lista = ttk.Treeview(marco_lista, columns=("marca", "nombre"), show="headings", selectmode="browse")
        self.lista.heading("marca", text="✓")
        self.lista.heading("nombre", text="Especificación alineada")
        self.lista.column("marca", width=32, minwidth=32, stretch=False, anchor="center")
        self.lista.column("nombre", width=430, minwidth=140)
        scroll = ttk.Scrollbar(marco_lista, orient="vertical", command=self.lista.yview)
        scroll_x = ttk.Scrollbar(marco_lista, orient="horizontal", command=self.lista.xview)
        self.lista.configure(yscrollcommand=scroll.set, xscrollcommand=scroll_x.set)
        marco_lista.rowconfigure(0, weight=1)
        marco_lista.columnconfigure(0, weight=1)
        self.lista.grid(row=0, column=0, sticky="nsew")
        scroll.grid(row=0, column=1, sticky="ns")
        scroll_x.grid(row=1, column=0, sticky="ew")
        self.lista.bind("<ButtonRelease-1>", self.clic_lista)
        self.lista.bind("<space>", self.marcar_actual)
        self.lista.bind("<Return>", self.marcar_actual)
        self.lista.bind("<<TreeviewSelect>>", self.seleccionar_fila)
        self.actualizar_boton = ttk.Button(self.busqueda_panel, text="Actualizar lista", command=self.recargar_catalogo)
        self.actualizar_boton.pack(fill="x", pady=(10, 0))
        self.contador = ttk.Label(self.busqueda_panel, wraplength=250)
        self.contador.pack(anchor="w", pady=(8, 0))
        encabezados = ttk.Frame(self.previa_panel)
        encabezados.pack(fill="x")
        self.botones_ampliar_vista = {}
        self.etiquetas_vista = {}
        for clave, titulo in (("tabla", "Vista previa · Tabla"), ("excel", "Vista previa · Excel")):
            encabezado = ttk.Frame(encabezados, padding=(4, 4))
            encabezado.pack(side="left", fill="x", expand=True)
            ttk.Button(encabezado, text="∅", width=3,
                       command=lambda c=clave: self.alternar_vista(c)).pack(side="right", padx=(4, 0))
            boton = ttk.Button(encabezado, text="+", width=3,
                               command=lambda c=clave: self.ampliar_vista(c))
            boton.pack(side="right", padx=(4, 0))
            self.botones_ampliar_vista[clave] = boton
            etiqueta = ttk.Label(encabezado, text=titulo)
            etiqueta.pack(side="left")
            self.etiquetas_vista[clave] = (etiqueta, titulo)
        self.paneles = ttk.Panedwindow(self.previa_panel, orient="horizontal")
        self.paneles.pack(fill="both", expand=True)
        self.izquierda = ttk.Frame(self.paneles)
        self.derecha = ttk.Frame(self.paneles)
        self.paneles.add(self.izquierda, weight=1)
        self.vaciar_previa()
        self.estado.configure(text="Cargando especificaciones…")
        self._programar("inicio", 0, self.recargar_catalogo)
        self._programar("posicionar", 100, self.posicionar_paneles)
        self._programar("zoom", 150, self._aplicar_zoom)
        self._programar("idiomas", 150, self._actualizar_idiomas)

    def cambiar_zoom(self, cambio):
        self.zoom_factor = round(min(1.6, max(.6, self.zoom_factor + cambio)), 1)
        self.zoom_porcentaje.configure(text=f"{round(self.zoom_factor * 100)}%")
        self.zoom_mas.configure(state="disabled" if self.zoom_factor >= 1.6 else "normal")
        self.zoom_menos.configure(state="disabled" if self.zoom_factor <= .6 else "normal")
        self._programar("zoom", 40, self._aplicar_zoom)

    def _aplicar_zoom(self):
        self.zoom_contenido.aplicar(self.zoom_factor)

    def _disponibles(self, origen):
        ruta = Path(origen).resolve()
        if getattr(self, "_catalogo_disponibilidad", None) is not self.catalogo:
            self._catalogo_disponibilidad = self.catalogo
            self._disponibilidad = disponibilidad_catalogo(self.catalogo.especificaciones)
        return self._disponibilidad.get(ruta, set())

    def _actualizar_idiomas(self, origen=None):
        rutas = set(self.marcadas)
        if origen or self.origen_actual:
            rutas.add(Path(origen or self.origen_actual))
        disponibles = set.intersection(*(self._disponibles(r) for r in rutas)) if rutas else set()
        if self.idioma_objetivo not in disponibles:
            self.idioma_objetivo = None
        for idioma, boton in self.botones_idioma.items():
            habilitado = idioma in disponibles
            boton.configure(
                state="normal" if habilitado else "disabled",
                fg_color=ui.paleta()["acento"] if habilitado else ui.paleta()["borde"])
        ambos = disponibles == {"ESPANOL", "INGLES"}
        self.ambos_check.configure(state="normal" if ambos else "disabled")
        if not ambos:
            self.ambos_idiomas.set(False)
        etiquetas = [etiqueta for idioma, etiqueta in (("ESPANOL", "Español"), ("INGLES", "Inglés")) if idioma in disponibles]
        self.etiqueta_idioma.configure(text="Disponibles: " + " · ".join(etiquetas) if etiquetas else
                                       ("Cada selección se guardará en su idioma." if rutas else "Selecciona una especificación."))
        self._cambiar_ambos()

    def _cambiar_ambos(self):
        self.guardar_boton.configure(text="Guardar ambas" if self.ambos_idiomas.get() else "Guardar HII")
        self._programar("zoom", 40, self._aplicar_zoom)

    def _cancelar_tarea(self, clave):
        identificador = self._tareas.pop(clave, None)
        if identificador is not None:
            try:
                self.ventana.after_cancel(identificador)
            except tk.TclError:
                pass

    def _programar(self, clave, demora, accion):
        self._cancelar_tarea(clave)
        if self._cerrada:
            return
        def ejecutar():
            self._tareas.pop(clave, None)
            if not self._cerrada:
                accion()
        self._tareas[clave] = self.ventana.after(demora, ejecutar)

    def _al_destruir(self, evento):
        if evento.widget != self.ventana:
            return
        self._cerrada = True
        for trabajo in self._trabajos_excel.values():
            trabajo["cancelar"].set()
        self._trabajos_excel.clear()
        self._pendientes_excel.put(None)
        self._cancelar_catalogo.set()
        self._acciones_tras_catalogo.clear()
        for clave in tuple(self._tareas):
            self._cancelar_tarea(clave)

    def posicionar_paneles(self):
        ancho = self.contenido.winfo_width()
        self.contenido.sashpos(0, 250 if self.compacta else round(ancho * 0.42))
        if all(self.vistas_visibles.values()):
            self._programar("paneles", 0, self._centrar_vistas)

    def _centrar_vistas(self):
        if all(self.vistas_visibles.values()):
            self.paneles.sashpos(0, self.paneles.winfo_width() // 2)

    def alternar_vista(self, clave):
        self._vista_ampliada = self._vistas_anteriores = None
        self.vistas_visibles[clave] = not self.vistas_visibles[clave]
        self._actualizar_vistas()

    def ampliar_vista(self, clave):
        if self._vista_ampliada == clave:
            self.vistas_visibles = self._vistas_anteriores.copy()
            self._vista_ampliada = self._vistas_anteriores = None
        else:
            if self._vista_ampliada is None:
                self._vistas_anteriores = self.vistas_visibles.copy()
            self._vista_ampliada = clave
            self.vistas_visibles = {c: c == clave for c in self.vistas_visibles}
        self._actualizar_vistas()

    def _invalidar_excel(self):
        self._revision_excel += 1
        self._revision_dibujo += 1
        self._cancelar_tarea("dibujo")
        self.excel = self._dibujo_excel = self._firma_excel = None
        for trabajo in self._trabajos_excel.values():
            if not trabajo["callbacks"]:
                trabajo["cancelar"].set()

    def _solicitar_excel(self, al_terminar=None):
        if self._cerrada:
            return
        try:
            if self.excel is not None:
                firma = getattr(self, "_firma_excel", None)
                if firma is not None and firma != _firma_plantilla(self.plantilla_actual):
                    self._invalidar_excel()
                elif al_terminar is not None:
                    al_terminar(self.excel)
                    return
                elif self._dibujo_excel is not None:
                    self.mostrar_excel(self.excel, self.datos)
                    return
            trabajo = self._trabajos_excel.get(self._revision_excel)
            if trabajo is None:
                trabajo = {"revision": self._revision_excel, "datos": copy.deepcopy(self.datos),
                           "plantilla": self.plantilla_actual, "firma": _firma_plantilla(self.plantilla_actual),
                           "contenido": self.excel, "dibujar": self.vistas_visibles["excel"],
                           "cancelar": Event(), "callbacks": []}
                self._trabajos_excel[self._revision_excel] = trabajo
                if al_terminar is not None:
                    trabajo["callbacks"].append(al_terminar)
                self._pendientes_excel.put(trabajo)
                if self._hilo_excel is None:
                    self._hilo_excel = threading.Thread(target=_procesar_trabajos_excel,
                        args=(self._pendientes_excel, self._cola_excel), daemon=True, name="preparar-hii")
                    self._hilo_excel.start()
            elif al_terminar is not None:
                trabajo["callbacks"].append(al_terminar)
            self._programar("excel", 40, self._recibir_excel)
        except Exception as exc:
            messagebox.showerror("No se pudo preparar HII", str(exc), parent=self.ventana)

    def _recibir_excel(self):
        if self._cerrada:
            return
        while True:
            if self._cerrada:
                return
            try:
                revision, correcto, resultado = self._cola_excel.get_nowait()
            except queue.Empty:
                break
            trabajo = self._trabajos_excel.pop(revision, None)
            if trabajo is None:
                continue
            actual = revision == self._revision_excel and not trabajo["cancelar"].is_set()
            if not correcto:
                if not isinstance(resultado, CancelledError) and (actual or trabajo["callbacks"]):
                    messagebox.showerror("No se pudo preparar HII", str(resultado), parent=self.ventana)
                if actual:
                    self._vistas_renderizadas.discard("excel")
                    self.estado.configure(text="No se pudo preparar HII.")
                continue
            contenido, dibujo = resultado
            if actual:
                self.excel, self._dibujo_excel = contenido, dibujo
                self._firma_excel = trabajo["firma"]
                if self.vistas_visibles["excel"]:
                    if dibujo is None:
                        self._solicitar_excel()
                    else:
                        try:
                            self.mostrar_excel(contenido, self.datos)
                        except Exception as exc:
                            messagebox.showerror("No se pudo dibujar Excel", str(exc), parent=self.ventana)
            for accion in trabajo["callbacks"]:
                if self._cerrada:
                    return
                accion(contenido)
        if self._trabajos_excel:
            self._programar("excel", 40, self._recibir_excel)

    def _actualizar_vistas(self):
        self._cancelar_tarea("paneles")
        for panel in self.paneles.panes():
            self.paneles.forget(panel)
        for clave, panel in (("tabla", self.izquierda), ("excel", self.derecha)):
            visible = self.vistas_visibles[clave]
            self.botones_ampliar_vista[clave].configure(text="−" if self._vista_ampliada == clave else "+")
            etiqueta, titulo = self.etiquetas_vista[clave]
            etiqueta.configure(text=titulo + (" (oculta)" if not visible else ""))
            if visible:
                self.paneles.add(panel, weight=1)
            else:
                for widget in panel.winfo_children():
                    widget.destroy()
                self._vistas_renderizadas.discard(clave)
                if clave == "excel":
                    self._invalidar_excel()
                    self.imagenes.clear()
        if any(self.vistas_visibles.values()):
            self.paneles.pack(fill="both", expand=True)
        else:
            self.paneles.pack_forget()
        try:
            self._renderizar_vistas()
        except Exception as exc:
            messagebox.showerror("No se pudo mostrar la vista previa", str(exc), parent=self.ventana)
        if all(self.vistas_visibles.values()):
            self._programar("paneles", 0, self._centrar_vistas)
        self._programar("zoom", 0, self._aplicar_zoom)

    def _renderizar_vistas(self):
        for clave, panel in (("tabla", self.izquierda), ("excel", self.derecha)):
            if not self.vistas_visibles[clave] or clave in self._vistas_renderizadas:
                continue
            if self.actual is None:
                mensaje = "Marca una especificación para ver su resumen." if clave == "tabla" else "Aquí verás la hoja Excel que se generará."
                ttk.Label(panel, text=mensaje, anchor="center", wraplength=350).pack(fill="both", expand=True, padx=12)
            elif clave == "tabla":
                self.mostrar_especificacion(self.actual)
            else:
                ttk.Label(panel, text="Preparando vista previa de Excel…", anchor="center").pack(fill="both", expand=True)
                self._solicitar_excel()
            self._vistas_renderizadas.add(clave)

    def alternar_lista(self):
        self.compacta = not self.compacta
        self.editar_seleccion.configure(text="Ampliar" if self.compacta else "Compactar")
        self.actualizar_lista()
        self.posicionar_paneles()

    def recargar_catalogo(self, al_terminar=None):
        if self._cerrada:
            return
        if al_terminar is not None:
            self._acciones_tras_catalogo.append(al_terminar)
        if self._cargando_catalogo:
            self._recarga_pendiente = True
            return
        self._cargando_catalogo = True
        self.actualizar_boton.configure(state="disabled")
        self.estado.configure(text="Cargando especificaciones…")
        # El trabajador solo recibe datos y una cola; no llama a Tk ni a SQLite.
        carpeta, cancelar, cola = self.carpeta_json, self._cancelar_catalogo, self._cola_catalogo
        def trabajar():
            try:
                catalogo = cargar_catalogo_hii(carpeta, cancelar=cancelar)
                cola.put((catalogo, None))
            except Exception as exc:
                cola.put((None, exc))
        threading.Thread(target=trabajar, daemon=True, name="catalogo-hii").start()
        self._programar("catalogo", 50, self._consultar_catalogo)

    def _consultar_catalogo(self):
        if self._cerrada:
            return
        try:
            catalogo, error = self._cola_catalogo.get_nowait()
        except queue.Empty:
            self._programar("catalogo", 50, self._consultar_catalogo)
            return
        self._cargando_catalogo = False
        if self._recarga_pendiente:
            self._recarga_pendiente = False
            self.recargar_catalogo()
            return
        self.actualizar_boton.configure(state="normal")
        acciones, self._acciones_tras_catalogo = self._acciones_tras_catalogo, []
        if error is not None:
            self.estado.configure(text="No se pudo actualizar la lista de especificaciones.")
            messagebox.showerror("No se pudo cargar la lista", str(error), parent=self.ventana)
            return
        self.catalogo = catalogo
        self.marcadas.intersection_update(e.ruta for e in self.catalogo.especificaciones)
        self._actualizar_idiomas()
        self.actualizar_lista()
        if acciones:
            for accion in acciones:
                accion()
        elif self.origen_actual and self.origen_actual not in self.marcadas:
            self.vaciar_previa()
        elif self.actual:
            self.cargar_previa(self.origen_actual or self.actual)
        else:
            self.estado.configure(text="Marca una especificación de la lista.")

    def programar_busqueda(self):
        self._cancelar_tarea("lista")
        self._programar("busqueda", 250, self.actualizar_lista)

    def actualizar_lista(self):
        if self._cerrada:
            return
        self._cancelar_tarea("busqueda")
        self._cancelar_tarea("lista")
        self._revision_lista += 1
        self.lista.delete(*self.lista.get_children())
        candidatas = self.catalogo.especificaciones
        if self.compacta and self.marcadas:
            candidatas = [e for e in candidatas if e.ruta in self.marcadas]
        self.visibles = buscar_especificaciones(candidatas, self.consulta.get())
        self._insertar_bloque(self._revision_lista, 0)
        ingles = sum(e.idioma == "INGLES" for e in self.catalogo.especificaciones)
        errores = f"\n{len(self.catalogo.errores)} archivos inválidos" if self.catalogo.errores else ""
        self.contador.configure(text=f"{len(self.catalogo.especificaciones)} especificaciones ({ingles} en inglés)\n"
                               f"{len(self.marcadas)} seleccionadas · {len(self.visibles)} visibles{errores}")
        self.lote_boton.configure(state="normal" if self.marcadas else "disabled")

    def _insertar_bloque(self, revision, inicio):
        if self._cerrada or revision != self._revision_lista:
            return
        fin = min(inicio + FILAS_POR_BLOQUE, len(self.visibles))
        for i in range(inicio, fin):
            especificacion = self.visibles[i]
            self.lista.insert("", "end", iid=f"{revision}:{i}", values=("☑" if especificacion.ruta in self.marcadas else "☐",
                especificacion.etiqueta + (" · English" if especificacion.idioma == "INGLES" else " · Español")))
        if fin < len(self.visibles):
            self._programar("lista", 1, lambda: self._insertar_bloque(revision, fin))

    def _indice_fila(self, fila):
        try:
            revision, indice = map(int, fila.split(":"))
            if revision == self._revision_lista and 0 <= indice < len(self.visibles):
                return indice
        except (ValueError, AttributeError):
            pass
        return None

    def clic_lista(self, evento):
        fila = self.lista.identify_row(evento.y)
        if fila and self.lista.identify_column(evento.x) == "#1":
            indice = self._indice_fila(fila)
            if indice is not None:
                self.marcar(indice)

    def marcar_actual(self, _evento=None):
        seleccion = self.lista.selection()
        if seleccion:
            indice = self._indice_fila(seleccion[0])
            if indice is not None:
                self.marcar(indice)
        return "break"

    def marcar(self, indice):
        especificacion = self.visibles[indice]
        if especificacion.ruta in self.marcadas:
            self.marcadas.remove(especificacion.ruta)
            if self.origen_actual == especificacion.ruta:
                siguiente = next((e for e in self.catalogo.especificaciones if e.ruta in self.marcadas), None)
                if siguiente:
                    self._actualizar_idiomas(siguiente.ruta)
                    self.cargar_previa(siguiente.ruta)
                else:
                    self.vaciar_previa()
        else:
            self.marcadas.add(especificacion.ruta)
            self._actualizar_idiomas(especificacion.ruta)
            self.cargar_previa(especificacion.ruta)
            self.compacta = True
            self.consulta.set("")
        if not self.marcadas:
            self.compacta = False
        self.editar_seleccion.configure(text="Ampliar" if self.compacta else "Compactar")
        self._actualizar_idiomas()
        self.actualizar_lista()
        self.posicionar_paneles()

    def seleccionar_fila(self, _evento=None):
        seleccion = self.lista.selection()
        if seleccion:
            indice = self._indice_fila(seleccion[0])
            if indice is None:
                return
            especificacion = self.visibles[indice]
            if especificacion.ruta in self.marcadas and especificacion.ruta != self.origen_actual:
                self.cargar_previa(especificacion.ruta)

    def vaciar_previa(self):
        self._invalidar_excel()
        self.excel = self.datos = self.actual = self.origen_actual = None
        self._vista_actual = None
        self.imagenes.clear()
        self.plantilla_actual = None
        self.guardar_boton.configure(state="disabled")
        self._vistas_renderizadas.clear()
        for panel in (self.izquierda, self.derecha):
            for widget in panel.winfo_children():
                widget.destroy()
        self._actualizar_vistas()
        self.estado.configure(text="Marca una especificación de la lista.")
        self._actualizar_idiomas()

    def abrir(self):
        ruta = filedialog.askopenfilename(parent=self.ventana, initialdir=self.carpeta_json,
                                          filetypes=[("JSON alineado", "*.json *.JSON")])
        if not ruta:
            return
        self.cargar_previa(Path(ruta))

    def cargar_previa(self, ruta):
        try:
            origen = Path(ruta)
            ruta = self.resolver_idioma(origen)
            if ruta is None:
                return
            datos, vista = preparar_seleccion(Path(ruta))
            plantilla = seleccionar_plantilla(self.formatos, datos.idioma)
            self._invalidar_excel()
            self._vista_actual = vista
            self.datos, self.excel = datos, None
            self.imagenes.clear()
            self.plantilla_actual = plantilla
            self.actual = Path(ruta)
            self.origen_actual = origen
            self._vistas_renderizadas.clear()
            for panel in (self.izquierda, self.derecha):
                for widget in panel.winfo_children():
                    widget.destroy()
            self._renderizar_vistas()
            self._actualizar_idiomas()
            self.guardar_boton.configure(state="normal")
            aviso = " | Sin fecha de reanálisis en el JSON" if not datos.reanalisis else ""
            self.estado.configure(text=f"{datos.codigo} | {datos.idioma}{aviso}")
            self._aplicar_zoom()
        except Exception as exc:
            self.vaciar_previa()
            messagebox.showerror("No se pudo generar HII", str(exc), parent=self.ventana)

    def generar_idioma(self, idioma):
        origen = self.origen_actual or self.actual
        if origen is None:
            origen = next((e.ruta for e in self.catalogo.especificaciones if e.ruta in self.marcadas), None)
        if origen:
            if idioma not in self._disponibles(origen):
                return
            self.idioma_objetivo = idioma
            self.cargar_previa(origen)
        else:
            self.estado.configure(text="Marca una especificación para generar el idioma elegido.")

    def resolver_idioma(self, origen):
        origen = Path(origen).resolve()
        if self.idioma_objetivo is None:
            return origen
        fuente = self._leer_entrada(origen)
        if fuente.idioma == self.idioma_objetivo:
            return origen
        opciones = candidatas_catalogo(self.catalogo.especificaciones, fuente, self.idioma_objetivo)
        if len(opciones) > 1:
            dialogo = ui.Ventana(self.ventana)
            dialogo.title("Elegir especificación traducida")
            dialogo.transient(self.ventana)
            ttk.Label(dialogo, text="Hay varias versiones con el mismo código, revisión y cliente. Elige la correspondiente.",
                      wraplength=650, padding=12).pack()
            lista = tk.Listbox(dialogo, width=90, height=min(12, len(opciones)), exportselection=False)
            for opcion in opciones:
                lista.insert("end", str(opcion.relative_to(self.carpeta_json.resolve())))
            lista.pack(fill="both", expand=True, padx=12, pady=8)
            elegida = []
            def aceptar():
                if lista.curselection():
                    elegida.append(opciones[lista.curselection()[0]])
                    dialogo.destroy()
            ttk.Button(dialogo, text="Usar esta especificación", command=aceptar).pack(pady=10)
            dialogo.grab_set()
            self.ventana.wait_window(dialogo)
            if not elegida:
                self.vaciar_previa()
                self.origen_actual = origen
                return None
            opciones = elegida
        if opciones:
            destino = opciones[0]
            elegida = self._leer_entrada(destino)
            if destino not in candidatas_catalogo([elegida], fuente, self.idioma_objetivo):
                raise ValueError("La especificación vinculada cambió. Actualiza la lista y vuelve a seleccionarla.")
            try:
                with conectar(self.base_datos) as con:
                    vincular(con, origen, destino, permitir_revision_distinta=True)
            except (OSError, ValueError, sqlite3.Error, RuntimeError) as exc:
                messagebox.showwarning("Versión disponible", f"Se usará la especificación traducida, pero no se pudo guardar el vínculo:\n{exc}", parent=self.ventana)
            self.catalogo = Catalogo(tuple(
                replace(e, versiones_idioma=tuple(dict.fromkeys((*e.versiones_idioma, destino if e.ruta == origen else origen))))
                if e.ruta in {origen, destino} else e for e in self.catalogo.especificaciones), self.catalogo.errores)
            return destino
        raise ValueError("No hay una versión disponible en el idioma solicitado. Actualiza la lista de especificaciones.")

    @staticmethod
    def _leer_entrada(ruta):
        catalogo = cargar_catalogo(ruta.parent, archivos=[ruta])
        if catalogo.errores:
            raise ValueError(catalogo.errores[0])
        return catalogo.especificaciones[0]

    def abrir_traduccion(self, origen, idioma):
        clave = (origen.resolve(), idioma)
        existente = self.editores_traduccion.get(clave)
        if existente is not None and existente.ventana.winfo_exists():
            existente.ventana.lift()
            return
        from interfaz_aplicacion_especificaciones import GeneradorEspecificacion
        from types import SimpleNamespace
        app = self.app or SimpleNamespace(raiz=self.ventana,
            rutas={"json_alineadas": self.carpeta_json, "formatos": self.formatos})
        sufijo = "ENG" if idioma == "INGLES" else "ESP"
        destino = self.carpeta_json / f"{origen.stem} {sufijo}.json"
        indice = 1
        while destino.exists():
            destino = self.carpeta_json / f"{origen.stem} {sufijo}_{indice}.json"
            indice += 1
        def guardada(ruta):
            datos = preparar_datos(ruta)
            if datos.idioma != idioma:
                raise ValueError("La traducción no tiene el idioma solicitado.")
            with conectar(self.base_datos) as con:
                vincular(con, origen, ruta, permitir_revision_distinta=True)
            self.recargar_catalogo(al_terminar=lambda: self.cargar_previa(origen))
        editor = GeneradorEspecificacion(app, archivo_inicial=origen, destino_inicial=destino,
            idioma_destino="en" if idioma == "INGLES" else "es", al_guardar=guardada)
        self.editores_traduccion[clave] = editor

    def registrar_hii(self, destino, especificacion, datos):
        with conectar(self.base_datos) as con:
            registrar(con, destino, especificacion, datos.idioma, datos.fecha)

    def guardar_archivo(self, destino, especificacion, datos, contenido):
        huella = huella_datos(datos)
        if huella_datos(preparar_datos(especificacion)) != huella:
            raise ValueError("La especificación cambió después de preparar la HII. Vuelve a cargar su vista previa.")
        guardar_con_historial(destino, especificacion, contenido, huella)
        try:
            self.registrar_hii(destino, especificacion, datos)
        except (sqlite3.Error, OSError, ValueError, RuntimeError) as exc:
            return f"{destino.name}: Excel e historial JSON guardados; registro en SQLite pendiente: {exc}"
        return None

    def mostrar_especificacion(self, ruta):
        from tkinter.scrolledtext import ScrolledText
        for widget in self.izquierda.winfo_children():
            widget.destroy()
        vista = self._vista_actual if ruta == self.actual and self._vista_actual is not None else preparar_vista_previa_alineada(ruta)
        ttk.Label(self.izquierda, text=vista.nombre_producto, foreground="#163b70",
                  font=("Segoe UI", 16, "bold"), wraplength=450).pack(anchor="w", padx=10, pady=5)
        separador = crear_divisor_redimensionable(self.izquierda, tk.VERTICAL)
        separador.pack(fill="both", expand=True, padx=8, pady=8)
        tabla_panel = ttk.Frame(separador)
        detalles_panel = ttk.Frame(separador)
        separador.add(tabla_panel, minsize=140, stretch="always")
        separador.add(detalles_panel, minsize=100, stretch="always")
        campos = encabezados_tabla(vista.idioma)
        tabla = ttk.Treeview(tabla_panel, columns=[c for _, c in campos], show="headings")
        for etiqueta, campo in campos:
            tabla.heading(campo, text=etiqueta)
            tabla.column(campo, width=210 if campo == "prueba" else 120, minwidth=80)
        for prueba in vista.pruebas:
            tabla.insert("", "end", values=[texto(prueba.get(c)) for _, c in campos])
        vertical = ttk.Scrollbar(tabla_panel, orient="vertical", command=tabla.yview)
        horizontal = ttk.Scrollbar(tabla_panel, orient="horizontal", command=tabla.xview)
        tabla.configure(yscrollcommand=vertical.set, xscrollcommand=horizontal.set)
        tabla_panel.rowconfigure(0, weight=1)
        tabla_panel.columnconfigure(0, weight=1)
        tabla.grid(row=0, column=0, sticky="nsew")
        vertical.grid(row=0, column=1, sticky="ns")
        horizontal.grid(row=1, column=0, sticky="ew")
        def fijar_divisor(evento):
            if evento.height > 100 and not getattr(separador, "iniciado", False):
                separador.sash_place(0, 0, round(evento.height * .7))
                separador.iniciado = True
        separador.bind("<Configure>", fijar_divisor)
        visor = ScrolledText(detalles_panel, wrap="word", font=("Segoe UI", 10))
        visor.pack(fill="both", expand=True, padx=8, pady=8)
        lineas = [f"{len(vista.pruebas)} pruebas · {vista.idioma}",
                  *[f"{k}: {v}" for k, v in vista.control], ""]
        lineas.extend(vista.notas)
        lineas.extend(vista.mensajes_astm_no_encontradas)
        visor.insert("1.0", "\n".join(lineas))
        visor.configure(state="disabled")

    def guardar_seleccionadas(self):
        carpeta = filedialog.askdirectory(parent=self.ventana, title="Guardar las HII seleccionadas")
        if not carpeta:
            return
        seleccionadas = [e for e in self.catalogo.especificaciones if e.ruta in self.marcadas]
        resueltas = []
        try:
            for especificacion in seleccionadas:
                rutas = self._rutas_guardado(especificacion.ruta, self.ambos_idiomas.get())
                if not rutas:
                    return
                resueltas.extend(ruta for ruta in rutas if ruta not in resueltas)
        except (OSError, ValueError) as exc:
            messagebox.showerror("No se pudo preparar el lote", str(exc), parent=self.ventana)
            return
        self._guardar_en_carpeta(resueltas, Path(carpeta))

    def _rutas_guardado(self, origen, ambos):
        if not ambos:
            ruta = self.resolver_idioma(origen)
            return [ruta] if ruta is not None else []
        if self._disponibles(origen) != {"ESPANOL", "INGLES"}:
            raise ValueError("Esta especificación no tiene ambas versiones disponibles. Actualiza la lista.")
        anterior = self.idioma_objetivo
        try:
            rutas = []
            for idioma in ("ESPANOL", "INGLES"):
                self.idioma_objetivo = idioma
                ruta = self.resolver_idioma(origen)
                if ruta is None:
                    return []
                rutas.append(ruta)
            return rutas
        finally:
            self.idioma_objetivo = anterior

    def _guardar_en_carpeta(self, rutas, carpeta):
        salidas = []
        try:
            for ruta in rutas:
                datos = preparar_datos(ruta)
                plantilla = seleccionar_plantilla(self.formatos, datos.idioma)
                destino = carpeta / f"{ruta.stem} HII {datos.idioma}{plantilla.suffix}"
                self._validar_destino(destino)
                salidas.append((ruta, destino, datos, plantilla))
            if len({p.resolve() for _, p, _, _ in salidas}) != len(salidas):
                raise ValueError("Dos especificaciones producirían el mismo nombre de archivo. Guárdalas por separado.")
        except (OSError, ValueError) as exc:
            messagebox.showerror("No se pudo preparar la salida", str(exc), parent=self.ventana)
            return
        existentes = [p for _, p, _, _ in salidas if p.exists()]
        if existentes and not messagebox.askyesno("Archivos existentes", f"Se reemplazarán {len(existentes)} archivos HII. ¿Continuar?", parent=self.ventana):
            return
        errores, avisos, guardadas = [], [], 0
        for especificacion, destino, datos, plantilla in salidas:
            try:
                aviso = self.guardar_archivo(destino, especificacion, datos, construir_excel(plantilla, datos))
                guardadas += 1
                if aviso:
                    avisos.append(aviso)
            except Exception as exc:
                errores.append(f"{especificacion.name}: {exc}")
        self.estado.configure(text=f"{guardadas} HII guardadas")
        if errores:
            messagebox.showerror("Algunas HII no se generaron", "\n".join(errores), parent=self.ventana)
        if avisos:
            messagebox.showwarning("HII guardadas con registro pendiente", "\n".join(avisos), parent=self.ventana)

    def _validar_destino(self, destino):
        if destino.resolve() in {p.resolve() for p in self.formatos.glob("*") if p.is_file()}:
            raise ValueError("Guarda la HII con otro nombre para conservar la plantilla.")

    def mostrar_excel(self, contenido, datos):
        """Dibuja el plano preparado en bloques breves dentro del hilo de Tk."""
        if self._dibujo_excel is None:
            raise ValueError("La vista previa todavía se está preparando.")
        self._revision_dibujo += 1
        revision = self._revision_dibujo
        for widget in self.derecha.winfo_children():
            widget.destroy()
        marco = ttk.Frame(self.derecha)
        marco.pack(fill="both", expand=True)
        canvas = tk.Canvas(marco, background="#e5e7eb")
        vertical = ttk.Scrollbar(marco, orient="vertical", command=canvas.yview)
        horizontal = ttk.Scrollbar(marco, orient="horizontal", command=canvas.xview)
        canvas.configure(yscrollcommand=vertical.set, xscrollcommand=horizontal.set)
        marco.rowconfigure(0, weight=1)
        marco.columnconfigure(0, weight=1)
        canvas.grid(row=0, column=0, sticky="nsew")
        vertical.grid(row=0, column=1, sticky="ns")
        horizontal.grid(row=1, column=0, sticky="ew")
        region, ordenes = self._dibujo_excel
        canvas._hii_region = region
        canvas._hii_factor = 1.0
        canvas._hii_textos = {}
        canvas._hii_imagenes_base = []
        canvas._hii_imagenes_zoom = []
        canvas.configure(scrollregion=region)
        self._dibujar_bloque_excel(canvas, ordenes, 0, revision)

    def _dibujar_bloque_excel(self, canvas, ordenes, inicio, revision):
        if self._cerrada or revision != self._revision_dibujo or not self.vistas_visibles["excel"]:
            return
        if not canvas.winfo_exists():
            return
        indice = inicio
        while indice < len(ordenes) and indice < inicio + 40:
            tipo, coordenadas, opciones = ordenes[indice]
            factor = canvas._hii_factor
            escaladas = tuple(c * factor for c in coordenadas)
            if tipo == "text":
                valores = dict(opciones, font=fuente_escalada(opciones["font"], factor), width=opciones["width"] * factor)
                item = canvas.create_text(*escaladas, **valores)
                canvas._hii_textos[item] = (opciones["font"], opciones["width"])
            elif tipo == "image":
                original = tk.PhotoImage(master=canvas, data=opciones["contenido"])
                original = original.subsample(max(1, math.ceil(original.width() / opciones["ancho"])))
                proporcion = Fraction(str(factor)).limit_denominator(10)
                imagen = original if factor == 1 else original.zoom(proporcion.numerator).subsample(proporcion.denominator)
                item = canvas.create_image(*escaladas, image=imagen, anchor="nw")
                self.imagenes.append(original)
                canvas._hii_imagenes_base.append((item, original))
                canvas._hii_imagenes_zoom.append(imagen)
            else:
                canvas.create_rectangle(*escaladas, **opciones)
            indice += 1
        aplicar_zoom_canvas(canvas, self.zoom_factor)
        if indice < len(ordenes):
            self._programar("dibujo", 1, lambda: self._dibujar_bloque_excel(canvas, ordenes, indice, revision))

    def guardar(self):
        if self.datos is None:
            return
        if self.ambos_idiomas.get():
            carpeta = filedialog.askdirectory(parent=self.ventana, title="Guardar ambas HII")
            if not carpeta:
                return
            try:
                rutas = self._rutas_guardado(self.origen_actual or self.actual, True)
            except (OSError, ValueError) as exc:
                messagebox.showerror("No se pudieron preparar ambas HII", str(exc), parent=self.ventana)
                return
            if rutas:
                self._guardar_en_carpeta(rutas, Path(carpeta))
            return
        try:
            plantilla = getattr(self, "plantilla_actual", None) or seleccionar_plantilla(self.formatos, self.datos.idioma)
        except (OSError, ValueError) as exc:
            messagebox.showerror("No se pudo guardar", str(exc), parent=self.ventana)
            return
        extension = plantilla.suffix
        ruta = filedialog.asksaveasfilename(parent=self.ventana, defaultextension=extension,
            initialfile=f"{self.datos.codigo} HII {self.datos.idioma}{extension}", filetypes=[("Formato HII vigente", "*" + extension)])
        if ruta:
            try:
                destino = Path(ruta)
                if destino.suffix.lower() != extension.lower():
                    raise ValueError(f"La plantilla vigente requiere la extensión {extension}.")
                self._validar_destino(destino)
                origen, datos = self.actual, copy.deepcopy(self.datos)
                self.plantilla_actual = plantilla
                def terminar(contenido):
                    try:
                        aviso = self.guardar_archivo(destino, origen, datos, contenido)
                        self.estado.configure(text=f"Guardado: {destino.name}")
                        if aviso:
                            messagebox.showwarning("HII guardada con registro pendiente", aviso, parent=self.ventana)
                    except Exception as exc:
                        messagebox.showerror("No se pudo guardar", str(exc), parent=self.ventana)
                self.estado.configure(text=f"Preparando: {destino.name}")
                self._solicitar_excel(terminar)
            except Exception as exc:
                messagebox.showerror("No se pudo guardar", str(exc), parent=self.ventana)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--formatos", type=lambda valor: resolver_ruta(valor, BASE))
    parser.add_argument("--json", type=lambda valor: resolver_ruta(valor, BASE))
    args = parser.parse_args()
    raiz = ui.Raiz()
    raiz.withdraw()
    generador = GeneradorHII(raiz, args.formatos, args.json)
    generador.ventana.protocol("WM_DELETE_WINDOW", raiz.destroy)
    raiz.mainloop()


if __name__ == "__main__":
    main()
