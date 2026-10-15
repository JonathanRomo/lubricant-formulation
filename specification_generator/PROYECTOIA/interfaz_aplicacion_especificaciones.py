#!/usr/bin/env python3
"""Interfaz principal para cargar, editar, seleccionar y comparar especificaciones."""
from __future__ import annotations
import estilo_interfaz as ui
from edicion_bilingue import BorradoresBilingues, guardar_pareja_json
from especificaciones_wip import es_wip, ruta_estado, serializar_estado, restaurar_estado
from traductor_especificaciones import (
    SERVICIO as SERVICIO_TRADUCTOR, actualizar_traductor,
    validar_registros as validar_registros_traducciones,
)
from rutas_proyecto import BASE, CARPETAS, resolver_ruta, ruta_portatil, rutas_configuradas, ruta_maestro, carpeta_listas
from generador_hii import conservar_historial
from cambios_alineacion import (cambios_documento, diferencias_filas, original_disponible,
                               comparar_valores_pruebas, ResaltadorCambios)


import json
import argparse
import copy
import contextlib
import difflib
import os
import queue
import re
import shutil
import sqlite3
import subprocess
import sys
import tempfile
import threading
import tkinter as tk
from datetime import date
from pathlib import Path
from tkinter import filedialog, messagebox
from estilo_interfaz import ttk

from tkinter.scrolledtext import ScrolledText
from typing import Any, Callable
from zipfile import BadZipFile

from docx.opc.exceptions import PackageNotFoundError
from alinear_especificaciones_json import (
    COLUMNAS_MAESTRO,
    MAESTRO,
    VERSION as VERSION_ALINEADOR,
    alternativas_unidades,
    cargar_maestro,
    metodo_maestro_canonico,
    metodo_maestro_para_idioma,
    metodo_referencias_canonico,
    metodos_maestros,
    mejor_coincidencia,
    normalizar_espacio_unidades,
    normalizar_especificacion,
    normalizar_formato_numerico,
    normalizar_terminos_especificacion,
    normalizar_valor_medicion,
    nombres_canonicos,
    normalizar,
    referencias_astm,
    unidad_original_esta_entre_alternativas,
)
from interfaz_seleccion_especificaciones import (
    OrdenadorColumnasTreeview,
    crear_divisor_redimensionable,
    encabezados_tabla,
    preparar_vista_previa,
    preparar_vista_previa_alineada,
)


VERSION = "1.1.4"
RUTAS_PREDETERMINADAS = {clave: BASE / carpeta for clave, carpeta in CARPETAS.items()}
ARCHIVO_CONFIGURACION = BASE / "configuracion_interfaz.json"
CAMPOS_TABLA = ("prueba", "unidades", "metodo", "especificacion", "valor_tipico")
CAMPO_PERIODICIDAD = "PERIODIDICAD DE ANALISIS"
SIMBOLOS_PERIODICIDAD = ("◊", "▶", "C", "I", "P", "♦", "Q", "●")
MAX_SIMBOLOS_PERIODICIDAD = 4
PERIODICIDAD_PREDETERMINADA = "▶"
DIGITOS_SUPERINDICE = "⁰¹²³⁴⁵⁶⁷⁸⁹"
DIGITOS_SUPERINDICE_A_ENTERO = str.maketrans(DIGITOS_SUPERINDICE, "0123456789")
LEYENDA_PERIODICIDAD = {
    "es": (
        ("◊", "Prueba crítica"),
        ("▶", "Prueba cada lote de producción"),
        ("C", "Prueba cada carrotanque"),
        ("I", "Control interno"),
        ("P", "Liberación con base a informe de calidad del proveedor"),
        ("♦", "Prueba cada 3 meses"),
        ("Q", "Prueba cada cambio de composición química"),
        ("●", "Prueba cada cambio de lote de materia prima"),
    ),
    "en": (
        ("◊", "Critical test"),
        ("▶", "Test every production batch"),
        ("C", "Test every tank truck"),
        ("I", "Internal control"),
        ("P", "Release based on supplier quality report"),
        ("♦", "Test every 3 months"),
        ("Q", "Test every chemical composition change"),
        ("●", "Test every raw material batch change"),
    ),
}
NOTAS_PLANTILLA = {
    "es": "La versión vigente del método aplicable corresponderá a la versión emitida proporcionada por el laboratorio.",
    "en": "The active version of the applicable method shall be the version issued and provided by the laboratory.",
}
VALORES_ELABORO_APROBO = {
    "es": "Investigación y desarrollo",
    "en": "Research and Development",
}
CATALOGO_PART_ID = BASE / "CODIGOS_200.xlsx"
EXTENSIONES_WORD = (".doc", ".docm", ".docx")
CAMPOS_CONTROL_GENERADOR = (
    ("elaboro_y_aprobo", "Elaboró y aprobó", "Prepared and approved by"),
    ("folio_apqp", "Folio APQP", "APQP number"),
    ("folio_vap", "Folio VAP", "VAP number"),
    ("codigo_unico_producto", "Código / Part ID", "Part ID"),
    ("fecha", "Fecha", "Date"),
    ("revision", "Revisión", "Revision"),
    ("nombre_cliente", "Nombre del cliente", "Customer's name"),
    ("revision_cliente", "Revisión del cliente", "Customer's revision number"),
)
ETIQUETA_PERIODICIDAD = {
    "es": "Periodicidad de análisis",
    "en": "Analysis periodicity",
}
ETIQUETAS_TABLA_GENERADOR = {
    "es": (
        "Prueba", "Unidades", "Método", "Especificación", "Valores típicos",
        ETIQUETA_PERIODICIDAD["es"],
    ),
    "en": (
        "Test", "Units", "Test method", "Specification", "Typical values",
        ETIQUETA_PERIODICIDAD["en"],
    ),
}


def _valores_control_predeterminados(fecha_actual: date | None = None) -> dict[str, str]:
    fecha = fecha_actual or date.today()
    valores = {clave: "" for clave, _, _ in CAMPOS_CONTROL_GENERADOR}
    valores.update(
        {
            "elaboro_y_aprobo": VALORES_ELABORO_APROBO["es"],
            "folio_vap": "-",
            "fecha": fecha.strftime("%d/%m/%y"),
            "revision": "00",
            "nombre_cliente": "-",
            "revision_cliente": "-",
        }
    )
    return valores


def _traducir_nombre_prueba(
    maestro: list[dict[str, str]],
    nombre: str,
    idioma_origen: str,
    idioma_destino: str,
    metodo: str = "",
) -> str:
    nombre_clave = normalizar(nombre)
    if not nombre_clave or idioma_origen == idioma_destino:
        return nombre
    idioma_alternativo = "en" if idioma_origen == "es" else "es"
    registros = [
        registro
        for registro in maestro
        if normalizar(_nombre_maestro_para_idioma(registro, idioma_origen))
        == nombre_clave
        or normalizar(_nombre_maestro_para_idioma(registro, idioma_alternativo))
        == nombre_clave
    ]
    referencias = set(referencias_astm(metodo))
    if referencias:
        relacionados = [
            registro
            for registro in registros
            if any(
                set(referencias_astm(metodo)) == referencias
                for metodo in metodos_maestros(registro)
            )
        ]
        if relacionados:
            registros = relacionados
    traducciones = tuple(
        dict.fromkeys(
            destino
            for registro in registros
            if (destino := _nombre_maestro_para_idioma(registro, idioma_destino))
        )
    )
    return traducciones[0] if len(traducciones) == 1 else nombre


def crear_json_temporal_comparacion(
    archivo_word: Path,
    carpeta_temporal: Path,
) -> Path:
    """Extrae Word con el gestor sin escribir en su registro ni en sus carpetas."""
    import gestor_especificaciones as gestor

    lectura, _ = gestor.preparar(archivo_word, carpeta_temporal)
    idioma = "en" if gestor.contiene_marcadores_ingles(lectura) else "es"
    datos = gestor.extraer(lectura, idioma)
    ruta_json = carpeta_temporal / "comparacion_temporal.json"
    ruta_json.write_text(
        json.dumps(datos, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    return ruta_json


def crear_documento_alineado_manual(datos: dict[str, Any]) -> dict[str, Any]:
    """Valida un JSON de especificación y crea los valores que consume la vista alineada."""
    resultado = dict(datos)
    producto = resultado.get("producto")
    if (
        not isinstance(producto, dict)
        or not isinstance(producto.get("nombre"), str)
        or not producto["nombre"].strip()
    ):
        raise ValueError("Se requiere 'producto.nombre' como texto.")
    pruebas = resultado.get("pruebas")
    if not isinstance(pruebas, list) or any(not isinstance(fila, dict) for fila in pruebas):
        raise ValueError("'pruebas' debe ser una lista de objetos.")

    alineacion_existente = resultado.get("alineacion_metodos_lab")
    detalle_existente = (
        alineacion_existente.get("detalle")
        if isinstance(alineacion_existente, dict)
        else None
    )
    if isinstance(alineacion_existente, dict) and isinstance(detalle_existente, list):
        tiene_snapshot = (
            len(detalle_existente) == len(pruebas)
            and all(
                isinstance(detalle, dict)
                and isinstance(detalle.get("valores_alineados"), dict)
                and all(campo in detalle["valores_alineados"] for campo in CAMPOS_TABLA)
                for detalle in detalle_existente
            )
        )
        if tiene_snapshot:
            return resultado
        raise ValueError(
            "La auditoría de este JSON no tiene valores_alineados completos. "
            "Edítalos en alineacion_metodos_lab.detalle antes de guardar para "
            "no reemplazarlos con valores extraídos."
        )

    detalle = []
    for indice, fila in enumerate(pruebas, 1):
        valores = {}
        for campo in CAMPOS_TABLA:
            valor = fila.get(campo)
            if valor is None:
                valor = ""
            elif not isinstance(valor, (str, int, float, bool)):
                raise ValueError(f"pruebas[{indice}].{campo} debe ser un valor simple.")
            valores[campo] = str(valor)
        detalle.append(
            {
                "numero_prueba": indice,
                "estado": "EDICION_MANUAL",
                "valores_alineados": valores,
            }
        )

    resultado["alineacion_metodos_lab"] = {
        "version_programa": "EDICION_MANUAL",
        "detalle": detalle,
        "mensajes_astm_no_encontradas": [],
    }
    return resultado


def _valor_reanalisis_editor(valor: Any) -> str:
    texto = str(valor if valor is not None else "").strip()
    coincidencia = re.fullmatch(r"([0-9]+)\s+meses", texto, flags=re.IGNORECASE)
    return coincidencia.group(1) if coincidencia else texto


def _valor_reanalisis_json(valor: Any) -> str:
    texto = _valor_reanalisis_editor(valor)
    return f"{texto} meses" if re.fullmatch(r"[0-9]+", texto) else texto


def crear_datos_especificacion_nueva(
    nombre_producto: str,
    idioma: str,
    control: dict[str, str],
    pruebas: list[dict[str, str]],
    notas: list[str],
    desarrollador: str = "",
    fecha_reanalisis: str = "",
) -> dict[str, Any]:
    """Construye el JSON alineado que usa la vista previa y la generación Word."""
    nombre = nombre_producto.strip()
    if not nombre:
        raise ValueError("Ingresa el nombre del producto.")
    if idioma not in {"es", "en"}:
        raise ValueError("Selecciona Español o English.")
    pruebas_activas = [
        fila
        for fila in pruebas
        if any(str(fila.get(campo, "")).strip() for campo in CAMPOS_TABLA)
        or str(fila.get("periodicidad", fila.get(CAMPO_PERIODICIDAD, ""))).strip()
    ]
    if not pruebas_activas:
        raise ValueError("Agrega al menos una prueba a la tabla.")

    pruebas_json = []
    for indice, fila in enumerate(pruebas_activas, 1):
        nombre_prueba = fila.get("prueba", "").strip()
        if not nombre_prueba:
            raise ValueError(f"Completa el nombre de la prueba en la fila {indice}.")
        periodicidad = _normalizar_periodicidad(
            str(fila.get("periodicidad", fila.get(CAMPO_PERIODICIDAD, ""))),
            indice,
        )
        pruebas_json.append(
            {
                "numero": indice,
                "prueba": nombre_prueba,
                "unidades": fila.get("unidades", "").strip(),
                "metodo": fila.get("metodo", "").strip(),
                "especificacion": fila.get("especificacion", "").strip(),
                "valor_tipico": fila.get("valor_tipico", "").strip(),
                CAMPO_PERIODICIDAD: periodicidad,
            }
        )

    notas_especificacion = [
        nota.strip() for nota in notas if nota.strip()
    ] or [NOTAS_PLANTILLA[idioma]]
    control_documental = {
        clave: valor.strip() for clave, valor in control.items()
    }
    control_documental["elaboro_y_aprobo"] = (
        control_documental.get("elaboro_y_aprobo")
        or VALORES_ELABORO_APROBO[idioma]
    )
    datos = {
        "esquema": "especificacion_producto_es_v3",
        "producto": {"nombre": nombre},
        "documento": {
            "tipo": "ESPECIFICACION",
            "idioma": idioma,
            "fecha_documento_normalizada": "",
        },
        "control_documental": control_documental,
        "FECHA DE REANALISIS": _valor_reanalisis_json(fecha_reanalisis),
        "Desarrollador": desarrollador.strip(),
        "pruebas": pruebas_json,
        "notas": notas_especificacion,
    }
    return crear_documento_alineado_manual(datos)


def alinear_filas_generador(
    pruebas: list[dict[str, str]],
    maestro: list[dict[str, str]],
    idioma: str,
    al_progreso: Callable[[int, int], None] | None = None,
) -> tuple[list[dict[str, str]], dict[str, int]]:
    """Aplica las reglas del alineador sin tocar notas ni comentarios de celda."""
    if idioma not in {"es", "en"}:
        raise ValueError("Selecciona Español o English.")
    if not maestro:
        raise ValueError("La lista maestra no contiene registros.")

    filas = copy.deepcopy(pruebas)
    filas_con_datos = [
        fila for fila in filas
        if any(str(fila.get(campo, "") or "").strip() for campo in CAMPOS_TABLA)
    ]
    total = len(filas_con_datos)
    conteo = {"alineados": 0, "revision": 0, "sin_coincidencia": 0}
    idioma_maestro = "INGLES" if idioma == "en" else "ESPANOL"
    referencias_maestras = {
        referencia
        for registro in maestro
        for metodo in metodos_maestros(registro)
        for referencia in referencias_astm(metodo)
    }
    procesadas = 0

    for fila in filas:
        if not any(
            str(fila.get(campo, "") or "").strip() for campo in CAMPOS_TABLA
        ):
            continue

        originales = {
            campo: str(fila.get(campo, "") or "").strip()
            for campo in ("prueba", "metodo", "unidades", "especificacion", "valor_tipico")
        }
        nombres_sin_comentario: dict[str, str] = {}
        for campo in ("prueba", "metodo"):
            coincidencia = re.search(
                rf"\s+([{DIGITOS_SUPERINDICE}]+)$",
                originales[campo],
            )
            nombres_sin_comentario[campo] = (
                originales[campo][:coincidencia.start()].rstrip()
                if coincidencia
                else originales[campo]
            )

        especificacion, ambiguos_especificacion = normalizar_valor_medicion(
            normalizar_terminos_especificacion(
                normalizar_especificacion(originales["especificacion"])
            )
        )
        valor_tipico, ambiguos_valor_tipico = normalizar_valor_medicion(
            normalizar_terminos_especificacion(originales["valor_tipico"])
        )
        match = mejor_coincidencia(
            nombres_sin_comentario["prueba"],
            nombres_sin_comentario["metodo"],
            maestro,
            idioma_maestro,
        )
        registro = match["registro"]
        nombre_maestro, unidad_maestra = nombres_canonicos(registro, idioma_maestro)
        metodo_maestro_fuente = metodo_maestro_para_idioma(registro, idioma_maestro)
        metodo_maestro = metodo_maestro_canonico(metodo_maestro_fuente)
        if len(referencias_astm(nombres_sin_comentario["metodo"])) > len(
            referencias_astm(metodo_maestro_fuente)
        ):
            metodo_maestro = metodo_referencias_canonico(
                nombres_sin_comentario["metodo"]
            )
        if match["estado"] == "ALINEADO" and not nombre_maestro:
            match["estado"] = "REVISION_REQUERIDA"
        if ambiguos_especificacion or ambiguos_valor_tipico:
            match["estado"] = "REVISION_REQUERIDA"

        alternativas = alternativas_unidades(unidad_maestra)
        unidades_alineadas = alternativas[0] if alternativas else unidad_maestra
        if match["estado"] == "ALINEADO":
            for campo, valor_base in (
                ("prueba", nombre_maestro),
                ("metodo", metodo_maestro),
            ):
                marcador = originales[campo][len(nombres_sin_comentario[campo]):].strip()
                fila[campo] = " ".join(parte for parte in (valor_base, marcador) if parte)
            if unidades_alineadas:
                if unidad_original_esta_entre_alternativas(
                    originales["unidades"], unidad_maestra
                ):
                    unidades_alineadas = originales["unidades"]
                fila["unidades"] = normalizar_formato_numerico(unidades_alineadas)
            conteo["alineados"] += 1
        elif match["estado"] == "REVISION_REQUERIDA":
            conteo["revision"] += 1
        else:
            conteo["sin_coincidencia"] += 1

        fila["especificacion"] = especificacion
        fila["valor_tipico"] = valor_tipico
        for campo in ("prueba", "metodo", "unidades"):
            fila[campo] = normalizar_espacio_unidades(fila.get(campo, ""))
        if al_progreso is not None:
            procesadas += 1
            al_progreso(procesadas, total)

    return filas, conteo


def _validar_pareja_editor(origen: dict, pareja: dict) -> None:
    """Comprueba la identidad de dos versiones que se editarán como pareja."""
    if {origen.get("idioma"), pareja.get("idioma")} != {"es", "en"}:
        raise ValueError("Las versiones deben estar en idiomas distintos: ESP y ENG.")
    diferencias = []
    for campo, etiqueta in (("codigo_unico_producto", "código"), ("nombre_cliente", "cliente"), ("revision", "revisión")):
        a, b = (contenido.get("control", {}).get(campo, "") for contenido in (origen, pareja))
        if a != b:
            diferencias.append(f"{etiqueta} ({a or 'vacío'} / {b or 'vacío'})")
    if diferencias:
        raise ValueError("No coinciden " + ", ".join(diferencias) + ".")


def _datos_generador_desde_json(datos: dict[str, Any]) -> dict[str, Any]:
    producto = datos.get("producto", {})
    documento = datos.get("documento", {})
    control = datos.get("control_documental", {})
    pruebas = datos.get("pruebas", [])
    if not isinstance(producto, dict) or not isinstance(documento, dict):
        raise ValueError("El JSON debe contener producto y documento como objetos.")
    if not isinstance(control, dict) or not isinstance(pruebas, list):
        raise ValueError("El JSON debe contener control documental y pruebas válidos.")
    if any(not isinstance(fila, dict) for fila in pruebas):
        raise ValueError("Cada prueba del JSON debe ser un objeto.")

    detalle = datos.get("alineacion_metodos_lab", {})
    detalle = detalle.get("detalle", []) if isinstance(detalle, dict) else []
    if not isinstance(detalle, list):
        detalle = []
    filas_generador: list[dict[str, str]] = []
    for indice, fila in enumerate(pruebas):
        alineados: dict[str, Any] = {}
        if indice < len(detalle) and isinstance(detalle[indice], dict):
            valores = detalle[indice].get("valores_alineados", {})
            if isinstance(valores, dict):
                alineados = valores
        fila_generador = {
            campo: str(alineados.get(campo, fila.get(campo, "")) or "")
            for campo in CAMPOS_TABLA
        }
        periodicidad = next(
            (str(objeto.get(campo)) for objeto in (fila, alineados)
             for campo in (CAMPO_PERIODICIDAD, "periodicidad", "periodicidad_analisis")
             if str(objeto.get(campo) or "").strip()),
            PERIODICIDAD_PREDETERMINADA,
        )
        fila_generador["periodicidad"] = periodicidad
        filas_generador.append(fila_generador)

    notas = datos.get("notas", [])
    if isinstance(notas, str):
        notas = [notas]
    if not isinstance(notas, list):
        raise ValueError("Las notas del JSON deben ser una lista de textos.")
    control_editor = {clave: str(control.get(clave, "") or "") for clave, _, _ in CAMPOS_CONTROL_GENERADOR}
    for clave in ("nombre_cliente", "revision_cliente"):
        if not control_editor[clave].strip():
            control_editor[clave] = "-"
    return {
        "nombre_producto": str(producto.get("nombre", "") or ""),
        "idioma": "en" if str(documento.get("idioma", "es")).casefold() == "en" else "es",
        "control": control_editor,
        "pruebas": filas_generador,
        "notas": [str(nota) for nota in notas if nota is not None],
        "fecha_reanalisis": _valor_reanalisis_editor(datos.get("FECHA DE REANALISIS", "")),
        "desarrollador": str(datos.get("Desarrollador", "") or ""),
    }


def _marcador_comentario_celda(valor: str) -> str | None:
    coincidencia = re.search(rf"\s+([{DIGITOS_SUPERINDICE}]+)$", valor)
    if coincidencia is None:
        return None
    digitos = coincidencia.group(1).translate(DIGITOS_SUPERINDICE_A_ENTERO)
    return str(int(digitos))


def _formatear_superindice(numero: int) -> str:
    return str(numero).translate(str.maketrans("0123456789", DIGITOS_SUPERINDICE))


def _actualizar_comentarios_celda(
    pruebas: list[dict[str, str]],
    texto_notas: str,
    campo_nuevo: str | None = None,
    indice_nuevo: int | None = None,
    remover_numero: str | set[str] | None = None,
) -> str:
    if campo_nuevo is not None:
        if campo_nuevo not in {"prueba", "metodo"}:
            raise ValueError("Los comentarios solo se pueden asociar a prueba o método.")
        if indice_nuevo is None or not 0 <= indice_nuevo < len(pruebas):
            raise ValueError("No se encontró la fila seleccionada para agregar el comentario.")
        if _marcador_comentario_celda(
            pruebas[indice_nuevo].get(campo_nuevo, "")
        ):
            raise ValueError("Esta celda ya tiene un comentario.")
        if not pruebas[indice_nuevo].get(campo_nuevo, "").strip():
            instruccion = (
                "Agrega primero un método"
                if campo_nuevo == "metodo"
                else "Agrega primero una prueba"
            )
            raise ValueError(f"{instruccion} antes de añadirle un comentario.")

    numeros_a_remover = (
        {remover_numero}
        if isinstance(remover_numero, str)
        else remover_numero or set()
    )
    comentarios_por_numero: dict[str, list[str]] = {}
    notas_base = []
    comentario_actual: str | None = None
    for linea in texto_notas.splitlines():
        coincidencia = re.match(
            rf"^([{DIGITOS_SUPERINDICE}]+)(?:\s+(.*))?$",
            linea,
        )
        if coincidencia:
            comentario_actual = str(
                int(coincidencia.group(1).translate(DIGITOS_SUPERINDICE_A_ENTERO))
            )
            comentarios_por_numero[comentario_actual] = (
                [coincidencia.group(2)] if coincidencia.group(2) else []
            )
        elif comentario_actual is None:
            notas_base.append(linea)
        else:
            comentarios_por_numero[comentario_actual].append(linea)

    filas_con_comentario: list[tuple[int, str, str]] = []
    for indice, prueba in enumerate(pruebas):
        for campo in ("prueba", "metodo"):
            marcador = _marcador_comentario_celda(prueba.get(campo, ""))
            if marcador is not None and marcador not in numeros_a_remover:
                filas_con_comentario.append((indice, campo, marcador))
    if campo_nuevo is not None and indice_nuevo is not None:
        filas_con_comentario.append((indice_nuevo, campo_nuevo, "nuevo"))
    filas_con_comentario.sort(
        key=lambda elemento: (elemento[0], 0 if elemento[1] == "prueba" else 1)
    )

    comentarios_ordenados = []
    for numero, (indice, campo, marcador_anterior) in enumerate(
        filas_con_comentario, 1
    ):
        texto_celda = pruebas[indice].get(campo, "").rstrip()
        texto_celda = re.sub(
            rf"\s+[{DIGITOS_SUPERINDICE}]+$",
            "",
            texto_celda,
        )
        marcador = _formatear_superindice(numero)
        pruebas[indice][campo] = f"{texto_celda} {marcador}"
        contenido = (
            []
            if marcador_anterior == "nuevo"
            else comentarios_por_numero.get(marcador_anterior, [])
        )
        comentarios_ordenados.append((marcador, contenido))

    lineas_resultado = list(notas_base)
    for marcador, contenido in comentarios_ordenados:
        lineas_resultado.append(
            marcador + (f" {contenido[0]}" if contenido else " ")
        )
        lineas_resultado.extend(contenido[1:])
    return "\n".join(lineas_resultado)


def _guardar_celda_con_comentario(
    pruebas: list[dict[str, str]],
    texto_notas: str,
    indice: int,
    campo: str,
    valor: str,
) -> str:
    marcador_anterior = _marcador_comentario_celda(
        pruebas[indice].get(campo, "")
    )
    marcador_editado = _marcador_comentario_celda(valor)
    if marcador_anterior is not None and marcador_editado == marcador_anterior:
        texto_celda = re.sub(
            rf"\s+[{DIGITOS_SUPERINDICE}]+$",
            "",
            valor,
        ).rstrip()
        valor = f"{texto_celda} {_formatear_superindice(int(marcador_anterior))}"
    elif marcador_anterior is not None:
        valor = re.sub(rf"\s+[{DIGITOS_SUPERINDICE}]+$", "", valor).rstrip()

    pruebas[indice][campo] = valor
    if marcador_anterior is not None and marcador_editado != marcador_anterior:
        texto_notas = _actualizar_comentarios_celda(
            pruebas,
            texto_notas,
            remover_numero=marcador_anterior,
        )
    return texto_notas


def _agregar_comentario_a_celda(
    pruebas: list[dict[str, str]],
    texto_notas: str,
    indice: int,
    campo: str,
    valor: str,
) -> str:
    marcador_previo = _marcador_comentario_celda(
        pruebas[indice].get(campo, "")
    )
    if marcador_previo is not None and _marcador_comentario_celda(valor) is None:
        valor = (
            f"{valor} {_formatear_superindice(int(marcador_previo))}"
            if valor
            else ""
        )
    pruebas[indice][campo] = valor
    return _actualizar_comentarios_celda(
        pruebas,
        texto_notas,
        campo,
        indice,
    )


def _integrar_edicion_json(
    original: dict[str, Any], actualizado: dict[str, Any]
) -> dict[str, Any]:
    resultado = copy.deepcopy(original)
    for clave in (
        "esquema",
        "producto",
        "documento",
        "control_documental",
        "pruebas",
        "notas",
        "alineacion_metodos_lab",
        "Desarrollador",
        "FECHA DE REANALISIS",
    ):
        valor_actualizado = actualizado.get(clave)
        valor_original = resultado.get(clave)
        if clave in {"producto", "documento", "control_documental"} and isinstance(
            valor_original, dict
        ) and isinstance(valor_actualizado, dict):
            fusionado = copy.deepcopy(valor_original)
            for campo, valor in valor_actualizado.items():
                if (
                    clave == "documento"
                    and campo == "fecha_documento_normalizada"
                    and not valor
                    and fusionado.get(campo)
                ):
                    continue
                fusionado[campo] = valor
            resultado[clave] = fusionado
        elif clave in actualizado:
            resultado[clave] = copy.deepcopy(valor_actualizado)
    fecha_actualizada = actualizado.get("control_documental", {}).get("fecha")
    if fecha_actualizada is not None and fecha_actualizada != original.get("control_documental", {}).get("fecha"):
        from gestor_especificaciones import parsear_fecha_documento
        fecha = parsear_fecha_documento(fecha_actualizada)
        resultado.setdefault("documento", {})["fecha_documento_normalizada"] = fecha.isoformat() if fecha else ""
    resultado.setdefault("FECHA DE REANALISIS", original.get("FECHA DE REANALISIS", ""))
    return resultado


def _opciones_maestro(
    maestro: list[dict[str, str]],
    campo: str,
    idioma: str,
    consulta: str = "",
    opciones_limitadas: tuple[str, ...] | None = None,
) -> tuple[str, ...]:
    es_ingles = idioma == "en"
    if campo == "prueba":
        opciones = (
            opciones_limitadas
            if opciones_limitadas is not None
            else tuple(
                _nombre_maestro_para_idioma(registro, idioma)
                for registro in maestro
            )
        )
    elif campo == "unidades":
        clave_idioma = "UNITS" if es_ingles else "UNIDADES"
        clave_alternativa = "UNIDADES" if es_ingles else "UNITS"
        opciones = tuple(
            registro.get(clave_idioma, "").strip()
            or registro.get(clave_alternativa, "").strip()
            for registro in maestro
        )
    elif campo == "metodo":
        opciones = tuple(
            metodo_maestro_canonico(
                metodo_maestro_para_idioma(registro, idioma)
            )
            for registro in maestro
        )
    else:
        return ()

    opciones_unicas = tuple(dict.fromkeys(opcion for opcion in opciones if opcion))
    texto_consulta = consulta.strip()
    if not texto_consulta:
        return opciones_unicas[:30]

    referencias = set(referencias_astm(texto_consulta)) if campo == "metodo" else set()
    if referencias:
        coincidencias_astm = tuple(
            opcion
            for opcion in opciones_unicas
            if (
                set(referencias_astm(opcion)) == referencias
                if len(referencias) > 1
                else referencias <= set(referencias_astm(opcion))
            )
        )
        if coincidencias_astm:
            return coincidencias_astm[:30]

    consulta_casefold = texto_consulta.casefold()
    consulta_normalizada = normalizar(texto_consulta)
    comienzan = tuple(
        opcion
        for opcion in opciones_unicas
        if opcion.casefold().startswith(consulta_casefold)
    )
    contienen = tuple(
        opcion
        for opcion in opciones_unicas
        if consulta_casefold in opcion.casefold() and opcion not in comienzan
    )
    normalizadas = (
        tuple(
            opcion
            for opcion in opciones_unicas
            if consulta_normalizada in normalizar(opcion)
            and opcion not in comienzan
            and opcion not in contienen
        )
        if consulta_normalizada
        else ()
    )
    return (comienzan + contienen + normalizadas)[:30]


def _normalizar_periodicidad(valor: str, numero_fila: int) -> str:
    texto = valor.strip()
    if not texto:
        return PERIODICIDAD_PREDETERMINADA
    if texto and not re.search(r"[\s,;/|]", texto):
        simbolos = list(texto)
    else:
        simbolos = re.split(r"[\s,;/|]+", texto)
    simbolos = [simbolo for simbolo in simbolos if simbolo]
    if not 1 <= len(simbolos) <= 4:
        raise ValueError(
            f"Selecciona de 1 a 4 símbolos de periodicidad en la fila {numero_fila}."
        )
    if any(simbolo not in SIMBOLOS_PERIODICIDAD for simbolo in simbolos):
        permitidos = " ".join(SIMBOLOS_PERIODICIDAD)
        raise ValueError(
            f"Usa únicamente estos símbolos en la fila {numero_fila}: {permitidos}."
        )
    if len(set(simbolos)) != len(simbolos):
        raise ValueError(f"No repitas símbolos de periodicidad en la fila {numero_fila}.")
    return "".join(simbolos)


def _part_id_existe(part_id: str, registros: list[dict[str, str]]) -> bool:
    clave = part_id.strip().casefold()
    return bool(clave) and any(
        str(registro.get("part_id", "")).strip().casefold() == clave
        for registro in registros
    )


def _sugerencias_part_id(
    registros: list[dict[str, str]], consulta: str, limite: int = 5
) -> tuple[str, ...]:
    texto = consulta.strip().casefold()
    if not texto:
        return ()
    coincidencias = []
    for registro in registros:
        part_id = str(registro.get("part_id", "")).strip()
        descripcion = str(registro.get("descripcion", "")).strip()
        if not part_id:
            continue
        id_normalizado = part_id.casefold()
        descripcion_normalizada = descripcion.casefold()
        if id_normalizado.startswith(texto):
            prioridad = 0
        elif texto in id_normalizado:
            prioridad = 1
        elif texto in descripcion_normalizada:
            prioridad = 2
        else:
            continue
        coincidencias.append((prioridad, part_id))
    coincidencias.sort(key=lambda item: item[0])
    return tuple(dict.fromkeys(codigo for _, codigo in coincidencias))[:limite]


def _part_id_canonico(part_id: str, registros: list[dict[str, str]]) -> str:
    clave = part_id.strip().casefold()
    return next(
        (
            str(registro.get("part_id", "")).strip()
            for registro in registros
            if str(registro.get("part_id", "")).strip().casefold() == clave
        ),
        "",
    )


def _sugerencias_nombre_producto(
    registros: list[dict[str, str]], consulta: str, limite: int = 5
) -> tuple[str, ...]:
    texto = normalizar(consulta.strip())
    if not texto:
        return ()
    coincidencias = []
    for registro in registros:
        descripcion = str(registro.get("descripcion", "")).strip()
        descripcion_normalizada = normalizar(descripcion)
        if not descripcion or not descripcion_normalizada:
            continue
        if descripcion_normalizada.startswith(texto):
            prioridad = 0
        elif texto in descripcion_normalizada:
            prioridad = 1
        else:
            continue
        coincidencias.append((prioridad, descripcion))
    coincidencias.sort(key=lambda item: item[0])
    return tuple(dict.fromkeys(nombre for _, nombre in coincidencias))[:limite]


def _part_id_por_nombre(
    nombre: str, registros: list[dict[str, str]]
) -> str:
    nombre_clave = normalizar(nombre.strip())
    if not nombre_clave:
        return ""
    codigos = tuple(
        dict.fromkeys(
            str(registro.get("part_id", "")).strip()
            for registro in registros
            if normalizar(str(registro.get("descripcion", ""))) == nombre_clave
            and str(registro.get("part_id", "")).strip()
        )
    )
    return codigos[0] if len(codigos) == 1 else ""


def _cargar_catalogo_part_ids() -> list[dict[str, str]]:
    import gestor_especificaciones as gestor

    registros, _ = gestor.leer_catalogo(CATALOGO_PART_ID)
    return [
        {"part_id": registro["part_id"], "descripcion": registro["descripcion"]}
        for registro in registros
        if registro.get("part_id")
    ]


def _prueba_cercana(
    maestro: list[dict[str, str]],
    idioma: str,
    consulta: str,
    opciones_limitadas: tuple[str, ...] | None = None,
) -> str:
    texto = normalizar(consulta.strip())
    if len(texto) < 2:
        return ""
    opciones = opciones_limitadas or tuple(
        dict.fromkeys(
            nombre
            for registro in maestro
            if (nombre := _nombre_maestro_para_idioma(registro, idioma))
        )
    )
    if not opciones:
        return ""
    candidata = max(
        opciones,
        key=lambda opcion: (
            normalizar(opcion).startswith(texto),
            difflib.SequenceMatcher(None, texto, normalizar(opcion)).ratio(),
        ),
    )
    similitud = difflib.SequenceMatcher(None, texto, normalizar(candidata)).ratio()
    return candidata if similitud >= 0.35 else ""


def _unidad_sugerida(
    maestro: list[dict[str, str]],
    prueba: str,
    metodo: str,
    idioma: str,
) -> str:
    nombre_clave = normalizar(prueba)
    if not nombre_clave:
        return ""
    relacionadas = [
        registro
        for registro in maestro
        if normalizar(_nombre_maestro_para_idioma(registro, idioma)) == nombre_clave
    ]
    referencias = set(referencias_astm(metodo))
    if referencias:
        por_metodo = [
            registro
            for registro in relacionadas
            if set(
                referencias_astm(metodo_maestro_para_idioma(registro, idioma))
            )
            == referencias
        ]
        if por_metodo:
            relacionadas = por_metodo
    if not relacionadas:
        return ""

    unidades_idioma = "UNITS" if idioma == "en" else "UNIDADES"
    unidades_alternativas = "UNIDADES" if idioma == "en" else "UNITS"
    unidades = tuple(
        dict.fromkeys(
            unidad.strip()
            for registro in relacionadas
            for columna in (unidades_idioma, unidades_alternativas)
            for unidad in re.split(r"\s*[,;/|]\s*", registro.get(columna, ""))
            if unidad.strip()
        )
    )
    es_temperatura = bool(
        re.search(
            r"\b(?:temperatura|temperature|flash point|punto de inflamacion|"
            r"pour point|punto de fluidez|cloud point|punto de nube|fire point)\b",
            normalizar(prueba),
            re.IGNORECASE,
        )
    )
    if es_temperatura:
        celsius = next(
            (
                unidad
                for unidad in unidades
                if unidad.strip().casefold().replace(" ", "")
                in {"°c", "ºc", "degc", "celsius", "centigrade"}
            ),
            "",
        )
        if celsius:
            return celsius
    return unidades[0] if len(unidades) == 1 else ""


def _valores_prueba_predeterminados(
    maestro: list[dict[str, str]], prueba: str, idioma: str
) -> dict[str, str]:
    prueba_clave = normalizar(prueba)
    relacionadas = [
        registro
        for registro in maestro
        if normalizar(_nombre_maestro_para_idioma(registro, idioma)) == prueba_clave
    ]
    if not prueba_clave or not relacionadas:
        return {}

    metodos = tuple(
        dict.fromkeys(
            metodo_maestro_canonico(
                metodo_maestro_para_idioma(registro, idioma)
            )
            for registro in relacionadas
            if metodo_maestro_canonico(
                metodo_maestro_para_idioma(registro, idioma)
            )
        )
    )
    valores: dict[str, str] = {}
    if len(metodos) == 1:
        valores["metodo"] = metodos[0]

    unidad = _unidad_sugerida(
        maestro,
        prueba,
        valores.get("metodo", ""),
        idioma,
    )
    registros_unidad = relacionadas
    if valores.get("metodo"):
        metodo_clave = normalizar(valores["metodo"])
        por_metodo = [
            registro
            for registro in relacionadas
            if normalizar(
                metodo_maestro_canonico(
                    metodo_maestro_para_idioma(registro, idioma)
                )
            )
            == metodo_clave
        ]
        if por_metodo:
            registros_unidad = por_metodo
    campo_unidad = "UNITS" if idioma == "en" else "UNIDADES"
    unidades_localizadas = tuple(
        dict.fromkeys(
            unidad_localizada.strip()
            for registro in registros_unidad
            for unidad_localizada in re.split(
                r"\s*[,;|]\s*|\s+/\s+", registro.get(campo_unidad, "")
            )
            if unidad_localizada.strip()
        )
    )
    if not unidad and len(unidades_localizadas) == 1:
        unidad = unidades_localizadas[0]
    if unidad:
        valores["unidades"] = unidad
    return valores


def _nombre_maestro_para_idioma(registro: dict[str, str], idioma: str) -> str:
    nombre, _ = nombres_canonicos(
        registro, "INGLES" if idioma == "en" else "ESPANOL"
    )
    if nombre:
        return nombre
    nombre_alternativo, _ = nombres_canonicos(
        registro, "ESPANOL" if idioma == "en" else "INGLES"
    )
    return nombre_alternativo


def _nombres_maestros_para_metodo(
    maestro: list[dict[str, str]], metodo: str, idioma: str
) -> tuple[str, ...]:
    referencias = set(referencias_astm(metodo))
    if not referencias:
        return ()
    return tuple(
        dict.fromkeys(
            nombre
            for registro in maestro
            if (
                set(
                    referencias_astm(
                        metodo_maestro_para_idioma(registro, idioma)
                    )
                )
                == referencias
                if len(referencias) > 1
                else referencias
                <= set(
                    referencias_astm(
                        metodo_maestro_para_idioma(registro, idioma)
                    )
                )
            )
            and (nombre := _nombre_maestro_para_idioma(registro, idioma))
        )
    )


def _procesar_carga_especificaciones(
    base: Path,
    archivos_seleccionados: list[Path] | None = None,
    al_mensaje: Callable[[str], None] | None = None,
    al_progreso: Callable[[str, int, int, str], None] | None = None,
) -> dict[str, Any]:
    import gestor_especificaciones as gestor
    from alinear_especificaciones_json import ejecutar as alinear_especificaciones

    resultado: dict[str, Any] = {}
    salida = (
        _SalidaCarga(al_mensaje, sys.stdout)
        if al_mensaje is not None
        else None
    )
    error = (
        _SalidaCarga(al_mensaje, sys.stderr)
        if al_mensaje is not None
        else None
    )
    try:
        if al_mensaje is not None:
            al_mensaje("Iniciando extracción de especificaciones.")
        if archivos_seleccionados is None:
            if salida is None:
                resultado["extraccion"] = gestor.procesar(base)
            else:
                with contextlib.redirect_stdout(salida), contextlib.redirect_stderr(error):
                    resultado["extraccion"] = gestor.procesar(base)
        else:
            json_generados: list[Path] = []
            extraer_seleccion = lambda: gestor.procesar(
                base,
                archivos_seleccionados=archivos_seleccionados,
                al_progreso=(
                    lambda completados, total, nombre: al_progreso(
                        "Extracción", completados, total, nombre
                    )
                    if al_progreso is not None
                    else None
                ),
                json_generados=json_generados,
            )
            if salida is None:
                resultado["extraccion"] = extraer_seleccion()
            else:
                with contextlib.redirect_stdout(salida), contextlib.redirect_stderr(error):
                    resultado["extraccion"] = extraer_seleccion()
            resultado["_json_generados"] = json_generados
    except Exception as exc:
        resultado["error_extraccion"] = str(exc)
        if al_mensaje is not None:
            al_mensaje(f"ERROR en extracción: {exc}")
    try:
        if al_mensaje is not None:
            al_mensaje("Iniciando alineación de especificaciones.")
        if archivos_seleccionados is not None:
            json_generados = resultado.get("_json_generados", [])
            if json_generados:
                alinear_seleccion = lambda: alinear_especificaciones(
                    base,
                    archivos_seleccionados=json_generados,
                    al_progreso=(
                        lambda completados, total, nombre: al_progreso(
                            "Alineación", completados, total, nombre
                        )
                        if al_progreso is not None
                        else None
                    ),
                )
                if salida is None:
                    resultado["alineacion"] = alinear_seleccion()
                else:
                    with contextlib.redirect_stdout(salida), contextlib.redirect_stderr(error):
                        resultado["alineacion"] = alinear_seleccion()
            else:
                resultado["alineacion"] = {
                    "archivos": 0,
                    "alineados": 0,
                    "revision": 0,
                    "sin_coincidencia": 0,
                    "errores": 0,
                }
        elif salida is None:
            resultado["alineacion"] = alinear_especificaciones(base)
        else:
            with contextlib.redirect_stdout(salida), contextlib.redirect_stderr(error):
                resultado["alineacion"] = alinear_especificaciones(base)
    except Exception as exc:
        resultado["error_alineacion"] = str(exc)
        if al_mensaje is not None:
            al_mensaje(f"ERROR en alineación: {exc}")
    resultado.pop("_json_generados", None)
    if salida is not None:
        salida.flush()
        error.flush()
    return resultado


class _SalidaCarga:
    def __init__(
        self,
        al_mensaje: Callable[[str], None],
        salida_original: Any,
    ):
        self.al_mensaje = al_mensaje
        self.salida_original = salida_original
        self.buffer = ""

    def write(self, texto: str) -> int:
        if self.salida_original is not None:
            self.salida_original.write(texto)
        self.buffer += texto
        while "\n" in self.buffer:
            linea, self.buffer = self.buffer.split("\n", 1)
            if linea.strip():
                self.al_mensaje(linea.rstrip())
        return len(texto)

    def flush(self) -> None:
        if self.salida_original is not None:
            self.salida_original.flush()
        if self.buffer.strip():
            self.al_mensaje(self.buffer.rstrip())
        self.buffer = ""


def _cargar_configuracion() -> dict[str, Path]:
    return rutas_configuradas(ARCHIVO_CONFIGURACION.parent, archivo=ARCHIVO_CONFIGURACION, estricto=True)


def _guardar_configuracion(rutas: dict[str, Path]) -> None:
    base = ARCHIVO_CONFIGURACION.parent
    rutas = {clave: resolver_ruta(ruta, base) for clave, ruta in rutas.items()}
    if "metodos_lab" in rutas:
        carpeta = Path(rutas["metodos_lab"])
        faltantes = [nombre for nombre in ("METODOS_LAB_UNIFICADO.xlsx", "TRADUCCIONES_ESPECIFICACIONES.xlsx")
                     if not (carpeta / nombre).is_file()]
        if faltantes:
            raise ValueError("La carpeta de métodos y traducciones debe contener:\n" + "\n".join(faltantes))
    temporal = ARCHIVO_CONFIGURACION.with_suffix(".tmp")
    temporal.write_text(
        json.dumps({clave: ruta_portatil(ruta, base) for clave, ruta in rutas.items()}, indent=2),
        encoding="utf-8",
    )
    temporal.replace(ARCHIVO_CONFIGURACION)


def _abrir_carpeta(ruta: Path) -> None:
    if not ruta.is_dir():
        raise FileNotFoundError(f"No existe la carpeta:\n{ruta}")
    abrir_en_windows = getattr(os, "startfile", None)
    if callable(abrir_en_windows):
        abrir_en_windows(str(ruta))
    elif sys.platform == "darwin":
        subprocess.Popen(["open", str(ruta)])
    else:
        subprocess.Popen(["xdg-open", str(ruta)])


def _consultar_tabla_historico(
    ruta: Path,
    tabla: str,
    busqueda: str = "",
    limite: int = 100,
    desplazamiento: int = 0,
) -> tuple[tuple[str, ...], list[tuple[str, ...]], int]:
    if limite < 1 or desplazamiento < 0:
        raise ValueError("El límite y el desplazamiento deben ser válidos.")
    if not ruta.is_file():
        raise FileNotFoundError(f"No se encontró la base de datos:\n{ruta}")

    conexion = sqlite3.connect(
        ruta.resolve().as_uri() + "?mode=ro",
        uri=True,
        timeout=30,
    )
    try:
        conexion.execute("PRAGMA busy_timeout=30000")
        tablas = {
            nombre
            for (nombre,) in conexion.execute(
                "SELECT name FROM sqlite_master "
                "WHERE type = 'table' AND name NOT LIKE 'sqlite_%'"
            )
        }
        if tabla not in tablas:
            raise ValueError(f"La tabla {tabla!r} no existe en la base de datos.")
        nombre_escapado = '"' + tabla.replace('"', '""') + '"'
        columnas = tuple(
            str(fila[1])
            for fila in conexion.execute(f"PRAGMA table_info({nombre_escapado})")
        )
        if not columnas:
            raise ValueError(f"No se pudieron leer las columnas de {tabla!r}.")

        clausula_busqueda = ""
        parametros: list[Any] = []
        if busqueda.strip():
            columnas_escapadas = [
                '"' + columna.replace('"', '""') + '"' for columna in columnas
            ]
            clausula_busqueda = " WHERE " + " OR ".join(
                f"CAST({columna} AS TEXT) LIKE ? COLLATE NOCASE"
                for columna in columnas_escapadas
            )
            patron = f"%{busqueda.strip()}%"
            parametros.extend(patron for _ in columnas)

        total = int(
            conexion.execute(
                f"SELECT COUNT(*) FROM {nombre_escapado}{clausula_busqueda}",
                parametros,
            ).fetchone()[0]
        )
        filas = [
            tuple("" if valor is None else str(valor) for valor in fila)
            for fila in conexion.execute(
                f"SELECT * FROM {nombre_escapado}{clausula_busqueda} "
                f"LIMIT {int(limite)} OFFSET {int(desplazamiento)}",
                parametros,
            )
        ]
        return columnas, filas, total
    finally:
        conexion.close()


def _exportar_tabla_historico_excel(
    ruta_base: Path,
    tabla: str,
    destino: Path,
    busqueda: str = "",
) -> int:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    columnas, _, total = _consultar_tabla_historico(
        ruta_base,
        tabla,
        busqueda,
        limite=1,
    )
    filas: list[tuple[str, ...]] = []
    desplazamiento = 0
    while desplazamiento < total:
        _, pagina, _ = _consultar_tabla_historico(
            ruta_base,
            tabla,
            busqueda,
            limite=1000,
            desplazamiento=desplazamiento,
        )
        if not pagina:
            break
        filas.extend(pagina)
        desplazamiento += len(pagina)

    libro = Workbook()
    hoja = libro.active
    nombre_hoja = re.sub(r"[\\/*?:\[\]]", "_", tabla)[:31] or "Historico"
    hoja.title = nombre_hoja
    hoja.append(list(columnas))
    for fila in filas:
        hoja.append(list(fila))
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1F4E78")
    for fila in hoja.iter_rows():
        for celda in fila:
            celda.data_type = "s"
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions
    for indice, columna in enumerate(columnas, start=1):
        ancho = min(40, max(12, len(columna) + 2))
        hoja.column_dimensions[hoja.cell(1, indice).column_letter].width = ancho
    libro.save(destino)
    return len(filas)


class VisorHistorico:
    TAMANO_PAGINA = 100

    def __init__(
        self,
        raiz: tk.Misc,
        ruta: Path,
        carpeta_historico: Path,
        contenedor: tk.Misc | None = None,
    ):
        self.ruta = ruta
        self.carpeta_historico = carpeta_historico
        self.desplazamiento = 0
        tablas = self._listar_tablas()
        self.ventana = contenedor if contenedor is not None else ui.Ventana(raiz)
        if contenedor is None:
            self.ventana.title("Histórico de especificaciones · Base de datos")
            self.ventana.geometry("1250x700")
            self.ventana.minsize(850, 450)

        cabecera = ttk.Frame(self.ventana, padding=(12, 10, 12, 6))
        cabecera.pack(fill="x")
        ttk.Label(
            cabecera,
            text="Tablas de especificaciones.db",
            font=("Segoe UI", 14, "bold"),
            foreground="#163b70",
        ).pack(side="left")
        ttk.Label(cabecera, text=str(ruta)).pack(side="right")
        ttk.Button(
            cabecera,
            text="Abrir carpeta de documentos históricos",
            command=self._abrir_carpeta_historico,
        ).pack(side="right", padx=(0, 12))

        controles = ttk.Frame(self.ventana, padding=(12, 4))
        controles.pack(fill="x")
        ttk.Label(controles, text="Tabla:").pack(side="left")
        self.tabla = ttk.Combobox(
            controles,
            state="readonly",
            width=24,
            values=tablas,
        )
        self.tabla.pack(side="left", padx=(6, 16))
        self.tabla.bind("<<ComboboxSelected>>", self._cambiar_tabla)
        ttk.Label(controles, text="Buscar en todos los campos:").pack(side="left")
        self.busqueda = tk.StringVar()
        entrada_busqueda = ttk.Entry(controles, textvariable=self.busqueda)
        entrada_busqueda.pack(side="left", fill="x", expand=True, padx=6)
        entrada_busqueda.bind("<Return>", self._buscar)
        ttk.Button(controles, text="Buscar", command=self._buscar).pack(
            side="left", padx=(0, 5)
        )
        ttk.Button(
            controles,
            text="Limpiar",
            command=self._limpiar_busqueda,
        ).pack(side="left")
        ttk.Button(
            controles,
            text="Exportar a Excel",
            command=self._exportar_excel,
        ).pack(side="right", padx=(8, 0))

        marco_tabla = ttk.Frame(self.ventana, padding=(12, 6))
        marco_tabla.pack(fill="both", expand=True)
        self.resultados = ttk.Treeview(marco_tabla, show="headings")
        barra_vertical = ttk.Scrollbar(
            marco_tabla,
            orient="vertical",
            command=self.resultados.yview,
        )
        barra_horizontal = ttk.Scrollbar(
            marco_tabla,
            orient="horizontal",
            command=self.resultados.xview,
        )
        self.resultados.configure(
            yscrollcommand=barra_vertical.set,
            xscrollcommand=barra_horizontal.set,
        )
        self.resultados.grid(row=0, column=0, sticky="nsew")
        barra_vertical.grid(row=0, column=1, sticky="ns")
        barra_horizontal.grid(row=1, column=0, sticky="ew")
        marco_tabla.rowconfigure(0, weight=1)
        marco_tabla.columnconfigure(0, weight=1)

        pie = ttk.Frame(self.ventana, padding=(12, 4, 12, 10))
        pie.pack(fill="x")
        self.estado = ttk.Label(pie, text="")
        self.estado.pack(side="left")
        self.boton_anterior = ttk.Button(
            pie,
            text="← Anterior",
            command=self._pagina_anterior,
        )
        self.boton_anterior.pack(side="right")
        self.boton_siguiente = ttk.Button(
            pie,
            text="Siguiente →",
            command=self._pagina_siguiente,
        )
        self.boton_siguiente.pack(side="right", padx=6)

        if self.tabla["values"]:
            self.tabla.current(0)
            self._cargar_pagina()
        else:
            self._mostrar_error("La base de datos no contiene tablas disponibles.")

    def _listar_tablas(self) -> tuple[str, ...]:
        if not self.ruta.is_file():
            raise FileNotFoundError(f"No se encontró la base de datos:\n{self.ruta}")
        conexion = sqlite3.connect(
            self.ruta.resolve().as_uri() + "?mode=ro",
            uri=True,
            timeout=30,
        )
        try:
            conexion.execute("PRAGMA busy_timeout=30000")
            return tuple(
                str(fila[0])
                for fila in conexion.execute(
                    "SELECT name FROM sqlite_master "
                    "WHERE type = 'table' AND name NOT LIKE 'sqlite_%' "
                    "ORDER BY name"
                )
            )
        finally:
            conexion.close()

    def _cargar_pagina(self) -> None:
        tabla = self.tabla.get()
        if not tabla:
            return
        try:
            columnas, filas, total = _consultar_tabla_historico(
                self.ruta,
                tabla,
                self.busqueda.get(),
                self.TAMANO_PAGINA,
                self.desplazamiento,
            )
        except (OSError, sqlite3.Error, ValueError) as exc:
            self._mostrar_error(str(exc))
            return

        self.resultados.delete(*self.resultados.get_children())
        self.resultados.configure(columns=columnas)
        for columna in columnas:
            self.resultados.heading(columna, text=columna)
            self.resultados.column(
                columna,
                width=min(280, max(110, len(columna) * 9)),
                minwidth=80,
                stretch=True,
            )
        for fila in filas:
            self.resultados.insert("", "end", values=fila)

        inicio = self.desplazamiento + 1 if filas else 0
        fin = self.desplazamiento + len(filas)
        self.estado.configure(
            text=f"{tabla}: {inicio}–{fin} de {total} registro(s)"
        )
        self.boton_anterior.configure(
            state="normal" if self.desplazamiento > 0 else "disabled"
        )
        self.boton_siguiente.configure(
            state=(
                "normal"
                if self.desplazamiento + len(filas) < total
                else "disabled"
            )
        )

    def _cambiar_tabla(self, _evento: tk.Event | None = None) -> None:
        self.desplazamiento = 0
        self._cargar_pagina()

    def _buscar(self, _evento: tk.Event | None = None) -> None:
        self.desplazamiento = 0
        self._cargar_pagina()

    def _limpiar_busqueda(self) -> None:
        self.busqueda.set("")
        self.desplazamiento = 0
        self._cargar_pagina()

    def _pagina_anterior(self) -> None:
        self.desplazamiento = max(0, self.desplazamiento - self.TAMANO_PAGINA)
        self._cargar_pagina()

    def _pagina_siguiente(self) -> None:
        self.desplazamiento += self.TAMANO_PAGINA
        self._cargar_pagina()

    def _exportar_excel(self) -> None:
        tabla = self.tabla.get()
        if not tabla:
            messagebox.showinfo(
                "Selecciona una tabla",
                "Selecciona una tabla antes de exportar.",
                parent=self.ventana,
            )
            return
        destino = filedialog.asksaveasfilename(
            parent=self.ventana,
            title="Exportar tabla del histórico a Excel",
            initialfile=f"{tabla}.xlsx",
            defaultextension=".xlsx",
            filetypes=(("Libro de Excel", "*.xlsx"),),
        )
        if not destino:
            return
        try:
            cantidad = _exportar_tabla_historico_excel(
                self.ruta,
                tabla,
                Path(destino),
                self.busqueda.get(),
            )
        except (OSError, PermissionError, sqlite3.Error, ValueError) as exc:
            messagebox.showerror(
                "No se pudo exportar a Excel",
                str(exc),
                parent=self.ventana,
            )
            return
        messagebox.showinfo(
            "Exportación completada",
            f"Se exportaron {cantidad} registro(s) a:\n{destino}",
            parent=self.ventana,
        )

    def _mostrar_error(self, mensaje: str) -> None:
        self.resultados.delete(*self.resultados.get_children())
        self.resultados.configure(columns=())
        self.estado.configure(text=mensaje)
        self.boton_anterior.configure(state="disabled")
        self.boton_siguiente.configure(state="disabled")

    def _abrir_carpeta_historico(self) -> None:
        try:
            self.carpeta_historico.mkdir(parents=True, exist_ok=True)
            _abrir_carpeta(self.carpeta_historico)
        except OSError as exc:
            messagebox.showerror(
                "No se pudo abrir el histórico",
                str(exc),
                parent=self.ventana,
            )


def mostrar_vista_json(contenedor: tk.Misc, ruta: Path, zoom: float = 1.0) -> dict | None:
    """Renderiza un JSON antiguo o alineado en un panel con zoom propio."""
    for widget in contenedor.winfo_children():
        widget.destroy()
    marco = ttk.Frame(contenedor)
    marco.pack(fill="both", expand=True)
    try:
        datos = json.loads(ruta.read_text(encoding="utf-8-sig"))
        alineado = (
            ruta.parent.name.casefold() == "especificaciones_json_alineadas".casefold()
            or isinstance(datos, dict) and isinstance(datos.get("alineacion_metodos_lab"), dict)
        )
        previa = (
            preparar_vista_previa_alineada(ruta)
            if alineado
            else preparar_vista_previa(ruta)
        )
    except (OSError, UnicodeError, json.JSONDecodeError, ValueError) as exc:
        ttk.Label(
            marco,
            text=f"No se pudo mostrar la especificación:\n{exc}",
            foreground="#9a3412",
            wraplength=440,
            justify="left",
        ).pack(anchor="w", padx=8, pady=8)
        return

    separador = crear_divisor_redimensionable(marco, tk.VERTICAL)
    separador.pack(fill="both", expand=True)
    area_tabla = ttk.Frame(separador, padding=8)
    area_detalles = ttk.Frame(separador)
    separador.add(area_tabla, minsize=120, stretch="always")
    separador.add(area_detalles, minsize=80, stretch="always")
    area_tabla.rowconfigure(2, weight=1)
    area_tabla.columnconfigure(0, weight=1)
    separador_inicializado = False

    def posicionar_separador(evento: tk.Event) -> None:
        nonlocal separador_inicializado
        if not separador_inicializado and evento.height > 100:
            separador_inicializado = True
            separador.sash_place(0, 0, round(evento.height * 0.68))

    separador.bind("<Configure>", posicionar_separador)
    ttk.Label(
        area_tabla,
        text=previa.nombre_producto,
        font=("Segoe UI", max(11, round(17 * zoom)), "bold"),
        foreground="#163b70",
        wraplength=680,
    ).grid(row=0, column=0, sticky="w", pady=(0, 4))
    estilo = ttk.Style(contenedor)
    sufijo = str(id(contenedor))
    estilo_filas = f"Preview{sufijo}.Treeview"
    estilo.configure(
        estilo_filas,
        font=("Segoe UI", max(7, round(9 * zoom))),
        rowheight=max(18, round(24 * zoom)),
    )
    estilo.configure(
        f"{estilo_filas}.Heading",
        font=("Segoe UI", max(7, round(9 * zoom)), "bold"),
    )
    ttk.Label(
        area_tabla,
        text=(
            f"Tests ({len(previa.pruebas)})"
            if previa.idioma == "INGLES"
            else f"Pruebas ({len(previa.pruebas)})"
        ),
        font=("Segoe UI", max(9, round(11 * zoom)), "bold"),
        foreground="#163b70",
    ).grid(row=1, column=0, sticky="w", pady=(0, 5))
    tabla = ttk.Treeview(
        area_tabla,
        columns=CAMPOS_TABLA,
        show="headings",
        height=min(max(len(previa.pruebas), 4), 16),
        style=estilo_filas,
    )
    for titulo, campo in encabezados_tabla(previa.idioma):
        tabla.heading(campo, text=titulo)
        tabla.column(campo, width=max(100, round(145 * zoom)), minwidth=70, stretch=True)
    barra_tabla_y = ttk.Scrollbar(area_tabla, orient="vertical", command=tabla.yview)
    barra_tabla_x = ttk.Scrollbar(area_tabla, orient="horizontal", command=tabla.xview)
    tabla.configure(yscrollcommand=barra_tabla_y.set, xscrollcommand=barra_tabla_x.set)
    tabla.grid(row=2, column=0, sticky="nsew")
    barra_tabla_y.grid(row=2, column=1, sticky="ns")
    barra_tabla_x.grid(row=3, column=0, sticky="ew")
    for prueba in previa.pruebas:
        tabla.insert("", tk.END, values=tuple(str(prueba.get(campo, "")) for campo in CAMPOS_TABLA))
    OrdenadorColumnasTreeview(
        tabla,
        {
            campo: titulo
            for titulo, campo in encabezados_tabla(previa.idioma)
            if campo in {"prueba", "metodo"}
        },
    )

    canvas = tk.Canvas(area_detalles, highlightthickness=0, background="#f4f6f8")
    vertical = ttk.Scrollbar(area_detalles, orient="vertical", command=canvas.yview)
    horizontal = ttk.Scrollbar(area_detalles, orient="horizontal", command=canvas.xview)
    canvas.configure(yscrollcommand=vertical.set, xscrollcommand=horizontal.set)
    canvas.grid(row=0, column=0, sticky="nsew")
    vertical.grid(row=0, column=1, sticky="ns")
    horizontal.grid(row=1, column=0, sticky="ew")
    area_detalles.rowconfigure(0, weight=1)
    area_detalles.columnconfigure(0, weight=1)
    pagina = ttk.Frame(canvas, padding=12)
    ventana = canvas.create_window((0, 0), window=pagina, anchor="nw")
    pagina.bind("<Configure>", lambda _e: canvas.configure(scrollregion=canvas.bbox("all")))
    canvas.bind("<Configure>", lambda e: canvas.itemconfigure(ventana, width=max(e.width, 640)))

    if previa.idioma == "INGLES":
        ttk.Label(
            pagina,
            text="English version",
            font=("Segoe UI", max(8, round(9 * zoom)), "italic"),
            foreground="#526174",
        ).pack(anchor="w", pady=(0, 4))
    ttk.Label(pagina, text=str(ruta), wraplength=680).pack(anchor="w", pady=(0, 8))
    for etiqueta, valor in previa.control:
        ttk.Label(
            pagina,
            text=f"{etiqueta}: {valor}",
            font=("Segoe UI", max(8, round(9 * zoom))),
            wraplength=680,
        ).pack(anchor="w")

    avisos = (*previa.mensajes_control, *previa.mensajes_astm_no_encontradas)
    for aviso in avisos:
        ttk.Label(
            pagina,
            text=aviso,
            foreground="#9a3412",
            font=("Segoe UI", max(8, round(9 * zoom))),
            wraplength=680,
        ).pack(anchor="w", pady=(5, 0))

    if previa.notas:
        ttk.Label(
            pagina,
            text=(
                "Notes: " if previa.idioma == "INGLES" else "Notas: "
            ) + " · ".join(previa.notas),
            wraplength=680,
            font=("Segoe UI", max(8, round(9 * zoom))),
        ).pack(anchor="w", pady=(8, 0))

    contenedor._ui_preservar_fuentes = True
    ventana_previa = contenedor.winfo_toplevel()
    if isinstance(ventana_previa, ui.VentanaAdaptable):
        ventana_previa._estilizar_hijos(contenedor, ajustar_fuentes=False)
    return {"tabla": tabla, "filas": tuple(tabla.get_children()), "pruebas": previa.pruebas,
            "idioma": "en" if previa.idioma == "INGLES" else "es", "alineada": alineado,
            "vertical": barra_tabla_y.set, "horizontal": barra_tabla_x.set}


def _leer_filas_maestro_editable(
    ruta: Path,
) -> list[tuple[int, tuple[str, ...]]]:
    from openpyxl import load_workbook

    libro = load_workbook(ruta, read_only=True, data_only=False)
    try:
        if "METODOS" not in libro.sheetnames:
            raise ValueError("El Excel maestro no contiene la hoja METODOS.")
        hoja = libro["METODOS"]
        encabezados = tuple(
            hoja.cell(1, indice).value
            for indice in range(1, len(COLUMNAS_MAESTRO) + 1)
        )
        if encabezados != COLUMNAS_MAESTRO:
            raise ValueError(
                "La hoja METODOS no tiene las siete columnas esperadas: "
                + ", ".join(COLUMNAS_MAESTRO)
            )
        return [
            (
                numero_fila,
                tuple("" if valor is None else str(valor) for valor in valores),
            )
            for numero_fila, valores in enumerate(
                hoja.iter_rows(
                    min_row=2,
                    max_col=len(COLUMNAS_MAESTRO),
                    values_only=True,
                ),
                start=2,
            )
            if any(valor is not None for valor in valores)
        ]
    finally:
        libro.close()


def _guardar_cambios_maestro(
    ruta: Path,
    cambios: dict[int, tuple[str, ...]],
) -> None:
    if not cambios:
        return
    from openpyxl import load_workbook

    libro = load_workbook(ruta)
    try:
        if "METODOS" not in libro.sheetnames:
            raise ValueError("El Excel maestro no contiene la hoja METODOS.")
        hoja = libro["METODOS"]
        encabezados = tuple(
            hoja.cell(1, indice).value
            for indice in range(1, len(COLUMNAS_MAESTRO) + 1)
        )
        if encabezados != COLUMNAS_MAESTRO:
            raise ValueError(
                "La hoja METODOS no tiene las siete columnas esperadas: "
                + ", ".join(COLUMNAS_MAESTRO)
            )
        hubo_cambios = False
        for numero_fila, valores in cambios.items():
            if numero_fila < 2 or numero_fila > hoja.max_row:
                raise ValueError(f"La fila {numero_fila} no existe en la lista maestra.")
            if len(valores) != len(COLUMNAS_MAESTRO):
                raise ValueError("Cada fila editada debe incluir los siete campos.")
            for indice_columna, valor in enumerate(valores, start=1):
                actual = hoja.cell(numero_fila, indice_columna).value
                hubo_cambios |= ("" if actual is None else str(actual).strip()) != ("" if valor is None else str(valor).strip())
                hoja.cell(numero_fila, indice_columna).value = valor

        if not hubo_cambios:
            return

        with tempfile.TemporaryDirectory(
            prefix=".edicion_metodos_",
            dir=ruta.parent,
        ) as directorio_temporal:
            temporal = Path(directorio_temporal) / ruta.name
            libro.save(temporal)
            temporal.replace(ruta)
    finally:
        libro.close()


class EditorListaMaestra:
    def __init__(self, raiz: tk.Misc, ruta: Path, al_guardar=None):
        ruta = resolver_ruta(ruta)
        self.ruta = ruta
        self.al_guardar = al_guardar
        self.filas = _leer_filas_maestro_editable(ruta)
        self.cambios: dict[int, tuple[str, ...]] = {}
        self.editor: ttk.Entry | None = None
        self.celda_editada: tuple[str, str] | None = None
        self.ventana = ui.Ventana(raiz)
        self.ventana.title("Editar lista maestra de métodos")
        self.ventana.geometry("1250x700")
        self.ventana.minsize(850, 450)
        self.ventana.transient(raiz)
        self.ventana.grab_set()

        ttk.Label(
            self.ventana,
            text="Lista maestra de métodos",
            font=("Segoe UI", 14, "bold"),
            foreground="#163b70",
        ).pack(anchor="w", padx=12, pady=(12, 4))
        ttk.Label(
            self.ventana,
            text="Haz doble clic en una celda para editarla. Se pueden modificar los siete campos.",
        ).pack(anchor="w", padx=12, pady=(0, 8))
        barra_busqueda = ttk.Frame(self.ventana)
        barra_busqueda.pack(fill="x", padx=12, pady=(0, 8))
        ttk.Label(barra_busqueda, text="Buscar:").pack(side="left")
        self.busqueda = tk.StringVar()
        entrada_busqueda = ttk.Entry(barra_busqueda, textvariable=self.busqueda)
        entrada_busqueda.pack(side="left", fill="x", expand=True, padx=(6, 0))
        entrada_busqueda.bind("<KeyRelease>", lambda _evento: self._filtrar())

        marco_tabla = ttk.Frame(self.ventana)
        marco_tabla.pack(fill="both", expand=True, padx=12)
        self.tabla = ttk.Treeview(
            marco_tabla,
            columns=COLUMNAS_MAESTRO,
            show="headings",
            selectmode="browse",
        )
        anchos = (210, 210, 270, 270, 110, 110, 170)
        for columna, ancho in zip(COLUMNAS_MAESTRO, anchos):
            self.tabla.heading(columna, text=columna)
            self.tabla.column(columna, width=ancho, minwidth=90, stretch=True)
        barra_vertical = ttk.Scrollbar(
            marco_tabla,
            orient="vertical",
            command=self.tabla.yview,
        )
        barra_horizontal = ttk.Scrollbar(
            marco_tabla,
            orient="horizontal",
            command=self.tabla.xview,
        )
        self.tabla.configure(
            yscrollcommand=barra_vertical.set,
            xscrollcommand=barra_horizontal.set,
        )
        self.tabla.grid(row=0, column=0, sticky="nsew")
        barra_vertical.grid(row=0, column=1, sticky="ns")
        barra_horizontal.grid(row=1, column=0, sticky="ew")
        marco_tabla.rowconfigure(0, weight=1)
        marco_tabla.columnconfigure(0, weight=1)
        self.tabla.bind("<Double-1>", self._iniciar_edicion)
        self._filtrar()

        acciones = ttk.Frame(self.ventana)
        acciones.pack(fill="x", padx=12, pady=10)
        ttk.Button(acciones, text="Cerrar", command=self._cerrar).pack(side="right")
        ttk.Button(
            acciones,
            text="Guardar cambios",
            command=self._guardar,
        ).pack(side="right", padx=(0, 8))
        self.ventana.protocol("WM_DELETE_WINDOW", self._cerrar)

    def _filtrar(self) -> None:
        consulta = self.busqueda.get().casefold().strip()
        for item in self.tabla.get_children():
            self.tabla.delete(item)
        for numero_fila, valores in self.filas:
            if consulta and consulta not in " ".join(valores).casefold():
                continue
            self.tabla.insert("", "end", iid=str(numero_fila), values=valores)

    def _iniciar_edicion(self, evento: tk.Event) -> None:
        item = self.tabla.identify_row(evento.y)
        columna = self.tabla.identify_column(evento.x)
        if not item or not columna:
            return
        caja = self.tabla.bbox(item, columna)
        if not caja:
            return
        if self.editor is not None:
            self._confirmar_edicion()
        numero_columna = int(columna[1:]) - 1
        valores = list(self.tabla.item(item, "values"))
        self.celda_editada = (item, columna)
        self.editor = ttk.Entry(self.tabla)
        self.editor.insert(0, valores[numero_columna])
        self.editor.select_range(0, "end")
        self.editor.place(x=caja[0], y=caja[1], width=caja[2], height=caja[3])
        self.editor.focus_set()
        self.editor.bind("<Return>", self._confirmar_edicion)
        self.editor.bind("<Escape>", self._cancelar_edicion)
        self.editor.bind("<FocusOut>", self._confirmar_edicion)

    def _confirmar_edicion(self, _evento: tk.Event | None = None) -> None:
        if self.editor is None or self.celda_editada is None:
            return
        item, columna = self.celda_editada
        numero_columna = int(columna[1:]) - 1
        valores = list(self.tabla.item(item, "values"))
        valores[numero_columna] = self.editor.get()
        self.tabla.item(item, values=valores)
        numero_fila = int(item)
        valores_fila = tuple(str(valor) for valor in valores)
        self.cambios[numero_fila] = valores_fila
        self.filas = [
            (fila, valores_fila if fila == numero_fila else existentes)
            for fila, existentes in self.filas
        ]
        editor = self.editor
        self.editor = None
        self.celda_editada = None
        editor.destroy()

    def _cancelar_edicion(self, _evento: tk.Event | None = None) -> None:
        editor = self.editor
        self.editor = None
        self.celda_editada = None
        if editor is not None:
            editor.destroy()

    def _guardar(self) -> bool:
        self._confirmar_edicion()
        filas_actualizadas = len(self.cambios)
        try:
            _guardar_cambios_maestro(self.ruta, self.cambios)
            if filas_actualizadas:
                self.filas = _leer_filas_maestro_editable(self.ruta)
        except (OSError, ValueError) as exc:
            messagebox.showerror(
                "No se pudieron guardar los cambios",
                str(exc),
                parent=self.ventana,
            )
            return False
        if filas_actualizadas:
            messagebox.showinfo(
                "Lista maestra actualizada",
                f"Se guardaron cambios en {filas_actualizadas} fila(s).",
                parent=self.ventana,
            )
            self.cambios.clear()
            self._filtrar()
            if self.al_guardar is not None:
                self.al_guardar()
        return True

    def _cerrar(self) -> None:
        if self.editor is not None:
            self._confirmar_edicion()
        if self.cambios:
            decision = messagebox.askyesnocancel(
                "Cambios sin guardar",
                "¿Quieres guardar los cambios de la lista maestra antes de cerrar?",
                parent=self.ventana,
            )
            if decision is None:
                return
            if decision and not self._guardar():
                return
        self.ventana.destroy()


class EditorTraducciones:
    def __init__(self, raiz, al_guardar=None):
        self.al_guardar = al_guardar
        self.filas = copy.deepcopy(SERVICIO_TRADUCTOR.estado["registros"])
        self.huella = SERVICIO_TRADUCTOR.estado["sha256_excel"]
        self._carpeta_traducciones = SERVICIO_TRADUCTOR.carpeta.resolve()
        self.modificado = False
        self.guardando = False
        self.cola = queue.Queue()
        self.ventana = ui.Ventana(raiz)
        self.ventana.title("Editar lista de traducciones")
        self.ventana.geometry("1050x700")
        self.ventana.transient(raiz)
        self.ventana.grab_set()
        ttk.Label(self.ventana, text="Lista de traducciones", font=("Segoe UI", 16, "bold")).pack(anchor="w", padx=14, pady=12)
        ttk.Label(self.ventana, text="Selecciona una fila para editarla. La primera variante española define la traducción inversa.").pack(anchor="w", padx=14)
        self.busqueda = tk.StringVar()
        entrada = ttk.Entry(self.ventana, textvariable=self.busqueda)
        entrada.pack(fill="x", padx=14, pady=8)
        entrada.bind("<KeyRelease>", lambda _e: self._mostrar())
        marco = ttk.Frame(self.ventana)
        marco.pack(fill="both", expand=True, padx=14)
        self.tabla = ttk.Treeview(marco, columns=("es", "en", "activo", "observaciones"), show="headings", selectmode="browse")
        for campo, etiqueta, ancho in (("es", "ESPAÑOL", 260), ("en", "INGLÉS", 260), ("activo", "ACTIVO", 80), ("observaciones", "OBSERVACIONES", 350)):
            self.tabla.heading(campo, text=etiqueta)
            self.tabla.column(campo, width=ancho)
        barra = ttk.Scrollbar(marco, orient="vertical", command=self.tabla.yview)
        self.tabla.configure(yscrollcommand=barra.set)
        self.tabla.pack(side="left", fill="both", expand=True)
        barra.pack(side="right", fill="y")
        self.tabla.bind("<<TreeviewSelect>>", self._seleccionar)
        formulario = ttk.Frame(self.ventana, padding=14)
        formulario.pack(fill="x")
        self.variables = {c: tk.StringVar(value="SI" if c == "activo" else "") for c in ("es", "en", "activo", "observaciones")}
        for columna, (campo, etiqueta) in enumerate(zip(self.variables, ("Español", "Inglés", "Activo", "Observaciones"))):
            ttk.Label(formulario, text=etiqueta).grid(row=0, column=columna, sticky="w")
            control = ttk.Combobox(formulario, textvariable=self.variables[campo], values=("SI", "NO"), state="readonly", width=6) if campo == "activo" else ttk.Entry(formulario, textvariable=self.variables[campo])
            control.grid(row=1, column=columna, sticky="ew", padx=(0, 8))
            formulario.columnconfigure(columna, weight=0 if campo == "activo" else 1)
        acciones = ttk.Frame(self.ventana, padding=14)
        acciones.pack(fill="x")
        for texto, funcion in (("Agregar fila", lambda: self._aplicar(nueva=True)), ("Aplicar a fila seleccionada", self._aplicar), ("Eliminar fila", self._eliminar)):
            ttk.Button(acciones, text=texto, command=funcion).pack(side="left", padx=(0, 6))
        self.boton_guardar = ttk.Button(acciones, text="Guardar cambios", command=self._guardar)
        self.boton_guardar.pack(side="right")
        ttk.Button(acciones, text="Cerrar", command=self._cerrar, variante="secundario").pack(side="right", padx=6)
        self.ventana.protocol("WM_DELETE_WINDOW", self._cerrar)
        self._mostrar()

    def _mostrar(self):
        self.tabla.delete(*self.tabla.get_children())
        consulta = self.busqueda.get().casefold().strip()
        for indice, fila in enumerate(self.filas):
            valores = (fila["es"], fila["en"], "SI" if fila["activo"] else "NO", fila.get("observaciones", ""))
            if not consulta or consulta in " ".join(valores).casefold():
                self.tabla.insert("", "end", iid=str(indice), values=valores)

    def _seleccionar(self, _evento=None):
        seleccion = self.tabla.selection()
        if seleccion:
            fila = self.filas[int(seleccion[0])]
            for campo, variable in self.variables.items():
                variable.set(("SI" if fila["activo"] else "NO") if campo == "activo" else fila.get(campo, ""))

    def _aplicar(self, nueva=False):
        if self.guardando:
            return False
        seleccion = self.tabla.selection()
        if not nueva and not seleccion:
            return False
        fila = {c: v.get().strip() for c, v in self.variables.items()}
        fila["activo"] = fila["activo"] == "SI"
        candidatas = copy.deepcopy(self.filas)
        if nueva:
            candidatas.append(fila)
        else:
            candidatas[int(seleccion[0])] = fila
        try:
            validar_registros_traducciones(candidatas)
        except ValueError as exc:
            messagebox.showerror("Revisa la traducción", str(exc), parent=self.ventana)
            return False
        self.filas = candidatas
        self.modificado = True
        self._mostrar()
        return True

    def _eliminar(self):
        seleccion = self.tabla.selection()
        if not self.guardando and seleccion and messagebox.askyesno("Eliminar traducción", "¿Eliminar la fila seleccionada?", parent=self.ventana):
            del self.filas[int(seleccion[0])]
            self.modificado = True
            self._mostrar()

    def _guardar(self):
        if self.guardando:
            return
        if self.tabla.selection() and not self._aplicar():
            return
        self.guardando = True
        self.boton_guardar.configure(state="disabled", text="Guardando…")
        filas = copy.deepcopy(self.filas)
        huella = self.huella
        def trabajo():
            try:
                with SERVICIO_TRADUCTOR.lock:
                    if SERVICIO_TRADUCTOR.carpeta.resolve() != self._carpeta_traducciones:
                        raise ValueError("La carpeta de las listas cambió. Vuelve a abrir el editor de traducciones antes de guardar.")
                    self.cola.put((True, SERVICIO_TRADUCTOR.guardar_registros(filas, huella)))
            except Exception as exc:
                self.cola.put((False, str(exc)))
        threading.Thread(target=trabajo, daemon=True).start()
        self.ventana.after(100, self._resultado)

    def _resultado(self):
        try:
            correcto, resultado = self.cola.get_nowait()
        except queue.Empty:
            self.ventana.after(100, self._resultado)
            return
        self.guardando = False
        self.boton_guardar.configure(state="normal", text="Guardar cambios")
        if correcto:
            self.huella = resultado["sha256_excel"]
            self.modificado = False
            if self.al_guardar is not None:
                self.al_guardar()
            for campo, variable in self.variables.items():
                variable.set("SI" if campo == "activo" else "")
            messagebox.showinfo("Traducciones guardadas", f"{resultado['cambios']} cambios registrados. Los editores abiertos recibirán las nuevas equivalencias.", parent=self.ventana)
        else:
            messagebox.showerror("No se guardaron las traducciones", resultado, parent=self.ventana)

    def _cerrar(self):
        if self.guardando:
            return
        formulario = {c: v.get().strip() for c, v in self.variables.items()}
        formulario["activo"] = formulario["activo"] == "SI"
        seleccion = self.tabla.selection()
        pendiente = formulario != self.filas[int(seleccion[0])] if seleccion else any(formulario[c] for c in ("es", "en", "observaciones"))
        if not (self.modificado or pendiente) or messagebox.askyesno("Cambios sin guardar", "¿Cerrar y descartar los cambios de la lista?", parent=self.ventana):
            self.ventana.destroy()


class GeneradorEspecificacion:
    def __init__(
        self,
        app: "AplicacionEspecificaciones",
        archivo_inicial: Path | None = None,
        destino_inicial: Path | None = None,
        idioma_destino: str | None = None,
        al_guardar: Callable[[Path], None] | None = None,
    ):
        self.app = app
        self.archivo_inicial = resolver_ruta(archivo_inicial) if archivo_inicial is not None else None
        self.destino_inicial = resolver_ruta(destino_inicial) if destino_inicial is not None else None
        self.al_guardar = al_guardar
        self.idioma_destino = idioma_destino
        self.datos_originales: dict[str, Any] | None = None
        self._traductor = SERVICIO_TRADUCTOR.diccionario
        self._cola_traductor = queue.Queue()
        self._cargando_traductor = False
        self._bilingue = None
        self._rutas_pareja_descartadas = set()
        self._originales_idioma = {}
        self._cambios_alineacion = {}
        self._rutas_idioma = {}
        self._marcas_pendientes = []
        self._redibujado_pendiente = None
        self.maestro = cargar_maestro(ruta_maestro(BASE))
        self.codigos_part_id = _cargar_catalogo_part_ids()
        self._editor_celda: ttk.Entry | ttk.Combobox | ttk.Frame | None = None
        self._celda_editada: tuple[str, str] | None = None
        self._nombres_sugeridos_por_fila: dict[str, tuple[str, ...]] = {}
        self._sugerencia_actual: tuple[str, str, str] | None = None
        self._codigo_validacion_activada = False
        self._celda_activa: tuple[str, str] | None = None
        self._periodicidad_editores: list[ttk.Combobox] = []
        self._periodicidad_marco: ttk.Frame | None = None
        self._periodicidad_boton_mas: ttk.Button | None = None
        self._comentario_marco: ttk.Frame | None = None
        self._comentario_boton_mas: ttk.Button | None = None
        self._popup_sugerencias: tk.Toplevel | None = None
        self._lista_sugerencias: tk.Listbox | None = None
        self._entrada_sugerencias: ttk.Combobox | None = None
        self._al_seleccionar_sugerencia: Any = None
        self._idioma_anterior = "es"
        self._nota_predeterminada_actual = ""
        self._elaboro_aprobo_predeterminado_actual = VALORES_ELABORO_APROBO["es"]
        self.ventana = ui.Ventana(app.raiz)
        self.ventana.title("Generar especificación")
        self.ventana.geometry("1050x760")
        self.ventana.minsize(760, 560)
        self.pruebas: list[dict[str, str]] = []
        self.variables_control = {
            clave: tk.StringVar(value=valor)
            for clave, valor in _valores_control_predeterminados().items()
        }
        self._etiquetas: dict[str, ttk.Label] = {}
        self._idioma = tk.StringVar(value="Español")
        self.desarrollador = tk.StringVar()
        self.fecha_reanalisis = tk.StringVar()

        encabezado = ttk.Frame(self.ventana, padding=(12, 4, 12, 2))
        encabezado.pack(fill="x")
        self.etiqueta_idioma = ttk.Label(encabezado)
        self.etiqueta_idioma.pack(side="left", padx=(0, 8))
        self.selector_idioma = ttk.Combobox(
            encabezado,
            textvariable=self._idioma,
            values=("Español", "English"),
            state="readonly",
            width=12,
        )
        self.selector_idioma.pack(side="left")
        self.selector_idioma.bind("<<ComboboxSelected>>", self._actualizar_idioma)
        self.aviso_pareja = ttk.Label(encabezado, foreground="#9a3412", wraplength=850, justify="left")

        identidad = ttk.Frame(self.ventana, padding=(12, 0, 12, 2))
        identidad.pack(fill="x")
        self.etiqueta_producto = ttk.Label(identidad)
        self.etiqueta_producto.grid(row=0, column=0, sticky="w", padx=(0, 8), pady=3)
        self.nombre_producto = tk.StringVar()
        self.entrada_nombre_producto = ttk.Combobox(
            identidad,
            textvariable=self.nombre_producto,
        )
        self.nombre_producto.trace_add(
            "write", self._sincronizar_codigo_desde_nombre
        )
        self.entrada_nombre_producto.grid(
            row=0, column=1, sticky="ew", pady=3
        )
        self.entrada_nombre_producto.bind(
            "<KeyRelease>", self._filtrar_nombres_producto
        )
        self.entrada_nombre_producto.bind(
            "<<ComboboxSelected>>", self._seleccionar_nombre_producto
        )
        self.entrada_nombre_producto.bind("<Down>", self._enfocar_sugerencias)
        identidad.columnconfigure(1, weight=1)
        identidad.columnconfigure(3, weight=1)

        titulo_control = self.encabezado_control = ttk.Frame(self.ventana)
        titulo_control.pack(fill="x", padx=12, pady=(0, 3))
        self.etiqueta_titulo_control = ttk.Label(titulo_control)
        self.etiqueta_titulo_control.pack(side="left")
        self.boton_control = ttk.Button(
            titulo_control, text="−", width=2, command=self._alternar_control_documental
        )
        self.boton_control.pack(side="left", padx=(6, 0))
        self.marco_control = ttk.Frame(self.ventana, padding=(10, 3), style="Borde.TFrame")
        self.marco_control.pack(fill="x", padx=12, pady=(0, 3))
        self._control_colapsado = False
        self._campos_control_ocultos: list[tk.Widget] = []
        self.entradas_control: dict[str, ttk.Entry | ttk.Combobox] = {}
        self.etiqueta_codigo_asterisco: ttk.Label | None = None
        self.etiqueta_codigo_invalido: ttk.Label | None = None
        for indice, (clave, _, _) in enumerate(CAMPOS_CONTROL_GENERADOR):
            fila, grupo = divmod(indice, 2)
            columna = grupo * 2
            padre_campo = identidad if clave == "codigo_unico_producto" else self.marco_control
            etiqueta = ttk.Label(padre_campo)
            etiqueta.grid(row=0 if clave == "codigo_unico_producto" else fila,
                          column=2 if clave == "codigo_unico_producto" else columna,
                          sticky="w", padx=(12 if clave == "codigo_unico_producto" else 0, 6), pady=3)
            self._etiquetas[clave] = etiqueta
            if clave == "codigo_unico_producto":
                contenedor_codigo = ttk.Frame(identidad)
                contenedor_codigo.grid(
                    row=0,
                    column=3,
                    sticky="ew",
                    padx=(0, 12 if grupo == 0 else 0),
                    pady=3,
                )
                entrada = ttk.Combobox(
                    contenedor_codigo,
                    textvariable=self.variables_control[clave],
                    width=24,
                )
                entrada.grid(row=0, column=0, sticky="ew")
                self.etiqueta_codigo_asterisco = ttk.Label(
                    contenedor_codigo,
                    text="*",
                    foreground="#c00000",
                )
                self.etiqueta_codigo_asterisco.grid(
                    row=0, column=1, sticky="w", padx=(4, 0)
                )
                contenedor_codigo.columnconfigure(0, weight=1)
                self.etiqueta_codigo_invalido = ttk.Label(
                    identidad,
                    foreground="#c00000",
                    wraplength=330,
                )
                self.etiqueta_codigo_invalido.grid(
                    row=1,
                    column=3,
                    sticky="w",
                    pady=(2, 0),
                )
                entrada.bind("<KeyRelease>", self._filtrar_part_ids)
                entrada.bind("<Down>", self._enfocar_sugerencias)
                entrada.bind(
                    "<<ComboboxSelected>>",
                    self._seleccionar_part_id,
                )
            elif clave == "revision":
                marco_revision = ttk.Frame(self.marco_control)
                marco_revision.grid(row=fila, column=columna + 1, sticky="ew", pady=3)
                entrada = ttk.Entry(marco_revision, textvariable=self.variables_control[clave], width=6)
                entrada.pack(side="left")
                self.etiqueta_reanalisis = ttk.Label(marco_revision)
                self.etiqueta_reanalisis.pack(side="left", padx=(12, 6))
                self.entrada_reanalisis = ttk.Entry(marco_revision, textvariable=self.fecha_reanalisis, width=5)
                self.entrada_reanalisis.pack(side="left")
                self.unidad_reanalisis = ttk.Label(marco_revision)
                self.unidad_reanalisis.pack(side="left", padx=(6, 0))
            else:
                entrada = ttk.Entry(
                    self.marco_control,
                    textvariable=self.variables_control[clave],
                    width=24,
                )
                entrada.grid(
                    row=fila,
                    column=columna + 1,
                    sticky="ew",
                    padx=(0, 12 if grupo == 0 else 0),
                    pady=3,
                )
            self.entradas_control[clave] = entrada
        self.marco_control.columnconfigure(1, weight=1)
        self.marco_control.columnconfigure(3, weight=1)
        self.etiqueta_desarrollador = ttk.Label(self.marco_control)
        self.etiqueta_desarrollador.grid(
            row=1,
            column=2,
            sticky="w",
            padx=(0, 6),
            pady=3,
        )
        self.entrada_desarrollador = ttk.Entry(
            self.marco_control,
            textvariable=self.desarrollador,
        )
        self.entrada_desarrollador.grid(
            row=1,
            column=3,
            sticky="ew",
            pady=3,
        )

        self.marco_pruebas = ttk.LabelFrame(self.ventana, padding=(8, 3))
        self.marco_pruebas.pack(fill="both", expand=True, padx=12, pady=(0, 4))
        self.tabla_pruebas = ttk.Treeview(
            self.marco_pruebas,
            columns=(*CAMPOS_TABLA, "periodicidad"),
            show="headings",
            selectmode="browse",
            height=8,
            style="Generador.Treeview",
        )
        estilo = ttk.Style(self.ventana)
        estilo.configure(
            "Generador.Treeview",
            bordercolor="#e1e5e9",
            darkcolor="#e1e5e9",
            lightcolor="#e1e5e9",
            borderwidth=1,
            relief="solid",
            rowheight=25,
            fieldbackground="#ffffff",
        )
        estilo.configure(
            "Generador.Treeview.Heading",
            bordercolor="#e1e5e9",
            borderwidth=1,
            relief="solid",
        )
        self.tabla_pruebas.tag_configure("fila_alterna", background="#f7f8fa")
        anchos = {
            "prueba": 250,
            "unidades": 100,
            "metodo": 160,
            "especificacion": 180,
            "valor_tipico": 130,
            "periodicidad": 245,
        }
        for campo in (*CAMPOS_TABLA, "periodicidad"):
            self.tabla_pruebas.column(
                campo,
                width=anchos[campo],
                minwidth=75,
                stretch=True,
            )
        self._ordenador_pruebas = OrdenadorColumnasTreeview(
            self.tabla_pruebas,
            {
                campo: encabezado
                for campo, encabezado in zip(
                    CAMPOS_TABLA, ETIQUETAS_TABLA_GENERADOR["es"]
                )
                if campo in {"prueba", "metodo"}
            },
        )
        self.tabla_pruebas.column(
            "periodicidad", width=245, minwidth=230, stretch=False
        )
        self.tabla_pruebas.bind("<Button-1>", self._editar_celda)
        self.tabla_pruebas.bind("<Configure>", self._reposicionar_editor)
        self.tabla_pruebas.bind("<Delete>", self._borrar_celda_activa)
        self.tabla_pruebas.bind("<KP_Delete>", self._borrar_celda_activa)
        barra_vertical = ttk.Scrollbar(
            self.marco_pruebas,
            orient="vertical",
            command=self.tabla_pruebas.yview,
        )
        barra_horizontal = ttk.Scrollbar(
            self.marco_pruebas,
            orient="horizontal",
            command=self.tabla_pruebas.xview,
        )
        self.tabla_pruebas.configure(
            yscrollcommand=lambda *args: (barra_vertical.set(*args), self._programar_resaltado()),
            xscrollcommand=lambda *args: (barra_horizontal.set(*args), self._programar_resaltado()),
        )
        self.tabla_pruebas.bind("<Configure>", lambda _e: self._programar_resaltado(), add="+")
        self.tabla_pruebas.grid(row=0, column=0, sticky="nsew")
        barra_vertical.grid(row=0, column=1, sticky="ns")
        barra_horizontal.grid(row=1, column=0, sticky="ew")
        self.marco_pruebas.rowconfigure(0, weight=1)
        self.marco_pruebas.columnconfigure(0, weight=1)
        self.etiqueta_sugerencia = ttk.Label(
            self.marco_pruebas,
            foreground="#78818c",
            cursor="hand2",
        )
        self.etiqueta_sugerencia.grid(
            row=2, column=0, columnspan=2, sticky="w", pady=(5, 0)
        )
        self.etiqueta_sugerencia.bind(
            "<Button-1>", self._aceptar_sugerencia
        )
        acciones_pruebas = ttk.Frame(self.marco_pruebas)
        acciones_pruebas.grid(row=3, column=0, columnspan=2, sticky="ew", pady=(3, 0))
        self.boton_ayuda_tabla = ttk.Button(
            acciones_pruebas,
            text="ⓘ",
            width=3,
            command=self._mostrar_ayuda_tabla,
        )
        self.boton_ayuda_tabla.pack(side="left")
        self.boton_eliminar = ttk.Button(
            acciones_pruebas,
            command=self._eliminar_seleccionada,
        )
        self.boton_eliminar.pack(side="right")
        self.boton_confirmar_traduccion = ttk.Button(
            acciones_pruebas, text="Confirmar traducción de celda",
            command=self._confirmar_celda_bilingue,
        )
        self.referencia_traduccion = ttk.Label(self.marco_pruebas, wraplength=900)
        self.referencia_traduccion.grid(row=4, column=0, columnspan=2, sticky="w")
        self.tabla_pruebas.bind("<<TreeviewSelect>>", self._actualizar_boton_eliminar)
        self.tabla_pruebas.bind("<<TreeviewSelect>>", lambda _e: self._programar_resaltado(), add="+")

        marco_notas = ttk.Frame(self.ventana, padding=8)
        marco_notas.pack(side="bottom", before=self.marco_pruebas, fill="x", padx=12, pady=(0, 8))
        self.etiqueta_notas = ttk.Label(marco_notas)
        self.etiqueta_notas.pack(anchor="w")
        from editor_comportamientos import NotasProtegidas
        self.notas = NotasProtegidas(marco_notas, height=3, wrap="word")
        self.notas.pack(fill="x", pady=(4, 0))
        self.boton_confirmar_notas = ttk.Button(
            marco_notas, text="Confirmar traducción de notas",
            command=lambda: self._confirmar_cabecera_bilingue("notas"),
        )
        self.referencia_notas = ttk.Label(marco_notas, wraplength=900)

        pie = ttk.Frame(self.ventana, padding=(12, 0, 12, 4))
        pie.pack(side="bottom", before=marco_notas, fill="x")
        self.boton_guardar = ttk.Button(pie, command=self.guardar)
        self.boton_guardar.pack(side="right")
        self.boton_guardar_wip = ttk.Button(pie, text="Guardar en progreso", command=self.guardar_en_progreso, variante="secundario")
        self.boton_cancelar = ttk.Button(pie, command=self.ventana.destroy)
        self.boton_cancelar.pack(side="right", padx=(0, 8))
        self.boton_alinear = ttk.Button(
            pie,
            command=self._alinear_especificaciones,
        )
        self.boton_alinear.pack(side="right", padx=(0, 8))
        self._actualizar_idioma()
        self._actualizar_tabla()
        self._actualizar_validacion_codigo()
        if archivo_inicial is not None:
            self._cargar_especificacion(archivo_inicial)
        if idioma_destino is not None:
            self._idioma.set("English" if idioma_destino == "en" else "Español")
            self._actualizar_idioma()
            self.traduccion_revisada = tk.BooleanVar(value=False)
            ttk.Checkbutton(self.ventana,
                text="He revisado y completado la traducción de pruebas, especificaciones, notas y control documental",
                variable=self.traduccion_revisada).pack(fill="x", padx=12, pady=6)
            self.ventana.title("Completar traducción de especificación")

        self._contenido_inicial_revision = self._contenido_para_revision()
        self._revision_inicial = self.variables_control["revision"].get()
        self._revision_aplicada = False
        self._cambio_alineacion_revision = False
        self._resultados_revision = queue.Queue()
        self._revisiones_historial = {}
        self._consultas_revision = set()
        self.ventana.after(250, self._vigilar_revision)
        poco_espacio = (self.ventana.winfo_screenheight() - 80) / self.ventana._get_window_scaling() <= 650
        if ui.PREFERENCIAS["control_compacto"] or (ui.PREFERENCIAS["adaptacion_auto"] and poco_espacio):
            self._alternar_control_documental()
        if idioma_destino is None and (self.datos_originales is None or self.datos_originales.get("alineacion_metodos_lab")):
            self._iniciar_bilingue()
            self._contenido_inicial_revision = self._contenido_para_revision()
        self._actualizar_traductor_editor(avisar=False)

    def _actualizar_traductor_editor(self, avisar=True):
        if self._cargando_traductor:
            return
        self._cargando_traductor = True
        cola = self._cola_traductor
        def cargar():
            try:
                resultado = actualizar_traductor()
                cola.put((True, resultado, SERVICIO_TRADUCTOR.diccionario, avisar))
            except Exception as exc:
                cola.put((False, str(exc), None, avisar))
        threading.Thread(target=cargar, daemon=True).start()
        self.ventana.after(100, self._recibir_actualizacion_traductor)

    def _recibir_actualizacion_traductor(self):
        if not self.ventana.winfo_exists():
            return
        try:
            correcto, resultado, diccionario, avisar = self._cola_traductor.get_nowait()
        except queue.Empty:
            self.ventana.after(100, self._recibir_actualizacion_traductor)
            return
        self._cargando_traductor = False
        if not correcto:
            messagebox.showwarning("No se actualizaron las traducciones", f"Se conserva la última versión válida.\n\n{resultado}", parent=self.ventana)
            return
        self._adoptar_diccionario_traductor(diccionario)
        if avisar:
            messagebox.showinfo("Traducciones actualizadas", f"{resultado['equivalencias']} equivalencias activas.\n{resultado['cambios']} cambios registrados.", parent=self.ventana)

    def _adoptar_diccionario_traductor(self, diccionario):
        if self._traductor is not diccionario:
            self._cerrar_editor_celda(guardar=True)
            if self._bilingue is not None:
                self._capturar_bilingue(self._idioma_anterior)
            self._traductor = diccionario
            if self._bilingue is not None:
                seleccion = self._seleccion()
                self._bilingue.actualizar_traducciones_automaticas()
                self._aplicar_borrador(self._idioma_anterior)
                self._actualizar_tabla()
                if seleccion is not None and seleccion < len(self.pruebas):
                    self.tabla_pruebas.selection_set(str(seleccion))
                    self.tabla_pruebas.see(str(seleccion))

    def _contenido_bilingue(self, idioma):
        return {
            "idioma": idioma, "nombre_producto": self.nombre_producto.get(),
            "control": {c: v.get() for c, v in self.variables_control.items()},
            "pruebas": copy.deepcopy(self.pruebas),
            "notas": self.notas.get("1.0", "end-1c").splitlines(),
            "desarrollador": self.desarrollador.get(),
            "fecha_reanalisis": self.fecha_reanalisis.get(),
        }

    def _traducir_campo_bilingue(self, campo, texto, origen, destino, fila):
        diccionario = getattr(self, "_traductor", SERVICIO_TRADUCTOR.diccionario)
        if campo == "notas":
            lineas = [(NOTAS_PLANTILLA[destino], True) if linea.strip() == NOTAS_PLANTILLA[origen].strip()
                      else diccionario.traducir(linea, origen, destino) for linea in texto.splitlines()]
            return "\n".join(t for t, _ in lineas), all(seguro for _, seguro in lineas)
        if campo == "elaboro_y_aprobo":
            if texto == VALORES_ELABORO_APROBO[origen]:
                return VALORES_ELABORO_APROBO[destino], True
            # Nombres de personas no requieren traducción.
            return texto, True
        if campo in {"prueba", "unidades", "metodo"}:
            registros = [r for r in self.maestro
                         if normalizar(_nombre_maestro_para_idioma(r, origen)) == normalizar(fila.get("prueba", ""))]
            referencias = set(referencias_astm(fila.get("metodo", "")))
            if referencias:
                registros = [r for r in registros if any(set(referencias_astm(m)) == referencias for m in metodos_maestros(r))]
            alternativas = set()
            for registro in registros:
                if campo == "prueba":
                    fuente = _nombre_maestro_para_idioma(registro, origen)
                    objetivo = nombres_canonicos(registro, "INGLES" if destino == "en" else "ESPANOL")[0]
                elif campo == "metodo":
                    fuente = metodo_maestro_para_idioma(registro, "INGLES" if origen == "en" else "ESPANOL")
                    objetivo = metodo_maestro_para_idioma(registro, "INGLES" if destino == "en" else "ESPANOL")
                    if destino == "en" and not (registro.get("MÉTODO INGLÉS") or registro.get("METODO INGLES")):
                        objetivo = ""
                else:
                    fuente = registro.get("UNITS" if origen == "en" else "UNIDADES", "")
                    objetivo = registro.get("UNITS" if destino == "en" else "UNIDADES", "")
                if normalizar(texto) == normalizar(fuente) and objetivo:
                    alternativas.add(objetivo)
            if len(alternativas) == 1:
                return alternativas.pop(), True
            if campo == "unidades" and (not texto or re.fullmatch(r"(?:N/A|[-−°%\d\s.,/*·²³^μµ]+|(?:mPa|Pa|kg|mg|g|mm|cm|m|ml|mL|L|s|cSt|SUS|K|F|C|ppm)(?:[-\d\s.,/*·²³^°%]|kg|mg|g|mm|cm|m|ml|mL|L|s|K|F|C|mol)*)", texto)):
                return texto, True
        return diccionario.traducir(texto, origen, destino)

    def _iniciar_bilingue(self):
        idioma = self._idioma_anterior
        modelo = BorradoresBilingues(self._contenido_bilingue(idioma), self._traducir_campo_bilingue)
        otro = "en" if idioma == "es" else "es"
        self._rutas_pareja_descartadas = set()
        self._originales_idioma.pop(otro, None)
        self._rutas_idioma.pop(otro, None)
        if hasattr(self, "aviso_pareja"):
            self.aviso_pareja.pack_forget()
        if self.datos_originales is not None:
            self._originales_idioma[idioma] = self.datos_originales
            self._rutas_idioma[idioma] = Path(self.destino_inicial or self.archivo_inicial).resolve()
            relaciones = self.datos_originales.get("versiones_idioma", {})
            referencias = relaciones.get("INGLES" if otro == "en" else "ESPANOL", []) if isinstance(relaciones, dict) else []
            if es_wip(self.archivo_inicial):
                referencias = []  # Los IDs y pendientes se recuperan del estado WIP.
            rutas = set()
            try:
                if not isinstance(referencias, list):
                    raise ValueError("Las referencias de idioma no tienen un formato válido.")
                for referencia in referencias:
                    if isinstance(referencia, dict) and isinstance(referencia.get("archivo_json"), str):
                        rutas.add((Path(self.archivo_inicial).parent / referencia["archivo_json"]).resolve())
                if len(rutas) > 1:
                    raise ValueError("Hay varias traducciones relacionadas.")
                if rutas:
                    ruta = next(iter(rutas))
                    if not ruta.is_file():
                        raise ValueError(f"No existe el JSON relacionado:\n{ruta}")
                    datos = json.loads(ruta.read_text(encoding="utf-8-sig"))
                    if not isinstance(datos, dict):
                        raise ValueError("La raíz del JSON relacionado debe ser un objeto.")
                    contenido = _datos_generador_desde_json(datos)
                    _validar_pareja_editor(modelo.borradores[idioma], contenido)
                    modelo.cargar_pareja(contenido)
                    self._originales_idioma[otro] = datos
                    self._rutas_idioma[otro] = ruta
            except (OSError, ValueError) as exc:
                self._rutas_pareja_descartadas = rutas - {self._rutas_idioma[idioma]}
                if hasattr(self, "aviso_pareja"):
                    self.aviso_pareja.configure(text=f"Borrador {'ENG' if otro == 'en' else 'ESP'} nuevo: "
                                               f"no se cargó la versión vinculada. {exc}")
                    self.aviso_pareja.pack(side="left", padx=(12, 0), fill="x", expand=True)
        self._bilingue = modelo
        if self.archivo_inicial is not None and es_wip(self.archivo_inicial):
            try:
                self._rutas_idioma = restaurar_estado(modelo, self.archivo_inicial)
                self._originales_idioma = {i: json.loads(r.read_text(encoding="utf-8-sig")) for i, r in self._rutas_idioma.items()}
                for lengua, datos in self._originales_idioma.items():
                    modelo.borradores[lengua].setdefault("fecha_reanalisis", _valor_reanalisis_editor(datos.get("FECHA DE REANALISIS", "")))
                self._aplicar_borrador(idioma)
            except (OSError, ValueError, KeyError, TypeError) as exc:
                messagebox.showerror("No se pudo recuperar el trabajo en progreso", str(exc), parent=self.ventana)
                self.boton_guardar.configure(state="disabled")
                return
        self.pruebas = copy.deepcopy(modelo.borradores[idioma]["pruebas"])
        for lengua, datos in self._originales_idioma.items():
            visibles = modelo.borradores[lengua]["pruebas"]
            original = original_disponible(datos, self._rutas_idioma.get(lengua), self.app.rutas.get("json_anteriores"))
            cambios = cambios_documento(datos, visibles, original)
            self._cambios_alineacion[lengua] = {f["_id_bilingue"]: c for f, c in zip(visibles, cambios)}
        self.boton_guardar_wip.pack(side="right", padx=(0, 8))
        self.ventana.title(f"Editar producto · {self.nombre_producto.get()} · ESP / ENG")
        self._programar_resaltado()

    def _capturar_bilingue(self, idioma):
        contenido = self._bilingue.actualizar(idioma, self._contenido_bilingue(idioma))
        self.pruebas = contenido["pruebas"]

    def _aplicar_borrador(self, idioma):
        contenido = copy.deepcopy(self._bilingue.borradores[idioma])
        self.pruebas = contenido["pruebas"]
        self.nombre_producto.set(contenido["nombre_producto"])
        self.desarrollador.set(contenido["desarrollador"])
        self.fecha_reanalisis.set(contenido.get("fecha_reanalisis", ""))
        for campo, valor in contenido["control"].items():
            self.variables_control[campo].set(valor)
        self.notas.delete("1.0", tk.END)
        self.notas.insert("1.0", "\n".join(contenido["notas"]))
        self.datos_originales = self._originales_idioma.get(idioma)
        self.destino_inicial = self._rutas_idioma.get(idioma)

    def _programar_resaltado(self):
        if (self._bilingue is not None or self._cambios_alineacion) and self._redibujado_pendiente is None:
            self._redibujado_pendiente = self.ventana.after_idle(self._resaltar_pendientes)

    def _resaltar_pendientes(self):
        self._redibujado_pendiente = None
        for marca in self._marcas_pendientes:
            marca.destroy()
        self._marcas_pendientes.clear()
        idioma = self._idioma_anterior
        pendientes = self._bilingue.pendientes[idioma] if self._bilingue is not None else set()
        cambios = getattr(self, "_cambios_alineacion", {}).get(idioma, {})
        estilo = ttk.Style(self.tabla_pruebas)
        nombre_estilo = self.tabla_pruebas.cget("style") or "Treeview"
        colores = ui.paleta()
        fondo = estilo.lookup(nombre_estilo, "fieldbackground") or colores["panel"]
        fuente = estilo.lookup(nombre_estilo, "font") or (ui.PREFERENCIAS["fuente"], 10)
        seleccionadas = set(self.tabla_pruebas.selection())
        rojo = "#ff6b6b" if colores["panel"].lower() not in {"#ffffff", "white"} else "#c62828"
        for indice, fila in enumerate(self.pruebas):
            for columna, campo in enumerate(CAMPOS_TABLA, 1):
                pendiente = (fila.get("_id_bilingue"), campo) in pendientes
                cambio = campo in cambios.get(fila.get("_id_bilingue", str(indice)), set())
                if not (pendiente or cambio) or self._celda_editada == (str(indice), campo):
                    continue
                bbox = self.tabla_pruebas.bbox(str(indice), f"#{columna}")
                if not bbox:
                    continue
                x, y, ancho, alto = bbox
                # Evitar cubrir encabezados y barras al desplazar horizontalmente.
                izquierda = max(0, x)
                ancho_visible = min(x + ancho, self.tabla_pruebas.winfo_width()) - izquierda
                if ancho_visible <= 0:
                    continue
                seleccionada = str(indice) in seleccionadas
                fondo_celda = (estilo.lookup(nombre_estilo, "background", ("selected",)) or colores["acento"]) if seleccionada else (
                    self.tabla_pruebas.tag_configure("fila_alterna").get("background") or colores["fondo"]
                    if indice % 2 else fondo
                )
                texto = fila.get(campo, "")
                marca = tk.Label(self.tabla_pruebas, text=texto if texto.strip() else "*", anchor="w",
                                 background=fondo_celda, foreground=("#ffb3b3" if seleccionada else rojo) if pendiente else ("#bfdbfe" if seleccionada else "#3b82f6"),
                                 font=fuente, borderwidth=0, relief="flat", padx=2)
                marca.place(x=izquierda, y=y, width=ancho_visible, height=alto)
                marca.bind("<Button-1>", lambda e, i=str(indice), c=columna: self._abrir_celda_marcada(i, c))
                self._marcas_pendientes.append(marca)
        if pendientes:
            self.boton_confirmar_traduccion.pack(side="left", padx=8)
        else:
            self.boton_confirmar_traduccion.pack_forget()
        notas_pendientes = ("cabecera", "notas") in pendientes
        self.notas.configure(background=colores["panel"], foreground=rojo if notas_pendientes else colores["texto"])
        self.notas.tag_configure("nota_formato", foreground=rojo if notas_pendientes else (
            colores["texto"] if getattr(self.notas, "_desbloqueadas", False) else "#9aa0a6"
        ))
        if notas_pendientes:
            self.boton_confirmar_notas.pack(anchor="w", pady=3)
            otro = "en" if idioma == "es" else "es"
            self.referencia_notas.configure(text=f"{'ENG' if otro == 'en' else 'ESP'} · " + "\n".join(self._bilingue.borradores[otro]["notas"]))
            self.referencia_notas.pack(anchor="w")
        else:
            self.boton_confirmar_notas.pack_forget()
            self.referencia_notas.pack_forget()
        if self._celda_activa is not None:
            iid, campo = self._celda_activa
            if iid.isdigit() and int(iid) < len(self.pruebas):
                self._mostrar_referencia_bilingue(iid, campo)

    def _abrir_celda_marcada(self, iid, columna):
        for marca in self._marcas_pendientes:
            marca.place_forget()
        bbox = self.tabla_pruebas.bbox(iid, f"#{columna}")
        if bbox:
            self._editar_celda(type("EventoCelda", (), {"x": max(1, bbox[0]) + 2, "y": bbox[1] + 2, "iid": iid, "columna": f"#{columna}"})())
        self._programar_resaltado()

    def _mostrar_referencia_bilingue(self, iid, campo):
        if self._bilingue is None:
            return
        fila = self.pruebas[int(iid)]
        otro = "en" if self._idioma_anterior == "es" else "es"
        origen = next((f for f in self._bilingue.borradores[otro]["pruebas"] if f["_id_bilingue"] == fila.get("_id_bilingue")), {})
        self.referencia_traduccion.configure(text=f"{'ENG' if otro == 'en' else 'ESP'} · {origen.get(campo, '')}")

    def _confirmar_celda_bilingue(self):
        self._cerrar_editor_celda(guardar=True)
        if self._celda_activa is None:
            return
        iid, campo = self._celda_activa
        if not iid.isdigit() or int(iid) >= len(self.pruebas):
            return
        fila = self.pruebas[int(iid)]
        if fila.get(campo, "").strip():
            self._bilingue.confirmar(self._idioma_anterior, (fila["_id_bilingue"], campo))
            self._programar_resaltado()

    def _confirmar_cabecera_bilingue(self, campo):
        self._capturar_bilingue(self._idioma_anterior)
        if self._bilingue._texto(self._bilingue.borradores[self._idioma_anterior], campo).strip():
            self._bilingue.confirmar(self._idioma_anterior, ("cabecera", campo))
            self._programar_resaltado()

    def _contenido_para_revision(self):
        if self._bilingue is not None:
            self._capturar_bilingue(self._idioma_anterior)
            return tuple((idioma, tuple(tuple(sorted((k, v) for k, v in f.items() if k != "_id_bilingue")) for f in c["pruebas"]), tuple(c["notas"]),
                          c.get("desarrollador", ""), c.get("fecha_reanalisis", ""),
                          c["control"].get("nombre_cliente", ""), c["control"].get("revision_cliente", ""))
                         for idioma, c in sorted(self._bilingue.borradores.items()))
        return (tuple(tuple(sorted(fila.items())) for fila in self.pruebas),
                self.notas.get("1.0", "end-1c"), self.desarrollador.get(), self.fecha_reanalisis.get(),
                self.variables_control["nombre_cliente"].get(), self.variables_control["revision_cliente"].get())

    def _actualizar_revision_automatica(self) -> bool:
        if not hasattr(self, "_contenido_inicial_revision") or (self.datos_originales is None and not self._originales_idioma):
            return True
        codigo = self.variables_control["codigo_unico_producto"].get().strip()
        if not codigo:
            return True
        while not self._resultados_revision.empty():
            clave, valor = self._resultados_revision.get_nowait()
            self._revisiones_historial[clave] = valor
        if codigo not in self._consultas_revision:
            self._consultas_revision.add(codigo)
            carpetas = [self.app.rutas.get(clave, RUTAS_PREDETERMINADAS[clave])
                        for clave in ("json_anteriores", "json_alineadas")]
            revision = self._revision_inicial
            resultados = self._resultados_revision
            def consultar():
                from editor_comportamientos import siguiente_revision
                try:
                    valor = siguiente_revision(codigo, carpetas, BASE / "especificaciones.db", revision)
                except Exception as exc:
                    valor = exc
                resultados.put((codigo, valor))
            threading.Thread(target=consultar, daemon=True).start()
        cambiado = self._cambio_alineacion_revision or self._contenido_para_revision() != self._contenido_inicial_revision
        if not cambiado or self._revision_aplicada:
            return True
        valor = self._revisiones_historial.get(codigo)
        if valor is None or isinstance(valor, Exception):
            return False
        self.variables_control["revision"].set(valor)
        self.variables_control["fecha"].set(date.today().strftime("%d/%m/%y"))
        self._revision_aplicada = True
        return True

    def _vigilar_revision(self):
        if self.ventana.winfo_exists():
            self._adoptar_diccionario_traductor(SERVICIO_TRADUCTOR.diccionario)
            self._actualizar_revision_automatica()
            self.ventana.after(250, self._vigilar_revision)

    def _alternar_control_documental(self) -> None:
        self._ocultar_lista_sugerencias()
        if self._control_colapsado:
            self.marco_control.pack(after=self.encabezado_control, fill="x", padx=12, pady=(0, 3))
            self.boton_control.configure(text="−")
        else:
            self.marco_control.pack_forget()
            self.boton_control.configure(text="+")
        self._control_colapsado = not self._control_colapsado

    def _cargar_especificacion(self, ruta: Path) -> None:
        try:
            datos = json.loads(ruta.read_text(encoding="utf-8-sig"))
            if not isinstance(datos, dict):
                raise ValueError("La raíz del JSON debe ser un objeto.")
            contenido = _datos_generador_desde_json(datos)
        except (OSError, UnicodeError, json.JSONDecodeError, ValueError) as exc:
            messagebox.showerror(
                "No se pudo abrir la especificación",
                str(exc),
                parent=self.ventana,
            )
            self.ventana.destroy()
            return

        self.datos_originales = datos
        self.nombre_producto.set(contenido["nombre_producto"])
        idioma = contenido["idioma"]
        self._idioma.set("English" if idioma == "en" else "Español")
        self._idioma_anterior = idioma
        self._actualizar_idioma()
        for clave, valor in contenido["control"].items():
            self.variables_control[clave].set(valor)
        self.desarrollador.set(contenido["desarrollador"])
        self.fecha_reanalisis.set(contenido["fecha_reanalisis"])
        self.pruebas = contenido["pruebas"]
        self.notas.delete("1.0", tk.END)
        self.notas.insert("1.0", "\n".join(contenido["notas"]))
        self._nota_predeterminada_actual = "\n".join(contenido["notas"])
        self._actualizar_tabla()
        self._actualizar_validacion_codigo()
        self.ventana.title(
            f"Editar especificación · {ruta.name}"
        )

    def _filtrar_part_ids(self, _evento: tk.Event | None = None) -> None:
        entrada = self.entradas_control["codigo_unico_producto"]
        if not isinstance(entrada, ttk.Combobox):
            return
        self._codigo_validacion_activada = True
        consulta = self.variables_control["codigo_unico_producto"].get()
        sugerencias = _sugerencias_part_id(self.codigos_part_id, consulta)
        entrada.configure(values=sugerencias)
        self._mostrar_lista_sugerencias(
            entrada,
            sugerencias,
            self._elegir_part_id,
        )
        self._actualizar_validacion_codigo()

    def _filtrar_nombres_producto(self, _evento: tk.Event | None = None) -> None:
        consulta = self.nombre_producto.get()
        nombres = _sugerencias_nombre_producto(self.codigos_part_id, consulta)
        self.entrada_nombre_producto.configure(values=nombres)
        self._mostrar_lista_sugerencias(
            self.entrada_nombre_producto,
            nombres,
            self._elegir_nombre_producto,
        )
        codigo = _part_id_por_nombre(consulta, self.codigos_part_id)
        if codigo:
            self.variables_control["codigo_unico_producto"].set(codigo)
            self._actualizar_validacion_codigo()
        else:
            codigo_actual = self.variables_control["codigo_unico_producto"].get()
            registro_actual = next(
                (
                    registro
                    for registro in self.codigos_part_id
                    if registro.get("part_id", "").casefold()
                    == codigo_actual.strip().casefold()
                ),
                None,
            )
            if registro_actual and normalizar(
                registro_actual.get("descripcion", "")
            ) != normalizar(consulta):
                self.variables_control["codigo_unico_producto"].set("")
                self._actualizar_validacion_codigo()

    def _sincronizar_codigo_desde_nombre(
        self, *_args: str
    ) -> None:
        codigo = _part_id_por_nombre(
            self.nombre_producto.get(), self.codigos_part_id
        )
        if codigo:
            self.variables_control["codigo_unico_producto"].set(codigo)

    def _seleccionar_part_id(self, _evento: tk.Event | None = None) -> None:
        self._elegir_part_id(
            self.variables_control["codigo_unico_producto"].get()
        )

    def _elegir_part_id(self, part_id: str) -> None:
        registro = next(
            (
                item
                for item in self.codigos_part_id
                if item.get("part_id", "").casefold() == part_id.casefold()
            ),
            None,
        )
        if registro is None:
            return
        self.variables_control["codigo_unico_producto"].set(registro["part_id"])
        self.nombre_producto.set(registro["descripcion"])
        self.entrada_nombre_producto.configure(values=())
        self._ocultar_lista_sugerencias()
        self._actualizar_validacion_codigo()

    def _seleccionar_nombre_producto(self, _evento: tk.Event | None = None) -> None:
        self._elegir_nombre_producto(self.nombre_producto.get())

    def _elegir_nombre_producto(self, descripcion: str) -> None:
        self.nombre_producto.set(descripcion)
        self._ocultar_lista_sugerencias()
        codigo = _part_id_por_nombre(descripcion, self.codigos_part_id)
        if codigo:
            self.variables_control["codigo_unico_producto"].set(codigo)
            self._actualizar_validacion_codigo()
        else:
            codigo_actual = self.variables_control["codigo_unico_producto"].get()
            registro_actual = next(
                (
                    registro
                    for registro in self.codigos_part_id
                    if registro.get("part_id", "").casefold()
                    == codigo_actual.strip().casefold()
                ),
                None,
            )
            if registro_actual and normalizar(
                registro_actual.get("descripcion", "")
            ) != normalizar(descripcion):
                self.variables_control["codigo_unico_producto"].set("")
                self._actualizar_validacion_codigo()
        self.entrada_nombre_producto.focus_set()

    def _ocultar_lista_sugerencias(self) -> None:
        popup = self._popup_sugerencias
        self._popup_sugerencias = None
        self._lista_sugerencias = None
        self._entrada_sugerencias = None
        self._al_seleccionar_sugerencia = None
        if popup is not None and popup.winfo_exists():
            popup.destroy()

    def _mostrar_lista_sugerencias(
        self,
        entrada: ttk.Combobox,
        opciones: tuple[str, ...],
        al_seleccionar: Any,
    ) -> None:
        opciones = tuple(opcion for opcion in opciones if opcion)
        if not opciones or not entrada.winfo_exists():
            self._ocultar_lista_sugerencias()
            return
        if (
            self._popup_sugerencias is not None
            and self._popup_sugerencias.winfo_exists()
            and self._entrada_sugerencias is entrada
        ):
            popup = self._popup_sugerencias
            lista = self._lista_sugerencias
            if lista is None:
                self._ocultar_lista_sugerencias()
                return
            lista.delete(0, tk.END)
            for opcion in opciones[:5]:
                lista.insert(tk.END, opcion)
        else:
            self._ocultar_lista_sugerencias()
            popup = ui.Ventana(self.ventana)
            popup.wm_overrideredirect(True)
            popup.transient(self.ventana)
            lista = tk.Listbox(
                popup,
                height=min(len(opciones), 5),
                activestyle="dotbox",
                exportselection=False,
                borderwidth=1,
                relief="solid",
                background="#ffffff",
                foreground="#202124",
                selectbackground="#dce8f8",
                selectforeground="#202124",
                highlightthickness=0,
            )
            lista.pack(fill="both", expand=True)
            lista.bind(
                "<ButtonRelease-1>",
                self._aceptar_sugerencia_popup,
            )
            lista.bind("<Return>", self._aceptar_sugerencia_popup)
            lista.bind("<Escape>", lambda _evento: entrada.focus_set())
            self._popup_sugerencias = popup
            self._lista_sugerencias = lista
            self._entrada_sugerencias = entrada
            self._al_seleccionar_sugerencia = al_seleccionar
            for opcion in opciones[:5]:
                lista.insert(tk.END, opcion)
        assert self._lista_sugerencias is not None
        assert self._popup_sugerencias is not None
        self._lista_sugerencias.configure(height=min(len(opciones), 5))
        ancho = max(entrada.winfo_width(), min(max(map(len, opciones[:5])) * 8 + 24, 520))
        alto = min(len(opciones), 5) * 22 + 2
        x = entrada.winfo_rootx()
        y = entrada.winfo_rooty() + entrada.winfo_height()
        self._popup_sugerencias.geometry(f"{ancho}x{alto}+{x}+{y}")
        self._popup_sugerencias.lift()
        entrada.focus_set()

    def _aceptar_sugerencia_popup(
        self, _evento: tk.Event | None = None
    ) -> str:
        lista = self._lista_sugerencias
        callback = self._al_seleccionar_sugerencia
        if lista is None or callback is None:
            return "break"
        seleccion = lista.curselection()
        if not seleccion:
            indice_activo = lista.index(tk.ACTIVE)
            if indice_activo < 0:
                return "break"
            seleccion = (indice_activo,)
        callback(lista.get(seleccion[0]))
        return "break"

    def _enfocar_sugerencias(self, _evento: tk.Event | None = None) -> str:
        lista = self._lista_sugerencias
        if lista is None or not lista.winfo_exists():
            return "break"
        if not lista.curselection() and lista.size():
            lista.selection_set(0)
            lista.activate(0)
        lista.focus_set()
        return "break"
    def _actualizar_validacion_codigo(self) -> bool:
        codigo = self.variables_control["codigo_unico_producto"].get()
        valido = _part_id_existe(codigo, self.codigos_part_id)
        codigo_canonico = _part_id_canonico(codigo, self.codigos_part_id)
        if valido and codigo != codigo_canonico:
            self.variables_control["codigo_unico_producto"].set(codigo_canonico)
            codigo = codigo_canonico
        if valido:
            registro = next(
                (
                    item
                    for item in self.codigos_part_id
                    if item.get("part_id", "").casefold()
                    == codigo.casefold()
                ),
                None,
            )
            if registro is not None:
                self.nombre_producto.set(registro.get("descripcion", ""))
        if self.etiqueta_codigo_asterisco is not None:
            self.etiqueta_codigo_asterisco.configure(
                text="" if valido or not self._codigo_validacion_activada else "*"
            )
        if self.etiqueta_codigo_invalido is not None:
            ingles = self._idioma.get() == "English"
            mensaje = ""
            if self._codigo_validacion_activada and not valido:
                if ingles and codigo.strip():
                    mensaje = "This Part ID is not in the catalog. Select a valid code."
                elif ingles:
                    mensaje = "Enter a Part ID from the catalog."
                elif codigo.strip():
                    mensaje = "El código no existe en el catálogo. Selecciona uno válido."
                else:
                    mensaje = "Ingresa un código del catálogo."
            self.etiqueta_codigo_invalido.configure(text=mensaje)
            if mensaje:
                self.etiqueta_codigo_invalido.grid()
            else:
                self.etiqueta_codigo_invalido.grid_remove()
        return valido

    def _mostrar_leyenda_periodicidad(self) -> None:
        idioma = "en" if self._idioma.get() == "English" else "es"
        lineas = [
            f"{simbolo} — {descripcion}"
            for simbolo, descripcion in LEYENDA_PERIODICIDAD[idioma]
        ]
        messagebox.showinfo(
            "Analysis periodicity" if idioma == "en" else "Periodicidad de análisis",
            "\n".join(lineas),
            parent=self.ventana,
        )

    def _mostrar_ayuda_tabla(self) -> None:
        ingles = self._idioma.get() == "English"
        mensaje = (
            "Click an empty cell to add a row; click any cell to edit it. "
            "Gray suggestions can be clicked to accept them. Enter 1–4 "
            "symbols with the periodicity dropdown. Press Delete to clear a cell."
            if ingles
            else "Haz clic en una celda vacía para agregar una fila y en "
            "cualquier celda para editarla. Puedes hacer clic en una sugerencia "
            "gris para aceptarla. Selecciona de 1 a 4 símbolos en el desplegable "
            "de periodicidad. Presiona Suprimir para borrar una celda."
        )
        messagebox.showinfo(
            "Table help" if ingles else "Ayuda de la tabla",
            mensaje,
            parent=self.ventana,
        )

    def _limpiar_sugerencia(self) -> None:
        self._sugerencia_actual = None
        self.etiqueta_sugerencia.configure(text="")
        self.etiqueta_sugerencia.grid_remove()

    def _mostrar_sugerencia(
        self, iid: str, campo: str, valor: str
    ) -> None:
        ingles = self._idioma.get() == "English"
        encabezados = {
            "prueba": "Suggested test" if ingles else "Prueba sugerida",
            "unidades": "Suggested unit" if ingles else "Unidad sugerida",
        }
        self._sugerencia_actual = (iid, campo, valor)
        self.etiqueta_sugerencia.grid()
        self.etiqueta_sugerencia.configure(
            text=f"{encabezados[campo]}: {valor}  ·  "
            + ("click to use" if ingles else "haz clic para usar")
        )

    def _actualizar_sugerencia_contextual(self, iid: str) -> None:
        if iid == "nueva":
            self._limpiar_sugerencia()
            return
        fila = self.pruebas[int(iid)]
        idioma = "en" if self._idioma.get() == "English" else "es"
        nombre = fila.get("prueba", "").strip()
        metodo = fila.get("metodo", "").strip()
        unidad = (
            _unidad_sugerida(self.maestro, nombre, metodo, idioma)
            if nombre and metodo
            else ""
        )
        if unidad and not fila.get("unidades", "").strip():
            self._mostrar_sugerencia(iid, "unidades", unidad)
            return
        nombres = _nombres_maestros_para_metodo(self.maestro, metodo, idioma)
        if len(nombres) == 1 and not nombre:
            self._mostrar_sugerencia(iid, "prueba", nombres[0])
            return
        self._limpiar_sugerencia()

    def _aceptar_sugerencia(self, _evento: tk.Event | None = None) -> None:
        sugerencia = self._sugerencia_actual
        if sugerencia is None:
            return
        iid, campo, valor = sugerencia
        if self._celda_editada == (iid, campo) and isinstance(
            self._editor_celda, ttk.Combobox
        ):
            self._editor_celda.set(valor)
            self._cerrar_editor_celda(guardar=True)
            self.tabla_pruebas.selection_set(iid)
            self._actualizar_sugerencia_contextual(iid)
            return
        fila = self.pruebas[int(iid)]
        fila[campo] = valor
        self._actualizar_tabla()
        self.tabla_pruebas.selection_set(iid)
        self._actualizar_sugerencia_contextual(iid)

    def _actualizar_idioma(self, _evento: tk.Event | None = None) -> None:
        ingles = self._idioma.get() == "English"
        idioma = "en" if ingles else "es"
        self._nombres_sugeridos_por_fila.clear()
        self._limpiar_sugerencia()
        if idioma != self._idioma_anterior:
            self._cerrar_editor_celda(guardar=True)
            if self._bilingue is not None:
                seleccion = self._seleccion()
                self._capturar_bilingue(self._idioma_anterior)
                self._bilingue.actualizar_traducciones_automaticas(idioma, completar_sin_traducir=True)
                self._aplicar_borrador(idioma)
            else:
                seleccion = None
                for fila in self.pruebas:
                    fila["prueba"] = _traducir_nombre_prueba(
                        self.maestro,
                        fila.get("prueba", ""),
                        self._idioma_anterior,
                        idioma,
                        fila.get("metodo", ""),
                    )
            self._idioma_anterior = idioma
            self._actualizar_tabla()
            if seleccion is not None and seleccion < len(self.pruebas):
                self.tabla_pruebas.selection_set(str(seleccion))
                self.tabla_pruebas.see(str(seleccion))
        elaborado = self.variables_control["elaboro_y_aprobo"].get().strip()
        if (
            self._bilingue is None and (not elaborado
            or elaborado == self._elaboro_aprobo_predeterminado_actual
        )):
            elaborado = VALORES_ELABORO_APROBO[idioma]
            self.variables_control["elaboro_y_aprobo"].set(elaborado)
        self._elaboro_aprobo_predeterminado_actual = elaborado
        notas_actuales = self.notas.get("1.0", "end-1c").strip()
        if (
            self._bilingue is None and (not notas_actuales
            or notas_actuales == self._nota_predeterminada_actual
        )):
            notas_actuales = NOTAS_PLANTILLA[idioma]
            self.notas.delete("1.0", tk.END)
            self.notas.insert("1.0", notas_actuales)
        self._nota_predeterminada_actual = notas_actuales
        self.ventana.title("New specification" if ingles else "Nueva especificación")
        self.etiqueta_idioma.configure(text="Language:" if ingles else "Idioma:")
        self.selector_idioma.configure(values=("English", "Español"))
        self.etiqueta_producto.configure(
            text="Product name:" if ingles else "Nombre del producto:"
        )
        self.etiqueta_desarrollador.configure(
            text="Developer *:" if ingles else "Desarrollador *:"
        )
        self.etiqueta_reanalisis.configure(text="Reanalysis date:" if ingles else "Fecha de reanálisis:")
        self.unidad_reanalisis.configure(text="months" if ingles else "meses")
        self.etiqueta_titulo_control.configure(
            text="Document control" if ingles else "Control documental"
        )
        self.marco_pruebas.configure(text="Tests" if ingles else "Pruebas")
        for (clave, etiqueta_es, etiqueta_en) in CAMPOS_CONTROL_GENERADOR:
            self._etiquetas[clave].configure(
                text=f"{etiqueta_en}:" if ingles else f"{etiqueta_es}:"
            )
        for campo, encabezado in zip(
            CAMPOS_TABLA, ETIQUETAS_TABLA_GENERADOR[idioma]
        ):
            self.tabla_pruebas.heading(campo, text=encabezado)
        self._ordenador_pruebas.establecer_encabezados(
            {
                campo: encabezado
                for campo, encabezado in zip(
                    CAMPOS_TABLA, ETIQUETAS_TABLA_GENERADOR[idioma]
                )
                if campo in {"prueba", "metodo"}
            }
        )
        self.tabla_pruebas.heading(
            "periodicidad",
            text=f"{ETIQUETA_PERIODICIDAD[idioma]} ⓘ",
            command=self._mostrar_leyenda_periodicidad,
        )
        if self.etiqueta_codigo_invalido is not None:
            self._actualizar_validacion_codigo()
        self.boton_eliminar.configure(text="Remove row" if ingles else "Eliminar fila")
        self.boton_alinear.configure(
            text="Align specifications" if ingles else "Alinear especificaciones"
        )
        self.etiqueta_notas.configure(text="Notes:" if ingles else "Notas:")
        self.boton_guardar.configure(text="Save JSON" if ingles else "Guardar JSON")
        self.boton_cancelar.configure(text="Cancel" if ingles else "Cancelar")
        if self._bilingue is not None:
            self.ventana.title(f"Editar producto · {self.nombre_producto.get()} · ESP / ENG")
            self._programar_resaltado()

    def _seleccion(self) -> int | None:
        seleccion = self.tabla_pruebas.selection()
        if not seleccion or seleccion[0] == "nueva":
            return None
        return int(seleccion[0])

    def _actualizar_tabla(self) -> None:
        self._cerrar_editor_celda(guardar=True)
        if self._bilingue is not None:
            self._capturar_bilingue(self._idioma_anterior)
        self._limpiar_sugerencia()
        filas_existentes = self.tabla_pruebas.get_children()
        if filas_existentes:
            self.tabla_pruebas.delete(*filas_existentes)
        for indice, fila in enumerate(self.pruebas):
            self.tabla_pruebas.insert(
                "",
                tk.END,
                iid=str(indice),
                values=tuple(
                    fila.get(campo, "")
                    for campo in (*CAMPOS_TABLA, "periodicidad")
                ),
                tags=("fila_alterna",) if indice % 2 else (),
            )
        self.tabla_pruebas.insert(
            "",
            tk.END,
            iid="nueva",
            values=("", "", "", "", "", ""),
        )
        self._ordenador_pruebas.registrar_orden_original(
            [*(str(indice) for indice in range(len(self.pruebas))), "nueva"]
        )
        self.boton_eliminar.configure(state="disabled")
        self._programar_resaltado()

    def _sugerencias(self, campo: str, texto_actual: str, iid: str) -> tuple[str, ...]:
        maestro = self.maestro
        if campo in {"unidades", "metodo"} and iid != "nueva":
            fila = self.pruebas[int(iid)]
            nombre_prueba = fila.get("prueba", "")
            if nombre_prueba:
                idioma_maestro = (
                    "en" if self._idioma.get() == "English" else "es"
                )
                relacionadas = [
                    registro
                    for registro in self.maestro
                    if normalizar(_nombre_maestro_para_idioma(registro, idioma_maestro))
                    == normalizar(nombre_prueba)
                ]
                if relacionadas:
                    maestro = relacionadas
                    metodo = fila.get("metodo", "").strip()
                    referencias = set(referencias_astm(metodo))
                    if campo == "unidades" and metodo:
                        if referencias:
                            por_metodo = [
                                registro
                                for registro in relacionadas
                                if set(
                                    referencias_astm(
                                        metodo_maestro_para_idioma(
                                            registro, idioma_maestro
                                        )
                                    )
                                )
                                == referencias
                            ]
                        else:
                            metodo_normalizado = normalizar(
                                metodo_maestro_canonico(metodo)
                            )
                            por_metodo = [
                                registro
                                for registro in relacionadas
                                if normalizar(
                                    metodo_maestro_canonico(
                                        metodo_maestro_para_idioma(
                                            registro, idioma_maestro
                                        )
                                    )
                                )
                                == metodo_normalizado
                            ]
                        if por_metodo:
                            maestro = por_metodo
        opciones = _opciones_maestro(
            maestro,
            campo,
            "en" if self._idioma.get() == "English" else "es",
            texto_actual,
            self._nombres_sugeridos_por_fila.get(iid),
        )
        return ("", *opciones)

    def _editar_celda(self, evento: tk.Event) -> str | None:
        iid = getattr(evento, "iid", None) or self.tabla_pruebas.identify_row(evento.y)
        columna = getattr(evento, "columna", None) or self.tabla_pruebas.identify_column(evento.x)
        if not iid or not columna.startswith("#"):
            return None
        indice_columna = int(columna[1:]) - 1
        campos_generador = (*CAMPOS_TABLA, "periodicidad")
        if not 0 <= indice_columna < len(campos_generador):
            return None
        campo = campos_generador[indice_columna]
        if self._celda_editada == (iid, campo):
            return "break"
        self._cerrar_editor_celda(guardar=True)
        if iid == "nueva":
            self.pruebas.append(
                {
                    nombre: PERIODICIDAD_PREDETERMINADA if nombre == "periodicidad" else ""
                    for nombre in (*CAMPOS_TABLA, "periodicidad")
                }
            )
            iid = str(len(self.pruebas) - 1)
            self._actualizar_tabla()
        self.tabla_pruebas.selection_set(iid)
        bbox = self.tabla_pruebas.bbox(iid, columna)
        if not bbox:
            return "break"
        x, y, ancho, alto = bbox
        valor = (
            self.pruebas[int(iid)].get(campo, "")
            if iid != "nueva"
            else ""
        )
        self._celda_activa = (iid, campo)
        self._mostrar_referencia_bilingue(iid, campo)
        sugerencias = self._sugerencias(campo, valor, iid)
        if campo == "periodicidad":
            editor = self._crear_editor_periodicidad(
                x, y, ancho, alto, valor
            )
        elif campo in {"prueba", "metodo"}:
            marco_comentario = ttk.Frame(self.tabla_pruebas)
            marco_comentario.place(x=x, y=y, width=ancho, height=alto)
            editor = ttk.Combobox(marco_comentario, values=sugerencias)
            editor.set(valor)
            editor.pack(side="left", fill="both", expand=True)
            self._comentario_marco = marco_comentario
            self._comentario_boton_mas = ttk.Button(
                marco_comentario,
                text="+",
                width=2,
                command=lambda fila=iid, celda=campo, entrada=editor: (
                    self._agregar_comentario_celda(fila, celda, entrada)
                ),
            )
            self._comentario_boton_mas.pack(side="right")
            editor.bind(
                "<KeyRelease>",
                lambda e, c=campo, r=iid: self._filtrar_sugerencias(e, c, r),
            )
            editor.bind("<Down>", self._enfocar_sugerencias)
            editor.bind(
                "<<ComboboxSelected>>",
                lambda _e: self._cerrar_editor_celda(guardar=True),
            )
        elif campo == "unidades":
            editor: ttk.Entry | ttk.Combobox = ttk.Combobox(
                self.tabla_pruebas,
                values=sugerencias,
            )
            editor.set(valor)
            editor.bind("<KeyRelease>", lambda e, c=campo, r=iid: self._filtrar_sugerencias(e, c, r))
            editor.bind("<Down>", self._enfocar_sugerencias)
            editor.bind("<<ComboboxSelected>>", lambda _e: self._cerrar_editor_celda(guardar=True))
        else:
            editor = ttk.Entry(self.tabla_pruebas)
            editor.insert(0, valor)
        if campo not in {"periodicidad", "prueba", "metodo"}:
            editor.place(x=x, y=y, width=ancho, height=alto)
            editor.focus_set()
            if valor:
                editor.icursor(tk.END)
        self._editor_celda = editor
        self._celda_editada = (iid, campo)
        self._actualizar_sugerencia_contextual(iid)
        controles = (
            [*self._periodicidad_editores]
            + ([self._periodicidad_boton_mas] if campo == "periodicidad" else [])
            if campo == "periodicidad"
            else [editor]
            + (
                [self._comentario_boton_mas]
                if self._comentario_boton_mas is not None
                else []
            )
            if campo in {"prueba", "metodo"}
            else [editor]
        )
        for control in controles:
            control.bind("<Return>", lambda _e: self._cerrar_editor_celda(guardar=True))
            control.bind("<Escape>", lambda _e: self._cerrar_editor_celda(guardar=False))
            control.bind("<FocusOut>", self._programar_cierre_editor)
            control.bind("<Delete>", self._borrar_valor_editor)
            control.bind("<KP_Delete>", self._borrar_valor_editor)
        return "break"

    def _filtrar_sugerencias(
        self, evento: tk.Event, campo: str, iid: str
    ) -> None:
        editor = self._editor_celda
        if not isinstance(editor, ttk.Combobox):
            return
        opciones = self._sugerencias(campo, editor.get(), iid)
        editor.configure(values=opciones)
        self._mostrar_lista_sugerencias(
            editor,
            tuple(opcion for opcion in opciones if opcion),
            lambda opcion: self._elegir_sugerencia_fila(iid, campo, opcion),
        )
        if campo == "prueba":
            cercana = _prueba_cercana(
                self.maestro,
                "en" if self._idioma.get() == "English" else "es",
                editor.get(),
                self._nombres_sugeridos_por_fila.get(iid),
            )
            if cercana and normalizar(cercana) != normalizar(editor.get()):
                self._mostrar_sugerencia(iid, "prueba", cercana)
            else:
                self._limpiar_sugerencia()
        elif campo == "unidades" and not editor.get().strip():
            self._actualizar_sugerencia_contextual(iid)

    def _elegir_sugerencia_fila(
        self, iid: str, campo: str, valor: str
    ) -> None:
        editor = self._editor_celda
        if not isinstance(editor, ttk.Combobox):
            return
        editor.set(valor)
        self._ocultar_lista_sugerencias()
        self._cerrar_editor_celda(guardar=True)
        self.tabla_pruebas.selection_set(iid)
        self._actualizar_sugerencia_contextual(iid)

    def _crear_editor_periodicidad(
        self,
        x: int,
        y: int,
        ancho: int,
        alto: int,
        valor: str,
    ) -> ttk.Frame:
        marco = ttk.Frame(self.tabla_pruebas)
        marco.place(x=x, y=y, width=ancho, height=alto)
        self._periodicidad_marco = marco
        self._periodicidad_editores = []
        self._periodicidad_boton_mas = ttk.Button(
            marco,
            text="+",
            width=2,
            command=self._agregar_selector_periodicidad,
        )
        valores = (
            re.split(r"[\s,;/|]+", valor.strip())
            if re.search(r"[\s,;/|]", valor.strip())
            else list(valor.strip())
        )
        valores = [simbolo for simbolo in valores if simbolo]
        if not valores:
            valores = [PERIODICIDAD_PREDETERMINADA]
        for simbolo in valores[:MAX_SIMBOLOS_PERIODICIDAD]:
            self._agregar_selector_periodicidad(simbolo)
        return marco

    def _agregar_selector_periodicidad(self, valor: str = "") -> None:
        marco = self._periodicidad_marco
        if marco is None or not marco.winfo_exists():
            return
        if len(self._periodicidad_editores) >= MAX_SIMBOLOS_PERIODICIDAD:
            return
        editor = ttk.Combobox(
            marco,
            values=SIMBOLOS_PERIODICIDAD,
            state="readonly",
            width=3,
        )
        editor.set(valor)
        editor.grid(
            row=0,
            column=len(self._periodicidad_editores),
            sticky="nsew",
            padx=(0, 2),
        )
        editor.bind("<<ComboboxSelected>>", self._actualizar_selectores_periodicidad)
        editor.bind(
            "<Return>", lambda _e: self._cerrar_editor_celda(guardar=True)
        )
        editor.bind(
            "<Escape>", lambda _e: self._cerrar_editor_celda(guardar=False)
        )
        editor.bind("<FocusOut>", self._programar_cierre_editor)
        editor.bind("<Delete>", self._borrar_valor_editor)
        editor.bind("<KP_Delete>", self._borrar_valor_editor)
        self._periodicidad_editores.append(editor)
        self._actualizar_selectores_periodicidad()

    def _actualizar_selectores_periodicidad(
        self, _evento: tk.Event | None = None
    ) -> None:
        simbolos = [editor.get() for editor in self._periodicidad_editores if editor.get()]
        if len(set(simbolos)) != len(simbolos):
            editor_cambiado = (
                _evento.widget
                if _evento is not None
                and isinstance(_evento.widget, ttk.Combobox)
                else None
            )
            if editor_cambiado is not None:
                editor_cambiado.set("")
            ingles = self._idioma.get() == "English"
            messagebox.showwarning(
                "Duplicate symbol" if ingles else "Símbolo duplicado",
                (
                    "Select each periodicity symbol only once."
                    if ingles
                    else "Selecciona cada símbolo de periodicidad una sola vez."
                ),
                parent=self.ventana,
            )
        self._mostrar_boton_mas_periodicidad()

    def _mostrar_boton_mas_periodicidad(self) -> None:
        boton = self._periodicidad_boton_mas
        marco = self._periodicidad_marco
        if boton is None or marco is None or not marco.winfo_exists():
            return
        if len(self._periodicidad_editores) < MAX_SIMBOLOS_PERIODICIDAD:
            boton.grid(
                row=0,
                column=len(self._periodicidad_editores),
                sticky="nsew",
                padx=(1, 0),
            )
        else:
            boton.grid_remove()

    def _programar_cierre_editor(self, _evento: tk.Event | None = None) -> None:
        editor = self._editor_celda
        if editor is None:
            return

        def cerrar_si_perdio_foco() -> None:
            if self._editor_celda is not editor:
                return
            foco = self.ventana.focus_get()
            marco = (
                self._comentario_marco
                if self._celda_editada is not None
                and self._celda_editada[1] in {"prueba", "metodo"}
                else None
            )
            if foco is not None and (
                str(foco).startswith(str(editor))
                or marco is not None and str(foco).startswith(str(marco))
            ):
                return
            self._cerrar_editor_celda(guardar=True)

        self.ventana.after(160, cerrar_si_perdio_foco)

    def _borrar_valor_editor(self, _evento: tk.Event | None = None) -> str:
        editor = self._editor_celda
        if editor is None:
            return "break"
        if isinstance(editor, ttk.Frame):
            for selector in self._periodicidad_editores:
                selector.set("")
        elif isinstance(editor, ttk.Combobox) and str(
            editor.cget("state")
        ) == "readonly":
            editor.set("")
        else:
            editor.delete(0, tk.END)
        return "break"

    def _borrar_celda_activa(self, _evento: tk.Event | None = None) -> str:
        if self._editor_celda is not None:
            return self._borrar_valor_editor()
        if self._celda_activa is None:
            return "break"
        iid, campo = self._celda_activa
        if iid == "nueva":
            return "break"
        marcador = _marcador_comentario_celda(
            self.pruebas[int(iid)].get(campo, "")
        ) if campo in {"prueba", "metodo"} else None
        self.pruebas[int(iid)][campo] = ""
        if marcador is not None:
            self.notas.delete("1.0", tk.END)
            self.notas.insert(
                "1.0",
                _actualizar_comentarios_celda(
                    self.pruebas,
                    self.notas.get("1.0", "end-1c"),
                    remover_numero=marcador,
                ),
            )
        self._actualizar_tabla()
        self.tabla_pruebas.selection_set(iid)
        self._actualizar_sugerencia_contextual(iid)
        return "break"

    def _reposicionar_editor(self, _evento: tk.Event) -> None:
        if self._editor_celda is None or self._celda_editada is None:
            return
        iid, campo = self._celda_editada
        columnas = {
            nombre: f"#{indice}"
            for indice, nombre in enumerate((*CAMPOS_TABLA, "periodicidad"), 1)
        }
        bbox = self.tabla_pruebas.bbox(iid, columnas[campo])
        if not bbox:
            self._cerrar_editor_celda(guardar=True)
            return
        marco = (
            self._comentario_marco
            if campo in {"prueba", "metodo"}
            else self._periodicidad_marco
            if campo == "periodicidad"
            else self._editor_celda
        )
        if marco is not None:
            marco.place(x=bbox[0], y=bbox[1], width=bbox[2], height=bbox[3])

    def _cerrar_editor_celda(self, guardar: bool) -> None:
        editor = self._editor_celda
        celda = self._celda_editada
        if editor is None or celda is None:
            return
        iid, campo = celda
        self._editor_celda = None
        self._celda_editada = None
        if campo == "periodicidad":
            valor = "".join(
                selector.get().strip()
                for selector in self._periodicidad_editores
                if selector.get().strip()
            )
        elif isinstance(editor, ttk.Frame):
            valor = ""
        else:
            valor = editor.get().strip()
        self._ocultar_lista_sugerencias()
        if campo in {"prueba", "metodo"} and self._comentario_marco is not None:
            self._comentario_marco.destroy()
        else:
            editor.destroy()
        self._periodicidad_editores = []
        self._periodicidad_marco = None
        self._periodicidad_boton_mas = None
        self._comentario_marco = None
        self._comentario_boton_mas = None
        if not guardar:
            return
        fila = self.pruebas[int(iid)]
        if campo in {"prueba", "metodo"}:
            notas_actualizadas = _guardar_celda_con_comentario(
                self.pruebas,
                self.notas.get("1.0", "end-1c"),
                int(iid),
                campo,
                valor,
            )
            if notas_actualizadas != self.notas.get("1.0", "end-1c"):
                self.notas.delete("1.0", tk.END)
                self.notas.insert("1.0", notas_actualizadas)
        else:
            fila[campo] = valor
        if campo == "prueba" and valor:
            idioma = "en" if self._idioma.get() == "English" else "es"
            for nombre_campo, valor_default in _valores_prueba_predeterminados(
                self.maestro, valor, idioma
            ).items():
                if not fila.get(nombre_campo, "").strip():
                    fila[nombre_campo] = valor_default
        if campo == "metodo":
            sugerencias = _nombres_maestros_para_metodo(
                self.maestro,
                valor,
                "en" if self._idioma.get() == "English" else "es",
            )
            self._nombres_sugeridos_por_fila[iid] = sugerencias
        self._actualizar_tabla()
        if iid != "nueva":
            self.tabla_pruebas.selection_set(iid)
            self._actualizar_sugerencia_contextual(iid)

    def _agregar_comentario_celda(
        self, iid: str, campo: str, editor: ttk.Combobox
    ) -> None:
        try:
            fila_indice = int(iid)
            texto_notas = self.notas.get("1.0", "end-1c")
            notas_actualizadas = _agregar_comentario_a_celda(
                self.pruebas,
                texto_notas,
                fila_indice,
                campo,
                editor.get().strip(),
            )
        except (IndexError, ValueError) as exc:
            ingles = self._idioma.get() == "English"
            messagebox.showwarning(
                "Method comment" if ingles else "Comentario de método",
                str(exc),
                parent=self.ventana,
            )
            return

        self.notas.delete("1.0", tk.END)
        self.notas.insert("1.0", notas_actualizadas)
        editor.set(self.pruebas[fila_indice][campo])
        self.notas.mark_set(tk.INSERT, tk.END)
        self.notas.focus_set()

    def _actualizar_boton_eliminar(self, _evento: tk.Event | None = None) -> None:
        seleccion = self._seleccion()
        self.boton_eliminar.configure(
            state="normal" if seleccion is not None else "disabled"
        )

    def _eliminar_seleccionada(self) -> None:
        self._cerrar_editor_celda(guardar=True)
        indice = self._seleccion()
        if indice is None:
            return
        ingles = self._idioma.get() == "English"
        confirmar = messagebox.askyesno(
            "Remove test" if ingles else "Eliminar prueba",
            "Remove the selected test?" if ingles else "¿Eliminar la prueba seleccionada?",
            parent=self.ventana,
        )
        if not confirmar:
            return
        marcadores_eliminados = {
            marcador
            for campo in ("prueba", "metodo")
            if (
                marcador := _marcador_comentario_celda(
                    self.pruebas[indice].get(campo, "")
                )
            )
        }
        del self.pruebas[indice]
        texto_notas = self.notas.get("1.0", "end-1c")
        if marcadores_eliminados:
            texto_notas = _actualizar_comentarios_celda(
                self.pruebas,
                texto_notas,
                remover_numero=marcadores_eliminados,
            )
        self.notas.delete("1.0", tk.END)
        self.notas.insert("1.0", texto_notas)
        self._nombres_sugeridos_por_fila.clear()
        self._actualizar_tabla()

    def _alinear_especificaciones(self) -> None:
        self._cerrar_editor_celda(guardar=True)
        ingles = self._idioma.get() == "English"
        idioma = "en" if ingles else "es"
        cantidad = sum(
            any(str(fila.get(campo, "") or "").strip() for campo in CAMPOS_TABLA)
            for fila in self.pruebas
        )
        if not cantidad:
            messagebox.showinfo(
                "Align specifications" if ingles else "Alinear especificaciones",
                (
                    "Add at least one test before aligning."
                    if ingles
                    else "Agrega al menos una prueba antes de alinear."
                ),
                parent=self.ventana,
            )
            return

        progreso = ui.Ventana(self.ventana)
        progreso.title("Aligning specifications" if ingles else "Alineando especificaciones")
        progreso.transient(self.ventana)
        progreso.resizable(False, False)
        progreso.protocol("WM_DELETE_WINDOW", lambda: None)
        contenido = ttk.Frame(progreso, padding=14)
        contenido.pack(fill="both", expand=True)
        estado = ttk.Label(
            contenido,
            text=(
                f"Aligning 0 of {cantidad} tests..."
                if ingles
                else f"Alineando 0 de {cantidad} pruebas..."
            ),
        )
        estado.pack(anchor="w", pady=(0, 8))
        barra = ttk.Progressbar(
            contenido,
            mode="determinate",
            maximum=cantidad,
            length=340,
        )
        barra.pack(fill="x")
        progreso.grab_set()
        progreso.update_idletasks()

        def actualizar_progreso(completadas: int, total: int) -> None:
            barra.configure(value=completadas)
            estado.configure(
                text=(
                    f"Aligning {completadas} of {total} tests..."
                    if ingles
                    else f"Alineando {completadas} de {total} pruebas..."
                )
            )
            progreso.update_idletasks()

        filas_antes = copy.deepcopy(self.pruebas)
        error_alineacion: str | None = None
        try:
            filas_alineadas, conteo = alinear_filas_generador(
                self.pruebas,
                self.maestro,
                idioma,
                al_progreso=actualizar_progreso,
            )
        except ValueError as exc:
            error_alineacion = str(exc)
        finally:
            progreso.grab_release()
            progreso.destroy()
        if error_alineacion is not None:
            messagebox.showerror(
                "Could not align specifications" if ingles else "No se pudieron alinear",
                error_alineacion,
                parent=self.ventana,
            )
            return

        self.pruebas = filas_alineadas
        self._cambios_alineacion[idioma] = {
            fila.get("_id_bilingue", str(indice)): campos
            for indice, (fila, campos) in enumerate(zip(filas_alineadas, diferencias_filas(filas_antes, filas_alineadas)))
        }
        if hasattr(self, "_cambio_alineacion_revision"):
            self._cambio_alineacion_revision = True
        self._nombres_sugeridos_por_fila.clear()
        self._actualizar_tabla()
        messagebox.showinfo(
            "Alignment complete" if ingles else "Alineación completada",
            (
                f"Aligned: {conteo['alineados']} | Review: {conteo['revision']} | "
                f"No match: {conteo['sin_coincidencia']}"
                if ingles
                else f"Alineadas: {conteo['alineados']} | Requieren revisión: "
                f"{conteo['revision']} | Sin coincidencia: "
                f"{conteo['sin_coincidencia']}"
            ),
            parent=self.ventana,
        )

    def guardar_en_progreso(self):
        self._cerrar_editor_celda(guardar=True)
        if self._bilingue is not None:
            self._guardar_bilingue(en_progreso=True)

    def _validar_campos_guardado(self, en_progreso=False):
        if not self.desarrollador.get().strip():
            messagebox.showerror("Desarrollador obligatorio", "Completa el campo Desarrollador para guardar.", parent=self.ventana)
            if getattr(self, "_control_colapsado", False):
                self._alternar_control_documental()
            self.entrada_desarrollador.focus_set()
            return False
        if not en_progreso and not re.fullmatch(r"[0-9]{2}", self.fecha_reanalisis.get().strip()):
            messagebox.showerror("Reanálisis obligatorio", "Ingresa exactamente dos dígitos para el reanálisis en meses (por ejemplo, 06 o 12).", parent=self.ventana)
            if getattr(self, "_control_colapsado", False):
                self._alternar_control_documental()
            self.entrada_reanalisis.focus_set()
            return False
        return True

    def _datos_en_progreso(self, contenido, idioma):
        datos = copy.deepcopy(self._originales_idioma.get(idioma, {}))
        datos.setdefault("esquema", "especificacion_producto_es_v3")
        datos.setdefault("producto", {})["nombre"] = contenido["nombre_producto"].strip() or "Sin nombre"
        documento = datos.setdefault("documento", {})
        documento.update({"tipo": "ESPECIFICACION", "idioma": idioma})
        documento.setdefault("fecha_documento_normalizada", "")
        datos["control_documental"] = copy.deepcopy(contenido["control"])
        datos["notas"] = list(contenido["notas"])
        datos["Desarrollador"] = contenido["desarrollador"]
        datos["FECHA DE REANALISIS"] = _valor_reanalisis_json(contenido.get("fecha_reanalisis", ""))
        datos["pruebas"] = [{"numero": indice, **{campo: fila.get(campo, "") for campo in CAMPOS_TABLA},
                             CAMPO_PERIODICIDAD: fila.get("periodicidad", "")}
                            for indice, fila in enumerate(contenido["pruebas"], 1)]
        datos.pop("alineacion_metodos_lab", None)
        return crear_documento_alineado_manual(datos)

    def _guardar_bilingue(self, en_progreso=False):
        if not self._validar_campos_guardado(en_progreso):
            return
        self._capturar_bilingue(self._idioma_anterior)
        if not en_progreso:
            try:
                _validar_pareja_editor(self._bilingue.borradores["es"], self._bilingue.borradores["en"])
            except ValueError as exc:
                messagebox.showerror("No se puede guardar la pareja", str(exc), parent=self.ventana)
                return
        conteos = {i: len(self._bilingue.pendientes[i]) for i in ("es", "en")}
        if any(conteos.values()) and not en_progreso:
            messagebox.showwarning("Completar ESP / ENG",
                f"ESP: {conteos['es']} pendientes · ENG: {conteos['en']} pendientes\n"
                "Completa o confirma las traducciones resaltadas antes de guardar ambos idiomas.", parent=self.ventana)
            idioma = self._idioma_anterior if conteos[self._idioma_anterior] else ("en" if self._idioma_anterior == "es" else "es")
            if idioma != self._idioma_anterior:
                self._idioma.set("English" if idioma == "en" else "Español")
                self._actualizar_idioma()
            pendientes = self._bilingue.pendientes[idioma]
            for indice, fila in enumerate(self.pruebas):
                campo = next((c for c in CAMPOS_TABLA if (fila["_id_bilingue"], c) in pendientes), None)
                if campo:
                    self.tabla_pruebas.see(str(indice))
                    self.tabla_pruebas.selection_set(str(indice))
                    self._celda_activa = (str(indice), campo)
                    self._mostrar_referencia_bilingue(str(indice), campo)
                    self._programar_resaltado()
                    return
            self.notas.focus_set()
            return
        documentos = {}
        try:
            for idioma, contenido in self._bilingue.borradores.items():
                datos = self._datos_en_progreso(contenido, idioma) if en_progreso else crear_datos_especificacion_nueva(contenido["nombre_producto"], idioma,
                    contenido["control"], contenido["pruebas"], contenido["notas"], contenido["desarrollador"], contenido.get("fecha_reanalisis", ""))
                if idioma in self._originales_idioma:
                    datos = _integrar_edicion_json(self._originales_idioma[idioma], datos)
                documentos[idioma] = datos
        except ValueError as exc:
            messagebox.showerror("No se puede guardar la pareja", str(exc), parent=self.ventana)
            return
        rutas = dict(self._rutas_idioma)
        if en_progreso:
            rutas = {i: r for i, r in rutas.items() if es_wip(r)}
        carpeta = self.app.rutas["json_alineadas"]
        carpeta.mkdir(parents=True, exist_ok=True)
        for idioma in ("es", "en"):
            if idioma not in rutas:
                nombre = re.sub(r'[<>:"/\\|?*]+', "_", self.nombre_producto.get()).strip(" .") or "ESPECIFICACION"
                directorio = next(iter(rutas.values())).parent if rutas else carpeta
                base_nombre = f"{nombre}_{'ESP' if idioma == 'es' else 'ENG'}{'_trabajo' if en_progreso else ''}"
                sugerido, indice = directorio / f"{base_nombre}.json", 1
                while sugerido.exists():
                    sugerido = directorio / f"{base_nombre}_{indice}.json"
                    indice += 1
                elegido = filedialog.asksaveasfilename(parent=self.ventana,
                    title=f"Guardar JSON alineado · {'ESP' if idioma == 'es' else 'ENG'}",
                    initialdir=directorio, initialfile=sugerido.name,
                    defaultextension=".json", filetypes=(("Archivos JSON", "*.json"),))
                if not elegido:
                    return
                rutas[idioma] = Path(elegido).resolve()
        if rutas["es"].resolve() == rutas["en"].resolve():
            messagebox.showerror("Rutas ESP / ENG", "Cada idioma necesita un archivo JSON diferente.", parent=self.ventana)
            return
        if any(Path(r).resolve() in getattr(self, "_rutas_pareja_descartadas", set()) and Path(r).exists()
               for r in rutas.values()):
            messagebox.showerror("Elige un archivo nuevo",
                "La versión vinculada que no se cargó se conserva por separado. Elige otro nombre para la nueva traducción.",
                parent=self.ventana)
            return
        if en_progreso and any(Path(r).resolve() == Path(original).resolve() and not es_wip(original)
                               for r in rutas.values() for original in self._rutas_idioma.values()):
            messagebox.showerror("Guardar en progreso", "Elige archivos nuevos para el borrador y conserva la especificación terminada.", parent=self.ventana)
            return
        if not messagebox.askyesno("Guardar pareja ESP / ENG",
            f"{self.nombre_producto.get()} · ESP / ENG\n" + (
                f"En progreso · ESP: {conteos['es']} pendientes · ENG: {conteos['en']} pendientes\n\n" if en_progreso
                else "ESP: Completo · ENG: Completo\n\n") +
            f"Se guardarán ambos JSON alineados (se reemplazarán si existen):\n\nESP: {rutas['es']}\nENG: {rutas['en']}\n\n" +
            ("¿Guardar el trabajo para continuarlo después?" if en_progreso else "¿Confirmas que revisaste ambas versiones y deseas guardarlas?"), parent=self.ventana):
            return
        for idioma, datos in documentos.items():
            if not en_progreso:
                datos["documento"]["edicion_bilingue"] = 1
            otro = "en" if idioma == "es" else "es"
            control = documentos[otro]["control_documental"]
            datos["versiones_idioma"] = {"INGLES" if otro == "en" else "ESPANOL": [{
                "archivo_json": Path(os.path.relpath(rutas[otro], rutas[idioma].parent)).as_posix(),
                "codigo": control.get("codigo_unico_producto", ""), "revision": control.get("revision", ""),
                "cliente": control.get("nombre_cliente", ""),
            }]}
        try:
            for idioma, datos in documentos.items():
                conservar_historial(datos, self._rutas_idioma.get(idioma, rutas[idioma]), rutas[idioma])
            archivos = {rutas[i]: documentos[i] for i in ("es", "en")}
            if en_progreso:
                archivos.update({ruta_estado(r): serializar_estado(self._bilingue, rutas, r) for r in rutas.values()})
            guardar_pareja_json(archivos)
            if not en_progreso:
                for r in rutas.values():
                    ruta_estado(r).unlink(missing_ok=True)
        except (OSError, ValueError) as exc:
            messagebox.showerror("No se pudo guardar la pareja", str(exc), parent=self.ventana)
            return
        self._rutas_idioma = rutas
        self._originales_idioma = documentos
        self.datos_originales = documentos[self._idioma_anterior]
        self.destino_inicial = rutas[self._idioma_anterior]
        if hasattr(self, "aviso_pareja"):
            self.aviso_pareja.pack_forget()
        self._contenido_inicial_revision = self._contenido_para_revision()
        self._revision_inicial = self.variables_control["revision"].get()
        self._revision_aplicada = False
        self._cambio_alineacion_revision = False
        self._consultas_revision.clear()
        self._revisiones_historial.clear()
        try:
            import generador_hii as hii
            if not en_progreso:
                with hii.conectar(BASE / "especificaciones.db") as con:
                    hii.vincular(con, rutas["es"], rutas["en"])
            if self.al_guardar is not None:
                for ruta in rutas.values():
                    self.al_guardar(ruta)
        except (OSError, ValueError, sqlite3.Error) as exc:
            messagebox.showwarning("Pareja guardada", f"Ambos JSON se guardaron, pero no se pudo actualizar el catálogo:\n{exc}", parent=self.ventana)
            return
        messagebox.showinfo("Trabajo en progreso guardado" if en_progreso else "Pareja guardada", f"ESP: {rutas['es']}\nENG: {rutas['en']}", parent=self.ventana)

    def guardar(self) -> None:
        if self.idioma_destino is not None:
            if not self.traduccion_revisada.get() or ("en" if self._idioma.get() == "English" else "es") != self.idioma_destino:
                messagebox.showerror("Traducción pendiente",
                    "Completa la traducción en el idioma solicitado y marca la casilla de revisión antes de guardar.",
                    parent=self.ventana)
                return
        self._cerrar_editor_celda(guardar=True)
        if not self._validar_campos_guardado():
            return
        ingles = self._idioma.get() == "English"
        idioma = "en" if ingles else "es"
        if not self._actualizar_revision_automatica():
            codigo = self.variables_control["codigo_unico_producto"].get().strip()
            error = self._revisiones_historial.get(codigo)
            messagebox.showerror("Revisión pendiente",
                str(error) if isinstance(error, Exception) else
                "Se está consultando el historial del producto. Espera un momento y vuelve a guardar.",
                parent=self.ventana)
            return
        self._codigo_validacion_activada = True
        if not self._actualizar_validacion_codigo():
            self.entradas_control["codigo_unico_producto"].focus_set()
            messagebox.showerror(
                "Invalid Part ID" if ingles else "Código inválido",
                (
                    "Select a Part ID from the catalog before saving."
                    if ingles
                    else "Selecciona un código Part ID del catálogo antes de guardar."
                ),
                parent=self.ventana,
            )
            return
        if self._bilingue is not None:
            self._guardar_bilingue()
            return
        try:
            datos = crear_datos_especificacion_nueva(
                self.nombre_producto.get(),
                idioma,
                {
                    clave: variable.get()
                    for clave, variable in self.variables_control.items()
                },
                self.pruebas,
                self.notas.get("1.0", "end-1c").splitlines(),
                self.desarrollador.get(),
                self.fecha_reanalisis.get(),
            )
            if self.datos_originales is not None:
                datos = _integrar_edicion_json(self.datos_originales, datos)
            self.app.rutas["json_alineadas"].mkdir(parents=True, exist_ok=True)
        except (OSError, ValueError) as exc:
            messagebox.showerror(
                "Cannot save specification" if ingles else "No se puede guardar",
                str(exc),
                parent=self.ventana,
            )
            return

        nombre = re.sub(
            r'[<>:"/\\|?*]+',
            "_",
            self.nombre_producto.get().strip(),
        ).strip(" .")
        if not nombre:
            nombre = "NEW_SPECIFICATION" if ingles else "ESPECIFICACION_NUEVA"
        if self.destino_inicial is not None:
            destino = self.destino_inicial
            if destino.exists() and not messagebox.askyesno(
                "Overwrite JSON" if ingles else "Reemplazar JSON",
                (
                    f"Save changes to the currently opened JSON?\n{destino}"
                    if ingles
                    else f"¿Guardar los cambios en el JSON abierto?\n{destino}"
                ),
                parent=self.ventana,
            ):
                return
        else:
            ruta = filedialog.asksaveasfilename(
                parent=self.ventana,
                title="Save specification JSON" if ingles else "Guardar JSON de especificación",
                initialdir=self.app.rutas["json_alineadas"],
                initialfile=nombre,
                defaultextension=".json",
                filetypes=(("Archivos JSON", "*.json"),),
            )
            if not ruta:
                return
            destino = Path(ruta)
        temporal = None
        try:
            origen_historial = self.archivo_inicial or destino
            if self.idioma_destino is not None:
                # Una traducción nueva no hereda las HII del idioma de origen.
                datos["hii_generadas"] = []
                origen_historial = destino
            conservar_historial(datos, Path(origen_historial), destino)
            with tempfile.NamedTemporaryFile(
                mode="w",
                encoding="utf-8",
                dir=destino.parent,
                prefix=f".{destino.name}.",
                suffix=".tmp",
                delete=False,
            ) as archivo:
                temporal = Path(archivo.name)
                json.dump(datos, archivo, ensure_ascii=False, indent=2)
                archivo.write("\n")
            os.replace(temporal, destino)
        except (OSError, ValueError) as exc:
            if temporal is not None:
                temporal.unlink(missing_ok=True)
            messagebox.showerror(
                "Cannot save JSON" if ingles else "No se pudo guardar el JSON",
                str(exc),
                parent=self.ventana,
            )
            return
        messagebox.showinfo(
            "JSON saved" if ingles else "JSON guardado",
            (
                f"The aligned JSON was saved at:\n{destino}"
                if ingles
                else f"El JSON alineado se guardó en:\n{destino}"
            ),
            parent=self.ventana,
        )
        if self.al_guardar is not None:
            self.al_guardar(destino)
        else:
            try:
                import generador_hii as hii
                with hii.conectar(BASE / "especificaciones.db") as con:
                    hii.sincronizar_idiomas(con, destino.parent)
            except (OSError, ValueError, sqlite3.Error) as exc:
                messagebox.showerror("Relaciones por idioma",
                    f"El JSON se guardó, pero no se pudieron actualizar sus relaciones: {exc}",
                    parent=self.ventana)


class EditorJson:
    def __init__(
        self,
        app: "AplicacionEspecificaciones",
        archivo_inicial: Path | None = None,
    ):
        self.app = app
        self.ventana = ui.Ventana(app.raiz)
        self.ventana.title("Generar o editar JSON alineado")
        self.ventana.geometry("1050x720")
        self.ventana.minsize(700, 480)
        barra = ttk.Frame(self.ventana, padding=8)
        barra.pack(fill="x")
        ttk.Button(barra, text="Nuevo", command=self.nuevo).pack(side="left", padx=3)
        ttk.Button(barra, text="Abrir JSON", command=self.abrir).pack(side="left", padx=3)
        ttk.Button(barra, text="Validar", command=self.validar).pack(side="left", padx=3)
        ttk.Button(barra, text="Guardar alineada", command=self.guardar).pack(side="left", padx=3)
        self.texto = ScrolledText(self.ventana, wrap="none", undo=True, font=("Consolas", 10))
        self.texto.pack(fill="both", expand=True, padx=8, pady=(0, 8))
        if archivo_inicial is None:
            self.nuevo()
        else:
            self._abrir_ruta(archivo_inicial)

    def nuevo(self) -> None:
        plantilla = {
            "esquema": "especificacion_producto_es_v3",
            "producto": {"nombre": ""},
            "documento": {"tipo": "ESPECIFICACION", "idioma": "es"},
            "control_documental": {"codigo_unico_producto": "", "revision": "00", "fecha": ""},
            "pruebas": [
                {
                    "prueba": "",
                    "unidades": "",
                    "metodo": "",
                    "especificacion": "",
                    "valor_tipico": "",
                }
            ],
            "notas": [],
        }
        self._poner_texto(plantilla)

    def _poner_texto(self, datos: dict[str, Any]) -> None:
        self.texto.delete("1.0", tk.END)
        self.texto.insert("1.0", json.dumps(datos, ensure_ascii=False, indent=2))

    def _leer(self) -> dict[str, Any]:
        try:
            datos = json.loads(self.texto.get("1.0", "end-1c"))
        except json.JSONDecodeError as exc:
            raise ValueError(f"JSON inválido en línea {exc.lineno}, columna {exc.colno}: {exc.msg}") from exc
        if not isinstance(datos, dict):
            raise ValueError("La raíz del JSON debe ser un objeto.")
        return crear_documento_alineado_manual(datos)

    def validar(self) -> None:
        try:
            self._leer()
        except ValueError as exc:
            messagebox.showerror("JSON no válido", str(exc), parent=self.ventana)
            return
        messagebox.showinfo("Validación", "El JSON es válido y sus pruebas se incluirán en detalle.valores_alineados.", parent=self.ventana)

    def abrir(self) -> None:
        ruta = filedialog.askopenfilename(
            parent=self.ventana,
            initialdir=self.app.rutas["json_alineadas"],
            filetypes=(("Archivos JSON", "*.json"), ("Todos los archivos", "*.*")),
        )
        if not ruta:
            return
        self._abrir_ruta(Path(ruta))

    def _abrir_ruta(self, ruta: Path) -> None:
        try:
            datos = json.loads(ruta.read_text(encoding="utf-8-sig"))
            if not isinstance(datos, dict):
                raise ValueError("La raíz del JSON debe ser un objeto.")
            self._poner_texto(datos)
        except (OSError, UnicodeError, json.JSONDecodeError, ValueError) as exc:
            messagebox.showerror("No se pudo abrir", str(exc), parent=self.ventana)

    def guardar(self) -> None:
        try:
            datos = self._leer()
        except ValueError as exc:
            messagebox.showerror("JSON no válido", str(exc), parent=self.ventana)
            return
        self.app.rutas["json_alineadas"].mkdir(parents=True, exist_ok=True)
        ruta = filedialog.asksaveasfilename(
            parent=self.ventana,
            initialdir=self.app.rutas["json_alineadas"],
            defaultextension=".json",
            filetypes=(("Archivos JSON", "*.json"),),
        )
        if not ruta:
            return
        destino = Path(ruta)
        temporal = destino.with_suffix(destino.suffix + ".tmp")
        try:
            temporal.write_text(json.dumps(datos, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
            temporal.replace(destino)
        except OSError as exc:
            temporal.unlink(missing_ok=True)
            messagebox.showerror("No se pudo guardar", str(exc), parent=self.ventana)
            return
        messagebox.showinfo("Guardado", f"JSON alineado guardado en:\n{destino}", parent=self.ventana)


class PanelComparacion:
    def __init__(self, comparador: "Comparador", indice: int):
        self.comparador = comparador
        self.indice = indice
        self.ruta: Path | None = None
        self.archivo_origen: Path | None = None
        self.directorio_temporal: tempfile.TemporaryDirectory[str] | None = None
        self.zoom = 1.44
        self.vista = None
        self.resaltador = None
        self.marco = ttk.LabelFrame(comparador.contenedor, text=f"Especificación {indice}", padding=8)
        comparador.contenedor.add(
            self.marco,
            minsize=240,
            stretch="always",
        )
        self.marco.rowconfigure(1, weight=1)
        self.marco.columnconfigure(0, weight=1)
        comandos = ttk.Frame(self.marco)
        comandos.grid(row=0, column=0, sticky="ew", pady=(0, 6))
        fila_archivo = ttk.Frame(comandos)
        fila_archivo.pack(fill="x", pady=(0, 6))
        fila_herramientas = ttk.Frame(comandos)
        fila_herramientas.pack(fill="x")
        ttk.Button(fila_archivo, text="Seleccionar archivo", width=16, command=self.seleccionar).pack(side="left")
        ttk.Button(
            fila_herramientas,
            text="ⓘ",
            width=3, variante="secundario", height=26,
            command=self.mostrar_ayuda_seleccion,
        ).pack(side="left", padx=(4, 0))
        self.boton_editar = ttk.Button(
            fila_archivo,
            text="Editar",
            command=self.editar, width=7, variante="secundario",
            state="disabled",
        )
        self.boton_editar.pack(side="left", padx=(6, 0))
        ttk.Button(fila_herramientas, text="−", width=3, height=26, variante="secundario", command=lambda: self.cambiar_zoom(-0.1)).pack(side="left", padx=(8, 4))
        ttk.Button(fila_herramientas, text="+", width=3, height=26, variante="secundario", command=lambda: self.cambiar_zoom(0.1)).pack(side="left")
        ttk.Button(fila_herramientas, text="Emergente", width=9, height=26, variante="secundario", command=self.comparador.abrir_emergente).pack(side="right")
        ttk.Button(fila_herramientas, text="×", width=3, height=26, variante="secundario", command=self.limpiar).pack(side="right", padx=(0, 6))
        self.previa = ttk.Frame(self.marco)
        self.previa.grid(row=1, column=0, sticky="nsew")
        self.etiqueta = ttk.Label(self.previa, text="Selecciona un JSON para mostrarlo.", anchor="center")
        self.etiqueta.pack(fill="both", expand=True)

    def mostrar_ayuda_seleccion(self) -> None:
        messagebox.showinfo(
            "Archivos para comparar",
            "Puedes seleccionar:\n"
            "• JSON limpios o alineados.\n"
            "• JSON viejos o extraídos.\n"
            "• Tus propias especificaciones Word (.doc, .docx o .docm).\n\n"
            "Los documentos Word se usan solo para esta comparación y no se "
            "registran en el historial.",
            parent=self.comparador.ventana,
        )

    def seleccionar(self) -> None:
        ruta = filedialog.askopenfilename(
            parent=self.comparador.ventana,
            initialdir=self.comparador.app.rutas["json_alineadas"],
            filetypes=(
                ("Especificaciones y documentos", "*.json *.doc *.docx *.docm"),
                ("Archivos JSON", "*.json"),
                ("Documentos Word", "*.doc *.docx *.docm"),
            ),
        )
        if ruta:
            self.cargar_archivo(Path(ruta))

    def cargar_archivo(self, archivo: Path) -> None:
        self.limpiar(actualizar=False)
        if archivo.suffix.casefold() == ".json":
            self.archivo_origen = archivo
            self.ruta = archivo
        elif archivo.suffix.casefold() in EXTENSIONES_WORD:
            temporal = tempfile.TemporaryDirectory(prefix="comparacion_especificacion_")
            self.directorio_temporal = temporal
            try:
                self.ruta = crear_json_temporal_comparacion(
                    archivo,
                    Path(temporal.name),
                )
            except (
                OSError,
                RuntimeError,
                ValueError,
                ImportError,
                subprocess.SubprocessError,
                BadZipFile,
                PackageNotFoundError,
            ) as exc:
                self.limpiar()
                messagebox.showerror(
                    "No se pudo preparar el documento Word",
                    f"No fue posible extraer la especificación para compararla:\n{exc}",
                    parent=self.comparador.ventana,
                )
                return
            self.archivo_origen = archivo
            self.marco.configure(text=f"Especificación {self.indice} · Temporal")
        else:
            messagebox.showerror(
                "Formato no compatible",
                "Selecciona un archivo JSON, DOC, DOCX o DOCM.",
                parent=self.comparador.ventana,
            )
            self.comparador.actualizar_resaltado()
            return
        self._mostrar_archivo()
        self.boton_editar.configure(state="normal")

    def _mostrar_archivo(self):
        self.vista = mostrar_vista_json(self.previa, self.ruta, self.zoom)
        self.resaltador = None
        self.comparador.actualizar_resaltado()

    def editar(self) -> None:
        if self.ruta is None:
            return
        destino = (
            self.archivo_origen
            if self.archivo_origen is not None
            and self.archivo_origen.suffix.casefold() == ".json"
            else None
        )
        self.comparador.app.editar_especificacion(self.ruta, destino)

    def limpiar(self, actualizar=True) -> None:
        self.vista = None
        self.resaltador = None
        self.ruta = None
        self.archivo_origen = None
        self.boton_editar.configure(state="disabled")
        self.marco.configure(text=f"Especificación {self.indice}")
        if self.directorio_temporal is not None:
            self.directorio_temporal.cleanup()
            self.directorio_temporal = None
        for widget in self.previa.winfo_children():
            widget.destroy()
        self.etiqueta = ttk.Label(
            self.previa,
            text="Selecciona un JSON para mostrarlo.",
            anchor="center",
        )
        self.etiqueta.pack(fill="both", expand=True)
        if actualizar:
            self.comparador.actualizar_resaltado()

    def cambiar_zoom(self, cambio: float) -> None:
        self.zoom = min(1.8, max(0.7, self.zoom + cambio))
        if self.ruta is not None:
            self._mostrar_archivo()


class Comparador:
    def __init__(self, app: "AplicacionEspecificaciones"):
        self.app = app
        self.ventana = ui.Ventana(app.raiz)
        self.ventana.title("Comparar especificaciones")
        self.ventana.geometry("1350x820")
        self.ventana.minsize(850, 550)
        encabezado = ttk.Frame(self.ventana)
        encabezado.pack(fill="x", padx=14, pady=(12, 6))
        ttk.Label(
            encabezado,
            text="Comparación de especificaciones",
            font=("Segoe UI", 15, "bold"),
            foreground="#163b70",
        ).pack(side="left", anchor="w")
        self.boton_agregar = ttk.Button(
            encabezado,
            text="＋ Agregar especificación",
            command=self.agregar_panel,
        )
        self.boton_agregar.pack(side="right")
        self.resaltado_activo = False
        self._cerrando = False
        self._maestro_comparacion = None
        self._ruta_maestro_comparacion = None
        self.boton_resaltar = ttk.Button(encabezado, text="RESALTAR", command=self.alternar_resaltado,
                                        variante="secundario", state="disabled")
        self.estado_resaltado = ttk.Label(self.ventana, text="")
        self.estado_resaltado.pack(fill="x", padx=14)
        self.contenedor = crear_divisor_redimensionable(self.ventana, tk.HORIZONTAL)
        self.contenedor.pack(fill="both", expand=True, padx=8, pady=5)
        separador_inicializado = False

        def centrar_separador(evento: tk.Event) -> None:
            nonlocal separador_inicializado
            if not separador_inicializado and evento.width > 100:
                separador_inicializado = True
                self.contenedor.after_idle(self._distribuir_paneles)

        self.contenedor.bind("<Configure>", centrar_separador)
        self.paneles: list[PanelComparacion] = []
        self.agregar_panel()
        self.agregar_panel()
        self.ventana.protocol("WM_DELETE_WINDOW", self.cerrar)

    def agregar_panel(self) -> None:
        if len(self.paneles) >= 3:
            messagebox.showinfo("Límite alcanzado", "La comparación permite hasta 3 especificaciones.", parent=self.ventana)
            return
        panel = PanelComparacion(self, len(self.paneles) + 1)
        self.paneles.append(panel)
        self.boton_agregar.configure(state="disabled" if len(self.paneles) == 3 else "normal")
        if len(self.paneles) == 3:
            self.contenedor.after_idle(self._distribuir_paneles)

    def _distribuir_paneles(self) -> None:
        cantidad = len(self.paneles)
        ancho = self.contenedor.winfo_width()
        if cantidad < 2 or ancho <= 1:
            return
        for indice in range(1, cantidad):
            self.contenedor.sash_place(
                indice - 1,
                round(ancho * indice / cantidad),
                0,
            )

    def alternar_resaltado(self):
        # FUNCIÓN EN ESTADO BETA: resalta diferencias entre pruebas comunes.
        self.resaltado_activo = not self.resaltado_activo
        self.actualizar_resaltado()

    def actualizar_resaltado(self):
        if self._cerrando:
            return
        paneles = [p for p in self.paneles if p.vista is not None]
        hay_alineada = any(p.vista["alineada"] for p in paneles)
        if hay_alineada:
            self.boton_resaltar.pack(side="right", padx=8)
        else:
            self.boton_resaltar.pack_forget()
        habilitado = hay_alineada and len(paneles) >= 2
        self.boton_resaltar.configure(state="normal" if habilitado else "disabled",
                                     text="Quitar resaltado" if self.resaltado_activo else "RESALTAR")
        if not habilitado or not self.resaltado_activo:
            for panel in paneles:
                if panel.resaltador is not None:
                    panel.resaltador.establecer_cambios({})
            self.estado_resaltado.configure(text="")
            return
        ruta_lista = ruta_maestro(BASE)
        if self._maestro_comparacion is None or getattr(self, "_ruta_maestro_comparacion", ruta_lista) not in (None, ruta_lista):
            try:
                self._maestro_comparacion = cargar_maestro(ruta_lista)
            except (OSError, ValueError, ImportError):
                self._maestro_comparacion = ()
            self._ruta_maestro_comparacion = ruta_lista
        cambios, resumen = comparar_valores_pruebas([p.vista for p in paneles], self._maestro_comparacion,
                                                   SERVICIO_TRADUCTOR.diccionario.traducir)
        for panel, filas in zip(paneles, cambios):
            vista = panel.vista
            celdas = {vista["filas"][indice]: campos for indice, campos in filas.items()}
            if panel.resaltador is None:
                panel.resaltador = ResaltadorCambios(vista["tabla"], celdas, vista["vertical"], vista["horizontal"])
            else:
                panel.resaltador.establecer_cambios(celdas)
        aviso = f" · {resumen['unidades_distintas']} grupos con unidades distintas" if resumen["unidades_distintas"] else ""
        self.estado_resaltado.configure(text=f"{resumen['comunes']} pruebas comunes · {resumen['celdas']} celdas diferentes{aviso}")

    def abrir_emergente(self) -> None:
        panel_seleccionado = self.ventana.focus_get()
        panel = next(
            (p for p in self.paneles if panel_seleccionado is not None and str(panel_seleccionado).startswith(str(p.marco))),
            None,
        )
        if panel is None:
            panel = next((p for p in self.paneles if p.ruta is not None), None)
        if panel is None or panel.ruta is None:
            messagebox.showinfo("Selecciona una especificación", "Primero carga un JSON en el panel que quieres abrir.", parent=self.ventana)
            return
        self.app.abrir_emergente(
            panel.ruta,
            panel.zoom,
            directorio_temporal=panel.directorio_temporal,
        )

    def cerrar(self) -> None:
        self._cerrando = True
        for panel in self.paneles:
            panel.limpiar()
        self.ventana.destroy()


class AplicacionEspecificaciones:
    def __init__(self, raiz: tk.Tk, zoom_aplicacion: float = 1.0):
        self.raiz = raiz
        self.zoom_aplicacion = zoom_aplicacion
        self._escala_tk_base = float(self.raiz.tk.call("tk", "scaling"))
        self.ventanas_flotantes: set[tk.Toplevel] = set()
        try:
            self.rutas = _cargar_configuracion()
        except (OSError, UnicodeError, json.JSONDecodeError, ValueError) as exc:
            messagebox.showerror("Configuración inválida", str(exc), parent=raiz)
            self.rutas = dict(RUTAS_PREDETERMINADAS)
        self.raiz.title("Gestor de especificaciones")
        self.raiz.geometry("1380x820")
        self.raiz.minsize(900, 600)
        self.raiz.configure(background=ui.paleta()["fondo"])
        self.zoom_aplicacion = ui.PREFERENCIAS["escala"]
        self._carga_en_proceso = False
        self._cola_carga: queue.Queue[tuple[str, str | int]] | None = None
        self._mensajes_carga: list[str] = []
        self._ventana_detalles_carga: tk.Toplevel | None = None
        self._texto_detalles_carga: ScrolledText | None = None
        self._reconstruccion_en_proceso = False
        self._acciones_tras_recuperacion = []
        self._crear_interfaz()
        from reconstruir_base_datos import necesita_recuperacion
        if necesita_recuperacion(BASE, rutas=self.rutas):
            self.reconstruir_base_datos(automatica=True)
        if ui.PREFERENCIAS["maximizar"]:
            self.raiz.after_idle(lambda: self.raiz.state("zoomed"))

    def _crear_interfaz(self) -> None:
        from menu_principal import crear_menu_principal
        crear_menu_principal(self)

    def reconstruir_base_datos(self, automatica=False):
        """Recuperación en segundo plano; todas las acciones de Tk se hacen en el hilo principal."""
        if self._reconstruccion_en_proceso or self._carga_en_proceso:
            messagebox.showinfo("Procesamiento en curso", "Espera a que termine el proceso actual.", parent=self.raiz)
            return
        from reconstruir_base_datos import reconstruir, asegurar_base
        self._reconstruccion_en_proceso = True
        rutas = dict(self.rutas)
        cola = queue.Queue()
        ventana = ui.Ventana(self.raiz)
        ventana.title("Recuperar base de datos")
        ventana.geometry("540x180")
        ventana.transient(self.raiz)
        ventana.grab_set()
        ventana.protocol("WM_DELETE_WINDOW", lambda: None)
        ttk.Label(ventana, text="Reconstruyendo desde los JSON existentes…", font=("Segoe UI", 13, "bold")).pack(padx=18, pady=(20, 10))
        estado = ttk.Label(ventana, text="Preparando recuperación…", wraplength=490)
        estado.pack(fill="x", padx=18, pady=8)
        barra = ttk.Progressbar(ventana, mode="indeterminate")
        barra.pack(fill="x", padx=18, pady=8)
        barra.start()

        def trabajar():
            try:
                recuperar = asegurar_base if automatica else reconstruir
                opciones = {} if automatica else {"reemplazar": True}
                resultado = recuperar(BASE, rutas=rutas, **opciones,
                    al_progreso=lambda actual, total, nombre: cola.put(("progreso", (actual, total, nombre))))
                cola.put(("terminado", resultado))
            except Exception as exc:
                cola.put(("error", str(exc)))

        def recibir():
            if not self.raiz.winfo_exists():
                return
            ultimo_progreso = None
            terminado = None
            while True:
                try:
                    tipo, dato = cola.get_nowait()
                except queue.Empty:
                    break
                if tipo == "progreso":
                    ultimo_progreso = dato
                else:
                    terminado = (tipo, dato)
            if ultimo_progreso:
                actual, total, nombre = ultimo_progreso
                estado.configure(text=f"{actual} de {total} JSON: {nombre}")
            if terminado is None:
                self.raiz.after(100, recibir)
                return
            barra.stop()
            ventana.grab_release()
            ventana.destroy()
            self._reconstruccion_en_proceso = False
            tipo, resultado = terminado
            if tipo == "error":
                self._acciones_tras_recuperacion.clear()
                messagebox.showerror("No se pudo reconstruir la base", resultado, parent=self.raiz)
                return
            menu = self.menu_principal
            if "Histórico" in menu.creadas:
                menu.reintentar_historial(menu.pestanas.tab("Histórico"))
            if resultado is not None:
                texto = (f"{resultado['documentos']} JSON registrados, {resultado['versiones']} versiones y "
                         f"{resultado['relaciones']} relaciones ESP/ENG.\n"
                         f"{resultado['wip']} borradores en progreso.\n\nSe conservaron los nombres y el contenido de los JSON.")
                if resultado["omitidos"]:
                    texto += f"\n\n{len(resultado['omitidos'])} archivos omitidos. Consulta REPORTE_RECONSTRUCCION_BASE.json."
                if resultado["respaldo"]:
                    texto += f"\n\nRespaldo anterior: {resultado['respaldo']}"
                messagebox.showinfo("Base reconstruida", texto, parent=self.raiz)
            acciones, self._acciones_tras_recuperacion = self._acciones_tras_recuperacion, []
            for accion in acciones:
                accion()

        threading.Thread(target=trabajar, daemon=True).start()
        self.raiz.after(100, recibir)

    def abrir_instrucciones(self) -> None:
        for nombre in ("README_ACTUALIZACION_CODIGOS_200.md", "README.md"):
            ruta = BASE / nombre
            if ruta.is_file():
                try:
                    abrir_en_windows = getattr(os, "startfile", None)
                    if callable(abrir_en_windows):
                        abrir_en_windows(str(ruta))
                    else:
                        _abrir_carpeta(BASE)
                except OSError as exc:
                    messagebox.showerror("No se pudo abrir", str(exc), parent=self.raiz)
                return
        messagebox.showerror("No encontrado", "No se encontró el archivo de instrucciones.", parent=self.raiz)

    def abrir_historico(self) -> None:
        try:
            VisorHistorico(
                self.raiz,
                BASE / "especificaciones.db",
                self.rutas["historico"],
            )
        except (OSError, sqlite3.Error) as exc:
            messagebox.showerror("No se pudo abrir el histórico", str(exc), parent=self.raiz)

    def subir_especificacion(self) -> None:
        if self._carga_en_proceso:
            messagebox.showinfo(
                "Procesamiento en curso",
                "Espera a que termine la carga anterior antes de iniciar otra.",
                parent=self.raiz,
            )
            return
        archivos = filedialog.askopenfilenames(
            parent=self.raiz,
            title="Selecciona especificaciones",
            filetypes=(
                ("Documentos Word", "*.doc *.docm *.docx"),
                ("Todos los archivos", "*.*"),
            ),
        )
        if not archivos:
            return
        destino = self.rutas["especificaciones_viejas"]
        self._carga_en_proceso = True
        cantidad_archivos = len(archivos)
        self._cola_carga = queue.Queue()
        self._mensajes_carga.clear()
        self._cerrar_detalles_carga()
        self._estado_progreso_carga.configure(
            text=f"Preparando {cantidad_archivos} especificación(es)..."
        )
        self._indicador_progreso_carga.configure(value=0)
        resultado: dict[str, Any] = {}
        terminado = threading.Event()

        def enviar(tipo: str, mensaje: str | int) -> None:
            if self._cola_carga is not None:
                self._cola_carga.put((tipo, mensaje))

        def actualizar_progreso(
            etapa: str,
            completados: int,
            total: int,
            nombre: str,
            inicio: int,
            fin: int,
        ) -> None:
            fraccion = completados / total if total else 1
            porcentaje = round(inicio + (fin - inicio) * fraccion)
            enviar("progreso", porcentaje)
            detalle = f": {Path(nombre).name}" if nombre else ""
            enviar("estado", f"{etapa} {completados}/{total}{detalle}")

        def procesar() -> None:
            try:
                import gestor_especificaciones as gestor

                destino.mkdir(parents=True, exist_ok=True)
                copiados = 0
                omitidos = 0
                bloqueados = 0
                errores_archivos = 0
                archivos_procesar: list[Path] = []
                for indice, archivo in enumerate(archivos, 1):
                    origen = Path(archivo)
                    enviar(
                        "estado",
                        f"Copiando especificación {indice}/{cantidad_archivos}: "
                        f"{origen.name}",
                    )
                    if origen.suffix.lower() not in EXTENSIONES_WORD:
                        omitidos += 1
                        enviar(
                            "mensaje",
                            f"OMITIDO (formato no compatible): {origen.name}. "
                            "Solo se procesan documentos Word.",
                        )
                        actualizar_progreso(
                            "Copiando",
                            indice,
                            cantidad_archivos,
                            origen.name,
                            0,
                            20,
                        )
                        continue
                    bloqueo = gestor.archivo_bloqueado_word(origen)
                    if bloqueo is not None:
                        bloqueados += 1
                        enviar(
                            "mensaje",
                            f"OMITIDO (Word bloqueado): {origen.name}; "
                            f"se encontró {bloqueo.name}. Cierra el documento y "
                            "vuelve a cargarlo cuando esté disponible.",
                        )
                        actualizar_progreso(
                            "Copiando",
                            indice,
                            cantidad_archivos,
                            origen.name,
                            0,
                            20,
                        )
                        continue
                    try:
                        candidato = destino / origen.name
                        numero = 1
                        while candidato.exists():
                            if candidato.read_bytes() == origen.read_bytes():
                                break
                            candidato = destino / (
                                f"{origen.stem}_{numero}{origen.suffix}"
                            )
                            numero += 1
                        if not candidato.exists():
                            shutil.copy2(origen, candidato)
                            copiados += 1
                            enviar("mensaje", f"COPIADO: {candidato.name}")
                        else:
                            enviar("mensaje", f"YA EXISTÍA: {candidato.name}")
                        archivos_procesar.append(candidato)
                    except (OSError, shutil.Error, ValueError) as exc:
                        errores_archivos += 1
                        enviar(
                            "mensaje",
                            f"ERROR al copiar {origen.name}: {exc}. "
                            "Se continuará con el siguiente archivo.",
                        )
                    actualizar_progreso(
                        "Copiando",
                        indice,
                        cantidad_archivos,
                        origen.name,
                        0,
                        20,
                    )
                resultado["copiados"] = copiados
                resultado["omitidos"] = omitidos
                resultado["bloqueados"] = bloqueados
                resultado["errores_archivos"] = errores_archivos
                if archivos_procesar:
                    resultado.update(
                        _procesar_carga_especificaciones(
                            BASE,
                            archivos_seleccionados=archivos_procesar,
                            al_mensaje=lambda mensaje: enviar("mensaje", mensaje),
                            al_progreso=lambda etapa, completados, total, nombre: (
                                actualizar_progreso(
                                    etapa,
                                    completados,
                                    total,
                                    nombre,
                                    20 if etapa == "Extracción" else 75,
                                    75 if etapa == "Extracción" else 100,
                                )
                            ),
                        )
                    )
                else:
                    resultado["extraccion"] = {
                        "nuevos": 0,
                        "actualizados": 0,
                        "ignorados": 0,
                        "errores": 0,
                    }
                    resultado["alineacion"] = {
                        "archivos": 0,
                        "alineados": 0,
                        "revision": 0,
                        "sin_coincidencia": 0,
                        "errores": 0,
                    }
                enviar("progreso", 100)
            except Exception as exc:
                resultado["error_flujo"] = str(exc)
                enviar("mensaje", f"ERROR del flujo de carga: {exc}")
            finally:
                terminado.set()

        def revisar_progreso() -> None:
            if self._cola_carga is not None:
                while True:
                    try:
                        tipo, mensaje = self._cola_carga.get_nowait()
                    except queue.Empty:
                        break
                    if tipo == "estado":
                        self._estado_progreso_carga.configure(text=str(mensaje))
                    elif tipo == "progreso":
                        self._indicador_progreso_carga.configure(value=int(mensaje))
                    else:
                        self._agregar_mensaje_carga(str(mensaje))
            if terminado.is_set():
                if self._cola_carga is not None:
                    while not self._cola_carga.empty():
                        tipo, mensaje = self._cola_carga.get_nowait()
                        if tipo == "estado":
                            self._estado_progreso_carga.configure(text=str(mensaje))
                        elif tipo == "progreso":
                            self._indicador_progreso_carga.configure(value=int(mensaje))
                        else:
                            self._agregar_mensaje_carga(str(mensaje))
                self._estado_progreso_carga.configure(text="Carga finalizada.")
                self._mostrar_resultado_carga(
                    resultado,
                    cantidad_archivos,
                    int(resultado.get("copiados", 0)),
                    destino,
                )
            else:
                self.raiz.after(100, revisar_progreso)

        threading.Thread(target=procesar, daemon=True).start()
        self.raiz.after(100, revisar_progreso)

    def _agregar_mensaje_carga(self, mensaje: str) -> None:
        self._mensajes_carga.append(mensaje)
        if self._carga_en_proceso and mensaje:
            self._estado_progreso_carga.configure(text=mensaje[:100])
        if self._texto_detalles_carga is not None:
            try:
                self._texto_detalles_carga.configure(state="normal")
                self._texto_detalles_carga.insert(tk.END, mensaje + "\n")
                self._texto_detalles_carga.see(tk.END)
                self._texto_detalles_carga.configure(state="disabled")
            except tk.TclError:
                self._texto_detalles_carga = None
                self._ventana_detalles_carga = None

    def _mostrar_detalles_carga(self) -> None:
        if (
            self._ventana_detalles_carga is not None
            and self._ventana_detalles_carga.winfo_exists()
        ):
            self._ventana_detalles_carga.deiconify()
            self._ventana_detalles_carga.lift()
            return
        ventana = ui.Ventana(self.raiz)
        ventana.title("Mensajes de carga")
        ventana.geometry("850x480")
        ventana.transient(self.raiz)
        texto = ScrolledText(
            ventana,
            wrap="word",
            state="normal",
            font=("Consolas", 9),
        )
        texto.pack(fill="both", expand=True, padx=8, pady=8)
        texto.insert("1.0", "\n".join(self._mensajes_carga))
        texto.see(tk.END)
        texto.configure(state="disabled")
        self._ventana_detalles_carga = ventana
        self._texto_detalles_carga = texto
        ventana.protocol("WM_DELETE_WINDOW", self._cerrar_detalles_carga)

    def _cerrar_detalles_carga(self) -> None:
        ventana = self._ventana_detalles_carga
        self._ventana_detalles_carga = None
        self._texto_detalles_carga = None
        if ventana is not None:
            try:
                if ventana.winfo_exists():
                    ventana.destroy()
            except tk.TclError:
                pass

    def _mostrar_resultado_carga(
        self,
        resultado: dict[str, Any],
        cantidad_archivos: int,
        cantidad_copiados: int,
        destino: Path,
    ) -> None:
        self._carga_en_proceso = False
        lineas = [
            f"Archivos recibidos: {cantidad_archivos}",
            f"Copias nuevas: {cantidad_copiados}",
            f"Destino: {destino}",
        ]
        omitidos = int(resultado.get("omitidos", 0))
        bloqueados = int(resultado.get("bloqueados", 0))
        errores_archivos = int(resultado.get("errores_archivos", 0))
        if omitidos:
            lineas.append(f"Formatos no compatibles omitidos: {omitidos}")
        if bloqueados:
            lineas.append(f"Documentos Word bloqueados omitidos: {bloqueados}")
        if errores_archivos:
            lineas.append(f"Errores al copiar archivos: {errores_archivos}")
        hay_errores = errores_archivos > 0
        error_flujo = resultado.get("error_flujo")
        if error_flujo:
            hay_errores = True
            lineas.append(f"Error del flujo: {error_flujo}")
        for etapa, etiqueta in (
            ("extraccion", "Extracción a JSON"),
            ("alineacion", "Alineación"),
        ):
            error = resultado.get(f"error_{etapa}")
            resumen = resultado.get(etapa)
            if error:
                hay_errores = True
                lineas.append(f"{etiqueta}: ERROR — {error}")
            elif resumen is not None:
                errores_etapa = int(resumen.get("errores", 0))
                hay_errores = hay_errores or errores_etapa > 0
                if etapa == "extraccion":
                    lineas.append(
                        f"{etiqueta}: {resumen.get('nuevos', 0)} nuevos, "
                        f"{resumen.get('actualizados', 0)} actualizados, "
                        f"{resumen.get('ignorados', 0)} sin cambios; "
                        f"{errores_etapa} error(es)."
                    )
                else:
                    lineas.append(
                        f"{etiqueta}: {resumen.get('archivos', 0)} archivos procesados; "
                        f"{errores_etapa} error(es)."
                    )
            else:
                hay_errores = True
                lineas.append(f"{etiqueta}: no se obtuvo resultado.")
        mostrar = messagebox.showwarning if hay_errores else messagebox.showinfo
        mostrar(
            "Carga terminada con errores" if hay_errores else "Carga terminada",
            "\n".join(lineas),
            parent=self.raiz,
        )

    def generar_hii(self) -> None:
        from generador_hii import GeneradorHII

        GeneradorHII(self.raiz, self.rutas["formatos"], self.rutas["json_alineadas"], app=self)

    def generar_especificacion(self) -> None:
        try:
            GeneradorEspecificacion(self)
        except (OSError, ValueError) as exc:
            messagebox.showerror(
                "No se pudo abrir el generador",
                f"No se pudo cargar la lista maestra o el catálogo CODIGOS_200.xlsx:\n{exc}",
                parent=self.raiz,
            )

    def editar_especificacion(
        self,
        archivo_inicial: Path | None = None,
        destino_inicial: Path | None = None,
    ) -> None:
        try:
            GeneradorEspecificacion(
                self,
                archivo_inicial=archivo_inicial,
                destino_inicial=destino_inicial,
            )
        except (OSError, ValueError) as exc:
            messagebox.showerror(
                "No se pudo abrir el generador",
                f"No se pudo cargar la lista maestra o el catálogo CODIGOS_200.xlsx:\n{exc}",
                parent=self.raiz,
            )

    def seleccionar_especificacion(self) -> None:
        from interfaz_seleccion_especificaciones import SelectorEspecificaciones

        def abrir_editor_json(ruta: Path) -> None:
            EditorJson(self, ruta)

        ventana = ui.Ventana(self.raiz)
        SelectorEspecificaciones(
            ventana,
            carpeta_json=self.rutas["json_alineadas"],
            carpeta_json_anterior=self.rutas["json_anteriores"],
            al_actualizar=abrir_editor_json,
            al_editar=lambda ruta: self.editar_especificacion(ruta, ruta),
            al_comparar=self.comparar_especificaciones,
        )

    def comparar_especificaciones(self, rutas: list[Path] | None = None) -> None:
        if rutas is None:
            Comparador(self)
            return
        if len(rutas) > 3:
            if not messagebox.askyesno("Abrir más de 3 especificaciones", f"Se abrirán {len(rutas)} especificaciones en ventanas flotantes. ¿Deseas continuar?", parent=self.raiz):
                return
            for ruta in rutas:
                self.abrir_emergente(ruta, 1.44)
            return
        comparador = Comparador(self)
        for panel, ruta in zip(comparador.paneles, rutas):
            panel.cargar_archivo(ruta)

    def actualizar_formatos(self) -> None:
        archivos = filedialog.askopenfilenames(
            parent=self.raiz,
            title="Selecciona plantillas Word para FORMATOS",
            filetypes=(("Plantillas Word y Excel", "*.doc *.docm *.docx *.xlsx *.xlsm"),),
        )
        if not archivos:
            return
        destino = self.rutas["formatos"]
        try:
            destino.mkdir(parents=True, exist_ok=True)
            nombres = []
            for archivo in archivos:
                origen = Path(archivo)
                if origen.suffix.lower() not in {*EXTENSIONES_WORD, ".xlsx", ".xlsm"}:
                    raise ValueError(f"Extensión no compatible: {origen.name}")
                shutil.copy2(origen, destino / origen.name)
                nombres.append(origen.name)
        except (OSError, ValueError) as exc:
            messagebox.showerror("No se pudieron actualizar los formatos", str(exc), parent=self.raiz)
            return
        messagebox.showinfo("Formatos actualizados", f"{len(nombres)} archivo(s) copiado(s) en:\n{destino}", parent=self.raiz)

    def editar_traducciones(self):
        self._cargar_traducciones_desde_metodos(lambda: EditorTraducciones(self.raiz, self._refrescar_fechas_listas))

    def _refrescar_fechas_listas(self):
        if hasattr(self, "menu_principal"):
            self.menu_principal.refrescar_fechas_metodos()

    def actualizar_traducciones(self):
        self._cargar_traducciones_desde_metodos()

    def _cargar_traducciones_desde_metodos(self, al_terminar=None, avisar=True):
        if getattr(self, "_traducciones_cargando", False):
            return
        self._traducciones_cargando = True
        cola = queue.Queue()
        def cargar():
            try:
                cola.put((True, actualizar_traductor()))
            except Exception as exc:
                cola.put((False, str(exc)))
        def recibir():
            if not self.raiz.winfo_exists():
                return
            try:
                correcto, resultado = cola.get_nowait()
            except queue.Empty:
                self.raiz.after(100, recibir)
                return
            self._traducciones_cargando = False
            if correcto:
                self._refrescar_fechas_listas()
            if not correcto:
                messagebox.showwarning("No se actualizaron las traducciones", f"Se conserva la última versión válida.\n\n{resultado}", parent=self.raiz)
            elif al_terminar is not None:
                al_terminar()
            elif avisar:
                messagebox.showinfo("Traducciones actualizadas", f"{resultado['equivalencias']} equivalencias activas.\n{resultado['cambios']} cambios registrados.\nLos editores abiertos recibirán las nuevas equivalencias.", parent=self.raiz)
        threading.Thread(target=cargar, daemon=True).start()
        self.raiz.after(100, recibir)

    def opciones_traducciones(self):
        from traductor_especificaciones import ARCHIVO_ESTADO, ARCHIVO_HISTORIAL
        CARPETA = carpeta_listas(BASE)
        ventana = ui.Ventana(self.raiz)
        ventana.title("Traducciones · ubicación e historial")
        ventana.transient(self.raiz)
        ventana.grab_set()
        ttk.Label(ventana, text="Lista de traducciones", font=("Segoe UI", 14, "bold")).pack(anchor="w", padx=16, pady=12)
        ttk.Label(ventana, text="Consulta la ubicación, el historial y la última versión de las traducciones.", wraplength=450).pack(anchor="w", padx=16, pady=8)
        def abrir(ruta):
            try:
                if not ruta.exists():
                    raise FileNotFoundError(f"No existe el archivo:\n{ruta}")
                if os.name == "nt":
                    os.startfile(str(ruta))
                else:
                    subprocess.Popen(["open" if sys.platform == "darwin" else "xdg-open", str(ruta)])
            except OSError as exc:
                messagebox.showerror("No se pudo abrir", str(exc), parent=ventana)
        for texto, ruta in (("Abrir carpeta de traducciones", CARPETA),
                            ("Consultar historial de cambios", CARPETA / ARCHIVO_HISTORIAL),
                            ("Consultar última versión registrada", CARPETA / ARCHIVO_ESTADO)):
            ttk.Button(ventana, text=texto, command=lambda r=ruta: abrir(r)).pack(fill="x", padx=16, pady=4)
        ttk.Button(ventana, text="Cerrar", command=ventana.destroy).pack(anchor="e", padx=16, pady=12)

    def actualizar_metodos(self) -> None:
        ventana = ui.Ventana(self.raiz)
        ventana.title("Lista maestra de métodos")
        ventana.transient(self.raiz)
        ventana.resizable(False, False)
        ventana.grab_set()
        ttk.Label(
            ventana,
            text="Consulta y edición manual de la lista maestra",
            font=("Segoe UI", 11, "bold"),
            wraplength=440,
        ).pack(anchor="w", padx=16, pady=(16, 12))
        ttk.Label(
            ventana,
            text=(
                "Puedes editar los siete campos aquí o abrir la ubicación del Excel "
                "para mantener la lista manualmente."
            ),
            wraplength=440,
            justify="left",
        ).pack(anchor="w", padx=16, pady=(0, 10))

        acciones = ttk.Frame(ventana)
        acciones.pack(fill="x", padx=16)

        def editar_aqui() -> None:
            ventana.destroy()
            try:
                EditorListaMaestra(self.raiz, ruta_maestro(BASE), self._refrescar_fechas_listas)
            except (OSError, ValueError, ImportError) as exc:
                messagebox.showerror(
                    "No se pudo abrir la lista maestra",
                    str(exc),
                    parent=self.raiz,
                )

        def abrir_ubicacion() -> None:
            ventana.destroy()
            ruta = ruta_maestro(BASE)
            if not ruta.is_file():
                messagebox.showerror(
                    "No se encontró el Excel",
                    f"No se encontró la lista maestra:\n{ruta}",
                    parent=self.raiz,
                )
                return
            try:
                if os.name == "nt":
                    subprocess.Popen(["explorer", f"/select,{ruta}"])
                elif sys.platform == "darwin":
                    subprocess.Popen(["open", "-R", str(ruta)])
                else:
                    subprocess.Popen(["xdg-open", str(ruta.parent)])
            except OSError as exc:
                messagebox.showerror(
                    "No se pudo abrir la ubicación",
                    str(exc),
                    parent=self.raiz,
                )

        ttk.Button(
            acciones,
            text="Editar la lista maestra aquí",
            command=editar_aqui,
        ).pack(fill="x", pady=4)
        ttk.Button(
            acciones,
            text="Ir a la ubicación del Excel",
            command=abrir_ubicacion,
        ).pack(fill="x", pady=4)
        ttk.Button(
            ventana,
            text="Cancelar",
            command=ventana.destroy,
        ).pack(anchor="e", padx=16, pady=12)

    def configuracion(self) -> None:
        ventana = ui.Ventana(self.raiz)
        ventana.title("Configuración de carpetas")
        ventana.transient(self.raiz)
        ventana.grab_set()
        campos: dict[str, tk.StringVar] = {}
        nombres = {
            "especificaciones_viejas": "Documentos Word recibidos",
            "json_anteriores": "JSON extraídos",
            "json_alineadas": "JSON alineados",
            "formatos": "Plantillas Word",
            "historico": "Documentos generados",
            "metodos_lab": "Lista de métodos y traducciones",
        }
        for fila, (clave, etiqueta) in enumerate(nombres.items()):
            ttk.Label(ventana, text=etiqueta).grid(row=fila, column=0, sticky="w", padx=10, pady=7)
            variable = tk.StringVar(value=ruta_portatil(self.rutas[clave], BASE))
            campos[clave] = variable
            ttk.Entry(ventana, textvariable=variable, width=64).grid(row=fila, column=1, padx=5, pady=7)
            ttk.Button(
                ventana,
                text="Examinar",
                command=lambda v=variable: self._elegir_carpeta(v, ventana),
            ).grid(row=fila, column=2, padx=5, pady=7)

        def guardar() -> None:
            nuevas = {clave: resolver_ruta(variable.get(), BASE) for clave, variable in campos.items()}
            try:
                _guardar_configuracion(nuevas)
            except (OSError, ValueError) as exc:
                messagebox.showerror("No se pudo guardar", str(exc), parent=ventana)
                return
            self.rutas = nuevas
            self._refrescar_fechas_listas()
            self._cargar_traducciones_desde_metodos(avisar=False)
            ventana.destroy()

        ttk.Button(ventana, text="Guardar", command=guardar).grid(row=len(nombres), column=1, sticky="e", padx=8, pady=12)
        ttk.Button(ventana, text="Cancelar", command=ventana.destroy).grid(row=len(nombres), column=2, padx=8, pady=12)

    @staticmethod
    def _elegir_carpeta(variable: tk.StringVar, parent: tk.Misc) -> None:
        ruta = filedialog.askdirectory(parent=parent, initialdir=str(resolver_ruta(variable.get() or BASE)))
        if ruta:
            variable.set(ruta)

    def cambiar_zoom(self, cambio: float) -> None:
        self.zoom_aplicacion = min(1.4, max(0.8, self.zoom_aplicacion + cambio))
        ui.aplicar_preferencias({**ui.PREFERENCIAS, "escala": self.zoom_aplicacion})

    def abrir_emergente(
        self,
        ruta: Path,
        zoom_inicial: float,
        *,
        directorio_temporal: tempfile.TemporaryDirectory[str] | None = None,
    ) -> None:
        vivas = {ventana for ventana in self.ventanas_flotantes if ventana.winfo_exists()}
        self.ventanas_flotantes = vivas
        if len(vivas) >= 10:
            messagebox.showinfo("Límite alcanzado", "Se permiten hasta 10 vistas emergentes.", parent=self.raiz)
            return
        temporal_emergente = None
        if directorio_temporal is not None:
            temporal_emergente = tempfile.TemporaryDirectory(
                prefix="comparacion_emergente_"
            )
            ruta_emergente = Path(temporal_emergente.name) / ruta.name
            try:
                shutil.copy2(ruta, ruta_emergente)
            except OSError as exc:
                temporal_emergente.cleanup()
                messagebox.showerror(
                    "No se pudo abrir la vista emergente",
                    str(exc),
                    parent=self.raiz,
                )
                return
            ruta = ruta_emergente
        ventana = ui.Ventana(self.raiz)
        ventana.title(f"Vista emergente · {ruta.name}")
        ventana.geometry("900x680")
        ventana.minsize(600, 400)
        self.ventanas_flotantes.add(ventana)
        if temporal_emergente is not None:
            ventana.bind(
                "<Destroy>",
                lambda evento, temporal=temporal_emergente: temporal.cleanup()
                if evento.widget is ventana
                else None,
                add="+",
            )
        zoom = [zoom_inicial]
        barra = ttk.Frame(ventana, padding=6)
        barra.pack(fill="x")
        marco = ttk.Frame(ventana)
        marco.pack(fill="both", expand=True)

        def actualizar(cambio: float) -> None:
            zoom[0] = min(1.8, max(0.7, zoom[0] + cambio))
            mostrar_vista_json(marco, ruta, zoom[0])

        ttk.Button(barra, text="−", width=3, command=lambda: actualizar(-0.1)).pack(side="left")
        ttk.Button(barra, text="+", width=3, command=lambda: actualizar(0.1)).pack(side="left", padx=4)
        ttk.Label(barra, text=str(ruta)).pack(side="left", padx=10)
        ventana.bind("<Destroy>", lambda evento: self.ventanas_flotantes.discard(ventana) if evento.widget is ventana else None)
        mostrar_vista_json(marco, ruta, zoom[0])


def main() -> None:
    parser = argparse.ArgumentParser(description="Gestor gráfico de especificaciones.")
    parser.add_argument(
        "--editar-json",
        type=Path,
        help="abre directamente el JSON indicado en el editor de especificaciones",
    )
    parser.add_argument(
        "--generar-json",
        type=Path,
        help="abre en el generador el JSON indicado",
    )
    args = parser.parse_args()
    raiz = ui.Raiz()
    app = AplicacionEspecificaciones(raiz)
    def abrir_inicial():
        if args.editar_json is not None:
            EditorJson(app, resolver_ruta(args.editar_json))
        if args.generar_json is not None:
            ruta = resolver_ruta(args.generar_json)
            app.editar_especificacion(ruta, ruta)
    if app._reconstruccion_en_proceso:
        app._acciones_tras_recuperacion.append(abrir_inicial)
    else:
        abrir_inicial()
    raiz.mainloop()


if __name__ == "__main__":
    main()
