#!/usr/bin/env python3
"""Genera especificaciones Word a partir de JSON y plantillas por idioma."""
from __future__ import annotations

import argparse
import copy
import json
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata
from datetime import date
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from rutas_proyecto import BASE, resolver_ruta, rutas_configuradas


EXTENSIONES_WORD = {".doc", ".docx", ".docm"}
MARCADOR_INGLES = re.compile(r"(?<![A-Za-z0-9])ENG(?![A-Za-z0-9])", re.IGNORECASE)
CODIGO_JSON_IGNORADO = "20000000"
TIPO_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"
IDIOMAS = {
    "en": "INGLES",
    "english": "INGLES",
    "ingles": "INGLES",
    "es": "ESPANOL",
    "espanol": "ESPANOL",
    "spanish": "ESPANOL",
}
COLUMNAS_PRUEBAS = {
    "PRUEBA": "prueba",
    "TEST": "prueba",
    "UNIDAD": "unidades",
    "UNIDADES": "unidades",
    "UNIT": "unidades",
    "UNITS": "unidades",
    "METODO": "metodo",
    "METHOD": "metodo",
    "TEST METHOD": "metodo",
    "ESPECIFICACION": "especificacion",
    "SPECIFICATION": "especificacion",
    "SPECIFICATIONS": "especificacion",
    "VALOR TIPICO": "valor_tipico",
    "VALORES TIPICOS": "valor_tipico",
    "TYPICAL VALUE": "valor_tipico",
    "TYPICAL VALUES": "valor_tipico",
}
ENCABEZADOS_FILA_PRUEBA = {"PRUEBA", "TEST"}
ENCABEZADOS_FILA_METODO = {
    "METODO",
    "METHOD",
    "PRUEBA METODO",
    "TEST METHOD",
}
ENCABEZADOS_FILA_PRUEBA_METODO = {"PRUEBA METODO", "TEST METHOD"}
ENCABEZADOS_FILA_ESPECIFICACION = {
    "ESPECIFICACION",
    "SPECIFICATION",
    "SPECIFICATIONS",
}
ETIQUETAS_CONTROL = {
    "ELABORO Y APROBO": "elaboro_y_aprobo",
    "ELABORADO Y APROBADO POR": "elaboro_y_aprobo",
    "ELABORATED AND APPROVED BY": "elaboro_y_aprobo",
    "PREPARED AND APPROVED BY": "elaboro_y_aprobo",
    "FOLIO APQP": "folio_apqp",
    "APQP FOLIO": "folio_apqp",
    "FOLIO VAP": "folio_vap",
    "VAP FOLIO": "folio_vap",
    "CODIGO UNICO DE PRODUCTO": "codigo_unico_producto",
    "CODIGO DE PRODUCTO": "codigo_unico_producto",
    "PART ID": "codigo_unico_producto",
    "PRODUCT CODE": "codigo_unico_producto",
    "UNIQUE PRODUCT CODE": "codigo_unico_producto",
    "FECHA": "fecha",
    "DATE": "fecha",
    "REVISION": "revision",
    "REV": "revision",
    "NOMBRE CLIENTE": "nombre_cliente",
    "CUSTOMER NAME": "nombre_cliente",
    "CUSTOMERS NAME": "nombre_cliente",
    "CLIENT NAME": "nombre_cliente",
    "REVISION CLIENTE": "revision_cliente",
    "CUSTOMER REVISION": "revision_cliente",
    "CUSTOMERS REVISION NUMBER": "revision_cliente",
    "CUSTOMER S REVISION NUMBER": "revision_cliente",
    "CLIENT REVISION": "revision_cliente",
}
def es_documento_word(path: Path) -> bool:
    return path.is_file() and path.suffix.lower() in EXTENSIONES_WORD


def es_formato_ingles(path: Path) -> bool:
    return es_documento_word(path) and bool(MARCADOR_INGLES.search(path.stem))


def clave_texto(valor: str) -> str:
    normalizado = "".join(
        caracter for caracter in unicodedata.normalize("NFD", valor.casefold())
        if unicodedata.category(caracter) != "Mn"
    ).upper()
    return re.sub(r"[^A-Z0-9]+", " ", normalizado).strip()


def detectar_formatos(carpeta: Path) -> list[Path]:
    if not carpeta.is_dir():
        raise FileNotFoundError(f"No existe la carpeta de formatos: {carpeta}")
    return sorted(
        (path for path in carpeta.rglob("*") if es_documento_word(path)),
        key=lambda path: str(path.relative_to(carpeta)).casefold(),
    )


def detectar_json(carpeta: Path) -> list[Path]:
    if not carpeta.is_dir():
        raise FileNotFoundError(f"No existe la carpeta de JSON: {carpeta}")
    return sorted(
        (path for path in carpeta.rglob("*") if path.is_file() and path.suffix.lower() == ".json"),
        key=lambda path: str(path.relative_to(carpeta)).casefold(),
    )


def leer_idioma_json(path: Path) -> str:
    with path.open(encoding="utf-8") as archivo:
        datos = json.load(archivo)
    if not isinstance(datos, dict):
        raise ValueError("la raiz del JSON debe ser un objeto")
    return leer_idioma_datos(datos)


def leer_datos_json(path: Path) -> dict:
    with path.open(encoding="utf-8") as archivo:
        datos = json.load(archivo)
    if not isinstance(datos, dict):
        raise ValueError("la raiz del JSON debe ser un objeto")
    leer_idioma_datos(datos)
    return datos


def leer_idioma_datos(datos: dict) -> str:
    documento = datos.get("documento")
    if not isinstance(documento, dict):
        raise ValueError("falta el objeto 'documento'")
    idioma = documento.get("idioma")
    if not isinstance(idioma, str) or not idioma.strip():
        raise ValueError("falta el campo de texto 'documento.idioma'")
    idioma_normalizado = clave_texto(idioma).replace(" ", "").casefold()
    try:
        return IDIOMAS[idioma_normalizado]
    except KeyError as exc:
        raise ValueError(f"idioma no compatible en 'documento.idioma': {idioma!r}") from exc


def seleccionar_plantilla(plantillas: list[Path], idioma: str) -> Path:
    coincidencias = [
        plantilla for plantilla in plantillas
        if es_formato_ingles(plantilla) == (idioma == "INGLES")
    ]
    if not coincidencias:
        raise FileNotFoundError(f"No hay una plantilla Word para el idioma {idioma}.")
    if len(coincidencias) > 1:
        nombres = ", ".join(plantilla.name for plantilla in coincidencias)
        raise ValueError(
            f"Hay varias plantillas para {idioma}; no se puede elegir una por idioma solamente: {nombres}"
        )
    return coincidencias[0]


def _convertir_docm_temporal(origen: Path, destino: Path) -> None:
    """Crea un paquete DOCX temporal sin VBA; el DOCM fuente no se modifica."""
    with ZipFile(origen) as entrada, ZipFile(destino, "w", ZIP_DEFLATED) as salida:
        for item in entrada.infolist():
            nombre = item.filename
            if (
                nombre.startswith("word/vbaProject")
                or nombre.startswith("word/vbaData")
                or nombre == "word/vbaData.xml"
            ):
                continue
            contenido = entrada.read(nombre)
            if nombre == "[Content_Types].xml":
                texto = contenido.decode("utf-8")
                texto = texto.replace(
                    "application/vnd.ms-word.document.macroEnabled.main+xml", TIPO_DOCX
                )
                texto = re.sub(
                    r'<Default\b(?=[^>]*\bExtension="bin")[^>]*/>', "", texto
                )
                texto = re.sub(
                    r'<Override\b(?=[^>]*\bPartName="/word/vba[^"]*")[^>]*/>', "", texto
                )
                contenido = texto.encode("utf-8")
            elif nombre == "word/_rels/document.xml.rels":
                texto = contenido.decode("utf-8")
                texto = re.sub(
                    r'<Relationship\b(?=[^>]*\bType="[^"]*/vbaProject")[^>]*/>', "", texto
                )
                contenido = texto.encode("utf-8")
            salida.writestr(item, contenido)


def preparar_plantilla(path: Path, temporal: Path) -> Path:
    if path.suffix.lower() == ".docx":
        return path
    destino = temporal / f"{path.stem}.docx"
    if path.suffix.lower() == ".docm":
        _convertir_docm_temporal(path, destino)
        return destino

    if path.suffix.lower() == ".doc":
        ejecutable = shutil.which("soffice") or shutil.which("libreoffice")
        if ejecutable:
            proceso = subprocess.run(
                [ejecutable, "--headless", "--convert-to", "docx", "--outdir", str(temporal), str(path)],
                capture_output=True,
                text=True,
                timeout=120,
                check=False,
            )
            if proceso.returncode == 0 and destino.exists():
                return destino
            detalle = proceso.stderr.strip() or proceso.stdout.strip()
            raise RuntimeError(f"No se pudo convertir la plantilla .doc: {detalle or proceso.returncode}")
        raise RuntimeError("Para usar una plantilla .doc se requiere LibreOffice (soffice).")
    raise ValueError(f"Formato de plantilla no compatible: {path.suffix}")


def _poner_texto(parrafo, texto: str) -> None:
    corridas = parrafo.runs
    if corridas:
        corridas[0].text = texto
        for corrida in corridas[1:]:
            corrida.text = ""
    else:
        parrafo.add_run(texto)


def _buscar_tabla_pruebas(documento):
    requeridas = {"prueba", "unidades", "metodo", "especificacion"}
    for tabla in documento.tables:
        if not tabla.rows:
            continue
        mapeo = [COLUMNAS_PRUEBAS.get(clave_texto(celda.text)) for celda in tabla.rows[0].cells]
        if requeridas.issubset(set(mapeo)):
            return tabla, mapeo
    raise ValueError("No se encontro una tabla de pruebas reconocible por sus encabezados.")


def _buscar_tabla_control(documento, tabla_pruebas):
    candidatas = []
    for tabla in documento.tables:
        if tabla is tabla_pruebas or len(tabla.columns) < 2:
            continue
        coincidencias = sum(
            ETIQUETAS_CONTROL.get(clave_texto(fila.cells[0].text)) is not None
            for fila in tabla.rows
        )
        if coincidencias:
            candidatas.append((coincidencias, tabla))
    if not candidatas:
        raise ValueError("No se encontro la tabla de control documental.")
    candidatas.sort(key=lambda item: item[0], reverse=True)
    if len(candidatas) > 1 and candidatas[0][0] == candidatas[1][0]:
        raise ValueError("Se encontraron varias tablas candidatas de control documental.")
    return candidatas[0][1]


def _es_fila_notas(fila) -> bool:
    return any(
        re.search(r"\b(?:NOTAS?|NOTES?)\b", clave_texto(celda.text))
        for celda in fila.cells
    )


def _llenar_pruebas(tabla, mapeo: list[str | None], pruebas: list[dict]) -> None:
    filas = list(tabla.rows)
    filas_notas = [fila for fila in filas[1:] if _es_fila_notas(fila)]
    if len(filas_notas) != 1:
        raise ValueError("La tabla de pruebas debe tener una sola fila de notas identificable.")
    fila_notas = filas_notas[0]
    filas_datos = [fila for fila in filas[1:] if fila is not fila_notas]

    while len(filas_datos) < len(pruebas):
        fila_nueva = copy.deepcopy(filas_datos[-1]._tr if filas_datos else filas[0]._tr)
        fila_notas._tr.addprevious(fila_nueva)
        filas_datos = [fila for fila in tabla.rows[1:] if fila._tr is not fila_notas._tr]

    for indice, fila in enumerate(filas_datos):
        prueba = pruebas[indice] if indice < len(pruebas) else {}
        for columna, campo in enumerate(mapeo):
            if campo is None or columna >= len(fila.cells):
                continue
            valor = prueba.get(campo, "") if isinstance(prueba, dict) else ""
            _poner_texto(fila.cells[columna].paragraphs[0], "" if valor is None else str(valor))


def _normalizar_nota(texto: str) -> str:
    texto = re.sub(r"^(?:Notas?|Notes?)\s*:\s*", "", texto.strip(), flags=re.IGNORECASE)
    return clave_texto(texto)


def _es_fila_encabezado_pruebas(prueba: dict) -> bool:
    encabezado_prueba = clave_texto(str(prueba.get("prueba", "") or ""))
    encabezado_metodo = clave_texto(str(prueba.get("metodo", "") or ""))
    encabezado_especificacion = clave_texto(
        str(prueba.get("especificacion", "") or "")
    )
    tiene_especificacion = encabezado_especificacion in ENCABEZADOS_FILA_ESPECIFICACION
    if not tiene_especificacion:
        return False
    if encabezado_prueba in ENCABEZADOS_FILA_PRUEBA_METODO:
        return True
    return (
        encabezado_prueba in ENCABEZADOS_FILA_PRUEBA
        and encabezado_metodo in ENCABEZADOS_FILA_METODO
    )


def _separar_notas_de_pruebas(pruebas: list[dict]) -> tuple[list[dict], list[str]]:
    pruebas_validas = []
    notas = []
    marcador_nota = re.compile(r"^(?:Notas?|Notes?)\s*:\s*(.*)$", re.IGNORECASE)
    campos = ("prueba", "unidades", "metodo", "especificacion", "valor_tipico")
    for prueba in pruebas:
        if _es_fila_encabezado_pruebas(prueba):
            continue
        valores = [
            str(prueba.get(campo, "") or "").strip()
            for campo in campos
            if str(prueba.get(campo, "") or "").strip()
        ]
        coincidencia = (
            marcador_nota.match(valores[0])
            if valores and len(set(valores)) == 1
            else None
        )
        if coincidencia:
            nota = coincidencia.group(1).strip()
            if nota:
                notas.append(nota)
        else:
            pruebas_validas.append(prueba)
    return pruebas_validas, notas


def _escribir_notas(tabla, notas: list) -> None:
    fila = next((fila for fila in tabla.rows if _es_fila_notas(fila)), None)
    if fila is None:
        raise ValueError("No se encontro la fila de notas en la tabla de pruebas.")
    celda = fila.cells[0]
    texto_plantilla = celda.text.strip()
    clave_plantilla = _normalizar_nota(texto_plantilla)
    valores = []
    claves_vistas = {clave_plantilla} if clave_plantilla else set()
    for nota in notas:
        if nota is None or not str(nota).strip():
            continue
        nota_texto = str(nota).strip()
        clave_nota = _normalizar_nota(nota_texto)
        if not clave_nota or clave_nota in claves_vistas:
            continue
        claves_vistas.add(clave_nota)
        valores.append(re.sub(r"^(?:Notas?|Notes?)\s*:\s*", "", nota_texto, flags=re.IGNORECASE))
    if not valores:
        return
    separador = "\n" if texto_plantilla else ""
    notas_adicionales = "\n".join(valores)
    _poner_texto(celda.paragraphs[0], f"{texto_plantilla}{separador}{notas_adicionales}")
    for parrafo in celda.paragraphs[1:]:
        _poner_texto(parrafo, "")


def preparar_control_generacion(control: dict, fecha_generacion: date | None = None) -> dict:
    resultado = dict(control)
    revision = resultado.get("revision")
    if isinstance(revision, bool) or not isinstance(revision, (str, int)):
        raise ValueError("control_documental.revision debe ser un entero o texto numerico.")
    revision_texto = str(revision).strip()
    if not re.fullmatch(r"[0-9]+", revision_texto):
        raise ValueError(f"revision no numerica en control_documental.revision: {revision!r}")

    fecha = fecha_generacion or date.today()
    resultado["fecha"] = fecha.strftime("%d/%m/%Y")
    resultado["revision"] = f"{int(revision_texto) + 1:02d}"
    return resultado


def _llenar_control(tabla, control: dict) -> None:
    for fila in tabla.rows:
        etiqueta = ETIQUETAS_CONTROL.get(clave_texto(fila.cells[0].text))
        if etiqueta is None or len(fila.cells) < 2:
            continue
        if etiqueta not in control:
            continue
        valor = control[etiqueta]
        _poner_texto(fila.cells[1].paragraphs[0], "" if valor is None else str(valor))


def _reemplazar_texto_parte(parte, buscar: str, reemplazo: str) -> bool:
    patron = re.compile(buscar, re.IGNORECASE)
    for elemento in parte._element.xpath('.//*[local-name()="p"]'):
        nodos = elemento.xpath('.//*[local-name()="t"]')
        texto = "".join(nodo.text or "" for nodo in nodos)
        coincidencia = patron.search(texto)
        if not coincidencia:
            continue
        texto_nuevo = texto[:coincidencia.start()] + reemplazo + texto[coincidencia.end():]
        if nodos:
            nodos[0].text = texto_nuevo
            for nodo in nodos[1:]:
                nodo.text = ""
            if texto_nuevo.startswith(" ") or texto_nuevo.endswith(" "):
                nodos[0].set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        return True
    return False


def _poner_nombre_producto(documento, nombre: str) -> None:
    encontrado = False
    for seccion in documento.sections:
        encontrado = _reemplazar_texto_parte(
            seccion.header, r"ESPECIFICACI[ÓO]N\s+NOMBRE", f"ESPECIFICACIÓN {nombre}"
        ) or encontrado
        if not encontrado:
            encontrado = _reemplazar_texto_parte(
                seccion.header, r"\bNOMBRE\b", nombre
            ) or encontrado
    if not encontrado:
        raise ValueError("No se encontro el campo del nombre del producto en el encabezado de la plantilla.")


def _poner_pie_pagina(documento, texto: str) -> None:
    if not texto.strip():
        return
    for seccion in documento.sections:
        parrafos = [
            (len("".join(nodo.text or "" for nodo in p.xpath('.//*[local-name()="t"]'))), p)
            for p in seccion.footer._element.xpath('.//*[local-name()="p"]')
        ]
        if not parrafos:
            continue
        _, parrafo = max(parrafos, key=lambda item: item[0])
        nodos = parrafo.xpath('.//*[local-name()="t"]')
        if nodos:
            nodos[0].text = texto
            for nodo in nodos[1:]:
                nodo.text = ""
        else:
            raise ValueError("El pie de pagina de la plantilla no tiene un campo de texto editable.")


def _nombre_destino(carpeta: Path, nombre: str) -> Path:
    destino = carpeta / f"{nombre}.docx"
    indice = 1
    while destino.exists():
        destino = carpeta / f"{nombre}_{indice:02d}.docx"
        indice += 1
    return destino


def generar_documento(datos: dict, plantilla: Path, destino: Path) -> None:
    leer_idioma_datos(datos)
    with tempfile.TemporaryDirectory(prefix="generador_especs_hii_") as directorio_temporal:
        archivo_plantilla = preparar_plantilla(plantilla, Path(directorio_temporal))
        documento = Document(archivo_plantilla)

    producto = datos.get("producto")
    nombre = producto.get("nombre") if isinstance(producto, dict) else None
    if not isinstance(nombre, str) or not nombre.strip():
        raise ValueError("falta el campo de texto 'producto.nombre'")
    pruebas = datos.get("pruebas")
    if not isinstance(pruebas, list) or any(not isinstance(item, dict) for item in pruebas):
        raise ValueError("el campo 'pruebas' debe ser una lista de objetos")
    pruebas, notas_en_pruebas = _separar_notas_de_pruebas(pruebas)
    control = datos.get("control_documental", {})
    if not isinstance(control, dict):
        raise ValueError("el campo 'control_documental' debe ser un objeto")
    control = preparar_control_generacion(control)

    tabla_pruebas, mapeo = _buscar_tabla_pruebas(documento)
    tabla_control = _buscar_tabla_control(documento, tabla_pruebas)
    _poner_nombre_producto(documento, nombre.strip())
    _llenar_pruebas(tabla_pruebas, mapeo, pruebas)
    notas = datos.get("notas", [])
    if not isinstance(notas, list):
        raise ValueError("el campo 'notas' debe ser una lista")
    _escribir_notas(tabla_pruebas, [*notas, *notas_en_pruebas])
    _llenar_control(tabla_control, control)

    institucionales = datos.get("datos_institucionales", {})
    if not isinstance(institucionales, dict):
        raise ValueError("el campo 'datos_institucionales' debe ser un objeto")
    pie = institucionales.get("pie_pagina")
    if pie is not None and not isinstance(pie, str):
        raise ValueError("el campo 'datos_institucionales.pie_pagina' debe ser texto")
    if isinstance(pie, str):
        _poner_pie_pagina(documento, pie)
    documento.save(destino)


def main(argv: list[str] | None = None) -> int:
    rutas = rutas_configuradas(BASE)
    parser = argparse.ArgumentParser(
        description=(
            "Genera un archivo DOCX por JSON usando la plantilla correspondiente a documento.idioma."
        )
    )
    parser.add_argument(
        "--formatos",
        type=lambda valor: resolver_ruta(valor, BASE),
        default=rutas["formatos"],
        help="carpeta a revisar (por defecto: FORMATOS junto al programa)",
    )
    parser.add_argument(
        "--json-dir",
        type=lambda valor: resolver_ruta(valor, BASE),
        default=rutas["json_anteriores"],
        help="carpeta de JSON (por defecto: ESPECIFICACIONES_JSON junto al programa)",
    )
    parser.add_argument(
        "--salida",
        type=lambda valor: resolver_ruta(valor, BASE),
        default=rutas["historico"],
        help="carpeta para los DOCX generados (por defecto: HISTORICO_ESPECIFICACIONES)",
    )
    args = parser.parse_args(argv)

    try:
        archivos_json = detectar_json(args.json_dir)
    except FileNotFoundError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1

    ignorados = [path for path in archivos_json if CODIGO_JSON_IGNORADO in path.name]
    archivos_json = [path for path in archivos_json if CODIGO_JSON_IGNORADO not in path.name]
    if not archivos_json:
        print(
            f"No hay JSON elegibles. Ignorados por contener {CODIGO_JSON_IGNORADO}: {len(ignorados)}"
        )
        return 0

    try:
        plantillas = detectar_formatos(args.formatos)
    except FileNotFoundError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1

    if not plantillas:
        print(f"ERROR: No se encontraron plantillas Word en: {args.formatos}", file=sys.stderr)
        return 1

    args.salida.mkdir(parents=True, exist_ok=True)
    errores = 0
    generados_esta_ejecucion = 0
    for archivo_json in archivos_json:
        if CODIGO_JSON_IGNORADO in archivo_json.name:
            print(f"IGNORADO\t{archivo_json.relative_to(args.json_dir)}")
            continue
        try:
            datos = leer_datos_json(archivo_json)
            idioma = leer_idioma_datos(datos)
            plantilla = seleccionar_plantilla(plantillas, idioma)
            destino = _nombre_destino(args.salida, archivo_json.stem)
            generar_documento(datos, plantilla, destino)
        except (OSError, json.JSONDecodeError, RuntimeError, ValueError) as exc:
            print(f"ERROR\t{archivo_json.relative_to(args.json_dir)}: {exc}", file=sys.stderr)
            errores += 1
            continue
        print(
            f"{idioma}\t{archivo_json.relative_to(args.json_dir)}"
            f"\t-> {plantilla.relative_to(args.formatos)}\t=> {destino.name}"
        )
        generados_esta_ejecucion += 1

    generados = len(list(args.salida.glob("*.docx")))
    print(
        f"\nJSON considerados: {len(archivos_json)} | Ignorados por nombre: {len(ignorados)}"
        f" | Generados en esta ejecucion: {generados_esta_ejecucion}"
        f" | DOCX en historico: {generados} | Errores: {errores}"
    )
    return 1 if errores else 0


if __name__ == "__main__":
    raise SystemExit(main())
