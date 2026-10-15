#!/usr/bin/env python3
"""Gestor de especificaciones SCHUTZ, Python 3.11."""
from __future__ import annotations

import argparse
import hashlib
import html
import json
import multiprocessing
import os
import queue
import re
import signal
import shutil
import sqlite3
import subprocess
import sys
import tempfile
import time
import unicodedata
from collections import defaultdict
from contextlib import contextmanager
from contextvars import ContextVar
from datetime import date, datetime
from pathlib import Path
from typing import Callable
from zipfile import ZipFile

from docx import Document
from openpyxl import load_workbook

APP_VERSION = "0.8.4"
PYTHON_REQUERIDO = (3, 11)
ENTRADA = "ESPECIFICACIONES_VIEJAS"
ENTRADA_ENG = "ESPECIFICACIONES_VIEJAS_ENG"
SALIDA = "ESPECIFICACIONES_JSON"
REVISAR = "REVISAR"
CATALOGO = "CODIGOS_200.xlsx"
DB = "especificaciones.db"
EXTENSIONES = {".doc", ".docx", ".docm"}
TIEMPO_LIMITE_CONVERSION_WORD = 30
_PROCESOS_WORD_PREEXISTENTES: set[int] | None = None
_ERROR_INICIO_WORD: str | None = None
_LECTURA_WORD = ContextVar("lectura_word_gestor", default=None)

COLUMNAS = {
    "PRUEBA": "prueba", "UNIDAD": "unidades", "UNIDADES": "unidades",
    "METODO": "metodo", "ESPECIFICACION": "especificacion",
    "VALOR TIPICO": "valor_tipico", "VALORES TIPICOS": "valor_tipico",
    "TEST": "prueba", "UNITS": "unidades", "UNIT": "unidades",
    "TEST METHOD": "metodo", "METHOD": "metodo",
    "SPECIFICATION": "especificacion", "SPECIFICATIONS": "especificacion",
    "TYPICAL VALUE": "valor_tipico", "TYPICAL VALUES": "valor_tipico",
}
CONTROL = {
    "ELABORO Y APROBO": "elaboro_y_aprobo",
    "ELABORADO Y APROBADO POR": "elaboro_y_aprobo",
    "PREPARED AND APPROVED BY": "elaboro_y_aprobo",
    "ELABORATED AND APPROVED BY": "elaboro_y_aprobo",
    "FOLIO APQP": "folio_apqp", "FOLIO VAP": "folio_vap",
    "APQP FOLIO": "folio_apqp", "VAP FOLIO": "folio_vap",
    "CODIGO UNICO DE PRODUCTO": "codigo_unico_producto",
    "CODIGO DE PRODUCTO": "codigo_unico_producto",
    "PART ID": "codigo_unico_producto",
    "PRODUCT CODE": "codigo_unico_producto",
    "UNIQUE PRODUCT CODE": "codigo_unico_producto",
    "PROTOTYPE BATCH NUMBER": "numero_lote_prototipo",
    "PROTOTYPE BATCH NO": "numero_lote_prototipo",
    "PROTOTYPE PART NUMBER": "codigo_unico_producto",
    "ID": "codigo_unico_producto",
    "FECHA": "fecha", "DATE": "fecha", "REVISION": "revision", "REV": "revision",
    "NOMBRE CLIENTE": "nombre_cliente", "REVISION CLIENTE": "revision_cliente",
    "NUMERO DE REVISION": "revision_cliente",
    "NAME CLIENT": "nombre_cliente", "CLIENT NAME": "nombre_cliente",
    "CUSTOMER NAME": "nombre_cliente", "CUSTOMERS NAME": "nombre_cliente",
    "REVISION CLIENT": "revision_cliente", "CLIENT REVISION": "revision_cliente",
    "CUSTOMER REVISION": "revision_cliente",
    "CUSTOMERS REVISION NUMBER": "revision_cliente",
    "CUSTOMER S REVISION NUMBER": "revision_cliente",
    "FECHA DE ELABORACION": "fecha", "FECHA ELABORACION": "fecha",
    "FECHA DE CREACION": "fecha", "FECHA CREACION": "fecha",
    "FECHA DE EMISION": "fecha", "FECHA EMISION": "fecha",
    "FECHA DE REVISION": "fecha", "FECHA REVISION": "fecha",
    "FECHA DE APROBACION": "fecha", "FECHA APROBACION": "fecha",
    "FECHA DE ACTUALIZACION": "fecha", "FECHA ACTUALIZACION": "fecha",
    "PREPARATION DATE": "fecha", "DATE PREPARED": "fecha",
    "DATE OF ISSUE": "fecha", "ISSUE DATE": "fecha",
    "DATE OF REVISION": "fecha", "REVISION DATE": "fecha",
    "DATE CREATED": "fecha", "CREATION DATE": "fecha",
    "DATE OF PREPARATION": "fecha", "DATE OF APPROVAL": "fecha",
    "APPROVAL DATE": "fecha", "UPDATE DATE": "fecha", "UPDATED ON": "fecha",
    "EFFECTIVE DATE": "fecha",
}

MESES = {
    "JAN": 1, "JANUARY": 1, "ENE": 1, "ENERO": 1,
    "FEB": 2, "FEBRUARY": 2, "FEBRERO": 2,
    "MAR": 3, "MARCH": 3, "MARZO": 3,
    "APR": 4, "APRIL": 4, "ABR": 4, "ABRIL": 4,
    "MAY": 5, "MAYO": 5,
    "JUN": 6, "JUNE": 6, "JUNIO": 6,
    "JUL": 7, "JULY": 7, "JULIO": 7,
    "AUG": 8, "AUGUST": 8, "AGO": 8, "AGOSTO": 8,
    "SEP": 9, "SEPT": 9, "SEPTEMBER": 9, "SET": 9,
    "SEPTIEMBRE": 9, "SETIEMBRE": 9,
    "OCT": 10, "OCTOBER": 10, "OCTUBER": 10, "OCTUBRE": 10,
    "NOV": 11, "NOVEMBER": 11, "NOVIEMBRE": 11,
    "DEC": 12, "DECEMBER": 12, "DIC": 12, "DICIEMBRE": 12,
}


def base_app() -> Path:
    from rutas_proyecto import base_programa
    return base_programa()


def limpiar(valor) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return re.sub(r"\s+", " ", str(valor).replace("\xa0", " ")).strip()


def clave(valor) -> str:
    texto = "".join(c for c in unicodedata.normalize("NFD", limpiar(valor).upper())
                    if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", re.sub(r"[^A-Z0-9]+", " ", texto)).strip()


def id_clave(valor) -> str:
    return re.sub(r"[^A-Z0-9]", "", limpiar(valor).upper())


def nombre_seguro(valor) -> str:
    texto = re.sub(r'[<>:"/\\|?*]+', " ", limpiar(valor).upper())
    return re.sub(r"\s+", " ", texto).strip(" ._")


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for bloque in iter(lambda: f.read(1024 * 1024), b""):
            h.update(bloque)
    return h.hexdigest()


def conectar(path: Path, *, recuperar: bool = True) -> sqlite3.Connection:
    from rutas_proyecto import resolver_ruta, ruta_portatil, reubicar_ruta_heredada
    path = resolver_ruta(path)
    if recuperar:
        from reconstruir_base_datos import asegurar_base
        asegurar_base(path.parent, destino=path)
    con = sqlite3.connect(path, timeout=30)
    try:
        con.row_factory = sqlite3.Row
        con.execute("PRAGMA busy_timeout=30000")
        con.execute("PRAGMA foreign_keys=ON")
        modo_journal = con.execute("PRAGMA journal_mode").fetchone()[0]
        if str(modo_journal).lower() != "wal":
            con.execute("PRAGMA journal_mode=WAL")
        con.executescript("""
    CREATE TABLE IF NOT EXISTS archivos(
      id INTEGER PRIMARY KEY, ruta_relativa TEXT UNIQUE NOT NULL,
      nombre_archivo TEXT NOT NULL, hash_actual TEXT NOT NULL,
      hash_catalogo TEXT, tamano_bytes INTEGER NOT NULL,
      fecha_modificacion_ns INTEGER NOT NULL, producto_clave TEXT,
      producto_nombre TEXT, nombre_fuente TEXT, resolucion_nombre TEXT,
      part_id_fuente TEXT, part_id_final TEXT,
      validacion_catalogo TEXT, revision_especificacion TEXT,
      codigo_formato TEXT, revision_formato TEXT, formato_origen TEXT,
      idioma TEXT, ultima_revision_json INTEGER, ultimo_json TEXT,
      primera_deteccion TEXT NOT NULL, ultima_revision_procesada TEXT NOT NULL,
      estado TEXT NOT NULL, ultimo_error TEXT);
    CREATE TABLE IF NOT EXISTS versiones(
      id INTEGER PRIMARY KEY, archivo_id INTEGER NOT NULL,
      producto_clave TEXT NOT NULL, revision_json INTEGER NOT NULL,
      hash_fuente TEXT NOT NULL, hash_catalogo TEXT,
      part_id_fuente TEXT, part_id_final TEXT, validacion_catalogo TEXT,
      nombre_fuente TEXT, resolucion_nombre TEXT,
      revision_especificacion TEXT, codigo_formato TEXT,
      revision_formato TEXT, formato_origen TEXT, idioma TEXT,
      fecha_documento TEXT,
      fecha_creacion TEXT NOT NULL, archivo_json TEXT UNIQUE NOT NULL,
      FOREIGN KEY(archivo_id) REFERENCES archivos(id),
      UNIQUE(producto_clave,revision_json));
    CREATE TABLE IF NOT EXISTS catalogo_codigos(
      id INTEGER PRIMARY KEY, hash_catalogo TEXT NOT NULL,
      hoja TEXT NOT NULL, fila_excel INTEGER NOT NULL,
      part_id TEXT NOT NULL, descripcion TEXT NOT NULL,
      part_id_clave TEXT NOT NULL, descripcion_clave TEXT NOT NULL,
      fecha_carga TEXT NOT NULL,
      UNIQUE(hash_catalogo,hoja,fila_excel));
    CREATE TABLE IF NOT EXISTS catalogo_fuentes(
      id INTEGER PRIMARY KEY, nombre_archivo TEXT NOT NULL,
      ruta_archivo TEXT NOT NULL, hash_catalogo TEXT UNIQUE NOT NULL,
      tamano_bytes INTEGER NOT NULL, fecha_modificacion_ns INTEGER NOT NULL,
      fecha_lectura TEXT NOT NULL, total_registros INTEGER NOT NULL);
    CREATE INDEX IF NOT EXISTS idx_archivos_hash_catalogo_estado
      ON archivos(hash_actual,hash_catalogo,estado);
    CREATE TABLE IF NOT EXISTS gestor_estado (
      clave TEXT PRIMARY KEY, valor TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS gestor_historial_cache (
      version_id INTEGER PRIMARY KEY REFERENCES versiones(id),
      tamano_bytes INTEGER NOT NULL, fecha_modificacion_ns INTEGER NOT NULL);
        """)
        # Migracion no destructiva desde versiones previas.
        requeridas_archivos = {
            "hash_catalogo": "TEXT", "part_id_fuente": "TEXT", "part_id_final": "TEXT",
            "nombre_fuente": "TEXT", "resolucion_nombre": "TEXT",
            "validacion_catalogo": "TEXT", "revision_especificacion": "TEXT",
            "codigo_formato": "TEXT", "revision_formato": "TEXT", "formato_origen": "TEXT",
            "idioma": "TEXT",
        }
        existentes = {r["name"] for r in con.execute("PRAGMA table_info(archivos)")}
        for nombre, tipo in requeridas_archivos.items():
            if nombre not in existentes:
                con.execute(f"ALTER TABLE archivos ADD COLUMN {nombre} {tipo}")
        requeridas_versiones = {k: v for k, v in requeridas_archivos.items() if k != "formato_origen"}
        requeridas_versiones["formato_origen"] = "TEXT"
        requeridas_versiones["fecha_documento"] = "TEXT"
        existentes = {r["name"] for r in con.execute("PRAGMA table_info(versiones)")}
        for nombre, tipo in requeridas_versiones.items():
            if nombre not in existentes:
                con.execute(f"ALTER TABLE versiones ADD COLUMN {nombre} {tipo}")
        if not con.execute("SELECT 1 FROM gestor_estado WHERE clave='rutas_catalogo_portatiles'").fetchone():
            for fila in con.execute("SELECT id, ruta_archivo FROM catalogo_fuentes").fetchall():
                ruta = reubicar_ruta_heredada(fila["ruta_archivo"], path.parent)
                con.execute("UPDATE catalogo_fuentes SET ruta_archivo=? WHERE id=?",
                            (ruta_portatil(ruta, path.parent), fila["id"]))
            con.execute("INSERT OR IGNORE INTO gestor_estado VALUES ('rutas_catalogo_portatiles', '1')")
        con.commit()
        return con
    except BaseException:
        con.close()
        raise


def leer_catalogo(path: Path, *, digest: str | None = None) -> tuple[list[dict], str]:
    """Lee CODIGOS_200.xlsx estrictamente en modo solo lectura.

    Esta funcion nunca llama save(), nunca cambia celdas y nunca crea una copia
    del libro. Toda la informacion derivada se devuelve en memoria y se guarda
    exclusivamente en SQLite.
    """
    if not path.exists():
        raise FileNotFoundError(f"No se encontro el catalogo obligatorio: {path.name}")
    digest = digest or sha256(path)
    wb = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    registros = []
    for ws in wb.worksheets:
        encabezado = None
        col_id = col_desc = None
        for n_fila, fila in enumerate(ws.iter_rows(values_only=True), 1):
            claves = [clave(x) for x in fila]
            if "ID" in claves and "DESCRIPTION" in claves:
                encabezado = n_fila
                col_id = claves.index("ID")
                col_desc = claves.index("DESCRIPTION")
                break
        if encabezado is None:
            continue
        for n_fila, fila in enumerate(ws.iter_rows(min_row=encabezado + 1, values_only=True), encabezado + 1):
            part_id = limpiar(fila[col_id] if col_id < len(fila) else "")
            descripcion = limpiar(fila[col_desc] if col_desc < len(fila) else "")
            if not part_id and not descripcion:
                continue
            registros.append({"hoja": ws.title, "fila_excel": n_fila,
                              "part_id": part_id, "descripcion": descripcion,
                              "part_id_clave": id_clave(part_id),
                              "descripcion_clave": clave(descripcion)})
    wb.close()
    if not registros:
        raise ValueError("CODIGOS_200.xlsx no contiene columnas ID y Description reconocibles.")
    return registros, digest


def obtener_catalogo(con: sqlite3.Connection, path: Path) -> tuple[list[dict], str]:
    """Reutiliza registros completos del mismo Excel, comprobando su contenido."""
    digest = sha256(path)
    fuente = con.execute("SELECT total_registros FROM catalogo_fuentes WHERE hash_catalogo=?", (digest,)).fetchone()
    if fuente:
        filas = con.execute("""SELECT hoja, fila_excel, part_id, descripcion,
            part_id_clave, descripcion_clave FROM catalogo_codigos
            WHERE hash_catalogo=? ORDER BY id""", (digest,)).fetchall()
        if filas and len(filas) == fuente["total_registros"]:
            return [dict(fila) for fila in filas], digest
    return leer_catalogo(path, digest=digest)


class CatalogoPreparado(list):
    """Índices de búsqueda con el mismo orden y deduplicación del validador."""
    def __init__(self, catalogo):
        super().__init__(catalogo)
        self.registros = list({(r["part_id_clave"], r["descripcion_clave"]): r
            for r in self if r["part_id_clave"] and r["descripcion_clave"]}.values())
        self.por_nombre = defaultdict(list)
        self.por_id = defaultdict(list)
        for registro in self.registros:
            self.por_nombre[registro["descripcion_clave"]].append(registro)
            self.por_id[registro["part_id_clave"]].append(registro)


@contextmanager
def lectura_word_compartida():
    """Caché por documento; también se limpia en errores y cambios de archivo."""
    token = _LECTURA_WORD.set({})
    try:
        yield
    finally:
        _LECTURA_WORD.reset(token)


def productos_historial_modificados(con, salida):
    afectados = set()
    for fila in con.execute("""SELECT v.id, v.producto_clave, v.archivo_json,
        c.tamano_bytes, c.fecha_modificacion_ns FROM versiones v
        LEFT JOIN gestor_historial_cache c ON c.version_id=v.id"""):
        try:
            estado = (salida / fila["archivo_json"]).stat()
            cambiado = (estado.st_size, estado.st_mtime_ns) != (fila["tamano_bytes"], fila["fecha_modificacion_ns"])
        except OSError:
            cambiado = True
        if cambiado:
            afectados.add(fila["producto_clave"])
    return afectados


def registrar_historial_revisado(con, salida, productos):
    for fila in con.execute("SELECT id, producto_clave, archivo_json FROM versiones").fetchall():
        if productos is not None and fila["producto_clave"] not in productos:
            continue
        try:
            estado = (salida / fila["archivo_json"]).stat()
        except OSError:
            continue
        con.execute("""INSERT INTO gestor_historial_cache VALUES (?, ?, ?)
            ON CONFLICT(version_id) DO UPDATE SET tamano_bytes=excluded.tamano_bytes,
            fecha_modificacion_ns=excluded.fecha_modificacion_ns""",
            (fila["id"], estado.st_size, estado.st_mtime_ns))


def ampliar_productos_afectados(con, salida, catalogo, productos):
    """Incluye historiales de destino cuando la normalización cambia el producto."""
    filas = con.execute("SELECT producto_clave, archivo_json FROM versiones").fetchall()
    procesados = set()
    while productos - procesados:
        grupo = productos - procesados
        procesados.update(grupo)
        for fila in filas:
            if fila["producto_clave"] not in grupo:
                continue
            try:
                datos = json.loads((salida / fila["archivo_json"]).read_text(encoding="utf-8"))
                producto, control = datos.get("producto", {}), datos.get("control_documental", {})
                nombre = producto.get("nombre_fuente") or producto.get("nombre", "")
                part_id = control.get("codigo_unico_producto_fuente", control.get("codigo_unico_producto", ""))
                contenido = dict(datos)
                contenido["producto"] = dict(producto, nombre=nombre)
                control_busqueda = dict(control)
                if "codigo_unico_producto_fuente" in control_busqueda:
                    control_busqueda["codigo_unico_producto"] = control_busqueda.pop("codigo_unico_producto_fuente")
                contenido["control_documental"] = control_busqueda
                validacion = validar_catalogo(nombre, part_id, catalogo,
                    contenido=json.dumps(contenido, ensure_ascii=False))
                if validacion.get("nombre_catalogo"):
                    productos.add(clave(validacion["nombre_catalogo"]))
            except (OSError, ValueError, AttributeError, TypeError):
                # El reindexador reporta el error y conserva el historial completo.
                continue
    return productos


def registrar_catalogo(con: sqlite3.Connection, registros: list[dict], digest: str,
                       fecha: str, archivo: Path) -> None:
    from rutas_proyecto import ruta_portatil
    base = Path(con.execute("PRAGMA database_list").fetchone()[2]).parent
    stat = archivo.stat()
    con.execute("""INSERT OR IGNORE INTO catalogo_fuentes
      (nombre_archivo,ruta_archivo,hash_catalogo,tamano_bytes,fecha_modificacion_ns,fecha_lectura,total_registros)
      VALUES(?,?,?,?,?,?,?)""",
      (archivo.name, ruta_portatil(archivo, base), digest, stat.st_size, stat.st_mtime_ns, fecha, len(registros)))
    con.execute("UPDATE catalogo_fuentes SET ruta_archivo=? WHERE hash_catalogo=?",
                (ruta_portatil(archivo, base), digest))
    existentes = con.execute("SELECT COUNT(*) FROM catalogo_codigos WHERE hash_catalogo=?", (digest,)).fetchone()[0]
    if existentes == len(registros):
        con.commit()
        return
    if existentes:
        # Reparar únicamente la caché derivada de esta huella si quedó incompleta.
        con.execute("DELETE FROM catalogo_codigos WHERE hash_catalogo=?", (digest,))
    con.execute("UPDATE catalogo_fuentes SET total_registros=? WHERE hash_catalogo=?", (len(registros), digest))
    con.executemany("""INSERT OR IGNORE INTO catalogo_codigos
      (hash_catalogo,hoja,fila_excel,part_id,descripcion,part_id_clave,descripcion_clave,fecha_carga)
      VALUES(?,?,?,?,?,?,?,?)""",
      [(digest, r["hoja"], r["fila_excel"], r["part_id"], r["descripcion"],
        r["part_id_clave"], r["descripcion_clave"], fecha) for r in registros])
    con.commit()


def archivo_bloqueado_word(path: Path) -> Path | None:
    nombre = path.name.casefold()
    for temporal in path.parent.glob("~$*"):
        nombre_temporal = temporal.name[2:].casefold()
        if (
            temporal.is_file()
            and temporal.suffix.casefold() == path.suffix.casefold()
            and len(nombre_temporal) >= 4
            and (
                nombre[2:].startswith(nombre_temporal)
                or nombre.startswith(nombre_temporal)
            )
        ):
            return temporal
    return None


def _convertir_word_com_worker(path: str, destino: str, resultados) -> None:
    import pythoncom
    import win32com.client
    import win32process

    resultados.put(("stage", "iniciar automatización COM"))
    pythoncom.CoInitialize()
    word = None
    documento = None
    try:
        resultados.put(("stage", "iniciar Microsoft Word"))
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        resultados.put(("stage", "abrir documento .doc"))
        documento = word.Documents.Open(
            path,
            ConfirmConversions=False,
            ReadOnly=True,
            AddToRecentFiles=False,
            Visible=False,
            OpenAndRepair=False,
            NoEncodingDialog=True,
        )
        # Hwnd pertenece a Window, no a Word.Application. El PID es auxiliar:
        # si no hay ventana disponible, el supervisor conserva su limpieza
        # por procesos de automatización y su límite de tiempo.
        try:
            hwnd = documento.Windows.Item(1).Hwnd
            _, pid_word = win32process.GetWindowThreadProcessId(hwnd)
            if pid_word:
                resultados.put(("word_pid", pid_word))
        except Exception:
            pass
        resultados.put(("stage", "guardar documento temporal .docx"))
        documento.SaveAs2(destino, FileFormat=16)
        resultados.put(("stage", "cerrar documento"))
        documento.Close(False)
        documento = None
        resultados.put(("stage", "cerrar Microsoft Word"))
        word.Quit()
        word = None
        resultados.put(("ok", None))
    except Exception as exc:
        resultados.put(("error", f"{type(exc).__name__}: {exc}"))
    finally:
        try:
            if documento is not None:
                resultados.put(("stage", "cerrar documento tras error"))
                try:
                    documento.Close(False)
                except Exception:
                    pass
            if word is not None:
                resultados.put(("stage", "cerrar Microsoft Word tras error"))
                try:
                    word.Quit()
                except Exception:
                    pass
        finally:
            resultados.put(("stage", "finalizar automatización COM"))
            pythoncom.CoUninitialize()


def _terminar_proceso_word(pid: int) -> None:
    try:
        os.kill(pid, signal.SIGTERM)
    except ProcessLookupError:
        pass


def _procesos_word_automatizacion() -> set[int]:
    consulta = (
        "$p = @(Get-CimInstance Win32_Process | "
        "Where-Object { $_.Name -eq 'WINWORD.EXE' -and "
        "$_.CommandLine -match '(?i)/automation' } | "
        "Select-Object -ExpandProperty ProcessId); "
        "ConvertTo-Json -InputObject $p -Compress"
    )
    resultado = subprocess.run(
        ["powershell.exe", "-NoProfile", "-NonInteractive", "-Command", consulta],
        capture_output=True,
        text=True,
        timeout=5,
        check=False,
    )
    if resultado.returncode:
        raise RuntimeError(
            "No se pudieron inspeccionar las instancias automatizadas de Word: "
            f"{resultado.stderr.strip() or resultado.stdout.strip()}"
        )
    datos = json.loads(resultado.stdout.strip() or "[]")
    if not isinstance(datos, list):
        datos = [datos]
    return {int(pid) for pid in datos}


def _limpiar_word_iniciado(
    procesos_preexistentes: set[int],
    pid_word: int | None,
) -> None:
    actuales = _procesos_word_automatizacion()
    procesos_nuevos = actuales - procesos_preexistentes
    if pid_word is not None and pid_word not in procesos_preexistentes:
        procesos_nuevos.add(pid_word)
    for pid in sorted(procesos_nuevos):
        _terminar_proceso_word(pid)


def _restablecer_error_inicio_word() -> None:
    global _ERROR_INICIO_WORD
    _ERROR_INICIO_WORD = None


def _convertir_word_com(
    path: Path,
    destino: Path,
    *,
    worker=_convertir_word_com_worker,
) -> None:
    global _ERROR_INICIO_WORD, _PROCESOS_WORD_PREEXISTENTES
    if _ERROR_INICIO_WORD is not None:
        raise RuntimeError(
            "Se omitió el intento con Microsoft Word porque ya falló su inicio "
            f"en esta carga: {_ERROR_INICIO_WORD}"
        )
    if _PROCESOS_WORD_PREEXISTENTES is None:
        _PROCESOS_WORD_PREEXISTENTES = _procesos_word_automatizacion()
    procesos_word_antes = set(_PROCESOS_WORD_PREEXISTENTES)
    contexto = multiprocessing.get_context("spawn")
    resultados = contexto.Queue()
    proceso = contexto.Process(
        target=worker,
        args=(str(path.resolve()), str(destino.resolve()), resultados),
    )
    pid_word = None
    resultado = None
    etapa = "iniciar el proceso de conversión"
    iniciado = False
    try:
        proceso.start()
        iniciado = True
        limite = time.monotonic() + TIEMPO_LIMITE_CONVERSION_WORD
        while resultado is None:
            restante = limite - time.monotonic()
            if restante <= 0:
                break
            try:
                tipo, valor = resultados.get(timeout=min(0.1, restante))
            except queue.Empty:
                if not proceso.is_alive():
                    break
                continue
            if tipo == "word_pid":
                pid_word = valor
            elif tipo == "stage":
                etapa = valor
            else:
                resultado = (tipo, valor)

        if resultado is None and time.monotonic() >= limite:
            proceso.terminate()
            proceso.join(timeout=5)
            if proceso.is_alive():
                proceso.kill()
                proceso.join()
            _limpiar_word_iniciado(procesos_word_antes, pid_word)
            mensaje = (
                f"Microsoft Word se bloqueó al intentar {etapa}; excedió el límite de "
                f"{TIEMPO_LIMITE_CONVERSION_WORD} segundos."
            )
            if etapa == "iniciar Microsoft Word":
                _ERROR_INICIO_WORD = mensaje
            raise TimeoutError(
                mensaje
            )

        proceso.join(timeout=5)
        if proceso.is_alive():
            proceso.terminate()
            proceso.join()
            if resultado is None or resultado[0] != "ok":
                if pid_word is not None:
                    _terminar_proceso_word(pid_word)
        if resultado is None:
            raise RuntimeError(
                f"El proceso de conversión de Word terminó sin resultado "
                f"(código {proceso.exitcode})."
            )
        if resultado[0] != "ok":
            if pid_word is not None:
                _terminar_proceso_word(pid_word)
            raise RuntimeError(f"Microsoft Word no pudo convertir el documento: {resultado[1]}")
    finally:
        if iniciado and proceso.is_alive():
            proceso.terminate()
            proceso.join()
        resultados.close()
        resultados.cancel_join_thread()


def convertir_word_legacy(path: Path, tmp: Path) -> Path:
    destino = tmp / f"{path.stem}.docx"
    error_word = None
    if sys.platform == "win32":
        try:
            _convertir_word_com(path, destino)
            if destino.exists():
                return destino
        except Exception as exc:
            error_word = str(exc)
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if exe:
        try:
            subprocess.run(
                [exe, "--headless", "--convert-to", "docx", "--outdir", str(tmp), str(path)],
                capture_output=True,
                timeout=TIEMPO_LIMITE_CONVERSION_WORD,
                check=False,
            )
        except subprocess.TimeoutExpired as exc:
            raise TimeoutError(
                f"LibreOffice no terminó la conversión en "
                f"{TIEMPO_LIMITE_CONVERSION_WORD} segundos."
            ) from exc
        if destino.exists():
            return destino
    detalle = f" Error de Microsoft Word: {error_word}" if error_word else ""
    raise RuntimeError(
        "No se pudo convertir el archivo .doc o .docm. "
        "Revisa Microsoft Word y pywin32, o usa LibreOffice como alternativa."
        f"{detalle}"
    )


def preparar(path: Path, tmp: Path) -> tuple[Path, bool]:
    extension = path.suffix.lower()
    if extension == ".docx":
        return path, False
    if extension in {".doc", ".docm"}:
        # El original se abre solo para lectura y se convierte a un DOCX temporal.
        # El archivo fuente nunca se guarda, modifica ni sobrescribe.
        return convertir_word_legacy(path, tmp), True
    raise ValueError(f"Formato Word no compatible: {extension}")


def validar_documento_word(path: Path, tmp: Path) -> None:
    if path.stat().st_size == 0:
        raise ValueError("el archivo esta vacio")
    lectura, _ = preparar(path, tmp)
    try:
        doc = Document(lectura)
    except Exception as exc:
        raise ValueError("el archivo no es un documento Word valido") from exc
    if not doc.tables:
        raise ValueError("el documento Word no contiene tablas")
    tiene_contenido = any(limpiar(p.text) for p in doc.paragraphs)
    tiene_contenido = tiene_contenido or any(
        limpiar(celda.text) for tabla in doc.tables for fila in tabla.rows for celda in fila.cells
    )
    if not tiene_contenido:
        raise ValueError("el documento Word no contiene texto")


def texto_ooxml(path: Path) -> str:
    cache = _LECTURA_WORD.get()
    clave_cache = str(path.resolve())
    if cache is not None and clave_cache in cache:
        return cache[clave_cache]
    partes = []
    with ZipFile(path) as z:
        for nombre in z.namelist():
            if nombre.startswith("word/") and nombre.endswith(".xml"):
                contenido = z.read(nombre).decode("utf-8", "ignore")
                partes.extend(re.findall(r"<w:t[^>]*>(.*?)</w:t>", contenido, re.S))
                partes.extend(re.findall(r"<a:t[^>]*>(.*?)</a:t>", contenido, re.S))
    resultado = limpiar(html.unescape(" ".join(re.sub(r"<[^>]+>", " ", x) for x in partes)))
    if cache is not None:
        cache[clave_cache] = resultado
    return resultado


def metadatos_formato(path: Path) -> dict:
    texto = texto_ooxml(path)
    datos = {}
    m = re.search(r"\b(R\d{2}-[A-Z]{2,}-[A-Z]{2,}-\d{2})\b", texto, re.I)
    if m:
        datos["codigo_formato"] = m.group(1).upper()
        cerca = texto[m.end():m.end() + 100]
        rev = re.search(r"\b(\d{1,2})\b", cerca)
        if rev: datos["revision_formato"] = rev.group(1).zfill(2)
    exe = shutil.which("tesseract")
    if not exe:
        return datos
    with tempfile.TemporaryDirectory() as td, ZipFile(path) as z:
        for i, nombre in enumerate(n for n in z.namelist() if n.startswith("word/media/")):
            img = Path(td) / f"{i}{Path(nombre).suffix}"
            img.write_bytes(z.read(nombre))
            try:
                r = subprocess.run([exe, str(img), "stdout", "-l", "spa+eng", "--psm", "6"],
                                   capture_output=True, text=True, timeout=20, check=False)
            except Exception:
                continue
            t = limpiar(r.stdout)
            cod = re.search(r"C[OÓ]DIGO\s*[:.]?\s*([A-Z0-9-]+)", t, re.I)
            rev = re.search(r"REV(?:ISI[OÓ]N)?\s*[.:]?\s*([A-Z0-9.-]+)", t, re.I)
            if cod: datos["codigo_formato"] = cod.group(1)
            if rev: datos["revision_formato"] = rev.group(1).strip(".").zfill(2)
    return datos


def contiene_marcadores_ingles(path: Path) -> bool:
    """Detecta idioma por etiquetas del documento, no por el formato visual usado."""
    texto = clave(texto_ooxml(path))
    marcadores_ingles = (
        "TEST", "UNITS", "TEST METHOD", "SPECIFICATION", "TYPICAL VALUES",
        "PART ID", "PRODUCT CODE", "CUSTOMER", "NAME", "DATE", "POSITION",
        "SIGNATURE", "COMPANY",
    )
    marcadores_espanol = (
        "PRUEBA", "UNIDAD", "METODO", "ESPECIFICACION", "VALOR TIPICO",
        "CODIGO UNICO DE PRODUCTO", "FECHA", "NOMBRE CLIENTE",
        "ELABORO Y APROBO",
    )
    ingles = sum(bool(re.search(rf"\b{re.escape(m)}\b", texto)) for m in marcadores_ingles)
    espanol = sum(bool(re.search(rf"\b{re.escape(m)}\b", texto)) for m in marcadores_espanol)
    return ingles > espanol


def menciones_yves_yoi(path: Path) -> list[str]:
    """Devuelve los marcadores que aparecen como palabras completas en el Word."""
    texto = texto_ooxml(path)
    return list(dict.fromkeys(re.findall(r"\b(?:YVES|YOI)\b", texto, flags=re.IGNORECASE)))


def destino_movimiento_sin_colision(destino: Path) -> Path:
    """Evita sobrescribir un archivo existente al mover una especificacion."""
    if not destino.exists():
        return destino
    indice = 1
    while True:
        candidato = destino.with_name(f"{destino.stem}_{indice:02d}{destino.suffix}")
        if not candidato.exists():
            return candidato
        indice += 1


def nombre_archivo_json(part_id: str, revision: str, fecha: datetime) -> str:
    codigo = nombre_seguro(part_id) or "SIN_CODIGO_200"
    revision_archivo = nombre_seguro(revision) or "SIN_REVISION"
    if revision_archivo.isdigit():
        revision_archivo = revision_archivo.zfill(2)
    return f"{codigo} {revision_archivo}_{fecha:%d_%m_%y}.JSON"


def parsear_fecha_documento(valor: str) -> date | None:
    texto = limpiar(valor)
    if not texto:
        return None
    texto = "".join(
        c for c in unicodedata.normalize("NFD", texto.upper())
        if unicodedata.category(c) != "Mn"
    )
    texto = re.sub(r"(\d+)(?:ST|ND|RD|TH)\b", r"\1", texto)
    texto = re.sub(r"(?<=/)F(?=FEB(?:RUARY)?\b)", "", texto)

    def fecha_segura(anio: int, mes: int, dia: int) -> date | None:
        if anio < 100:
            anio += 2000 if anio <= 69 else 1900
        try:
            return date(anio, mes, dia)
        except ValueError:
            return None

    compacto_largo = re.search(r"\b(\d{8})\b", texto)
    if compacto_largo:
        digitos = compacto_largo[1]
        primero = int(digitos[:4])
        if 1900 <= primero <= 2100:
            return fecha_segura(primero, int(digitos[4:6]), int(digitos[6:]))
        return fecha_segura(int(digitos[4:]), int(digitos[2:4]), int(digitos[:2]))

    compacto = re.search(r"\b(\d{2})(\d{2})(\d{2})\b", texto)
    if compacto:
        dia, mes, anio = map(int, compacto.groups())
        return fecha_segura(anio, mes, dia)

    iso = re.search(r"\b(\d{4})\s*[-/.,\s]\s*(\d{1,2})\s*[-/.,\s]\s*(\d{1,2})\b", texto)
    if iso:
        return fecha_segura(int(iso[1]), int(iso[2]), int(iso[3]))

    numerica = re.search(r"\b(\d{1,2})\s*[-/.,\s]\s*(\d{1,2})\s*[-/.,\s]\s*(\d{2,4})\b", texto)
    if numerica:
        primero, segundo, anio = map(int, numerica.groups())
        if primero > 12:
            dia, mes = primero, segundo
        elif segundo > 12:
            mes, dia = primero, segundo
        else:
            # Las fechas numéricas ambiguas usan día/mes; es el patrón predominante
            # en las especificaciones existentes, también en varios documentos ingleses.
            dia, mes = primero, segundo
        return fecha_segura(anio, mes, dia)

    tokens = re.findall(r"[A-Z]+|\d+", texto)
    for indice, token in enumerate(tokens):
        mes = MESES.get(token)
        if mes is None:
            continue
        numeros = [(i, int(t)) for i, t in enumerate(tokens) if t.isdigit() and i != indice]
        anios = [(i, n) for i, n in numeros if n >= 1000]
        if not anios:
            anios = [(i, n) for i, n in numeros if len(tokens[i]) == 2]
        if not anios:
            continue
        indice_anio, anio = max(anios, key=lambda par: par[0])
        dias = [(i, n) for i, n in numeros if i != indice_anio and 1 <= n <= 31]
        dia = min(dias, key=lambda par: abs(par[0] - indice))[1] if dias else 1
        return fecha_segura(anio, mes, dia)

    mes_anio = re.search(r"\b(\d{1,2})\s*[-/.]\s*(\d{4})\b", texto)
    if mes_anio:
        return fecha_segura(int(mes_anio[2]), int(mes_anio[1]), 1)
    anio_mes = re.search(r"\b(\d{4})\s*[-/.]\s*(\d{1,2})\b", texto)
    if anio_mes:
        return fecha_segura(int(anio_mes[1]), int(anio_mes[2]), 1)
    return None


def reindexar_revisiones(
    con: sqlite3.Connection,
    salida: Path,
    catalogo: list[dict],
    producto_claves: set[str] | None = None,
) -> dict[str, int]:
    filtro = ""
    parametros: tuple[str, ...] = ()
    if producto_claves is not None:
        claves_ordenadas = sorted(producto_claves)
        if claves_ordenadas:
            placeholders = ",".join("?" for _ in claves_ordenadas)
            filtro = f"WHERE v.producto_clave IN ({placeholders})"
            parametros = tuple(claves_ordenadas)
        else:
            filtro = "WHERE 0"
    filas = con.execute(
        f"""SELECT v.id, v.archivo_id, v.producto_clave, v.revision_json,
                  v.archivo_json, v.fecha_creacion, v.fecha_documento, a.ruta_relativa,
                  v.part_id_fuente, v.part_id_final, v.validacion_catalogo,
                  v.nombre_fuente, v.resolucion_nombre
           FROM versiones v JOIN archivos a ON a.id=v.archivo_id
           {filtro}
           ORDER BY v.producto_clave, v.revision_json, v.id""",
        parametros,
    ).fetchall()
    totales_por_producto: dict[str, int] = defaultdict(int)
    productos_con_json_invalido: set[str] = set()
    errores_json: list[tuple[str, str]] = []
    versiones_cargadas: list[dict] = []
    for fila in filas:
        totales_por_producto[fila["producto_clave"]] += 1
        ruta_json = salida / fila["archivo_json"]
        try:
            datos = json.loads(ruta_json.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as exc:
            productos_con_json_invalido.add(fila["producto_clave"])
            errores_json.append((fila["archivo_json"], str(exc)))
            continue
        if not isinstance(datos, dict):
            productos_con_json_invalido.add(fila["producto_clave"])
            errores_json.append((fila["archivo_json"], "el contenido JSON no es un objeto"))
            continue
        json_original = json.dumps(datos, ensure_ascii=False, sort_keys=True)
        control = datos.get("control_documental")
        documento = datos.get("documento")
        if not isinstance(control, dict) or not isinstance(documento, dict):
            productos_con_json_invalido.add(fila["producto_clave"])
            errores_json.append((fila["archivo_json"], "faltan metadatos documentales válidos"))
            continue
        fecha = parsear_fecha_documento(str(control.get("fecha") or ""))
        fecha_iso = fecha.isoformat() if fecha else None
        documento["fecha_documento_normalizada"] = fecha_iso or ""
        versiones_cargadas.append(
            {"fila": fila, "ruta": ruta_json, "datos": datos,
             "fecha": fecha, "fecha_iso": fecha_iso,
             "json_original": json_original,
             "producto_clave": fila["producto_clave"],
             "producto_clave_original": fila["producto_clave"],
             "nombre_fuente": fila["nombre_fuente"] or "",
             "resolucion_nombre": fila["resolucion_nombre"] or "",
             "part_id_fuente": fila["part_id_fuente"] or "",
             "part_id_final": fila["part_id_final"] or "",
             "archivo_json": fila["archivo_json"],
             "validacion_catalogo": fila["validacion_catalogo"] or ""})

    nombres_normalizados = 0
    nombres_ambiguos = 0
    ids_catalogo = {r["part_id_clave"] for r in catalogo if r["part_id_clave"]}
    # Avoid merging product histories while an unreadable file prevents a
    # complete view of that history.
    if not errores_json:
        for version in versiones_cargadas:
            datos = version["datos"]
            producto = datos.get("producto")
            control = datos.get("control_documental")
            if not isinstance(producto, dict) or not isinstance(control, dict):
                continue
            nombre_actual = str(producto.get("nombre") or "")
            nombre_fuente = str(producto.get("nombre_fuente") or version["nombre_fuente"] or nombre_actual)
            if "codigo_unico_producto_fuente" in control:
                part_id_fuente = str(control["codigo_unico_producto_fuente"] or "")
            elif version["part_id_fuente"]:
                part_id_fuente = str(version["part_id_fuente"])
            else:
                part_id_fuente = str(control.get("codigo_unico_producto") or "")

            contenido_busqueda = dict(datos)
            producto_busqueda = dict(producto)
            control_busqueda = dict(control)
            if "nombre_fuente" in producto_busqueda:
                producto_busqueda["nombre"] = producto_busqueda.pop("nombre_fuente")
            if "codigo_unico_producto_fuente" in control_busqueda:
                control_busqueda["codigo_unico_producto"] = control_busqueda.pop(
                    "codigo_unico_producto_fuente"
                )
            contenido_busqueda["producto"] = producto_busqueda
            contenido_busqueda["control_documental"] = control_busqueda
            contenido_busqueda.pop("validacion_codigos_200", None)
            validacion = validar_catalogo(
                nombre_fuente,
                part_id_fuente,
                catalogo,
                contenido=json.dumps(contenido_busqueda, ensure_ascii=False),
            )
            datos["validacion_codigos_200"] = validacion
            version["validacion_catalogo"] = validacion["estado"]
            version["nombre_fuente"] = nombre_fuente if nombre_fuente != nombre_actual else ""
            version["resolucion_nombre"] = validacion.get("resolucion_nombre", "")
            version["part_id_fuente"] = part_id_fuente
            version["part_id_final"] = validacion.get("part_id_final", part_id_fuente)
            if validacion.get("nombre_catalogo"):
                nombre_canonico = validacion["nombre_catalogo"]
                if nombre_canonico != nombre_actual:
                    nombres_normalizados += 1
                    producto["nombre_fuente"] = nombre_fuente
                    producto["nombre"] = nombre_canonico
                if id_clave(version["part_id_final"]) != id_clave(part_id_fuente):
                    control["codigo_unico_producto_fuente"] = part_id_fuente
                if version["part_id_final"]:
                    control["codigo_unico_producto"] = version["part_id_final"]
                version["producto_clave"] = clave(nombre_canonico)
                version["nombre_fuente"] = nombre_fuente if nombre_fuente != nombre_canonico else ""
            elif validacion["estado"] == "CONFLICTO_NOMBRE_AMBIGUO":
                nombres_ambiguos += 1

    nombres_por_cambiar = set()
    if not errores_json:
        for version in versiones_cargadas:
            archivo_json = version["archivo_json"]
            part_id_final = version["part_id_final"]
            codigo_valido = id_clave(part_id_final) in ids_catalogo
            if archivo_json.startswith(("SIN_CODIGO_200", "DUPLICADO")) and (
                (codigo_valido and not archivo_json.startswith(nombre_seguro(part_id_final) + " "))
                or (not codigo_valido and archivo_json.startswith("DUPLICADO"))
            ):
                nombres_por_cambiar.add(archivo_json)
    nombres_reservados = {
        version["archivo_json"]
        for version in versiones_cargadas
        if version["archivo_json"] not in nombres_por_cambiar
    }
    for version in versiones_cargadas:
        archivo_json = version["archivo_json"]
        if archivo_json not in nombres_por_cambiar:
            continue
        part_id_final = version["part_id_final"]
        codigo_valido = id_clave(part_id_final) in ids_catalogo
        prefijo_anterior = "SIN_CODIGO_200" if archivo_json.startswith("SIN_CODIGO_200") else "DUPLICADO"
        sufijo = archivo_json[len(prefijo_anterior):]
        destino_base = f"{nombre_seguro(part_id_final)}{sufijo}" if codigo_valido else f"SIN_CODIGO_200{sufijo}"
        destino = destino_base
        indice = 1
        while (
            destino in nombres_reservados
            or (salida / destino).exists()
        ):
            ruta_destino = Path(destino_base)
            destino = f"{ruta_destino.stem}_{indice:02d}{ruta_destino.suffix}"
            indice += 1
        nombres_reservados.add(destino)
        version["archivo_json"] = destino
        version["ruta_destino"] = salida / destino
        version["datos"]["archivo_json"] = destino

    asignaciones = []
    cantidad_cambiada = 0
    fechas_no_reconocidas = 0
    por_producto: dict[str, list[dict]] = defaultdict(list)
    for version in versiones_cargadas:
        por_producto[version["producto_clave"]].append(version)
    for versiones in por_producto.values():
        if any(v["fila"]["producto_clave"] in productos_con_json_invalido for v in versiones):
            continue
        versiones.sort(key=lambda v: (
            v["fecha"] is None,
            v["fecha"] or date.max,
            v["producto_clave_original"].casefold(),
            v["fila"]["revision_json"],
            v["fila"]["fecha_creacion"],
            v["fila"]["ruta_relativa"].casefold(),
            v["fila"]["archivo_json"].casefold(),
        ))
        for revision, version in enumerate(versiones):
            fila = version["fila"]
            revision_texto = f"{revision:02d}"
            version["revision_normalizada"] = revision_texto
            if version["fecha"] is None:
                fechas_no_reconocidas += 1
            control = version["datos"].get("control_documental")
            if isinstance(control, dict):
                revision_fuente = str(control.get("revision") or "").strip()
                if re.fullmatch(r"-{1,3}", revision_fuente):
                    control.setdefault("revision_fuente", revision_fuente)
                    control["revision"] = revision_texto
            version["datos"]["revision_json"] = revision_texto
            json_actual = json.dumps(version["datos"], ensure_ascii=False, sort_keys=True)
            version["dirty"] = json_actual != version["json_original"]
            if (fila["revision_json"] != revision or fila["fecha_documento"] != version["fecha_iso"]
                    or version["dirty"] or version["producto_clave"] != fila["producto_clave"]
                    or version["validacion_catalogo"] != (fila["validacion_catalogo"] or "")
                    or version["part_id_final"] != (fila["part_id_final"] or "")
                    or version["part_id_fuente"] != (fila["part_id_fuente"] or "")
                    or version["nombre_fuente"] != (fila["nombre_fuente"] or "")
                    or version["resolucion_nombre"] != (fila["resolucion_nombre"] or "")):
                cantidad_cambiada += 1
            asignaciones.append((fila["id"], fila["archivo_id"], revision, version["fecha_iso"], version))

    nombres_revision_a_normalizar = [
        version for version in versiones_cargadas
        if "revision_normalizada" in version
        and re.search(r"(?<= )-{1,3}(?=_)", version["archivo_json"])
    ]
    nombres_reservados.difference_update(
        version["archivo_json"] for version in nombres_revision_a_normalizar
    )
    for version in nombres_revision_a_normalizar:
        archivo_json = version["archivo_json"]
        revision_texto = version["revision_normalizada"]
        destino_base = re.sub(
            r"(?<= )-{1,3}(?=_)", revision_texto, archivo_json, count=1
        )
        destino = destino_base
        indice = 1
        while destino in nombres_reservados or (salida / destino).exists():
            ruta_destino = Path(destino_base)
            destino = f"{ruta_destino.stem}_{indice:02d}{ruta_destino.suffix}"
            indice += 1
        nombres_reservados.add(destino)
        version["archivo_json"] = destino
        version["ruta_destino"] = salida / destino
        version["datos"]["archivo_json"] = destino
        version["dirty"] = (
            json.dumps(version["datos"], ensure_ascii=False, sort_keys=True)
            != version["json_original"]
        )

    for _, _, _, _, version in asignaciones:
        if version["dirty"]:
            version["ruta"].write_text(
                json.dumps(version["datos"], ensure_ascii=False, indent=2), encoding="utf-8"
            )

    renombres_aplicados = []
    try:
        for _, _, _, _, version in asignaciones:
            ruta_destino = version.get("ruta_destino")
            if ruta_destino is not None:
                version["ruta"].rename(ruta_destino)
                renombres_aplicados.append((ruta_destino, version["ruta"]))
    except OSError:
        for ruta_actual, ruta_original in reversed(renombres_aplicados):
            ruta_actual.rename(ruta_original)
        raise

    try:
        con.execute("BEGIN")
        con.executemany(
            "UPDATE versiones SET revision_json=-id, producto_clave='__reindex__'||id WHERE id=?",
            [(version_id,) for version_id, _, _, _, _ in asignaciones],
        )
        con.executemany(
            """UPDATE versiones SET producto_clave=?, revision_json=?, fecha_documento=?,
                      part_id_fuente=?, part_id_final=?, validacion_catalogo=?,
                      nombre_fuente=?, resolucion_nombre=?, archivo_json=?
               WHERE id=?""",
            [(version["producto_clave"], revision, fecha_iso,
              version["part_id_fuente"], version["part_id_final"],
              version["validacion_catalogo"], version["nombre_fuente"],
              version["resolucion_nombre"], version["archivo_json"], version_id)
             for version_id, _, revision, fecha_iso, version in asignaciones],
        )
        ultima_version_archivo = {}
        for _, archivo_id, _, _, version in asignaciones:
            if (archivo_id not in ultima_version_archivo
                    or version["fila"]["id"] > ultima_version_archivo[archivo_id]["fila"]["id"]):
                ultima_version_archivo[archivo_id] = version
        for archivo_id, version in ultima_version_archivo.items():
            producto = version["datos"].get("producto") or {}
            con.execute(
                """UPDATE archivos SET producto_clave=?, producto_nombre=?,
                       nombre_fuente=?, resolucion_nombre=?, part_id_fuente=?,
                       part_id_final=?, validacion_catalogo=?,
                       ultima_revision_json=(
                         SELECT v.revision_json FROM versiones v
                         WHERE v.archivo_id=archivos.id
                         ORDER BY v.revision_json DESC LIMIT 1
                       ),
                       ultimo_json=(
                         SELECT v.archivo_json FROM versiones v
                         WHERE v.archivo_id=archivos.id
                         ORDER BY v.revision_json DESC LIMIT 1
                       )
                   WHERE id=?""",
                (version["producto_clave"], producto.get("nombre", ""),
                 version["nombre_fuente"], version["resolucion_nombre"],
                 version["part_id_fuente"], version["part_id_final"],
                 version["validacion_catalogo"], archivo_id),
            )
        con.commit()
    except sqlite3.Error:
        con.rollback()
        for ruta_actual, ruta_original in reversed(renombres_aplicados):
            ruta_actual.rename(ruta_original)
        raise

    for nombre, error in errores_json:
        print(f"ADVERTENCIA: no se reordenó el historial con JSON ilegible {nombre}: {error}")
    return {
        "revisiones_reordenadas": cantidad_cambiada,
        "revisiones_sin_fecha_reconocida": fechas_no_reconocidas,
        "nombres_catalogo_normalizados": nombres_normalizados,
        "nombres_catalogo_ambiguos": nombres_ambiguos,
        "archivos_renombrados_codigo_catalogo": len(renombres_aplicados),
        "revisiones_historial_omitidas": sum(
            totales_por_producto[producto] for producto in productos_con_json_invalido
        ),
    }


def extraer(path: Path, idioma: str = "es") -> dict:
    doc = Document(path)
    encabezados = [
        limpiar(p.text) for s in doc.sections for p in s.header.paragraphs
        if limpiar(p.text) and clave(p.text) not in {"ESPECIFICACION", "SPECIFICATION"}
    ]
    producto = encabezados[0] if encabezados else ""
    if not producto:
        candidatos = [
            limpiar(p.text) for p in doc.paragraphs
            if limpiar(p.text) and clave(p.text) not in {"ESPECIFICACION", "SPECIFICATION"}
        ]
        producto = candidatos[0] if candidatos else ""
    if not producto:
        raise ValueError("No se encontro el nombre del producto.")

    tablas = []
    for i, tabla in enumerate(doc.tables):
        if tabla.rows:
            enc = [clave(c.text) for c in tabla.rows[0].cells]
            tablas.append((sum(x in COLUMNAS for x in enc), i, tabla, enc))
    if not tablas:
        raise ValueError("El documento no contiene tablas.")
    puntos, indice, tabla, enc = max(tablas, key=lambda x: x[0])
    if puntos < 3:
        raise ValueError("No se encontro la tabla de pruebas.")
    mapa = [COLUMNAS.get(x) for x in enc]
    pruebas, notas = [], []
    for fila_origen, fila in enumerate(tabla.rows[1:], 2):
        valores = [limpiar(c.text) for c in fila.cells]
        no_vacios = [v for v in valores if v]
        if not no_vacios: continue
        if len({clave(v) for v in no_vacios}) == 1 and re.search(r"\b(?:NOTA|NOTES?)\b", clave(no_vacios[0])):
            notas.append(re.sub(r"^(?:Notas?|Notes?)\s*:\s*", "", no_vacios[0], flags=re.I)); continue
        item = {"numero": len(pruebas) + 1, "fila_origen": fila_origen}
        for j, valor in enumerate(valores):
            if j < len(mapa) and mapa[j]: item[mapa[j]] = valor
        for campo in ("prueba", "unidades", "metodo", "especificacion", "valor_tipico"):
            item.setdefault(campo, "")
        pruebas.append(item)

    control, mejor, puntos_mejor = {}, None, 0
    for i, candidata in enumerate(doc.tables):
        if i == indice or len(candidata.columns) < 2: continue
        puntos = sum(clave(r.cells[0].text) in CONTROL for r in candidata.rows)
        if puntos > puntos_mejor: mejor, puntos_mejor = candidata, puntos
    if mejor:
        for fila in mejor.rows:
            etiqueta = CONTROL.get(clave(fila.cells[0].text))
            if etiqueta: control[etiqueta] = limpiar(fila.cells[1].text)

    pie = " ".join(limpiar(p.text) for s in doc.sections for p in s.footer.paragraphs if limpiar(p.text))
    return {"esquema": "especificacion_producto_es_v3", "producto": {"nombre": producto},
            "documento": {"tipo": "ESPECIFICACION", "idioma": idioma, **metadatos_formato(path)},
            "control_documental": control, "pruebas": pruebas, "notas": notas,
            "datos_institucionales": {"pie_pagina": pie},
            "estructura_detectada": {"numero_tablas": len(doc.tables), "tabla_pruebas_indice": indice,
                                      "filas_pruebas": len(pruebas),
                                      "encabezados": [limpiar(c.text) for c in tabla.rows[0].cells]}}


def validar_catalogo(
    nombre: str,
    part_id: str,
    catalogo: list[dict],
    *,
    contenido: str = "",
    permitir_coincidencia_part_id: bool = False,
) -> dict:
    nk, ik = clave(nombre), id_clave(part_id)
    # Collapse repeated spreadsheet rows, but keep distinct IDs as real candidates.
    if isinstance(catalogo, CatalogoPreparado):
        registros = catalogo.registros
        por_nombre = catalogo.por_nombre.get(nk, [])
        por_id = catalogo.por_id.get(ik, []) if ik else []
    else:
        registros = list({(r["part_id_clave"], r["descripcion_clave"]): r
            for r in catalogo if r["part_id_clave"] and r["descripcion_clave"]}.values())
        por_nombre = [r for r in registros if r["descripcion_clave"] == nk]
        por_id = [r for r in registros if ik and r["part_id_clave"] == ik]
    texto_contenido = f" {clave(contenido)} " if contenido else ""
    coincidencias_contenido = [
        r for r in registros
        if len(r["descripcion_clave"]) >= 8
        and f" {r['descripcion_clave']} " in texto_contenido
    ]
    if coincidencias_contenido:
        largo_maximo = max(len(r["descripcion_clave"]) for r in coincidencias_contenido)
        candidatos = [
            r for r in coincidencias_contenido
            if len(r["descripcion_clave"]) == largo_maximo
        ]
        fuente_coincidencia = "CONTENIDO_JSON"
    elif por_nombre:
        candidatos = por_nombre
        fuente_coincidencia = "NOMBRE_PRODUCTO"
    else:
        candidatos = []
        fuente_coincidencia = ""

    ids_candidatos = {r["part_id_clave"] for r in candidatos}
    ids_200 = {
        r["part_id_clave"] for r in candidatos
        if re.fullmatch(r"2\d{7}", r["part_id_clave"])
    }
    resultado = {"estado": "", "prefijo_archivo": "", "nombre_fuente": nombre,
                 "part_id_fuente": part_id, "part_id_final": part_id,
                 "coincidencias_nombre": len(por_nombre),
                 "coincidencias_part_id": len(por_id),
                 "coincidencias_contenido": [
                     {"part_id": r["part_id"], "descripcion": r["descripcion"]}
                     for r in candidatos
                 ],
                 "candidatos_codigo_200": sorted(ids_200)}

    seleccionado = None
    metodo = ""
    if ik and por_id:
        # A real catalog ID remains authoritative when its description is among
        # the name candidates; it must not silently change to another ID family.
        candidatos_id = [r for r in por_id if r in candidatos]
        if candidatos_id:
            seleccionado = candidatos_id[0]
            metodo = "CODIGO_CATALOGO_EXISTENTE"
        elif not candidatos and len({r["descripcion_clave"] for r in por_id}) == 1:
            seleccionado = por_id[0]
            metodo = "CODIGO_CATALOGO_EXISTENTE"
        elif not candidatos:
            resultado.update(estado="CONFLICTO_NOMBRE_PART_ID", prefijo_archivo="--- ",
                             registro_por_part_id=por_id[0])
        else:
            resultado.update(
                estado="CONFLICTO_NOMBRE_PART_ID",
                prefijo_archivo="--- ",
                registro_por_nombre=candidatos[0],
                registro_por_part_id=por_id[0],
            )
    elif len({r["descripcion_clave"] for r in candidatos}) > 1:
        resultado.update(estado="CONFLICTO_NOMBRE_AMBIGUO", prefijo_archivo="--- ")
    elif len(ids_200) == 1:
        id_200 = next(iter(ids_200))
        seleccionado = next(r for r in candidatos if r["part_id_clave"] == id_200)
        metodo = "CODIGO_200_POR_NOMBRE" if fuente_coincidencia == "NOMBRE_PRODUCTO" else "CODIGO_200_EN_CONTENIDO"
    elif len(ids_200) > 1:
        resultado.update(estado="CONFLICTO_NOMBRE_AMBIGUO", prefijo_archivo="--- ")
    elif len(ids_candidatos) == 1:
        seleccionado = candidatos[0]
        metodo = "COINCIDENCIA_UNICA_DE_CATALOGO"
    elif len(ids_candidatos) > 1:
        resultado.update(estado="CONFLICTO_NOMBRE_AMBIGUO", prefijo_archivo="--- ")
    elif ik and permitir_coincidencia_part_id and por_id:
        seleccionado = por_id[0]
        metodo = "CODIGO_CATALOGO_INGLES"

    if seleccionado:
        nombre_catalogo = seleccionado["descripcion"]
        nombre_clave_catalogo = seleccionado["descripcion_clave"]
        nombre_clave_fuente = nk
        if metodo == "CODIGO_CATALOGO_EXISTENTE":
            estado = "COINCIDENCIA_EXACTA" if nombre_clave_fuente == nombre_clave_catalogo else "NOMBRE_CANONICO_POR_CODIGO"
        elif metodo in {"CODIGO_200_POR_NOMBRE", "CODIGO_200_EN_CONTENIDO"}:
            estado = "NOMBRE_ANTIGUO_MAPEADO_CODIGO_200"
        elif metodo == "COINCIDENCIA_UNICA_DE_CATALOGO":
            estado = "COINCIDENCIA_UNICA_EN_CATALOGO" if por_nombre else "NOMBRE_RECUPERADO_DESDE_CONTENIDO"
        else:
            estado = "COINCIDENCIA_PART_ID"
        resultado.update(
            estado=estado,
            prefijo_archivo="",
            part_id_final=seleccionado["part_id"],
            nombre_catalogo=nombre_catalogo,
            resolucion_nombre=metodo,
            registro_catalogo=seleccionado,
        )
        if not ik and metodo == "COINCIDENCIA_UNICA_DE_CATALOGO":
            resultado["prefijo_archivo"] = "&&& "
    elif not resultado["estado"]:
        resultado.update(estado="NO_ENCONTRADO_EN_CATALOGO", prefijo_archivo="--- ")
    return resultado


def procesar(
    base: Path,
    archivos_seleccionados: list[Path] | None = None,
    al_progreso: Callable[[int, int, str], None] | None = None,
    json_generados: list[Path] | None = None,
) -> dict:
    from rutas_proyecto import resolver_ruta, ruta_portatil, rutas_configuradas
    base = resolver_ruta(base)
    _restablecer_error_inicio_word()
    rutas = rutas_configuradas(base)
    entrada, salida = rutas["especificaciones_viejas"], rutas["json_anteriores"]
    entrada_eng = entrada.parent / ENTRADA_ENG
    entrada.mkdir(parents=True, exist_ok=True)
    entrada_eng.mkdir(parents=True, exist_ok=True)
    if archivos_seleccionados is not None:
        archivos_seleccionados = [resolver_ruta(ruta, base) for ruta in archivos_seleccionados]
    if archivos_seleccionados is None:
        archivos = sorted(
            ((p, raiz == entrada_eng) for raiz in (entrada, entrada_eng)
             for p in raiz.rglob("*")
             if p.is_file() and p.suffix.lower() in EXTENSIONES and not p.name.startswith("~$")),
            key=lambda item: str(item[0]),
        )
    else:
        archivos = sorted(
            (
                (path, path.resolve().is_relative_to(entrada_eng.resolve()))
                for path in archivos_seleccionados
                if path.is_file()
                and path.suffix.lower() in EXTENSIONES
                and not path.name.startswith("~$")
            ),
            key=lambda item: str(item[0]),
        )
    total_archivos = len(archivos)

    salida.mkdir(parents=True, exist_ok=True)
    con = conectar(base / DB)
    ahora_inicio = datetime.now().isoformat(timespec="seconds")
    catalogo, hash_catalogo = obtener_catalogo(con, base / CATALOGO)
    ids_catalogo = {r["part_id_clave"] for r in catalogo if r["part_id_clave"]}
    registrar_catalogo(con, catalogo, hash_catalogo, ahora_inicio, base / CATALOGO)
    catalogo = CatalogoPreparado(catalogo)
    huella_historial = f"{APP_VERSION}:{hash_catalogo}"
    estado_historial = con.execute("SELECT valor FROM gestor_estado WHERE clave='historial_revisado'").fetchone()
    revisar_todo = estado_historial is None or estado_historial["valor"] != huella_historial
    resumen = {"nuevos": 0, "actualizados": 0, "ignorados": 0, "bloqueados": 0,
               "movidos_eng": 0,
               "movidos_revisar": 0, "conflictos": 0, "completados": 0,
               "yves_yoi_a_revisar": 0, "errores": 0, "revisiones_reordenadas": 0}
    productos_afectados = productos_historial_modificados(con, salida) if not revisar_todo else set()
    archivos_json_reindexados: list[int] = []

    def informar_progreso(indice: int, ruta: str) -> None:
        if al_progreso is not None:
            al_progreso(indice, total_archivos, ruta)

    for indice_archivo, (origen, ya_esta_en_ingles) in enumerate(archivos, 1):
        rel = ruta_portatil(origen, base)
        ahora = datetime.now(); fecha_iso = ahora.isoformat(timespec="seconds")
        digest = ""
        stat = None
        print(f"PROCESANDO: {rel}", flush=True)
        try:
            bloqueo = archivo_bloqueado_word(origen)
            if bloqueo is not None:
                resumen["bloqueados"] += 1
                print(
                    f"OMITIDO (Word bloqueado): {rel}; archivo de bloqueo: "
                    f"{bloqueo.name}",
                    flush=True,
                )
                informar_progreso(indice_archivo, rel)
                continue
            digest = sha256(origen)
            stat = origen.stat()
            registrado = con.execute(
                """SELECT id FROM archivos
                   WHERE hash_actual=? AND hash_catalogo=? AND estado='OK'
                   ORDER BY id DESC LIMIT 1""",
                (digest, hash_catalogo),
            ).fetchone()
            if registrado:
                resumen["ignorados"] += 1
                print(f"IGNORADO (sin cambios): {rel}", flush=True)
                informar_progreso(indice_archivo, rel)
                continue
            with tempfile.TemporaryDirectory(prefix="espec_") as td, lectura_word_compartida():
                lectura, convertido = preparar(origen, Path(td))
                menciones = menciones_yves_yoi(lectura)
                if menciones:
                    ruta_original = rel
                    destino_revisar = destino_movimiento_sin_colision(
                        base / REVISAR / origen.name
                    )
                    destino_revisar.parent.mkdir(exist_ok=True)
                    shutil.move(str(origen), str(destino_revisar))
                    origen = destino_revisar
                    rel = ruta_portatil(origen, base)
                    detalle = ", ".join(menciones)
                    con.execute(
                        """UPDATE archivos SET ruta_relativa=?
                           WHERE ruta_relativa=?
                           AND NOT EXISTS (SELECT 1 FROM archivos WHERE ruta_relativa=?)""",
                        (rel, ruta_original, rel),
                    )
                    con.execute(
                        """INSERT INTO archivos
                          (ruta_relativa,nombre_archivo,hash_actual,hash_catalogo,tamano_bytes,
                           fecha_modificacion_ns,primera_deteccion,ultima_revision_procesada,
                           estado,ultimo_error)
                          VALUES(?,?,?,?,?,?,?,?,?,?)
                          ON CONFLICT(ruta_relativa) DO UPDATE SET
                           nombre_archivo=excluded.nombre_archivo,hash_actual=excluded.hash_actual,
                           hash_catalogo=excluded.hash_catalogo,
                           tamano_bytes=excluded.tamano_bytes,
                           fecha_modificacion_ns=excluded.fecha_modificacion_ns,
                           ultima_revision_procesada=excluded.ultima_revision_procesada,
                           estado=excluded.estado,ultimo_error=excluded.ultimo_error""",
                        (rel, origen.name, digest, hash_catalogo, stat.st_size, stat.st_mtime_ns,
                         fecha_iso, fecha_iso, "REVISAR_YVES_YOI",
                         f"El documento contiene la(s) palabra(s): {detalle}"),
                    )
                    con.commit()
                    resumen["movidos_revisar"] += 1
                    resumen["yves_yoi_a_revisar"] += 1
                    print(f"MOVIDO A REVISAR (Yves/YOI): {rel}", flush=True)
                    informar_progreso(indice_archivo, rel)
                    continue
                idioma = "en" if ya_esta_en_ingles or contiene_marcadores_ingles(lectura) else "es"
                datos = extraer(lectura, idioma)
            if idioma == "en" and not ya_esta_en_ingles:
                destino_eng = destino_movimiento_sin_colision(entrada_eng / origen.name)
                shutil.move(str(origen), str(destino_eng))
                origen = destino_eng
                rel = ruta_portatil(origen, base)
                con.commit()
                resumen["movidos_eng"] += 1
                print(f"MOVIDO ENG: {rel}")
            anterior = con.execute("SELECT * FROM archivos WHERE ruta_relativa=?", (rel,)).fetchone()
            if anterior is None and idioma == "en":
                anterior_movido = con.execute(
                    """SELECT id FROM archivos
                       WHERE nombre_archivo=? AND hash_actual=? AND estado='MOVIDO_ENG'
                       ORDER BY id DESC LIMIT 1""",
                    (origen.name, digest),
                ).fetchone()
                if anterior_movido:
                    con.execute("UPDATE archivos SET ruta_relativa=? WHERE id=?", (rel, anterior_movido["id"]))
                    con.commit()
                    anterior = con.execute("SELECT * FROM archivos WHERE ruta_relativa=?", (rel,)).fetchone()
            nombre = datos["producto"]["nombre"]
            part_fuente = datos["control_documental"].get("codigo_unico_producto", "")
            fecha_documento = parsear_fecha_documento(
                str(datos["control_documental"].get("fecha") or "")
            )
            fecha_documento_iso = fecha_documento.isoformat() if fecha_documento else None
            datos["documento"]["fecha_documento_normalizada"] = fecha_documento_iso or ""
            validacion = validar_catalogo(
                nombre, part_fuente, catalogo, contenido=json.dumps(datos, ensure_ascii=False),
                permitir_coincidencia_part_id=idioma == "en"
            )
            part_final = validacion["part_id_final"]
            nombre_final = validacion.get("nombre_catalogo", nombre)
            if nombre_final != nombre:
                datos["producto"]["nombre_fuente"] = nombre
                datos["producto"]["nombre"] = nombre_final
            if part_final and id_clave(part_final) != id_clave(part_fuente):
                datos["control_documental"]["codigo_unico_producto_fuente"] = part_fuente
            if part_final:
                datos["control_documental"]["codigo_unico_producto"] = part_final
            datos["validacion_codigos_200"] = validacion
            pk = clave(nombre_final)
            productos_afectados.add(pk)
            if anterior and anterior["producto_clave"]:
                productos_afectados.add(anterior["producto_clave"])
            ultima = con.execute("SELECT MAX(revision_json) m FROM versiones WHERE producto_clave=?", (pk,)).fetchone()["m"]
            revision = 0 if ultima is None else ultima + 1
            rev_esp = datos["control_documental"].get("revision", "")
            prefijo = validacion["prefijo_archivo"]
            codigo_archivo = part_final if id_clave(part_final) in ids_catalogo else ""
            nombre_base = nombre_archivo_json(codigo_archivo, rev_esp, ahora)
            archivo_json = nombre_base
            sufijo = 1
            while (salida / archivo_json).exists() or con.execute(
                "SELECT 1 FROM versiones WHERE archivo_json=?", (archivo_json,)
            ).fetchone():
                archivo_json = f"{Path(nombre_base).stem}_{sufijo:02d}.JSON"
                sufijo += 1
            destino = salida / archivo_json
            datos.update(archivo_origen=rel, hash_sha256_origen=digest, hash_sha256_catalogo=hash_catalogo,
                         formato_origen=origen.suffix[1:].lower(), conversion_temporal=convertido,
                         revision_json=f"{revision:02d}", fecha_creacion_json=fecha_iso,
                         archivo_json=archivo_json, extractor_version=APP_VERSION)
            destino.write_text(json.dumps(datos, ensure_ascii=False, indent=2), encoding="utf-8")
            cod_fmt = datos["documento"].get("codigo_formato", "")
            rev_fmt = datos["documento"].get("revision_formato", "")
            con.execute("""INSERT INTO archivos
              (ruta_relativa,nombre_archivo,hash_actual,hash_catalogo,tamano_bytes,fecha_modificacion_ns,
               producto_clave,producto_nombre,nombre_fuente,resolucion_nombre,part_id_fuente,part_id_final,validacion_catalogo,
               revision_especificacion,codigo_formato,revision_formato,formato_origen,idioma,
               ultima_revision_json,ultimo_json,primera_deteccion,ultima_revision_procesada,estado,ultimo_error)
              VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,'OK',NULL)
              ON CONFLICT(ruta_relativa) DO UPDATE SET nombre_archivo=excluded.nombre_archivo,
               hash_actual=excluded.hash_actual,hash_catalogo=excluded.hash_catalogo,tamano_bytes=excluded.tamano_bytes,
               fecha_modificacion_ns=excluded.fecha_modificacion_ns,producto_clave=excluded.producto_clave,
               producto_nombre=excluded.producto_nombre,nombre_fuente=excluded.nombre_fuente,
               resolucion_nombre=excluded.resolucion_nombre,part_id_fuente=excluded.part_id_fuente,
               part_id_final=excluded.part_id_final,validacion_catalogo=excluded.validacion_catalogo,
               revision_especificacion=excluded.revision_especificacion,codigo_formato=excluded.codigo_formato,
               revision_formato=excluded.revision_formato,formato_origen=excluded.formato_origen,
               idioma=excluded.idioma,ultima_revision_json=excluded.ultima_revision_json,ultimo_json=excluded.ultimo_json,
               ultima_revision_procesada=excluded.ultima_revision_procesada,estado='OK',ultimo_error=NULL""",
              (rel, origen.name, digest, hash_catalogo, stat.st_size, stat.st_mtime_ns, pk, nombre_final,
               nombre if nombre_final != nombre else "", validacion.get("resolucion_nombre", ""),
               part_fuente, part_final, validacion["estado"], rev_esp, cod_fmt, rev_fmt,
               origen.suffix[1:].lower(), idioma.upper(), revision, archivo_json, fecha_iso, fecha_iso))
            archivo_id = con.execute("SELECT id FROM archivos WHERE ruta_relativa=?", (rel,)).fetchone()["id"]
            con.execute("""INSERT INTO versiones
              (archivo_id,producto_clave,revision_json,hash_fuente,hash_catalogo,part_id_fuente,
               part_id_final,validacion_catalogo,nombre_fuente,resolucion_nombre,revision_especificacion,codigo_formato,
               revision_formato,formato_origen,idioma,fecha_documento,fecha_creacion,archivo_json)
              VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
              (archivo_id, pk, revision, digest, hash_catalogo, part_fuente, part_final,
               validacion["estado"], nombre if nombre_final != nombre else "",
               validacion.get("resolucion_nombre", ""), rev_esp, cod_fmt, rev_fmt, origen.suffix[1:].lower(),
               idioma.upper(), fecha_documento_iso, fecha_iso, archivo_json))
            con.commit()
            archivos_json_reindexados.append(archivo_id)
            resumen["actualizados" if anterior else "nuevos"] += 1
            if prefijo == "--- ": resumen["conflictos"] += 1
            if prefijo == "&&& ": resumen["completados"] += 1
            print(f"GENERADO: {archivo_json} [{validacion['estado']}]")
            informar_progreso(indice_archivo, rel)
        except Exception as exc:
            con.rollback()
            ruta_original = rel
            try:
                carpeta_revisar = base / REVISAR
                carpeta_revisar.mkdir(exist_ok=True)
                if origen.exists():
                    destino_revisar = destino_movimiento_sin_colision(carpeta_revisar / origen.name)
                    shutil.move(str(origen), str(destino_revisar))
                    origen = destino_revisar
                    rel = ruta_portatil(origen, base)
                    resumen["movidos_revisar"] += 1
                    print(f"MOVIDO A REVISAR: {rel}", flush=True)
            except OSError as error_movimiento:
                exc = RuntimeError(f"{exc}; no se pudo mover a {REVISAR}: {error_movimiento}")

            if ruta_original != rel:
                con.execute(
                    """UPDATE archivos SET ruta_relativa=?
                       WHERE ruta_relativa=?
                       AND NOT EXISTS (SELECT 1 FROM archivos WHERE ruta_relativa=?)""",
                    (rel, ruta_original, rel),
                )
            con.execute("""INSERT INTO archivos
              (ruta_relativa,nombre_archivo,hash_actual,hash_catalogo,tamano_bytes,fecha_modificacion_ns,
               primera_deteccion,ultima_revision_procesada,estado,ultimo_error)
              VALUES(?,?,?,?,?,?,?,?, 'ERROR',?)
              ON CONFLICT(ruta_relativa) DO UPDATE SET hash_actual=excluded.hash_actual,
               hash_catalogo=excluded.hash_catalogo,ultima_revision_procesada=excluded.ultima_revision_procesada,
               estado='ERROR',ultimo_error=excluded.ultimo_error""",
              (rel, origen.name, digest or "ERROR", hash_catalogo,
               stat.st_size if stat else 0, stat.st_mtime_ns if stat else 0,
               fecha_iso, fecha_iso, str(exc)))
            con.commit(); resumen["errores"] += 1; print(f"ERROR: {rel}: {exc}")
            informar_progreso(indice_archivo, rel)
    try:
        if not revisar_todo:
            productos_afectados = ampliar_productos_afectados(con, salida, catalogo, productos_afectados)
        resumen_revisiones = reindexar_revisiones(
            con,
            salida,
            catalogo,
            None if revisar_todo else productos_afectados,
        )
        if resumen_revisiones["revisiones_historial_omitidas"] == 0:
            import generador_hii as hii
            with hii.conectar(base / DB) as relaciones:
                hii.sincronizar_idiomas(relaciones, salida)
            registrar_historial_revisado(con, salida, None if revisar_todo else productos_afectados)
            con.execute("""INSERT INTO gestor_estado VALUES ('historial_revisado', ?)
                ON CONFLICT(clave) DO UPDATE SET valor=excluded.valor""", (huella_historial,))
            con.commit()
        resumen.update(resumen_revisiones)
        resumen["errores"] += resumen_revisiones["revisiones_historial_omitidas"]
        if json_generados is not None and archivos_json_reindexados:
            placeholders = ",".join("?" for _ in archivos_json_reindexados)
            filas_generadas = con.execute(
                f"""SELECT ultimo_json FROM archivos WHERE id IN ({placeholders})
                    AND ultimo_json IS NOT NULL""",
                tuple(archivos_json_reindexados),
            ).fetchall()
            json_generados[:] = [salida / fila["ultimo_json"] for fila in filas_generadas]
    finally:
        con.close()
    return resumen


def ver_registro(base: Path) -> None:
    con = conectar(base / DB)
    for r in con.execute("""SELECT estado,ruta_relativa,ultima_revision_json,producto_nombre,
      nombre_fuente,resolucion_nombre,part_id_fuente,part_id_final,validacion_catalogo,codigo_formato,revision_formato,
      revision_especificacion,idioma FROM archivos ORDER BY ruta_relativa"""):
        rev = "--" if r["ultima_revision_json"] is None else f"{r['ultima_revision_json']:02d}"
        print(f"{r['estado']} | idioma={r['idioma'] or '--'} | JSON {rev} | {r['validacion_catalogo'] or '--'} | "
              f"nombre fuente={r['nombre_fuente'] or '--'} método={r['resolucion_nombre'] or '--'} | "
              f"Part ID fuente={r['part_id_fuente'] or '--'} final={r['part_id_final'] or '--'} | "
              f"formato={r['codigo_formato'] or '--'} rev={r['revision_formato'] or '--'} | {r['ruta_relativa']}")
    con.close()


def main() -> int:
    if sys.version_info[:2] != PYTHON_REQUERIDO:
        raise SystemExit(f"Se requiere Python 3.11. Detectado {sys.version_info.major}.{sys.version_info.minor}.")
    parser = argparse.ArgumentParser()
    parser.add_argument("--ver-registro", action="store_true")
    args = parser.parse_args(); base = base_app()
    if args.ver_registro: ver_registro(base); return 0
    try:
        resumen = procesar(base)
    except Exception as exc:
        print(f"ERROR: {exc}")
        return 1
    print("\nRESUMEN")
    for k, v in resumen.items(): print(f"  {k}: {v}")
    return 1 if resumen["errores"] else 0


if __name__ == "__main__":
    multiprocessing.freeze_support()
    raise SystemExit(main())
